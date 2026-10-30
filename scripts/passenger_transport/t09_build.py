"""T09 — ساخت فرم‌های عملیاتی پیش‌نویس (Word، PDF) و گزارش داخلی.

اجرا (پایتونی که python-docx، openpyxl، pywin32 و PyMuPDF دارد):
    py -3.14 scripts/passenger_transport/t09_build.py            # همه‌چیز
    py -3.14 scripts/passenger_transport/t09_build.py --md-only  # فقط گزارش داخلی

ورودی‌ها فقط خوانده می‌شوند:
- کتاب‌کار T08: data/deliverables/passenger_transport/event_relationship_design_v0_1.xlsx
- پایگاه دفاتر: data/work/taxi_terminals/tt_v21.sqlite (mode=ro، کپی در حافظه)
- محتوای فرم‌ها: t09_content.py
خروجی‌ها:
- data/deliverables/passenger_transport/operational_forms_draft_v0_1.docx و .pdf
- docs/project/passenger_transport_form_design_v0_1.md
- data/work/passenger_transport/t09/ (کنترل‌ها، تصویر صفحه‌ها)
"""
from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import sys
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import t09_content as C  # noqa: E402

ROOT = HERE.parents[1]
T08 = ROOT / "data/deliverables/passenger_transport/event_relationship_design_v0_1.xlsx"
DB = ROOT / "data/work/taxi_terminals/tt_v21.sqlite"
OUT_DOCX = ROOT / "data/deliverables/passenger_transport/operational_forms_draft_v0_1.docx"
OUT_PDF = OUT_DOCX.with_suffix(".pdf")
OUT_MD = ROOT / "docs/project/passenger_transport_form_design_v0_1.md"
WORK = ROOT / "data/work/passenger_transport/t09"
VISUAL = WORK / "visual_check.json"

DATE_FA = "۱۴۰۵/۰۸/۰۸"
FA_DIG = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


def fa(x) -> str:
    """عدد یا متن با ارقام فارسی؛ عدد صحیح با جداکنندهٔ هزارگان."""
    if isinstance(x, int):
        # جداکنندهٔ «٬» در B Nazanin با صفر فارسی اشتباه می‌شود؛ ویرگول لاتین خواناتر است.
        return f"{x:,}".translate(FA_DIG)
    return str(x).translate(FA_DIG)


def sha16(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


# ---------------------------------------------------------------------------
# T08
# ---------------------------------------------------------------------------

def load_t08():
    wb = openpyxl.load_workbook(T08, read_only=True)
    fm = OrderedDict()
    for r in wb["فیلدهای_فرم"].iter_rows(min_row=4, values_only=True):
        if r[0]:
            fm[r[0]] = dict(form=r[1], form_title=r[2], label=r[4], el=r[5], req=r[11], cond=r[12], dec=r[16])
    els = OrderedDict()
    for r in wb["عناصر_داده"].iter_rows(min_row=4, values_only=True):
        if r[0]:
            els[r[0]] = dict(name=r[1], dec=r[10])
    pv = []
    for r in wb["نیاز_داده_مسیرها"].iter_rows(min_row=4, values_only=True):
        if r[0]:
            pv.append(dict(id=r[0], flow=r[1], lp=r[2].split(" — ")[0], stage=r[3], ev=r[4], sec=r[5], el=r[6],
                           label=r[7], req=r[10], cond=r[11], prof=r[16]))
    lps = OrderedDict()
    for r in wb["الگوهای_چرخه"].iter_rows(min_row=4, values_only=True):
        if r[0]:
            lps[r[0]] = dict(flow=r[1], stage=r[3], ev=r[4])
    wb.close()
    return fm, els, pv, lps


FM, ELS, PV, LPS = load_t08()


def items(form):
    for sec in form["sections"]:
        for it in sec["items"]:
            yield sec, it


def all_items():
    for form in C.FORMS:
        for sec, it in items(form):
            yield form, sec, it


def lp_of(it):
    out = []
    for g in it["lp"]:
        out += C.LPG.get(g, [g])
    return sorted(set(out))


def pv_of(it):
    L = set(lp_of(it))
    E = set(it["el"])
    doc = it["kind"] == "check" or it.get("vtype") == "up"
    return [p for p in PV if p["lp"] in L and (p["el"] in E or (doc and p["el"] == "—"))]


def t08_reqs(it):
    s = set()
    for f in it["fm"]:
        s.add(FM[f]["req"])
    for p in pv_of(it):
        for part in str(p["req"]).split("/"):
            s.add(part)
    return s


def disp_fid(fid: str) -> str:
    head, num = fid.split("-")
    head = {"FC": "ک"}.get(head, head.replace("F", ""))
    return fa(f"{head}-{num}")


# ---------------------------------------------------------------------------
# کنترل‌ها
# ---------------------------------------------------------------------------

F6_SYSTEM_ELS = {"EL-01", "EL-13", "EL-70", "EL-71", "EL-68", "EL-69", "EL-72", "EL-73", "EL-74", "EL-75", "EL-76", "EL-40"}


def run_controls():
    res = []

    def add(cid, title, ok, detail):
        res.append(dict(id=cid, title=title, status="PASS" if ok else "FAIL", detail=detail))

    # K-T01 شناسه‌های ارجاع‌شده وجود دارند
    bad = []
    for form, sec, it in all_items():
        for f in it["fm"]:
            if f not in FM:
                bad.append(f"{it['fid']}:{f}")
        for e in it["el"]:
            if e not in ELS:
                bad.append(f"{it['fid']}:{e}")
        for lp in lp_of(it):
            if lp not in LPS:
                bad.append(f"{it['fid']}:{lp}")
    add("K-T01", "همهٔ شناسه‌های FM، EL و LP ارجاع‌شده در T08 وجود دارند", not bad, "; ".join(bad) or "۰ ارجاع نامعتبر")

    # K-T02 هر فیلد به T08 یا دلیل طراحی نگاشت دارد
    orphan = [it["fid"] for _, _, it in all_items() if not it["fm"] and not pv_of(it) and not it.get("why")]
    n = sum(1 for _ in all_items())
    add("K-T02", "هر فیلد به فیلد T08، ردیف نیاز مسیر یا دلیل طراحی صریح نگاشت دارد", not orphan,
        f"{n} قلم؛ بی‌نگاشت: {', '.join(orphan) or '۰'}")

    # K-T03 سطح نیاز سخت‌تر از پیشنهاد T08 نیست
    viol = []
    for _, _, it in all_items():
        need, s, why = it["need"], t08_reqs(it), it.get("why")
        if need in ("start", "review") and "لازم" not in s and not why:
            viol.append(f"{it['fid']}({need}:{'/'.join(sorted(s)) or '—'})")
        if need == "cond" and not (s & {"شرطی", "لازم"}) and not why:
            viol.append(f"{it['fid']}(cond:{'/'.join(sorted(s)) or '—'})")
        if need == "dec" and "منوط به تصمیم" not in s and not why:
            viol.append(f"{it['fid']}(dec:{'/'.join(sorted(s)) or '—'})")
        if need == "opt" and "اختیاری" not in s and not why:
            viol.append(f"{it['fid']}(opt:{'/'.join(sorted(s)) or '—'})")
        if s == {"منوط به تصمیم"} and need not in ("dec", "auto") and not why:
            viol.append(f"{it['fid']}(stricter than decision-only)")
    add("K-T03", "سطح نیاز هر فیلد از «پیشنهاد — الزام» T08 سخت‌تر نیست (یا دلیل طراحی دارد)", not viol,
        "; ".join(viol) or "۰ تخطی")

    # K-T04 پوشش ردیف‌های نیاز مسیر در مرحله‌های در دامنه
    cover = defaultdict(set)
    for form, sec, it in all_items():
        for p in pv_of(it):
            if form["id"] == "F6" and p["el"] not in F6_SYSTEM_ELS:
                continue
            cover[(p["lp"], p["el"])].add(it["fid"])
    groups = OrderedDict()
    for p in PV:
        if p["lp"] in C.IN_SCOPE_LP:
            groups.setdefault((p["lp"], p["el"]), []).append(p)
    uncovered = [(k, v) for k, v in groups.items() if k not in cover]
    unjust = [k for k, _ in uncovered if k not in C.UNCOVERED_WHY]
    unc_lines = [f"{lp}/{el}" + ("" if (lp, el) in C.UNCOVERED_WHY else " (بی‌دلیل)") for (lp, el), v in uncovered]
    add("K-T04", "ردیف‌های «نیاز_داده_مسیرها» در مرحله‌های در دامنه پوشش دارند (فرم‌های ۱ تا ۵ و کارت؛ فرم ۶ فقط برای عناصر بسته و سطر) یا با دلیل بیرون مانده‌اند",
        not unjust, f"{len(groups)} گروه (مرحله × عنصر)، {sum(len(v) for v in groups.values())} ردیف PV؛ "
        f"پوشش‌دار {len(groups) - len(uncovered)}؛ بی‌فیلد با دلیل: " + ("; ".join(unc_lines) or "۰"))

    # K-T05 فیلدهای T08 یا نگاشت دارند یا در فهرست بیرون از دامنه‌اند
    used = {f for _, _, it in all_items() for f in it["fm"]}
    omitted = [f for f in FM if f not in used]
    allowed_prefix = ("FM09-", "FM10-", "FM11-")
    bad_om = [f for f in omitted if not (f.startswith(allowed_prefix) or f == "FM06-10")]
    add("K-T05", "هر فیلد T08 یا در فرم‌های T09 نگاشت دارد یا با دلیل بیرون مانده است", not bad_om,
        f"نگاشت‌شده {len(used)} از {len(FM)}؛ بیرون {len(omitted)}؛ بیرونِ بی‌دلیل: {', '.join(bad_om) or '۰'}")

    # K-T06 هیچ نیاز «شروع/بازبینی» فقط بر ردیف دارای پروفایل تقریبی تکیه ندارد
    approx = []
    for _, _, it in all_items():
        if it["need"] in ("start", "review") and not it["fm"] and not it.get("why"):
            ps = [p for p in pv_of(it) if p["req"] and "لازم" in str(p["req"])]
            if ps and all("تقریبی" in str(p["prof"]) for p in ps):
                approx.append(it["fid"])
    add("K-T06", "هیچ فیلد «برای شروع» یا «پیش از بازبینی» فقط بر ردیف دارای پروفایل تقریبی تکیه ندارد", not approx,
        ", ".join(approx) or "۰ فیلد؛ سطح نیاز فقط از ستون «پیشنهاد — الزام» خوانده شد، نه از ستون پروفایل")
    return res, uncovered, omitted


# ---------------------------------------------------------------------------
# شواهد دفتر
# ---------------------------------------------------------------------------

def open_db():
    src = sqlite3.connect(f"file:{DB.as_posix()}?mode=ro", uri=True)
    con = sqlite3.connect(":memory:")
    src.backup(con)
    src.close()
    for s in ("create index ix_doc on sanad(doc_key)", "create index ix_ov on sanad(org, vahed)",
              "create index ix_oh on sanad(org, hcode)", "create index ix_le on line_event(sid)"):
        con.execute(s)
    return con


def mask_digits(s: str, keep: int = 3) -> str:
    s = str(s)
    return "•••" + fa(s[-keep:]) if len(s) > keep else fa(s)


def mask_hcode(h: str) -> str:
    """کد حساب بانکی بدون شمارهٔ حساب: 111-2-2-4695607 → 111-2-2 (سه رقم آخر در عنوان می‌آید)."""
    parts = h.split("-")
    if parts and parts[0] == "111" and len(parts) == 4 and len(parts[3]) > 3:
        return "-".join(parts[:3])
    return h


def bank_suffix(h: str) -> str:
    parts = h.split("-")
    if parts and parts[0] == "111" and len(parts) == 4 and len(parts[3]) > 3:
        return f" — حساب با پایان {fa(parts[3][-3:])} (شماره پوشانده شد)"
    return ""


def mask_title(t: str) -> str:
    return re.sub(r"\d{5,}", lambda m: f"حساب با پایان {fa(m.group(0)[-3:])}", t or "")


def title_of(con, org, fy, hcode):
    r = con.execute("select name from hesab where org=? and fy=? and hcode=?", (org, fy, hcode)).fetchone()
    if not r:
        return "—"
    t = r[0]
    if hcode.startswith("111-"):
        # عنوان حساب بانکی شمارهٔ حساب و شبا دارد؛ فقط نام بانک می‌ماند.
        t = re.sub(r"\s+", " ", re.sub(r"(حساب\s*)?\d{5,}|IR\d+", " ", t)).strip()
        return t + bank_suffix(hcode)
    return mask_title(t)


def sample_a(con):
    q = lambda s, *a: con.execute(s, a).fetchall()  # noqa: E731
    series_docs = [d for (d,) in q("""select doc_key from sanad where org='Taxi' and fy='1404' and hcode='612-5-14-0'
        and cast(bed as real)>0 and sharh like '%انتظامات%' group by doc_key order by min(date), cast(min(sanad_num) as int)""")]
    series = []
    for d in series_docs:
        ls = q("""select s.hcode, s.bed, s.bes, s.person, s.markaz, s.date from sanad s where s.doc_key=? and s.sharh like '%انتظامات%'
                  order by s.src_row""", d)
        g = sum(b for h, b, c, p, m, dt in ls if h == "612-5-14-0")
        m114 = sum(b for h, b, c, p, m, dt in ls if h == "114-9-1-0")
        ins = sum(c for h, b, c, p, m, dt in ls if h == "212-6-4-0")
        gp = sum(c for h, b, c, p, m, dt in ls if h == "212-6-3-0")
        pay = sum(c for h, b, c, p, m, dt in ls if h == "211-1-3-0")
        bank_any = q("select count(*) from sanad where doc_key=? and hcode like '111-%' and cast(bes as real)>0", d)[0][0]
        codes = sorted({p for h, b, c, p, m, dt in ls if h in ("212-6-4-0", "211-1-3-0")})
        acts = sorted({m for h, b, c, p, m, dt in ls if h == "612-5-14-0"})
        series.append(dict(doc=d, date=ls[0][5], gross=g, m114=m114, ins=ins, gp=gp, pay=pay, bank=bank_any, codes=codes, acts=acts))
    period = [s for s in series if s["pay"] > 0]
    assert len(period) == 1, "انتظار یک سند با بستانکار پرداختنی"
    per = period[0]
    cp = per["codes"][0]
    pay_doc = q("""select doc_key from sanad where org='Taxi' and hcode='211-1-3-0' and person=? and cast(bed as real)=?
                  and date>=? order by date limit 1""", cp, per["pay"], per["date"])[0][0]
    rows = []
    for d in (per["doc"], pay_doc):
        for (src_row, dt, h, b, c, p, m) in q("""select src_row, date, hcode, bed, bes, person, markaz from sanad where doc_key=? order by src_row""", d):
            rows.append(dict(doc=d, row=src_row, date=dt, hcode=mask_hcode(h), title=title_of(con, "Taxi", "1404", h),
                             bed=int(b or 0), bes=int(c or 0), person=p, markaz=m))
    pay_row = [r for r in rows if r["doc"] == pay_doc and r["hcode"].startswith("111-")][0]
    f1 = q("select f1 from sanad where doc_key=? and hcode like '111-%'", pay_doc)[0][0]
    exp = [r for r in rows if r["doc"] == per["doc"] and r["hcode"] == "612-5-14-0"]
    # همان کد طرف‌حساب در اسناد هزینهٔ دیگر با شرحی جز «انتظامات»
    other = q("""select count(distinct s.doc_key) from sanad s where s.org='Taxi' and s.fy='1404' and s.person=?
                 and s.hcode in ('211-1-3-0','212-6-4-0') and cast(s.bes as real)>0 and s.sharh not like '%انتظامات%'
                 and s.doc_key in (select doc_key from sanad where org='Taxi' and hcode='612-5-14-0')""", cp)[0][0]
    return dict(series=series, per=per, pay_doc=pay_doc, rows=rows, pay_row=pay_row, cheque_mask=mask_digits(f1),
                cp=cp, exp=exp, other_docs=other, bank_title=pay_row["title"])


def sample_b(con):
    q = lambda s, *a: con.execute(s, a).fetchall()  # noqa: E731
    rows = q("""select s.vahed, s.hcode, e.own_event, e.lifecycle_action, s.person, case when cast(s.bed as real)>0 then 'D' else 'C' end
        from sanad s join line_event e on e.sid=s.sid join doc_class c on c.doc_key=s.doc_key
        where s.org='Terminals' and s.fy='1404' and s.vahed like '___-001-___' and c.kind='REGULAR'""")
    st = defaultdict(set)
    persons = defaultdict(set)
    for v, h, oe, la, p, dc in rows:
        persons[v].add(p)
        if h.startswith("113-1-4") and dc == "D" and oe == "TT-NE-15":
            st[v].add("claim")
        if h.startswith("113-1-4") and dc == "C":
            st[v].add("claim_settle")
        if la == "RECEIVE_DEPOSIT":
            st[v].add("deposit_in")
        if la == "REFUND_DEPOSIT":
            st[v].add("deposit_out")
        if la == "RECEIVE_INSTRUMENT":
            st[v].add("chq_recv")
        if la == "SEND_FOR_COLLECTION":
            st[v].add("chq_send")
        if la == "CLEAR_COLLECTED":
            st[v].add("chq_clear")
    full = sorted(v for v in st if len(st[v]) == 7 and len(persons[v]) == 1)
    V = full[0]
    name_v = q("select distinct name_vahed from sanad where org='Terminals' and vahed=?", V)[0][0]
    person = next(iter(persons[V]))

    def lines(where, *a):
        return q(f"""select s.doc_key, s.src_row, s.date, s.hcode, s.bed, s.bes, e.own_event, e.lifecycle_action, s.f1, s.f2, s.f3, s.vahed, c.kind
            from sanad s join line_event e on e.sid=s.sid join doc_class c on c.doc_key=s.doc_key
            where s.org='Terminals' and {where} order by s.date, cast(s.sanad_num as int), s.src_row""", *a)

    claims = [r for r in lines("s.fy='1404' and s.vahed=? and s.hcode='113-1-4-0' and e.own_event='TT-NE-15'", V)]
    recv = [r for r in lines("s.fy='1404' and s.vahed=? and e.lifecycle_action='RECEIVE_INSTRUMENT' and s.hcode like '113-5-%'", V)]
    recv_cr = [r for r in lines("s.fy='1404' and s.vahed=? and e.lifecycle_action='RECEIVE_INSTRUMENT' and s.hcode='213-2-4-2'", V)]
    send = [r for r in lines("s.fy='1404' and s.vahed=? and e.lifecycle_action='SEND_FOR_COLLECTION' and s.hcode like '113-6-%'", V)]
    clear = [r for r in lines("s.fy='1404' and s.vahed=? and e.lifecycle_action='CLEAR_COLLECTED'", V)]
    dep_in = [r for r in lines("s.fy='1404' and s.vahed=? and e.lifecycle_action='RECEIVE_DEPOSIT'", V)]
    dep_kind = []
    for r in dep_in:
        prev = q("select hcode, bed from sanad where doc_key=? and src_row=?", r[0], r[1] - 1)
        h, b = (prev[0] if prev else ("", 0))
        kind = "bank" if h.startswith("111-") and int(b or 0) == int(r[5]) else ("auction" if h == "212-6-2-0" and int(b or 0) == int(r[5]) else "other")
        dep_kind.append(kind)
    dep_out = [r for r in lines("s.fy='1404' and s.vahed=? and e.lifecycle_action='REFUND_DEPOSIT'", V)]
    opening = q("""select sum(bes)-sum(bed) from sanad s join doc_class c on c.doc_key=s.doc_key where s.org='Terminals' and s.fy='1404'
                  and s.vahed=? and s.hcode='212-6-5-0' and c.kind='OPENING'""", V)[0][0]
    closing = q("""select sum(bed)-sum(bes) from sanad s join doc_class c on c.doc_key=s.doc_key where s.org='Terminals' and s.fy='1404'
                  and s.vahed=? and s.hcode='212-6-5-0' and c.kind like 'CLOSING%'""", V)[0][0]
    chq_amt = int(recv[0][4])
    first_full = [r for r in claims if int(r[4]) == chq_amt][0]
    claim_month = first_full[2][:7]
    cheque = [r for r in recv if str(r[9]).startswith(claim_month)][0]
    chq_no = cheque[8]
    clear_doc = [r for r in clear if r[8] == chq_no][0]
    triplet = q("""select s.src_row, s.hcode, s.bed, s.bes, e.own_event, e.lifecycle_action from sanad s join line_event e on e.sid=s.sid
                   where s.doc_key=? and s.vahed=? order by s.src_row""", clear_doc[0], V)
    send_this = [r for r in send if r[8] == chq_no][0]
    claim_doc_cr = q("""select count(*), count(distinct vahed), group_concat(distinct hcode) from sanad where doc_key=? and hcode like '211-1-1-%'""",
                     first_full[0])[0]
    claim_doc_lines = q("select count(*), count(distinct vahed) from sanad where doc_key=? and hcode='113-1-4-0'", first_full[0])[0]
    # سطرهای پس از سطر استرداد همین غرفه در همان سند تا سطر ۲۱۲-۶-۵ بعدی
    refund = dep_out[0]
    after = q("select src_row, hcode, bed, bes, vahed from sanad where doc_key=? and src_row>? order by src_row", refund[0], refund[1])
    refund_follow = []
    for r in after:
        if r[1] == "212-6-5-0":
            break
        refund_follow.append(dict(row=r[0], hcode=mask_hcode(r[1]), title=title_of(con, "Terminals", "1404", r[1]), bed=int(r[2]), bes=int(r[3]), vahed=r[4]))
    # سپردهٔ شرکت در مزایده با مبلغ برابر در همان سند نخستین سپرده
    d0 = [r for r, k in zip(dep_in, dep_kind) if k == "bank"][0]
    auction_eq = q("""select count(*) from sanad where doc_key=? and hcode='212-6-2-0' and cast(bes as real)=?""", d0[0], d0[5])[0][0]
    auction_vahed = q("""select distinct vahed from sanad where doc_key=? and hcode='212-6-2-0'""", d0[0])
    n_claim_docs = len(claims)
    amounts = OrderedDict()
    for r in claims:
        amounts.setdefault(int(r[4]), []).append(r)
    return dict(dep_kind=dep_kind, d0=d0, V=V, name_v=name_v, person=person, stages=sorted(st[V]), n_full=len(full), claims=claims, amounts=amounts,
                recv=recv, recv_cr=recv_cr, send=send, clear=clear, dep_in=dep_in, dep_out=dep_out, opening=int(opening or 0),
                closing=int(closing or 0), chq_amt=chq_amt, first_full=first_full, cheque=cheque, chq_mask=mask_digits(chq_no),
                clear_doc=clear_doc, triplet=triplet, send_this=send_this, claim_doc_cr=claim_doc_cr, claim_doc_lines=claim_doc_lines,
                refund=refund, refund_follow=refund_follow, auction_eq=auction_eq, auction_vahed=[a[0] for a in auction_vahed],
                n_claim_docs=n_claim_docs, title_113_6=title_of(con, "Terminals", "1404", "113-6-6-0"))


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------

def build_docx(SA, SB):
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
    from docx.shared import Cm, Pt, RGBColor

    FONT = "B Nazanin"
    NAVY, GREY, RED = "1F3864", "595959", "C00000"
    LATIN = re.compile(r"[A-Za-z0-9][A-Za-z0-9\-./•_:]*[A-Za-z0-9•]|[A-Za-z0-9]")

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(11)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    szcs = OxmlElement("w:szCs")
    szcs.set(qn("w:val"), "22")
    rpr.append(szcs)
    ppr = st.element.get_or_add_pPr()
    ppr.append(OxmlElement("w:bidi"))

    def set_bidi(p):
        pPr = p._p.get_or_add_pPr()
        if pPr.find(qn("w:bidi")) is None:
            pPr.insert(0, OxmlElement("w:bidi"))

    def style_run(run, size, bold, color, rtl):
        rPr = run._r.get_or_add_rPr()
        f = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            f.set(qn(a), FONT)
        rPr.insert(0, f)
        run.font.size = Pt(size)
        s = OxmlElement("w:szCs")
        s.set(qn("w:val"), str(int(size * 2)))
        rPr.append(s)
        if bold:
            run.bold = True
            rPr.append(OxmlElement("w:bCs"))
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if rtl:
            rPr.append(OxmlElement("w:rtl"))

    def add_text(p, text, size=11, bold=False, color=None):
        # B Nazanin نشانهٔ U+0654 ندارد؛ «هٔ» با «ۀ» نوشته می‌شود. «☐» به واژهٔ بعد می‌چسبد.
        text = str(text).replace("\u0647\u0654", "\u06c0").replace("☐ ", "☐\u00a0")
        lines = text.split("\n")
        for i, line in enumerate(lines):
            pos = 0
            for m in LATIN.finditer(line):
                if m.start() > pos:
                    style_run(p.add_run(line[pos:m.start()]), size, bold, color, True)
                # LRE … PDF: کد لاتین (حساب، سند، واحد) در متن راست‌به‌چپ وارونه نشود، حتی با «•••».
                style_run(p.add_run("\u200e\u202a" + m.group(0) + "\u202c\u200e"), size, bold, color, False)
                pos = m.end()
            if pos < len(line):
                style_run(p.add_run(line[pos:]), size, bold, color, True)
            if i < len(lines) - 1:
                r = p.add_run()
                r.add_break()

    def para(container, text="", size=11, bold=False, color=None, align=None, after=4, before=0, keep=False, page_break=False,
             shade=None, style=None):
        p = container.add_paragraph(style=style) if style else container.add_paragraph()
        set_bidi(p)
        if align is not None:
            p.alignment = align
        pf = p.paragraph_format
        pf.space_after = Pt(after)
        pf.space_before = Pt(before)
        pf.keep_with_next = keep
        pf.page_break_before = page_break
        if text:
            add_text(p, text, size, bold, color)
        elif size < 6:
            p._p.get_or_add_pPr().append(parse_xml(f'<w:rPr {nsdecls("w")}><w:sz w:val="{int(size * 2)}"/><w:szCs w:val="{int(size * 2)}"/></w:rPr>'))
            pf.line_spacing = Pt(size + 1)
        if shade:
            pPr = p._p.get_or_add_pPr()
            pPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{shade}"/>'))
        return p

    def heading(text, level, page_break=False):
        size = {1: 16, 2: 13, 3: 11.5}[level]
        p = para(doc, "", after=6 if level < 3 else 3, before=6 if level > 1 else 0, keep=True, page_break=page_break,
                 style=f"Heading {level}")
        add_text(p, text, size, True, NAVY)
        return p

    def shade_cell(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))

    def cell_write(cell, text, size=9, bold=False, color=None, shade=None, center=False):
        p = cell.paragraphs[0]
        set_bidi(p)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if text:
            add_text(p, text, size, bold, color)
        if shade:
            shade_cell(cell, shade)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        return p

    def cell_add(cell, text, size=8, bold=False, color=GREY):
        p = cell.add_paragraph()
        set_bidi(p)
        p.paragraph_format.space_after = Pt(0)
        add_text(p, text, size, bold, color)
        return p

    def table(container, widths, header=None, head_fill=NAVY, font=8.5):
        t = container.add_table(rows=0, cols=len(widths))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblPr = t._tbl.tblPr
        tblPr.append(parse_xml(f'<w:bidiVisual {nsdecls("w")}/>'))
        tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))
        tblPr.append(parse_xml(f'<w:tblCellMar {nsdecls("w")}><w:top w:w="40" w:type="dxa"/><w:left w:w="70" w:type="dxa"/>'
                               f'<w:bottom w:w="40" w:type="dxa"/><w:right w:w="70" w:type="dxa"/></w:tblCellMar>'))
        t._widths = widths
        if header:
            r = add_row(t, header, size=font, bold=True, color="FFFFFF", fill=head_fill)
            trPr = r._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:tblHeader"))
            keep_row(r)
        return t

    def keep_row(r):
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True

    def add_row(t, values, size=8.5, bold=False, color=None, fill=None, height=None, fills=None):
        r = t.add_row()
        trPr = r._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if height:
            trPr.append(parse_xml(f'<w:trHeight {nsdecls("w")} w:val="{int(height * 567)}" w:hRule="atLeast"/>'))
        for i, (c, w) in enumerate(zip(r.cells, t._widths)):
            c.width = Cm(w)
            v = values[i] if i < len(values) else ""
            f = (fills[i] if fills else None) or fill
            cell_write(c, v, size=size, bold=bold, color=color, shade=f)
        return r

    def set_margins(sec, landscape=False):
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(1.5)
        sec.left_margin = sec.right_margin = Cm(1.6)
        sec.header_distance = sec.footer_distance = Cm(0.7)

    def fill_header(sec, watermark=False):
        sec.header.is_linked_to_previous = False
        hdr = sec.header
        p = hdr.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        set_bidi(p)
        if watermark:
            add_text(p, "پیش‌نویس — قابل ارجاع نهایی نیست", 10, True, RED)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            wm = (f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office">'
                  '<w:pict><v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800" path="m@7,l@8,m@5,21600l@6,21600e">'
                  '<v:formulas><v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/><v:f eqn="sum 21600 0 @1"/><v:f eqn="sum 0 0 @2"/>'
                  '<v:f eqn="sum 21600 0 @3"/><v:f eqn="if @0 @3 0"/><v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/>'
                  '<v:f eqn="if @0 @4 21600"/><v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/><v:f eqn="mid @7 @8"/><v:f eqn="mid @6 @7"/>'
                  '<v:f eqn="sum @6 0 @5"/></v:formulas><v:path textpathok="t" o:connecttype="custom" o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800" o:connectangles="270,180,90,0"/>'
                  '<v:textpath on="t" fitshape="t"/><v:handles><v:h position="#0,bottomRight" xrange="6629,14971"/></v:handles>'
                  '<o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype>'
                  '<v:shape id="PowerPlusWaterMarkObject1" o:spid="_x0000_s1025" type="#_x0000_t136" '
                  'style="position:absolute;margin-left:0;margin-top:0;width:640pt;height:96pt;rotation:340;z-index:-251657216;'
                  'mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin" '
                  'o:allowincell="f" fillcolor="#c00000" stroked="f"><v:fill opacity=".16"/>'
                  '<v:textpath style="font-family:&quot;B Nazanin&quot;;font-size:1pt" string="پیش‌نویس — قابل ارجاع نهایی نیست"/></v:shape></w:pict></w:r>')
            p._p.append(parse_xml(wm))
        else:
            add_text(p, "فرم‌های عملیاتی پیشنهادی — قرارداد خدمات و اجارهٔ غرفه یا دفتر شرکت | پیش‌نویس برای بازبینی، نسخهٔ ۰٫۱", 8.5, False, GREY)

    sec0 = doc.sections[0]
    set_margins(sec0)
    fill_header(sec0)
    fp = sec0.footer.paragraphs[0]
    set_bidi(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(fp, "صفحهٔ ", 8.5, False, GREY)
    fp._p.append(parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple>'))

    W_FIELD = [1.1, 6.2, 2.7, 4.6, 3.2]

    # ---------------- جلد و راهنما ----------------
    para(doc, "", after=40)
    para(doc, "سازمان حمل‌ونقل مسافر — تاکسیرانی و پایانه‌ها", 13, True, GREY, WD_ALIGN_PARAGRAPH.CENTER, after=6)
    para(doc, "فرم‌های عملیاتی پیشنهادی", 22, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=4)
    para(doc, "قرارداد خدمات یا حجمی، و اجارهٔ غرفه یا دفتر شرکت در پایانه", 15, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=24)
    box = table(doc, [17.8])
    add_row(box, ["پیش‌نویس برای بازبینی\n"
                  "این فرم‌ها قاعدهٔ مصوب، الزام نرم‌افزاری یا مجوز ارجاع نهایی به حسابداری نیستند. هیچ فیلدی به‌طور رسمی اجباری نشده است.\n"
                  "داده‌های دو بخش تاکسیرانی و پایانه‌ها جدا می‌ماند. موارد منوط به تصمیم باز مانده‌اند.\n"
                  "نمونه‌ها با پروندهٔ واقعی تکمیل نشده‌اند؛ فقط شاهد دفتر برای دو نمونه آمده است (بخش ۷)."],
            size=11, bold=False, color="7F0000", fill="FBE9E7")
    para(doc, "", after=18)
    para(doc, f"نسخهٔ ۰٫۱ — تاریخ {DATE_FA}", 11, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, "ورودی بازنگری بستهٔ جلسه؛ جایگزین هیچ فرم یا سند پیشین نیست.", 10, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, after=24)
    para(doc, "فهرست", 12, True, NAVY, after=4)
    for line in ["۱. راهنمای کوتاه تکمیل", "۲. فرم‌ها و ترتیب استفاده", "۳. مسیر الف — قرارداد خدمات یا حجمی (فرم ۱ و ۲)",
                 "۴. مسیر ب — اجارهٔ غرفه یا دفتر شرکت (فرم ۳ و ۴)", "۵. فرم و کارت مشترک (فرم ۵ و کارت معرفی طرف)",
                 "۶. پیش‌نمایش برای بازبینی مالی و حسابداری (فرم ۶)", "۷. وضعیت شواهد دو نمونه", "۸. موارد باز که روی فرم‌ها اثر دارند"]:
        para(doc, line, 10.5, after=1)

    heading("۱. راهنمای کوتاه تکمیل", 1, page_break=True)
    guide = [
        ("هر داده یک بار", "مشخصات قرارداد و محل یک بار در فرم ۱ یا ۳ ثبت می‌شود. فرم‌های دوره‌ای آن‌ها را «پیش‌پر» نشان می‌دهند. اگر مقدار پیش‌پر با سند دوره فرق دارد، فقط با ذکر دلیل تغییر دهید."),
        ("اگر مقدار ندارید", "یکی از سه گزینه را بزنید: «دریافت نشده» (سند هست ولی به دست ما نرسیده)، «موضوعیت ندارد» (این مورد در این قرارداد وجود ندارد) یا «در انتظار تصمیم». خانهٔ خالی یا عدد صفر به جای این گزینه‌ها ننویسید. صفر فقط وقتی نوشته می‌شود که خود سند صفر را نشان دهد."),
        ("ستون «نوع و نیاز»", "سطر اول نوع مقدار است: «ورود از سند»، «انتخاب از فهرست»، «پیش‌پر (مرجع)»، «محاسبهٔ پیشنهادی» (سامانه حساب می‌کند و باید پذیرفته شود)، «تأیید کارکنان» یا «خودکار (سامانه)». سطر دوم می‌گوید فیلد کی لازم می‌شود: «برای شروع ثبت»، «پیش از بازبینی»، «شرطی»، «پس از تصمیم مرجع» یا «اختیاری». این‌ها پیشنهادند و الزام رسمی نیستند. ارجاع نهایی فقط پس از انتشار قاعدهٔ مصوب ممکن است."),
        ("خانهٔ خاکستری", "کاربر پر نمی‌کند: یا سامانه پر می‌کند یا مقدار پس از تصمیم مرجع می‌آید."),
        ("نقش طرف‌ها جداست", "طرف قرارداد، طرف‌حساب (کسی که بدهکار یا بستانکار می‌شود)، پرداخت‌کنندهٔ واقعی و ذی‌نفع (دریافت‌کنندهٔ وجه) هر کدام فیلد جدا دارند، حتی اگر یک نفر باشند."),
        ("بخش‌ها جدا", "هر برگه فقط برای یک بخش است: تاکسیرانی یا پایانه‌ها. داده‌ها تلفیق نمی‌شوند و کد حساب همیشه با بخش و سال خوانده می‌شود."),
        ("نرخ و مبلغ", "نرخ کسور و مبلغ‌ها فقط از متن قرارداد، قانون یا سند دوره می‌آید. نرخ آمادهٔ سامانه به کار نمی‌رود و هیچ مبلغی برآورد نمی‌شود. مبالغ به «واحد منبع» است، چون واحد پول دفاتر هنوز تأیید نشده است."),
        ("پیوند مراحل", "دریافت، چک و پرداخت فقط با تخصیص صریح به مطالبه یا صورت‌وضعیت وصل می‌شوند. هم‌مبلغی یا هم‌نامی فقط پیشنهاد است و تا تأیید کاربر پیوند قطعی نیست."),
        ("اطلاعات شخصی", "کد ملی، شمارهٔ حساب افراد و فهرست نفرات فقط در پیوست‌ها و با دسترسی محدود نگهداری می‌شود. در نمونه‌های این سند نام اشخاص و شمارهٔ حساب‌ها پوشانده شده است."),
        ("نقش‌های پیشنهادی", "نقش تکمیل‌کننده و بازبین هر فرم پیشنهادی است و پس از معرفی مسئولان سازمان تأیید می‌شود."),
    ]
    t = table(doc, [4.0, 13.8], ["موضوع", "راهنما"])
    for a, b in guide:
        add_row(t, [a, b], size=9.5, fills=["EEF2F8", None])

    # ---------------- فهرست فرم‌ها ----------------
    heading("۲. فرم‌ها و ترتیب استفاده", 1, page_break=True)
    para(doc, "شش فرم و یک کارت مشترک. فرم‌های ۱ و ۳ یک بار برای هر قرارداد پر می‌شوند؛ فرم‌های ۲ و ۴ برای هر دوره؛ فرم ۵ برای هر سپرده، تضمین یا چک؛ کارت طرف فقط برای طرف تازه؛ فرم ۶ را سامانه برای بازبینی حسابداری می‌سازد.", 10.5)
    t = table(doc, [2.2, 5.2, 4.2, 6.2], ["فرم", "کاربرد", "تکمیل‌کننده (پیشنهادی)", "زمان استفاده"])
    inv = [
        ("فرم ۱", C.FORM1["title"], "امور قراردادها", "یک بار پس از امضا؛ با هر الحاقیه نسخهٔ تازه"),
        ("فرم ۲", C.FORM2["title"], "ناظر قرارداد؛ حسابداری؛ خزانه", "برای هر دورهٔ صورت‌وضعیت؛ بخش پرداخت پس از پرداخت"),
        ("فرم ۳", C.FORM3["title"], "واحد درآمد و اماکن پایانه‌ها؛ حقوقی", "یک بار پس از امضا یا مزایده؛ با هر تمدید نسخهٔ تازه"),
        ("فرم ۴", C.FORM4["title"], "واحد درآمد و اماکن؛ خزانه؛ حسابداری", "مطالبه در آغاز هر دوره؛ رسید هنگام هر دریافت؛ تخصیص پس از دریافت"),
        ("فرم ۵", C.FORM5["title"], "خزانه و صندوق؛ امور قراردادها یا واحد درآمد", "هنگام دریافت هر سپرده، تضمین یا چک و هر تغییر وضعیت آن"),
        ("کارت", C.CARD["title"], "مسئول اطلاعات پایه", "فقط وقتی طرف در فهرست نیست"),
        ("فرم ۶", C.FORM6["title"], "سامانه؛ بازبینی: حسابداری", "پس از تکمیل یک بسته؛ ارجاع نیست"),
    ]
    for r in inv:
        add_row(t, list(r), size=9.5, fills=["EEF2F8", None, None, None])
    heading("چرا شش فرم و یک کارت؟", 3)
    for b in [
        "مشخصات قرارداد و محل یک بار ثبت می‌شود و فرم‌های دوره‌ای فقط آن را انتخاب می‌کنند. به همین دلیل محل اجاره در همان فرم قرارداد اجاره آمده است.",
        "مطالبه، دریافت و تسویه در یک فرم با سه بخش جدا و امضای جدا آمده‌اند. این سه مرحله با هم دیده می‌شوند، ولی ادغام نمی‌شوند.",
        "سپرده، تضمین و چک در هر دو مسیر عمر جدای خود را دارند و از چند دوره عبور می‌کنند؛ پس یک فرم مشترک با ماهیت حقوقی جدا دارند.",
        "نیرو و کارکرد بخش شرطی صورت‌وضعیت است، نه فرم جدا؛ فقط برای قرارداد بر پایهٔ نیرو باز می‌شود.",
        "سطر طبقه‌بندی و بسته را کاربر عملیاتی نمی‌سازد؛ فقط در پیش‌نمایش حسابداری دیده می‌شوند.",
        "فهرست حقوق، اعلامیهٔ فی‌مابین و دریافت خدمات تاکسیرانی به این دو مسیر مربوط نیستند و در این سند نیامده‌اند.",
    ]:
        para(doc, "• " + b, 10, after=2)

    def path_table(stages):
        t = table(doc, [1.0, 4.4, 4.2, 4.4, 3.8], ["#", "مرحله", "چه کسی (پیشنهادی)", "سند یا منبع اطلاعات", "فرم و بخش"])
        for s in stages:
            add_row(t, list(s), size=9, fills=["EEF2F8", None, None, None, None])

    def guide_card(form):
        c = form["card"]
        t = table(doc, [4.0, 13.8] if not form.get("landscape") else [5.0, 20.4])
        for k, lab in [("who", "چه کسی و چه زمانی"), ("source", "منبع اطلاعات (سند یا شخص)"), ("prefill", "از ثبت‌های قبلی پیش‌پر می‌شود"),
                       ("reviewer", "بازبین و تأییدکننده"), ("attach", "پیوست‌ها (فقط در صورت شرط)"), ("open", "هنوز گردآوری‌نشده یا منوط به تصمیم")]:
            add_row(t, [lab, c[k]], size=9, fills=["EEF2F8", None])

    def field_rows(form, sec_items, widths):
        t = table(doc, widths, ["#", "فیلد و راهنما", "نوع و نیاز", "مقدار", "اگر مقدار ندارد"])
        for it in sec_items:
            need = C.NEED[it["need"]]
            if it["need"] == "cond" and it.get("cond"):
                need += "\n(" + it["cond"] + ")"
            r = add_row(t, [disp_fid(it["fid"]), "", C.VTYPE[it["vtype"]] + "\n" + need, "", ""], size=8.5, height=0.95,
                        fills=["EEF2F8", None, None, None, None])
            c1 = r.cells[1]
            cell_write(c1, it["label"], 9.5, True)
            if it.get("guide"):
                cell_add(c1, it["guide"], 8)
            cv = r.cells[3]
            if it.get("locked"):
                txt = it.get("opts") or ("در انتظار تصمیم مرجع — کاربر پر نمی‌کند" if it["need"] == "dec" else "سامانه پر می‌کند")
                if it["need"] == "dec" and it.get("opts") is None:
                    txt = "در انتظار تصمیم مرجع — کاربر پر نمی‌کند"
                cell_write(cv, txt, 8.5, False, GREY, shade="E7E6E6")
            elif it.get("opts"):
                cell_write(cv, it["opts"], 8.5)
            miss = "—" if (it.get("locked") or it["vtype"] in ("sys", "calc")) else C.MISSING_OPTS
            cell_write(r.cells[4], miss, 7.5, False, GREY)
        return t

    def render_form(form, page_break=True):
        land = form.get("landscape")
        usable = 25.4 if land else 17.8
        widths = [w * usable / 17.8 for w in W_FIELD]
        heading(f"فرم {form['no']} — {form['title']}" if form["id"] != "FC" else form["title"], 2, page_break=page_break)
        if form.get("banner"):
            para(doc, form["banner"], 14, True, RED, WD_ALIGN_PARAGRAPH.CENTER, after=4, shade="FBE9E7")
        para(doc, "پیش‌نویس برای بازبینی — الزام رسمی نیست", 9, True, RED, after=2)
        para(doc, form["purpose"], 10, after=4)
        guide_card(form)
        para(doc, "", after=2)
        for sec in form["sections"]:
            heading(sec["title"], 3)
            if sec.get("cond"):
                para(doc, "بخش شرطی — " + sec["cond"], 9, True, "7F6000", after=2, keep=True, shade="FFF4CE")
            buf = []

            def flush():
                if buf:
                    field_rows(form, list(buf), widths)
                    para(doc, "", 2, after=0)
                    buf.clear()

            for it in sec["items"]:
                if it["kind"] == "field":
                    buf.append(it)
                    continue
                flush()
                if it["kind"] == "table":
                    head = f"{disp_fid(it['fid'])} — {it['label']}"
                    para(doc, head, 9.5, True, keep=True, after=1)
                    meta = f"{C.VTYPE[it['vtype']]} • {C.NEED[it['need']]}" + (f" ({it['cond']})" if it.get("cond") else "")
                    para(doc, it["guide"] + "\n" + meta if it.get("guide") else meta, 8.5, False, GREY, keep=True, after=2)
                    ws = it["widths"]
                    if land and sum(ws) < 20:
                        ws = [w * usable / sum(ws) for w in ws]
                    t = table(doc, ws, it["cols"])
                    last_missing = it["cols"][-1] == "اگر مقدار ندارد"
                    keep_all = (not it["empty_rows"] and len(it["rows"]) <= 12) or (not it["rows"] and it["empty_rows"] <= 5)
                    for k, row in enumerate(it["rows"]):
                        vals = list(row) + [""] * (len(it["cols"]) - len(row))
                        if last_missing:
                            vals[-1] = C.MISSING_OPTS
                        r = add_row(t, vals, size=8.5, height=0.85, fills=["EEF2F8"] + [None] * (len(it["cols"]) - 1))
                        if last_missing:
                            for p in r.cells[-1].paragraphs:
                                for run in p.runs:
                                    run.font.size = Pt(7.5)
                        if keep_all and k < len(it["rows"]) - 1:
                            keep_row(r)
                    for k2 in range(it["empty_rows"]):
                        vals = [""] * len(it["cols"])
                        fills = [None] * len(it["cols"])
                        for ci, txt in it.get("fixed", {}).items():
                            vals[ci], fills[ci] = txt, "E7E6E6"
                        if last_missing:
                            vals[-1] = C.MISSING_OPTS
                        r = add_row(t, vals, size=7.5, height=0.85, fills=fills)
                        if keep_all and k2 < it["empty_rows"] - 1:
                            keep_row(r)
                    if it.get("note"):
                        para(doc, it["note"], 8.5, False, GREY, after=4, before=2)
                    else:
                        para(doc, "", 2, after=0)
                elif it["kind"] == "check":
                    para(doc, f"{disp_fid(it['fid'])} — {it['label']}", 9.5, True, keep=True, after=2)
                    t = table(doc, [w * usable / 17.8 for w in [5.6, 6.4, 5.8]], ["پیوست", "شرط", "وضعیت"])
                    for a, b in it["items"]:
                        add_row(t, [a, b, "☐ پیوست شد  ☐ دریافت نشده  ☐ موضوعیت ندارد"], size=8.5, height=0.75, fills=["EEF2F8", None, None])
                    para(doc, "", 2, after=0)
                elif it["kind"] == "sign":
                    para(doc, f"{disp_fid(it['fid'])} — {it['label']}", 9.5, True, keep=True, after=2)
                    if it.get("note"):
                        para(doc, it["note"], 9, keep=True, after=4)
                    t = table(doc, [w * usable / 17.8 for w in [5.6, 3.4, 2.4, 3.0, 3.4]], ["نقش", "نام", "تاریخ", "امضا", "نتیجه"])
                    for k, role in enumerate(it["roles"]):
                        r = add_row(t, [role, "", "", "", "☐ تأیید\n☐ برگشت با دلیل"], size=8.5, height=1.1, fills=["EEF2F8", None, None, None, None])
                        if k < len(it["roles"]) - 1:
                            keep_row(r)
                    para(doc, "", 2, after=0)
            flush()
            para(doc, "", after=2)

    # ---------------- مسیر الف ----------------
    heading("۳. مسیر الف — قرارداد خدمات یا حجمی", 1, page_break=True)
    para(doc, "از ثبت قرارداد تا صورت‌وضعیت دوره‌ای، تأیید و ثبت مرجع پرداخت. برای انتظامات، نظافت، تأسیسات یا تأمین نیرو در هر دو بخش.", 10.5)
    path_table(C.PATH_A_STAGES)
    para(doc, "واریز کسور به امور مالیاتی و تأمین اجتماعی تجمیعی است و چند قرارداد را با هم می‌پوشاند؛ در این فرم‌ها نیامده است.", 9, False, GREY, before=4)
    render_form(C.FORM1)
    render_form(C.FORM2)

    # ---------------- مسیر ب ----------------
    heading("۴. مسیر ب — اجارهٔ غرفه یا دفتر شرکت در پایانه", 1, page_break=True)
    para(doc, "از مزایده یا ثبت قرارداد تا مطالبه، دریافت، تسویه و پیگیری سپرده و چک. مطالبه، دریافت و تسویه سه مرحلهٔ جدا هستند.", 10.5)
    path_table(C.PATH_B_STAGES)
    para(doc, "عوارض ۵٪ فروش بلیط، سهم کارگزاری و تهاتر با شهرداری به این مسیر مربوط نیستند و تا تصمیم دربارهٔ نقش پایانه‌ها باز می‌مانند.", 9, False, GREY, before=4)
    render_form(C.FORM3)
    render_form(C.FORM4)

    # ---------------- مشترک ----------------
    heading("۵. فرم و کارت مشترک", 1, page_break=True)
    para(doc, "فرم ۵ در هر دو مسیر به کار می‌رود. وضعیت‌های مجاز هر ماهیت:", 10.5)
    t = table(doc, [4.0, 13.8], ["ماهیت", "وضعیت‌های مجاز (به ترتیب)"])
    for a, b in C.STATUS_BY_NATURE:
        add_row(t, [a, b], size=9, fills=["EEF2F8", None])
    render_form(C.FORM5, page_break=False)
    render_form(C.CARD)

    # ---------------- فرم ۶ (افقی، با واترمارک) ----------------
    sec6 = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(sec6, landscape=True)
    fill_header(sec6, watermark=True)
    heading("۶. پیش‌نمایش برای بازبینی مالی و حسابداری", 1)
    render_form(C.FORM6, page_break=False)

    # ---------------- نمونه‌ها ----------------
    sec7 = doc.add_section(WD_SECTION.NEW_PAGE)
    set_margins(sec7)
    fill_header(sec7)
    heading("۷. وضعیت شواهد دو نمونه", 1)
    para(doc, "سه دستهٔ جدا به کار رفته است و با هم مخلوط نمی‌شوند:", 10.5)
    for b in ["«نمونهٔ تکمیل‌شده با مدارک واقعی»: فرم با پروندهٔ واقعی سازمان پر شده و هر مقدار به مدرک مشخص ارجاع دارد.",
              "«شاهد دفتر؛ پروندهٔ قرارداد در دسترس نیست»: فقط ردیف‌های دفاتر ۱۴۰۴ که در اختیار ماست. از این ردیف‌ها نمی‌توان قرارداد، صورت‌وضعیت یا اعلامیهٔ مشخصی را نتیجه گرفت.",
              "«دریافت نشده»: مدرکی که برای تکمیل نمونه لازم است و هنوز نرسیده است."]:
        para(doc, "• " + b, 10, after=2)
    t = table(doc, [2.8, 2.4, 2.8, 5.0, 4.8], ["نمونه", "وضعیت", "نمونهٔ تکمیل‌شده با مدارک واقعی", "شاهد دفتر (پروندهٔ قرارداد در دسترس نیست)", "دریافت نشده"])
    add_row(t, ["الف — قرارداد خدمات", "تکمیل‌نشده", "ندارد",
                f"سری «خدمات انتظامات» دفتر تاکسیرانی ۱۴۰۴: {fa(len(SA['series']))} سند شناسایی هزینه و {fa(1)} سند پرداخت",
                "قرارداد، صورت‌وضعیت دوره، پیوست کارکرد، تأییدها، صورت‌حساب مالیاتی، دستور پرداخت (جدول ۷-الف-۴)"],
            size=9, fills=["EEF2F8", "FBE9E7", None, None, None])
    add_row(t, ["ب — اجارهٔ غرفه", "تکمیل‌نشده", "ندارد",
                f"غرفهٔ {fa(SB['V'].split('-')[2].lstrip('0'))} پایانهٔ کاوه در دفتر پایانه‌ها ۱۴۰۴: مطالبه، سپرده، چک و تسویه",
                "قرارداد اجاره، اسناد مزایده، اعلامیهٔ مطالبه، رسیدها، فهرست چک‌ها (جدول ۷-ب-۴)"],
            size=9, fills=["EEF2F8", "FBE9E7", None, None, None])
    para(doc, "هیچ‌یک از دو نمونه «تکمیل‌شده» نیست، چون پروندهٔ واقعی قرارداد و مدارک دوره هنوز نرسیده است. شمارهٔ قرارداد، نرخ، تاریخ و طرف قرارداد ساخته نشده است.", 10, True, "7F0000", before=4)

    def origin(o):
        return {"L": "شاهد دفتر؛ پروندهٔ قرارداد در دسترس نیست", "N": "دریافت نشده", "D": "در انتظار تصمیم", "C": "پیوند نامزد؛ تأییدشده نیست"}[o]

    def sample_form_table(rows):
        t = table(doc, [5.0, 7.2, 5.6], ["فیلد", "مقدار", "منشأ یا وضعیت"])
        for f, v, o in rows:
            fill = {"L": "E2EFDA", "N": "FBE9E7", "D": "E7E6E6", "C": "FFF4CE"}[o]
            add_row(t, [f, v, origin(o)], size=8.5, fills=["EEF2F8", None, fill])

    def ledger_table(rows):
        t = table(doc, [2.8, 1.5, 1.9, 2.7, 3.9, 2.5, 2.5], ["سند منبع", "ردیف منبع", "تاریخ", "حساب", "عنوان حساب", "بدهکار", "بستانکار"])
        for r in rows:
            add_row(t, [r["doc"], fa(r["row"]), fa(r["date"]), r["hcode"], r["title"], fa(r["bed"]) if r["bed"] else "—",
                        fa(r["bes"]) if r["bes"] else "—"], size=8)

    # ---- نمونهٔ الف
    per, cp = SA["per"], SA["cp"]
    e1, e2 = SA["exp"][0], SA["exp"][1]
    m114 = per["m114"]
    heading("۷-الف. نمونهٔ الف — قرارداد خدمات (شاهد دفتر)", 2, page_break=True)
    para(doc, "شاهد دفتر؛ پروندهٔ قرارداد در دسترس نیست", 13, True, "375623", WD_ALIGN_PARAGRAPH.CENTER, shade="E2EFDA", after=6)
    para(doc, "روش انتخاب: در دفتر تاکسیرانی ۱۴۰۴ سطرهای «خدمات پیمانکاری» که شرحشان واژهٔ «انتظامات» دارد یک سری ماهانه می‌سازند. "
              "این واژه فقط نشانه است و قرارداد را ثابت نمی‌کند. از این سری، تنها سندی انتخاب شد که شناسایی هزینه را به بستانکار پرداختنی برده و پرداخت در سند جدا آمده است؛ "
              "پس مرحلهٔ صورت‌وضعیت و مرحلهٔ پرداخت جدا دیده می‌شوند.", 10)
    heading("۷-الف-۱. پیش‌نمایش جزئی فرم ۱ و فرم ۲ با شاهد دفتر", 3)
    sample_form_table([
        ("۱-۰۱ بخش سازمان", "تاکسیرانی", "L"),
        ("۱-۰۲ و ۱-۰۳ شماره و تاریخ قرارداد", "—", "N"),
        ("۱-۰۴ نوع خدمت", "— (شرح دفتر به «انتظامات» اشاره دارد؛ نشانه است، نه مدرک)", "N"),
        ("۱-۰۸ طرف قرارداد", f"— (در دفتر فقط کد طرف‌حساب {fa(cp)} روی سطرهای کسر و پرداختنی دیده می‌شود)", "N"),
        ("۱-۰۹ امضاکنندهٔ سمت کارفرما", "—", "N"),
        ("۱-۱۴ فعالیت پیش‌فرض", f"کد فعالیت منبع روی سطرهای هزینهٔ این سری: {'، '.join(fa(a) for a in sorted({a for s in SA['series'] for a in s['acts']}))} (در ماه‌های مختلف متفاوت)", "L"),
        ("۱-۱۷ تا ۱-۲۰ مدت، مبلغ کل و مبنای مبلغ", "—", "N"),
        ("۱-۲۴ کسور و تضمین قراردادی", "— (بخش ۷-الف-۳: در دفتر فقط «سپرده بیمه» دیده می‌شود)", "N"),
        ("۲-۰۴ شماره و تاریخ صورت‌وضعیت", "— (شمارهٔ صورت‌وضعیت در ستون جدا ثبت نشده است)", "N"),
        ("۲-۰۵ دورهٔ خدمت", "— (شرح سند نام ماه را دارد؛ نشانه است)", "N"),
        ("۲-۰۷ مبلغ ناخالص", f"جمع دو سطر «خدمات پیمانکاری»: {fa(e1['bed'])} + {fa(e2['bed'])} = {fa(per['gross'])} واحد منبع؛ برابری با مبلغ صورت‌وضعیت تأیید نشده است", "L"),
        ("۲-۰۹ ارزش افزوده", f"سطر «مالیات» با کد 114-9-1-0 به مبلغ {fa(m114)}؛ اینکه ارزش افزوده است، بدون صورت‌حساب مالیاتی معلوم نیست", "L"),
        ("۲-۲۳ کسر یا سپردهٔ بیمه", f"{fa(per['ins'])} واحد منبع؛ ذی‌نفع و نرخ قراردادی دریافت نشده است", "L"),
        ("۲-۲۳ سپردهٔ حسن انجام کار و مالیات تکلیفی", "در این سند سطری دیده نشد. لازم بودن آن‌ها فقط از متن قرارداد معلوم می‌شود؛ این خانه صفر نیست", "N"),
        ("۲-۲۴ خالص قابل پرداخت", f"بستانکار پرداختنی: {fa(per['pay'])} = {fa(per['gross'])} + {fa(m114)} − {fa(per['ins'])}", "L"),
        ("بخش ۴ فرم ۲ — نیرو و کارکرد", "— (مبنای مبلغ معلوم نیست؛ پیوست کارکرد نرسیده است)", "N"),
        ("۲-۲۵ تأییدهای صورت‌وضعیت", "—", "N"),
        ("۲-۲۷ شمارهٔ دستور پرداخت", "—", "N"),
        ("۲-۲۸ تاریخ و مبلغ پرداخت", f"{fa(SA['pay_row']['date'])}، {fa(SA['pay_row']['bes'])} واحد منبع؛ سند {SA['pay_doc']}", "L"),
        ("۲-۳۰ حساب بانکی سازمان", f"{SA['pay_row']['hcode']} — {SA['bank_title']}", "L"),
        ("۲-۳۱ روش و مرجع پرداخت", f"مرجع سطر بانک پوشانده شد ({SA['cheque_mask']})؛ شرح سند به چک اشاره دارد", "L"),
        ("۲-۲۹ دریافت‌کنندهٔ وجه", "— (در دفتر فقط بدهکار شدن کد طرف‌حساب پرداختنی دیده می‌شود)", "N"),
        ("۲-۳۲ پیوند پرداخت با صورت‌وضعیت", "هم‌کد طرف‌حساب و هم‌مبلغ با پرداختنی سند شناسایی", "C"),
        ("نقش سازمان در هزینه‌های پایانه", "موضوعیت ندارد (بخش تاکسیرانی)", "L"),
    ])
    heading("۷-الف-۲. ردیف‌های دفتر منبع (شاهد تاریخی؛ پیشنهاد ثبت نیست)", 3)
    ledger_table(SA["rows"])
    para(doc, "شمارهٔ حساب بانکی و مرجع بانکی پوشانده شده است. شرح ردیف‌ها و نام طرف‌ها نقل نشده است.", 8.5, False, GREY, before=2)
    heading("۷-الف-۳. سری ۱۴۰۴ (انتخاب با نشانهٔ شرح؛ پیوند به یک قرارداد تأیید نشده است)", 3)
    t = table(doc, [2.7, 1.9, 2.8, 2.1, 2.3, 1.8, 1.6, 2.6],
              ["سند", "تاریخ", "جمع سطرهای هزینه", "سطر «مالیات»", "سپرده بیمه", "نسبت سپرده به هزینه", "کد فعالیت", "تسویه"])
    for s in SA["series"]:
        ratio = fa(f"{s['ins'] / s['gross'] * 100:.1f}") + "٪" if s["gross"] else "—"
        how = "بستانکار پرداختنی؛ پرداخت در سند جدا" if s["pay"] else ("خروج بانک در همان سند" if s["bank"] else "—")
        add_row(t, [s["doc"], fa(s["date"]), fa(s["gross"]), fa(s["m114"]), fa(s["ins"]), ratio, "، ".join(fa(a) for a in s["acts"]), how], size=8)
    para(doc, f"در هر {fa(len(SA['series']))} سند، «سپرده بیمه» دقیقاً ۵٫۰٪ جمع سطرهای هزینه است و «سپردهٔ حسن انجام کار» دیده نمی‌شود. این نسبت مشاهده است، نه نرخ قرارداد. "
              f"همان کد طرف‌حساب در {fa(SA['other_docs'])} سند هزینهٔ دیگر با شرحی جز «انتظامات» هم آمده است؛ پس کد طرف‌حساب هویت قرارداد نیست و فیلد قرارداد جدا لازم است.",
         9, before=3)
    heading("۷-الف-۴. مدارک لازم برای تکمیل نمونهٔ الف", 3)
    t = table(doc, [7.6, 6.4, 3.8], ["مدرک", "فیلدهایی که کامل می‌کند", "وضعیت"])
    for a, b in [
        ("متن کامل قرارداد امضاشده با پیوست‌ها و الحاقیه‌ها (همان قرارداد سری بالا یا قرارداد دیگری که سازمان انتخاب کند)", "فرم ۱ همهٔ بخش‌ها؛ نرخ کسور و تضمین"),
        ("صورت‌وضعیت یک دورهٔ کامل با تأییدهای ناظر، مدیر و حراست (برای همین سری: دورهٔ سند انتخاب‌شده)", "فرم ۲ بخش‌های ۱، ۲، ۵ و ۶"),
        ("پیوست کارکرد و حضور و فهرست پرداخت نفرات همان دوره (فقط اگر قرارداد بر پایهٔ نیرو است؛ با پوشاندن اطلاعات شخصی)", "فرم ۲ بخش ۴"),
        ("صورت‌حساب مالیاتی همان دوره (اگر پیمانکار مشمول است)", "۲-۰۹؛ روشن شدن ماهیت سطر «مالیات»"),
        ("دستور پرداخت و مدرک پرداخت (چک یا انتقال)", "۲-۲۷ تا ۲-۳۱"),
        ("تأیید سازمان که ردیف‌های دفتر بالا به همین قرارداد و دوره مربوط‌اند", "پیوند شاهد دفتر با پرونده"),
        ("ضمانت‌نامهٔ انجام تعهدات (اگر قرارداد خواسته است)", "فرم ۵"),
        ("ردیف بودجه یا نامهٔ تأمین اعتبار (اگر هست)", "۱-۲۵؛ ۲-۳۳"),
        ("معرفی فرد توضیح‌دهندهٔ گردش کار و فرم‌های فعلی همین فرایند", "نقش‌ها و مراحل؛ بازبینی چیدمان فرم"),
        ("واحد پول مبالغ دفاتر", "نمایش قطعی مبالغ"),
    ]:
        add_row(t, [a, b, "دریافت نشده"], size=8.5, fills=[None, None, "FBE9E7"])

    # ---- نمونهٔ ب
    V = SB["V"]
    booth_no = fa(V.split("-")[2].lstrip("0"))
    heading("۷-ب. نمونهٔ ب — اجارهٔ غرفه در پایانه (شاهد دفتر)", 2, page_break=True)
    para(doc, "شاهد دفتر؛ پروندهٔ قرارداد در دسترس نیست", 13, True, "375623", WD_ALIGN_PARAGRAPH.CENTER, shade="E2EFDA", after=6)
    para(doc, f"روش انتخاب: از {fa(SB['n_full'])} غرفهٔ دفتر پایانه‌ها ۱۴۰۴ که هر هفت مرحلهٔ مطالبه، تسویهٔ مطالبه، دریافت و استرداد سپرده، و دریافت، ارسال و وصول چک را دارند و فقط یک کد طرف‌حساب دارند، "
              f"غرفه‌ای که کوچک‌ترین کد واحد را دارد انتخاب شد: {V}. نخستین ماهی که مبلغ مطالبه با مبلغ چک‌های تازه برابر شد، به‌عنوان دورهٔ نمونه آمده است.", 10)
    ff = SB["first_full"]
    clr = SB["clear_doc"]
    trip = SB["triplet"]
    sent = SB["send_this"]
    chq = SB["cheque"]
    dep_total = sum(int(r[5]) for r in SB["dep_in"])
    kinds = list(zip(SB["dep_in"], SB["dep_kind"]))
    bank_d = [r for r, k in kinds if k == "bank"]
    auc_d = [r for r, k in kinds if k == "auction"]
    oth_d = [r for r, k in kinds if k == "other"]
    dep_txt = f"{fa(len(bank_d))} سطر سپرده هم‌مبلغ با سطر واریز بانکی مجاور ({fa(bank_d[0][2])} تا {fa(bank_d[-1][2])}) به جمع {fa(sum(int(r[5]) for r in bank_d))}"
    if auc_d:
        dep_txt += f"؛ {fa(len(auc_d))} سطر انتقال از «سپرده شرکت در مزایده» در {fa(auc_d[0][2])} به مبلغ {fa(sum(int(r[5]) for r in auc_d))}"
    if oth_d:
        dep_txt += f"؛ {fa(len(oth_d))} سطر دیگر به جمع {fa(sum(int(r[5]) for r in oth_d))}"
    refund = SB["refund"]
    heading("۷-ب-۱. پیش‌نمایش جزئی فرم ۳، ۴ و ۵ با شاهد دفتر", 3)
    amt_rows = "؛ ".join(f"{fa(a)} در {fa(len(v))} سند ({fa(v[0][2])} تا {fa(v[-1][2])})" for a, v in SB["amounts"].items())
    sample_form_table([
        ("۳-۰۱ بخش سازمان", "پایانه‌ها", "L"),
        ("۳-۰۲ پایانه", "کاوه", "L"),
        ("۳-۰۳ و ۳-۰۴ نوع و شمارهٔ زیرمحل", f"غرفه، شمارهٔ {booth_no} (عنوان واحد در دفتر)", "L"),
        ("۳-۰۵ کد قدیمی واحد در دفتر", V, "L"),
        ("۳-۰۶ و ۳-۰۷ متراژ، کاربری و دورهٔ واگذاری", "—", "N"),
        ("۳-۰۸ شماره و تاریخ قرارداد", "—", "N"),
        ("۳-۱۱ شیوهٔ واگذاری", f"— (در سند نخستین سپردهٔ بانکی، {fa(SB['auction_eq'])} سطر «سپرده شرکت در مزایده» با مبلغ برابر در سطح پایانه دیده می‌شود و یک سطر سپرده از مزایده منتقل شده است؛ نشانه است، نه مدرک)" if SB['dep_kind'].count('auction') else f"— (در سند نخستین سپردهٔ بانکی، {fa(SB['auction_eq'])} سطر «سپرده شرکت در مزایده» با مبلغ برابر در سطح پایانه دیده می‌شود؛ نشانه است، نه مدرک)", "N"),
        ("۳-۱۲ طرف قرارداد", f"— (کد طرف‌حساب منبع {fa(SB['person'])}؛ نام شخص پوشانده شد)", "N"),
        ("۳-۱۳ امضاکنندهٔ سمت واگذارکننده", "—", "N"),
        ("۳-۱۴ نقش سازمان در این درآمد", "—", "D"),
        ("۳-۱۸ اقلام و مبلغ هر دوره", f"مبلغ قراردادی دریافت نشده است. مبلغ مطالبهٔ ماهانهٔ این غرفه در دفتر: {amt_rows}", "L"),
        ("۳-۱۸ تفکیک قلم (اجاره، حق شارژ)", f"از دفتر قابل تشخیص نیست: سطر طرف مقابل مطالبه در سطح پایانه تجمیع شده است ({fa(SB['claim_doc_cr'][0])} سطر برای {fa(SB['claim_doc_cr'][1])} واحد در سند ماهانه)", "N"),
        ("۳-۲۰ سپردهٔ انجام تعهدات (قرارداد)", "مبلغ قراردادی دریافت نشده است. در دفتر: " + dep_txt + f"؛ ماندهٔ آغاز ۱۴۰۴ {fa(SB['opening'])} و پایان ۱۴۰۴ {fa(SB['closing'])}", "L"),
        ("۵-۱۷ تا ۵-۲۱ پایان سپردهٔ قبلی", f"{fa(refund[2])}: {fa(int(refund[4]))}، سند {refund[0]}؛ از آن {fa(sum(r['bes'] for r in SB['refund_follow'] if not r['hcode'].startswith('111-')))} به حساب دیگری رفته و بقیه از بانک خارج شده است (ترتیب سطر؛ پیوند قطعی نیست)", "C"),
        ("۵-۰۱ تا ۵-۱۰ چک‌های اقساط", f"{fa(len(SB['recv']))} چک هر یک {fa(SB['chq_amt'])}، بانک صادرکننده «{SB['recv'][0][10]}»، دریافت {fa(SB['recv'][0][2])}؛ سررسید ماهانه از {fa(min(r[9] for r in SB['recv']))} تا {fa(max(r[9] for r in SB['recv']))}", "L"),
        ("۵-۱۶ تاریخچهٔ چک دورهٔ نمونه", f"چک {fa(SB['chq_mask'])}: در صندوق {fa(chq[2])} ← نزد بانک {fa(sent[2])} ← وصول‌شده {fa(clr[2])}", "L"),
        ("۴-۰۲ شمارهٔ اعلامیهٔ مطالبه", "— (در دفتر شمارهٔ اعلامیه وجود ندارد)", "N"),
        ("۴-۰۳ و ۴-۰۷ تاریخ و دورهٔ مطالبه", f"سند ماهانهٔ درآمدی {ff[0]} به تاریخ {fa(ff[2])}؛ دوره از شرح سند (نشانه)", "L"),
        ("۴-۰۸ مبلغ مطالبه", f"{fa(int(ff[4]))} (بدهکار «{title_b(SB)}»)", "L"),
        ("۴-۱۷ تا ۴-۲۴ رسید دریافت", f"وصول چک {fa(SB['chq_mask'])} در {fa(clr[2])}: {fa(int(clr[5]))}", "L"),
        ("۴-۲۰ پرداخت‌کنندهٔ واقعی", "— (صادرکنندهٔ چک در دفتر نیامده است)", "N"),
        ("۴-۲۹ تخصیص به مطالبه", f"در همان سند وصول، «پیش‌دریافت» بدهکار و حساب بهره‌بردار بستانکار شده است؛ این تسویه {fa(clr[2])} و مطالبهٔ ماه {fa(ff[2])} ثبت شده است. پیوند با مطالبهٔ مشخص فقط از هم‌غرفه‌گی و هم‌مبلغی برمی‌آید", "C"),
    ])
    heading("۷-ب-۲. ردیف‌های دفتر منبع (شاهد تاریخی؛ پیشنهاد ثبت نیست)", 3)
    rows_b = []
    for r in SB["dep_in"]:
        rows_b.append(dict(doc=r[0], row=r[1], date=r[2], hcode=r[3], title="سپرده انجام تعهدات (غرفات)", bed=int(r[4]), bes=int(r[5])))
    rows_b.append(dict(doc=refund[0], row=refund[1], date=refund[2], hcode=refund[3], title="سپرده انجام تعهدات (غرفات)", bed=int(refund[4]), bes=int(refund[5])))
    for f in SB["refund_follow"]:
        rows_b.append(dict(doc=refund[0], row=f["row"], date=refund[2], hcode=f["hcode"], title=f["title"], bed=f["bed"], bes=f["bes"]))
    rows_b.append(dict(doc=chq[0], row=chq[1], date=chq[2], hcode=chq[3], title=f"اسناد دریافتنی — چک {SB['chq_mask']}", bed=int(chq[4]), bes=int(chq[5])))
    cr = SB["recv_cr"][0]
    rows_b.append(dict(doc=cr[0], row=cr[1], date=cr[2], hcode=cr[3], title=f"پیش دریافت — جمع {fa(len(SB['recv']))} چک", bed=int(cr[4]), bes=int(cr[5])))
    rows_b.append(dict(doc=sent[0], row=sent[1], date=sent[2], hcode=sent[3], title=f"{SB['title_113_6']} — چک {SB['chq_mask']}", bed=int(sent[4]), bes=int(sent[5])))
    for (row, h, b, c, oe, la) in trip:
        rows_b.append(dict(doc=clr[0], row=row, date=clr[2], hcode=h, title={"113-6-6-0": SB["title_113_6"], "213-2-4-2": "پیش دریافت", "113-1-4-0": "بدهکاران — بهره‌برداران"}.get(h, h),
                           bed=int(b), bes=int(c)))
    rows_b.append(dict(doc=ff[0], row=ff[1], date=ff[2], hcode=ff[3], title="بدهکاران — بهره‌برداران (مطالبهٔ ماه)", bed=int(ff[4]), bes=int(ff[5])))
    ledger_table(rows_b)
    para(doc, "فقط سطرهای همین غرفه آمده است؛ سطر طرف مقابل سپرده و مطالبه در سطح پایانه یا بانک تجمیع شده است. شمارهٔ چک و حساب بانکی پوشانده شده و نام بهره‌بردار نقل نشده است.", 8.5, False, GREY, before=2)
    heading("۷-ب-۳. آنچه این شاهد برای فرم‌ها نشان می‌دهد", 3)
    for b in [
        "وصول چک و تسویهٔ حساب بهره‌بردار پیش از ثبت مطالبهٔ همان ماه آمده است؛ پس رسید و مطالبه در زمان‌های جدا رخ می‌دهند و پیوند آن‌ها باید در بخش ج فرم ۴ صریح ثبت شود.",
        "شمارهٔ چک تنها پیوند قطعی میان دریافت، ارسال و وصول همان چک است (فرم ۵، شمارهٔ ابزار).",
        "مبلغ مطالبهٔ ماهانه در میانهٔ سال تغییر کرده و هم‌زمان سپردهٔ تازه دریافت و سپردهٔ قبلی مسترد شده است. اینکه این تغییر تمدید یا قرارداد تازه است، فقط از پروندهٔ قرارداد معلوم می‌شود؛ فرم ۳ برای هر تمدید نسخهٔ تازه دارد.",
        "قلم مطالبه (اجاره یا حق شارژ) در سطر غرفه دیده نمی‌شود؛ جدول اقلام فرم ۳ و فرم ۴ همین کمبود را جبران می‌کند.",
    ]:
        para(doc, "• " + b, 9.5, after=2)
    heading("۷-ب-۴. مدارک لازم برای تکمیل نمونهٔ ب", 3)
    t = table(doc, [7.6, 6.4, 3.8], ["مدرک", "فیلدهایی که کامل می‌کند", "وضعیت"])
    for a, b in [
        ("قرارداد امضاشدهٔ اجارهٔ همین غرفه یا قرارداد دیگری که سازمان انتخاب کند: طرفین، امضاکنندهٔ سمت واگذارکننده، محل، مدت، اقلام و سپرده", "فرم ۳ همهٔ بخش‌ها"),
        ("آگهی و صورت‌جلسهٔ مزایده (اگر واگذاری با مزایده بوده) و رسید سپردهٔ شرکت در مزایده", "۳-۱۱؛ فرم ۵ (سپردهٔ مزایده)"),
        ("رسیدهای سپردهٔ انجام تعهدات", "فرم ۵ بخش ۱"),
        ("جدول اقساط و فهرست چک‌های دریافتی", "۳-۱۸؛ ۵-۰۸ تا ۵-۱۴"),
        ("اعلامیهٔ مطالبهٔ یک دورهٔ کامل", "فرم ۴ بخش الف"),
        ("رسید یا اعلامیهٔ وصول همان دوره", "فرم ۴ بخش ب"),
        ("مدارک استرداد سپردهٔ قبلی و مبنای مبلغ کسرشده از آن", "فرم ۵ بخش ۳"),
        ("مبنای حق شارژ (اگر مطالبه می‌شود)", "۳-۱۸"),
        ("پاسخ سازمان دربارهٔ طرف قرارداد، وصول‌کننده و مالک درآمد اجارهٔ غرفه", "۳-۱۳؛ ۳-۱۴؛ ۳-۲۳"),
        ("فهرست کاری پایانه‌ها و زیرمحل‌ها؛ معرفی فرد توضیح‌دهندهٔ گردش کار؛ فرم‌های فعلی", "فرم ۳ بخش ۱؛ نقش‌ها و مراحل"),
    ]:
        add_row(t, [a, b, "دریافت نشده"], size=8.5, fills=[None, None, "FBE9E7"])

    # ---------------- موارد باز ----------------
    heading("۸. موارد باز که روی فرم‌ها اثر دارند", 1, page_break=True)
    para(doc, "تا پاسخ هر مورد، فرم همان رفتاری را دارد که در ستون آخر آمده است. هیچ‌کدام با فرض کاری بسته نشده است.", 10.5)
    t = table(doc, [5.6, 4.6, 7.6], ["موضوع باز", "فیلدهای متأثر", "رفتار فرم تا تصمیم"])
    for row in OPEN_ITEMS_WORD:
        add_row(t, list(row), size=8.5, fills=["EEF2F8", None, None])

    doc.core_properties.title = "فرم‌های عملیاتی پیشنهادی — پیش‌نویس برای بازبینی"
    doc.core_properties.subject = "T09"
    doc.save(OUT_DOCX)


def title_b(SB):
    return "بدهکاران — بهره‌برداران"


OPEN_ITEMS_WORD = [
    ("نقش حسابداری پایانه‌ها در اجاره، عوارض و هزینه‌ها (مالک درآمد و هزینه)", "۱-۱۰؛ ۳-۱۴؛ ۳-۲۳", "امضاکنندهٔ روی سند ثبت می‌شود، ولی نقش سازمان و بُعد درآمد خاکستری می‌مانند."),
    ("شمارهٔ یکتای اعلامیهٔ مطالبه و اصلاحیه", "۴-۰۲؛ ۴-۱۱", "سامانه شماره پیشنهاد می‌کند؛ تا تصمیم، پیشنهاد است."),
    ("روش تخصیص دریافت و چک به مطالبه", "۴-۲۹؛ ۵-۱۴", "تخصیص صریح با تأیید کاربر؛ بدون ترتیب پیش‌فرض."),
    ("جدایی مرحلهٔ شناسایی از پرداخت یا وصول", "فرم ۲ بخش ۸؛ فرم ۴ بخش‌های ب و ج", "مرحله‌ها جدا ثبت می‌شوند؛ اثر حسابداری باز است."),
    ("وصول در برابر تهاتر با پیش‌دریافت چک", "۴-۲۹", "منبع تسویه «پیش‌دریافت چک» جدا از «رسید» انتخاب می‌شود."),
    ("تفکیک صورت‌وضعیت چند محل یا چند سازمان", "۱-۱۲؛ ۱-۱۳؛ ۲-۱۴", "فقط با مدرک کمّی؛ بدون مدرک «در انتظار مبنا»."),
    ("منبع نرخ کسور", "۱-۲۴؛ ۲-۲۳", "فقط متن قرارداد یا قانون؛ نرخ آمادهٔ سامانه به کار نمی‌رود."),
    ("نوع صورت‌وضعیت (موقت، قطعی، تعدیل) و مرز پیش‌پرداخت", "۲-۰۶؛ ۱-۲۳؛ ۲-۲۳", "آنچه روی سند آمده ثبت می‌شود؛ اثر آن باز است."),
    ("ارزش افزودهٔ خرید (قابل استرداد بودن)", "۲-۰۹", "مبلغ و شمارهٔ صورت‌حساب ثبت می‌شود؛ طبقه‌بندی باز است."),
    ("جایگاه تضمین‌ها و رفتار پس از برگشت چک", "فرم ۵", "تاریخچهٔ وضعیت ثبت می‌شود؛ رویداد مالی آن باز است."),
    ("فعالیت و مرکز هزینهٔ اختصاصی؛ اتصال درآمد به مرکز", "۱-۱۴؛ ۱-۱۵؛ ۲-۱۰؛ ۲-۱۱؛ ۳-۲۲؛ ۳-۲۳", "فعالیت و مرکز با برچسب «پیشنهادی»؛ بُعد درآمد خاکستری."),
    ("نمودار حساب مقصد", "۶-۱۱", "حساب مقصد «حل‌نشده»؛ حساب دفتر قدیم فقط برای ردیابی."),
    ("کاربرد کد یکتا و بخش‌های نامعلوم آن", "۶-۱۳؛ ۱-۲۵", "بخش نامعلوم خالی می‌ماند؛ مقدار پیش‌فرض درج نمی‌شود."),
    ("دامنهٔ بسته و گزینهٔ ارجاع به دفتر؛ واحد پول", "۶-۰۵؛ ۶-۱۲", "بسته «یک قرارداد × یک دوره» پیشنهادی؛ جدول پیوند تا ارجاع رسمی خالی؛ مبالغ به واحد منبع."),
    ("ثبت قرارداد درآمدی در نرم‌افزار", "فرم ۳ و ۴", "فرم مستقل از مدل نرم‌افزاری طراحی شده است."),
    ("مرجع تأیید قواعد", "همهٔ ستون‌های «نوع و نیاز»", "هیچ فیلدی الزام رسمی ندارد؛ ارجاع نهایی ممکن نیست."),
]


# ---------------------------------------------------------------------------
# PDF و تصویر صفحه‌ها
# ---------------------------------------------------------------------------

def export_pdf():
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        d = word.Documents.Open(str(OUT_DOCX), ReadOnly=True, AddToRecentFiles=False)
        d.ExportAsFixedFormat(str(OUT_PDF), 17, OpenAfterExport=False, OptimizeFor=0, CreateBookmarks=1)
        n = d.ComputeStatistics(2)
        d.Close(False)
    finally:
        word.Quit()
    return n


def render_png(dpi=80):
    import fitz
    out = WORK / "png"
    out.mkdir(parents=True, exist_ok=True)
    for f in out.glob("*.png"):
        f.unlink()
    d = fitz.open(OUT_PDF)
    for i, page in enumerate(d):
        page.get_pixmap(dpi=dpi).save(out / f"p{i + 1:02d}.png")
    n = d.page_count
    d.close()
    return n


def scan_outputs():
    """کنترل محتوای خروجی‌ها: حریم شخصی، مقدار ساختگی و کد داخلی در Word."""
    from docx import Document
    d = Document(OUT_DOCX)
    texts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                texts.append(c.text)
    for s in d.sections:
        texts += [p.text for p in s.header.paragraphs]
    body = re.sub("[\u200e\u202a\u202c]", "", "\n".join(texts))
    md = OUT_MD.read_text(encoding="utf-8") if OUT_MD.exists() else ""
    res = []
    long_digits = re.findall(r"[0-9۰-۹]{7,}", body)
    res.append(("K-T07", "Word رشتهٔ عددی بلند (شمارهٔ حساب، چک یا شناسهٔ ملی) ندارد", not long_digits, ", ".join(long_digits[:5]) or "۰"))
    names = ["بهادري", "بهادری", "شهروكار", "شهروکار", "استيلا", "استیلا", "مرشدي", "صديقي", "رسالت ۴۶"]
    hits = [n for n in names if n in body or n in md]
    res.append(("K-T08", "نام اشخاص و طرف‌های نمونه در Word و گزارش داخلی نقل نشده است", not hits, ", ".join(hits) or "۰"))
    synth = [s for s in ("00000000", "۰۰۰۰۰۰۰۰") if s in body] + re.findall(r"\b\d{2}-\d{2}-\d{3}-\d{8}\b", body)
    res.append(("K-T09", "مقدار پیش‌فرض یا ساختگی کد یکتا در Word درج نشده است", not synth, ", ".join(synth) or "۰"))
    codes = sorted(set(re.findall(r"\b(?:DR|DE|DG|CD|T-D|FM|EL|LP|PV|PFE|TT-NE|FE|SP|RL|PO|YK)-?\d", body)))
    res.append(("K-T10", "شناسهٔ تصمیم، پروفایل یا عنصر داخلی در Word نمایش داده نشده است", not codes, ", ".join(codes) or "۰"))
    wm = any("قابل ارجاع نهایی نیست" in t for t in texts)
    res.append(("K-T11", "پیش‌نمایش برچسب «پیش‌نویس — قابل ارجاع نهایی نیست» دارد", wm, "سربرگ بخش افقی و بنر فرم ۶" if wm else "یافت نشد"))
    return [dict(id=a, title=b, status="PASS" if c else "FAIL", detail=d_) for a, b, c, d_ in res]


# ---------------------------------------------------------------------------
# گزارش داخلی
# ---------------------------------------------------------------------------

REUSE = [
    ("گام ۱۰۱ «HR Initialization» — فرم پایهٔ قرارداد حجمی (عنوان، پیمانکار، مبلغ کل، شروع و پایان، دسته)",
     "`scripts/seed_volume_contracts_workflow.py:232-250`؛ `app/routers/civil.py:351-419`", "فعال",
     "نیازمند تغییر پارامتر یا عنوان", "فرم ۱",
     "شمارهٔ قرارداد قانونی، بخش، جهت وجه، محل‌ها، مبنای مبلغ و شرایط کسور را ندارد؛ هویت قرارداد در `CivilProject` است نه `Contract`."),
    ("مدل `Contract` (شماره، عنوان، پیمانکار، مبلغ، ردیف بودجه، درخواست اعتبار، فعالیت، تاریخ‌ها)",
     "`app/models.py:1085-1145`", "فعال (مسیر عمومی قرارداد)", "قابل استفاده با همان عنوان (فقط مسیر الف)", "فرم ۱",
     "برای قرارداد اجاره قابل استفاده نیست: `contractor_id` اجباری است (`app/models.py:1100`)."),
    ("رجیستری پیمانکار `Contractor` (نوع، کد ملی، نام، شمارهٔ ثبت، کد اقتصادی، شبا)",
     "`app/models.py:1000-1044`", "فعال (صفحهٔ ثبت‌نام عمومی ذخیره نمی‌کند)", "نیازمند تغییر پارامتر یا عنوان", "کارت طرف",
     "فیلدها بازاستفاده می‌شوند، ولی کارت باید نقش‌بی‌طرف باشد (بهره‌بردار، پرداخت‌کننده، ذی‌نفع، نهاد شهرداری)."),
    ("گام ۱۰۲ «Contractor Upload» — شمارهٔ صورت‌وضعیت، دوره، مبلغ ناخالص، فایل صورت‌وضعیت، فایل ۶ ستونی، صورت‌حساب ارزش افزوده",
     "`scripts/seed_volume_contracts_workflow.py:251-270`", "فعال", "نیازمند تغییر پارامتر یا عنوان", "فرم ۲ بخش‌های ۱، ۲، ۴ و ۷",
     "تاریخ سند، مبلغ تأییدشده، تفکیک محل و جدول کسور با ذی‌نفع افزوده می‌شود؛ مبلغ ناخالص نباید از بزرگ‌ترین خانهٔ فایل حدس زده شود (ممیزی فنی §۰)."),
    ("قالب ۱۸ ستونی کارکرد `AttendanceRecord` (نام، کد ملی، روز کارکرد، غیبت، اضافه‌کار و …)",
     "`app/models/attendance.py:140-175`", "فعال", "نیازمند تغییر پارامتر یا عنوان", "۲-۱۶ تا ۲-۱۸",
     "`WorkPeriod` فقط با (سال، ماه) یکتاست و کلید قرارداد ندارد (`app/models/attendance.py:24-37`)؛ «منطقه» جای پایانه یا واحد نشسته است. برای قرارداد سازمان، کلید قرارداد و محل لازم است."),
    ("قالب ۶ ستونی `ContractorStatementItem` (کد ملی، نام، ناخالص، مالیات، بیمه، خالص)",
     "`app/models/attendance.py:218-240`", "فعال", "قابل استفاده با همان عنوان", "۲-۱۹", "فقط در قرارداد بر پایهٔ نیرو؛ دادهٔ شخصی در پیوست می‌ماند."),
    ("تطبیق حضور: جمع قانونی مناطق در برابر جمع اعلامی پیمانکار منهای جریمه",
     "`app/services/attendance_service.py:608-663`", "فعال", "نیازمند تغییر پارامتر یا عنوان", "۲-۲۰",
     "تلورانس و مقایسهٔ تعداد نفر ندارد؛ ارزش قانونی از ثابت‌های قیمت `app/utils/pricing_engine.py:5-15` می‌آید، نه از نرخ قرارداد. در فرم فقط «محاسبهٔ پیشنهادی» است."),
    ("جریمه `PenaltyRecord` (شرح، تعداد، مبلغ واحد، جمع)", "`app/models/attendance.py:249-263`", "فعال",
     "قابل استفاده با همان عنوان", "۲-۲۱", "ذی‌نفع جریمه در جدول کسور (۲-۲۳) ثبت می‌شود."),
    ("زنجیرهٔ تأیید برگهٔ حضور (ایجادکننده ← نمایندهٔ شرکت ← حراست ← مدیرعامل)", "`app/core/workflow.py:10-65`", "فعال",
     "نیازمند تغییر پارامتر یا عنوان", "۲-۲۲؛ ۲-۲۵", "عنوان نقش‌ها به ناظر، مدیر پایانه یا واحد، حراست و حسابداری تغییر می‌کند."),
    ("`ProgressStatement` (شماره، نوع، دوره، ناخالص، چهار ستون کسر، پیش‌پرداخت، تأیید چهارسطحی)",
     "`app/models.py:1182-1269`؛ `app/services/statement_service.py:37-57`", "فعال (مسیر عمومی)", "نیازمند تغییر پارامتر یا عنوان",
     "فرم ۲", "کسرها ذی‌نفع ندارند؛ `line_items` محل ندارد و در API نیامده است (`app/schemas/statement.py:16-45`)."),
    ("گام ۱۰۳ ورودی‌های کد یکتا (منطقه، معاونت، قسمت، بودجه، مرکز، اقدام مستمر، فعالیت ویژه، ذی‌نفع، رویداد، سال)",
     "`scripts/seed_volume_contracts_workflow.py:271-306`؛ `frontend/src/components/volume_contracts/YektaIssuancePanel.tsx:226`",
     "فعال", "نیازمند تغییر پارامتر یا عنوان", "۶-۱۳",
     "بخش نامعلوم باید خالی بماند: سال پیش‌فرض فرانت ۱۴۰۵ و پیش‌فرض‌های سطح پایین (صفرها، 1403، 001 و «Unknown») در `app/services/yekta_codes.py:73-106`."),
    ("گام ۱۰۴ تولید خودکار ماتریس و کد با نرخ ثابت (ارزش افزوده ۱۰٪، مالیات ۳٪، بیمه ۷٪ و ۲۳٪، سپردهٔ تأمین اجتماعی ۵٪، نگهداری ۱۰٪، حسن انجام کار ۵٪) و تراکنش آمادهٔ پرداخت",
     "`app/services/yekta_service.py:63-115,132-184,286-347`", "فعال", "خارج از دامنهٔ این دو مسیر", "—",
     "فرم‌ها تا «در انتظار مرجع» پیش می‌روند و نرخ از متن قرارداد می‌آید (۱-۲۴). شاهد دفتر نمونهٔ الف هم نسبت دیگری نشان می‌دهد (فقط سپردهٔ بیمه ۵٪، بی حسن انجام کار)."),
    ("`ContractGuarantee` (حسن انجام کار، پیش‌پرداخت، مناقصه؛ فعال، آزاد، منقضی، ضبط)", "`app/models.py:1274-1301`",
     "فقط مدل (API و صفحه ندارد)", "نیازمند تغییر پارامتر یا عنوان", "فرم ۵ (تضمین غیرنقدی)",
     "ضمانت‌نامهٔ بانکی، سفته و چک تضمین، تمدید و تاریخچهٔ وضعیت لازم است."),
    ("سپردهٔ نقدی، چک دریافتی و تاریخچهٔ وضعیت ابزار", "ندارد", "—", "نیازمند بخش تازه", "فرم ۵", "در زیرساخت فعلی مسیری ندارد."),
    ("`TransactionAllocation` (واحد، مرکز، ردیف بودجه، مبلغ، کسور، جریمه، خالص، کد یکتا)", "`app/models/finance.py:87-113`",
     "فعال فقط در مالی حضور", "نیازمند تغییر پارامتر یا عنوان", "۲-۱۴؛ ۶-۱۰؛ ۶-۱۲", "مبنا و مدرک تفکیک ندارد."),
    ("مرجع پرداخت: `Transaction.tracking_number` و کد یکتا؛ تاریخ پرداخت و پیگیری خزانه فقط در گام کهنهٔ ۱۱۰",
     "`app/models.py:361-461`؛ `frontend/src/config/volumeContractWorkflow.ts:80-87`", "فعال / کهنه", "نیازمند بخش تازه", "فرم ۲ بخش ۸",
     "گام ۱۰۴ پایانی است و مرجع پرداخت بانکی ندارد."),
    ("سازوکار پیوست و فرم پویا (`EntityDocument`)", "`app/routers/volume_contracts.py:132-169`؛ `frontend/src/components/volume_contracts/VolumeContractFormRenderer.tsx:50-218`",
     "فعال", "قابل استفاده با همان عنوان", "پیوست همهٔ فرم‌ها", "داده در JSON با نوع رشته‌ای است؛ برای فیلدهای کلیدی، نوع صریح لازم است (ممیزی فنی §۰)."),
    ("قرارداد اجاره، مطالبه، رسید و تخصیص", "ندارد", "—", "نیازمند بخش تازه", "فرم ۳ و ۴", "فقط عناصر عمومی طرف، تاریخ و پیوست بازاستفاده می‌شوند."),
    ("پیکربندی کهنه: طرح ۱۰ مرحله‌ای ۱۰۱ تا ۱۱۰ (بازنویسی می‌شود)، متادیتای فرانت ۱۰۵ تا ۱۱۰، کدهای `CONTRACTORS_101..104` در `config_master_v5.json` (کد فعالیت‌اند نه مرحله)",
     "`scripts/seed_volume_contracts_workflow.py:139-230`؛ `frontend/src/config/volumeContractWorkflow.ts:16-87`؛ `app/config/config_master_v5.json:2629-2674`",
     "کهنه", "خارج از دامنهٔ این دو مسیر", "—", "مبنای طراحی نیست؛ بارگذار پیکربندی به‌طور پیش‌فرض `config_master.json` را می‌خواند."),
]

DECISIONS_MD = [
    ("T-D2، DE-01، CD-01", "نقش حسابداری پایانه‌ها (اصیل یا عامل)", "۱-۱۰، ۳-۱۴، ۳-۲۳؛ فرم ۵ برای صاحب بدهی سپرده", "امضاکنندهٔ روی سند (۱-۰۹، ۳-۱۳) واقعیت سند است و جدا ثبت می‌شود؛ نقش و بُعد درآمد قفل‌اند."),
    ("T-D5", "نمونهٔ ب و طرف حقوقی قرارداد؛ اجاره در تاکسیرانی", "فرم ۳ و ۴؛ بخش سازمان", "فرم‌ها برای پایانه‌ها؛ تاکسیرانی فقط پس از تصمیم (LP-F2-13 بیرون)."),
    ("T-D4", "پروندهٔ واقعی نمونهٔ الف", "بخش ۷ Word", "نمونه «تکمیل‌نشده»؛ شاهد دفتر جدا برچسب خورده است."),
    ("T-D3", "نمودار حساب مقصد", "۶-۱۱", "حساب مقصد «حل‌نشده»؛ فقط کلید (بخش، سال، کد) دفتر قدیم."),
    ("T-D6، DG-02، CD-03", "مرکز هزینه و اتصال درآمد به مرکز", "۱-۱۵، ۲-۱۱، ۳-۲۳، ۶-۰۹", "مرکز هزینه با برچسب «پیشنهادی»؛ بُعد درآمد قفل."),
    ("T-D7، DR-12", "کاربرد کد یکتا و بخش‌های نامعلوم", "۶-۱۳، ۱-۲۵، ۲-۳۳", "ورودی‌ها پیش‌بینی شدند؛ بخش نامعلوم خالی؛ هیچ مقدار پیش‌فرض."),
    ("T-D8، CD-02", "گزینهٔ ارجاع و واحد پول", "۶-۰۸، ۶-۱۲؛ همهٔ مبالغ", "پیش‌نمایش «ارجاع نیست»؛ جدول پیوند تا ارجاع رسمی خالی؛ «واحد منبع»."),
    ("DR-01", "هویت طرف در برابر کد طرف‌حساب منبع", "کارت طرف (ک-۰۶)، ۳-۱۵، ۴-۰۶، ۴-۲۱", "کارت نقش‌بی‌طرف؛ کدهای قدیم نگاشت چند به یک. شاهد نمونهٔ الف: یک کد در دو سری خدمت."),
    ("DR-02", "شمارهٔ یکتای مطالبه", "۴-۰۲، ۴-۱۱", "شمارهٔ سامانه پیشنهادی است. شاهد نمونهٔ ب: دفتر شمارهٔ اعلامیه ندارد."),
    ("DR-03، DE-10", "تخصیص دریافت و چک به مطالبه؛ وصول یا تهاتر", "۴-۲۹، ۵-۱۴", "تخصیص صریح، بدون ترتیب پیش‌فرض؛ منبع تسویه «پیش‌دریافت چک» جدا. شاهد نمونهٔ ب: تسویه پیش از مطالبهٔ همان ماه."),
    ("DE-02 (DR-04)", "جدایی شناسایی، پرداخت و وصول", "فرم ۲ بخش ۸، فرم ۴ بخش‌های ب و ج", "مرحله‌ها جدا ثبت می‌شوند؛ ماهیت سطر بانک باز است."),
    ("DG-06 (DR-05)", "تفکیک صورت‌وضعیت چند محل یا سازمان", "۱-۱۲، ۱-۱۳، ۲-۱۴", "فقط با مدرک کمّی؛ بدون مدرک «در انتظار مبنا» (SP-02/SP-04)."),
    ("DR-06", "دامنهٔ بسته", "۶-۰۵", "«یک قرارداد × یک دوره» پیشنهادی و قفل."),
    ("DR-07", "ثبت قرارداد درآمدی در نرم‌افزار", "فرم ۳", "فرم مستقل از مدل؛ هر دو گزینهٔ DR-07 با آن سازگارند."),
    ("DR-08، DE-21 (DR-09)", "ثبت ابزار واحد؛ جایگاه تضمین‌ها", "فرم ۵", "یک فرم با سه ماهیت حقوقی جدا؛ تضمین صادره و کالای امانی بیرون (LP-F4-07/08)."),
    ("DR-13، DE-28", "فعالیت پیش‌فرض قرارداد", "۱-۱۴، ۲-۱۰، ۳-۲۲، ۴-۱۳", "پیش‌فرض از قرارداد، تغییر در دوره با دلیل. شاهد نمونهٔ الف: کد فعالیت منبع در ماه‌های مختلف عوض شده است."),
    ("DR-14، DE-17", "منبع نرخ کسور؛ کدهای هم‌شکل", "۱-۲۴، ۲-۲۳، ۶-۱۱", "نرخ از متن قرارداد یا قانون؛ کد حساب همیشه با بخش و سال."),
    ("DE-14", "صورت‌وضعیت در برابر صورت‌حساب؛ نوع", "۲-۰۶", "آنچه روی سند آمده ثبت می‌شود؛ اثر آن باز."),
    ("DE-15", "کار سرمایه‌ای پیمانکار", "۱-۱۶، ۲-۱۲", "پروژه شرطی."),
    ("DE-18", "ارزش افزودهٔ خرید", "۲-۰۹", "مبلغ و شمارهٔ صورت‌حساب؛ طبقه‌بندی باز."),
    ("DE-29", "مرز پیش‌پرداخت و علی‌الحساب", "۱-۲۳، ۲-۲۳", "شرطی و پیوند صریح به پرداخت قبلی."),
    ("DE-11", "چک برگشتی", "۵-۱۶، ۵-۱۷", "وضعیت‌های پس از «برگشتی» ثبت می‌شوند؛ رویداد آن باز."),
    ("T-D1", "هویت حقوقی و کلید بخش", "۱-۰۱، ۲-۰۱، ۳-۰۱، ۴-۰۱، ۵-۰۲، ۶-۰۲", "یک برگه فقط برای یک بخش؛ بدون تلفیق."),
    ("DG-17", "واحدهای ستادی در برابر پایانه‌ها", "۱-۱۱", "محل خدمت می‌تواند پایانه یا واحد ستادی باشد؛ جدا از مرکز هزینه."),
    ("هـ-۴ (G-04)", "مرجع تأیید قواعد", "همهٔ ستون‌های «نوع و نیاز»", "هیچ فیلدی الزام رسمی ندارد (ADR-0019)."),
]


STAGE_FA = {"claim": "صدور مطالبه", "claim_settle": "تسویهٔ مطالبه", "deposit_in": "دریافت سپرده", "deposit_out": "استرداد سپرده",
            "chq_recv": "دریافت چک", "chq_send": "ارسال چک", "chq_clear": "وصول چک"}


def build_md(SA, SB, ctrl, uncovered, omitted, visual, n_pages_word, n_pages_pdf):
    L = []
    w = L.append
    w("# طراحی فرم‌های عملیاتی پیش‌نویس — قرارداد خدمات و اجارهٔ غرفه یا دفتر شرکت در پایانه (T09)")
    w("")
    w(f"**نسخه:** ۰٫۱ | **تاریخ:** {DATE_FA} | **وضعیت:** آماده برای بازبینی فرم‌ها؛ آمادهٔ پایلوت یا ارجاع نهایی **نیست**. شمار قواعد تأییدشده: ۰.")
    w("")
    w("سند داخلی تیم. فرم‌ها در `data/deliverables/passenger_transport/operational_forms_draft_v0_1.docx` و PDF همان آمده‌اند. "
      f"مبنا: T08 در `data/deliverables/passenger_transport/event_relationship_design_v0_1.xlsx` (هش `{sha16(T08)}`)، به‌ویژه شیت‌های «عناصر_داده»، «نیاز_داده_مسیرها»، «فیلدهای_فرم» و «تصمیم‌ها»؛ "
      "سند دامنه (ویرایش ۲)؛ درخواست اطلاعات T03 (هنوز ارسال نشده)؛ T05 و T07. فرهنگ دادهٔ T08 اینجا تکرار نشده و فقط با شناسه ارجاع شده است.")
    w("")
    w("## ۱. خلاصه")
    w("")
    n_fields = sum(1 for _ in all_items())
    by_kind = defaultdict(int)
    for _, _, it in all_items():
        by_kind[it["kind"]] += 1
    w(f"- **شمار فرم:** ۶ فرم و ۱ کارت مشترک، با {fa(n_fields)} قلم ({fa(by_kind['field'])} فیلد، {fa(by_kind['table'])} جدول تکرارشونده، {fa(by_kind['check'])} فهرست پیوست و {fa(by_kind['sign'])} جدول تأیید).")
    w("- **پوشش:** مسیر الف با فرم ۱، ۲، ۵ و ۶؛ مسیر ب با فرم ۳، ۴، ۵ و ۶؛ کارت طرف در هر دو.")
    w(f"- **نگاشت:** {fa(len(FM) - len(omitted))} از {fa(len(FM))} فیلد T08 در فرم‌ها آمده‌اند؛ {fa(len(omitted))} فیلد (FM06-10، FM-09، FM-10، FM-11) بیرون از این دو مسیرند (بخش ۴).")
    w("- **نمونه‌ها:** هیچ نمونه‌ای «تکمیل‌شده با مدارک واقعی» نیست. برای هر مسیر یک «شاهد دفتر؛ پروندهٔ قرارداد در دسترس نیست» و فهرست مدارک لازم آمده است (بخش ۷).")
    w("- **بازاستفاده:** گردش کار فعال قرارداد حجمی (گام‌های ۱۰۱ تا ۱۰۴، قالب‌های ۱۸ و ۶ ستونی، جریمه و تطبیق حضور) برای مسیر الف بازاستفاده‌پذیر است، با تغییر پارامتر. برای مسیر ب زیرساختی نیست (بخش ۵).")
    w("- **هیچ کد برنامه، منبع، تحویل پیشین یا فایل Word و PowerPoint جلسه‌های قبل تغییر نکرد.** نامه‌ای ارسال نشد.")
    w("")
    w("## ۲. فهرست فرم‌ها، نقش تکمیل‌کننده و مرحلهٔ استفاده")
    w("")
    w("| فرم | عنوان | مسیر | تکمیل‌کننده و مرحله (پیشنهادی) | منبع اطلاعات | بازبین |")
    w("|---|---|---|---|---|---|")
    for f in C.FORMS:
        path = {"A": "الف", "B": "ب", "AB": "هر دو"}[f["path"]]
        w(f"| {f['id']} ({'فرم ' + f['no'] if f['id'] != 'FC' else 'کارت'}) | {f['title']} | {path} | {f['card']['who']} | {f['card']['source']} | {f['card']['reviewer']} |")
    w("")
    w("### ۲-۱. دلیل شمار و چیدمان")
    w("")
    for b in [
        "۱۳ فرم یا موجودیت T08 مکانیکی به ۱۳ فرم تبدیل نشد. معیار: هر داده یک بار و نزد کسی که سند اصلی را دارد.",
        "**قرارداد** به دو فرم جهت‌دار تقسیم شد (فرم ۱ خدمات، فرم ۳ اجاره)، چون مسیر و تکمیل‌کننده و منبع متفاوت دارند؛ هر دو فیلدهای FM-03 را با همان عنصر داده دارند.",
        "**محل اجاره (FM-02)** در فرم ۳ آمد، چون زیرمحل و قرارداد هم‌زمان از یک سند ثبت می‌شوند. فهرست پایانه‌ها و واحدها دادهٔ پایهٔ مشترک و فقط انتخاب‌کردنی است.",
        "**مطالبه، دریافت و تخصیص (FM-06، FM-07)** یک فرم با سه بخش جدا و امضای جدا شدند تا مرزشان حفظ شود ولی پرونده یکجا دیده شود. شاهد نمونهٔ ب نشان می‌دهد دریافت ممکن است پیش از مطالبه برسد؛ پس هر بخش مستقل پر می‌شود.",
        "**ابزار و سپرده (FM-04، FM-08)** یک فرم مشترک شدند؛ ماهیت حقوقی (سپردهٔ نقدی، تضمین غیرنقدی، چک پرداخت) فیلد اول است و فهرست وضعیت مجاز را تعیین می‌کند.",
        "**نیرو و کارکرد** بخش شرطی فرم ۲ است، نه فرم جدا؛ کلید آن «مبنای مبلغ» در فرم ۱ است.",
        "**سطر طبقه‌بندی و بسته (FM-12، FM-13)** فرم کاربر عملیاتی نیستند؛ فقط در پیش‌نمایش (فرم ۶) با برچسب «پیش‌نویس — قابل ارجاع نهایی نیست» دیده می‌شوند.",
        "**رجیستری طرف (FM-01)** کارت مشترک شد که فقط برای طرف تازه پر می‌شود.",
    ]:
        w(f"- {b}")
    w("")
    w("### ۲-۲. تبار ادغام فرم‌های T08")
    w("")
    w("| فرم T08 | عنوان | جای آن در T09 | توضیح |")
    w("|---|---|---|---|")
    for r in C.CONSOLIDATION:
        w(f"| {r[0]} | {r[1]} | {r[2]} | {r[3]} |")
    w("")
    w("## ۳. نگاشت فیلدها به عناصر داده و نیازهای T08")
    w("")
    w("شمارهٔ فیلد در Word با ارقام فارسی آمده است (مثلاً «۲-۰۷» = `F2-07`؛ «ک-۰۳» = `FC-03`). ستون «ردیف‌های نیاز مسیر» ردیف‌های شیت «نیاز_داده_مسیرها» را نشان می‌دهد که همان (مرحله، عنصر) را دارند. "
      "«نیاز T09» سطح پیشنهادی در فرم است؛ «الزام T08» ستون «پیشنهاد — الزام» همان ردیف‌ها و فیلدهای T08 است. ستون پروفایل T08 در تعیین نیاز به کار نرفت (کنترل K-T06).")
    w("")
    for form in C.FORMS:
        w(f"### {form['id']} — {form['title']}")
        w("")
        w("| فیلد | عنوان | فیلد T08 | عنصر داده | ردیف‌های نیاز مسیر | نیاز T09 | الزام T08 | دلیل طراحی |")
        w("|---|---|---|---|---|---|---|---|")
        for sec, it in items(form):
            pvs = pv_of(it)
            if pvs:
                byl = OrderedDict()
                for p in pvs:
                    byl.setdefault(p["lp"], []).append(p["id"])
                pvtxt = "؛ ".join(f"{lp}: {len(v)} ({v[0]}…{v[-1]})" if len(v) > 1 else f"{lp}: {v[0]}" for lp, v in byl.items())
            else:
                pvtxt = "—"
            req = "/".join(sorted(t08_reqs(it))) or "—"
            w(f"| {it['fid']} | {it['label']} | {'، '.join(it['fm']) or '—'} | {'، '.join(it['el']) or '—'} | {pvtxt} | {C.NEED[it['need']]} | {req} | {it.get('why') or '—'} |")
        w("")
    w("## ۴. فیلدها و مرحله‌های T08 بیرون از این دو مسیر")
    w("")
    groups = OrderedDict([("FM06-10", "گزارش فروش بلیط مبنای عوارض ۵٪ است، نه قرارداد اجاره؛ مالک درآمد عوارض منوط به T-D2 است و سند دامنه آن را «نیازمند تصمیم» دانسته است."),
                          ("FM09-", "دریافت خدمات تاکسیرانی (مسیر F5)؛ فرم و پایلوت آن به فاز بعد رفته است (سند دامنه، ردیف ۱۰)."),
                          ("FM10-", "فهرست حقوق (مسیر F6)؛ انتساب حقوق به فاز بعد رفته است (CD-08)."),
                          ("FM11-", "اعلامیهٔ فی‌مابین (مسیر F6)؛ تهاتر با شهرداری «نیازمند تصمیم» است (DE-01، DE-24).")])
    w("| فیلدهای T08 | شمار | دلیل |")
    w("|---|---|---|")
    for pre, why in groups.items():
        ids = [f for f in omitted if f.startswith(pre)]
        w(f"| {'، '.join(ids)} | {fa(len(ids))} | {why} |")
    w("")
    w("**مرحله‌های T08 بیرون از این دو مسیر:**")
    w("")
    w("| مرحله | عنوان | دلیل |")
    w("|---|---|---|")
    for lp, why in C.OUT_OF_SCOPE_LP.items():
        w(f"| {lp} | {LPS[lp]['stage']} | {why} |")
    for lp, why in C.PARTIAL_LP.items():
        w(f"| {lp} (پوشش جزئی) | {LPS[lp]['stage']} | {why} |")
    w("")
    if uncovered:
        w("**ردیف‌های در دامنه که فیلد جدا نگرفتند:**")
        w("")
        w("| مرحله | عنصر | الزام T08 | ردیف‌ها | دلیل |")
        w("|---|---|---|---|---|")
        for (lp, el), v in uncovered:
            w(f"| {lp} | {el} | {'/'.join(sorted({x['req'] for x in v}))} | {', '.join(x['id'] for x in v)} | {C.UNCOVERED_WHY.get((lp, el), 'بررسی شود')} |")
        w("")
    w("## ۵. بازاستفاده از گردش کار قرارداد حجمی موجود")
    w("")
    w("وضعیت‌ها: «قابل استفاده با همان عنوان»، «نیازمند تغییر پارامتر یا عنوان»، «نیازمند بخش تازه»، «خارج از دامنهٔ این دو مسیر». ستون «فعال یا کهنه» گردش کار در حال اجرا را از پیکربندی کهنه جدا می‌کند. "
      "شواهد از کد فعلی مخزن خوانده شد و کدی تغییر نکرد. گردش کار فعال حجمی چهار گام دارد (طرح ۱۰ مرحله‌ای در همان اسکریپت بازنویسی می‌شود) و گام ۱۰۳→۱۰۴ کد یکتا و تراکنش می‌سازد.")
    w("")
    w("| قلم | محل در کد | فعال یا کهنه | وضعیت | جای آن در فرم‌ها | یادداشت |")
    w("|---|---|---|---|---|---|")
    for r in REUSE:
        w("| " + " | ".join(r) + " |")
    w("")
    w("نقص‌های ممیزی فنی (`docs/contractor-hr/TECHNICAL_AUDIT_CONTRACTOR_HR_WORKFLOW.md` §۰ و §۲٫۳) مانند سه ماشین حالت ناهمگون و صدور غیراتمیک کد، پیش‌شرط پایلوت نرم‌افزاری‌اند، نه پیش‌شرط بازبینی فرم.")
    w("")
    w("## ۶. تصمیم‌های باز و اثر ساختاری بر فرم‌ها")
    w("")
    w("| تصمیم | موضوع | فیلدهای متأثر | رفتار فرم تا پاسخ |")
    w("|---|---|---|---|")
    for r in DECISIONS_MD:
        w("| " + " | ".join(r) + " |")
    w("")
    w("رابطهٔ چند به چند سطر طبقه‌بندی و سطر سند (PO-04 در RL-03) در فرم ۶ بخش ۴ فقط «گزینهٔ طراحی پیشنهادی» است و تا T-D8 خالی می‌ماند.")
    w("")
    w("## ۷. مدارک لازم برای تکمیل دو نمونهٔ مرجع")
    w("")
    w("هیچ‌یک از دو نمونهٔ آقای کلانی «تکمیل‌شده» نیست. درخواست اطلاعات T03 هنوز ارسال نشده است (همهٔ ردیف‌ها «آمادهٔ درخواست»). جست‌وجوی تازه در `H:/Freelancing_Project/KalaniProject` (۲۰۲۶-۱۰-۲۹، نام فایل) پروندهٔ تازه‌ای نیافت.")
    w("")
    w("| نمونه | مدرک | بند درخواست T03 | فیلدهایی که کامل می‌کند |")
    w("|---|---|---|---|")
    for r in [
        ("الف", "متن قرارداد امضاشده با پیوست‌ها و الحاقیه‌ها", "ب-۱", "فرم ۱"),
        ("الف", "یک دورهٔ کامل صورت‌وضعیت با تأییدها، کسور، سند حسابداری و مدرک پرداخت", "ب-۲", "فرم ۲ بخش‌های ۱ تا ۸"),
        ("الف", "پیوست کارکرد و حضور همان دوره (شرطی)", "ب-۳", "فرم ۲ بخش ۴"),
        ("الف", "ردیف بودجه یا اعتبار", "ب-۴", "۱-۲۵، ۲-۳۳"),
        ("الف", f"پیشنهاد افزوده: پروندهٔ همان قراردادی که سری «انتظامات» دفتر تاکسیرانی ۱۴۰۴ ({SA['per']['doc']} و {SA['pay_doc']}) به آن مربوط است، تا شاهد دفتر به پرونده وصل شود", "ب-۱، ب-۲ (انتخاب نمونه)", "پیوند بخش ۷-الف Word"),
        ("ب", "قرارداد اجارهٔ غرفه یا استقرار شرکت با طرفین و امضاکنندگان", "ج-۱", "فرم ۳"),
        ("ب", "یک دورهٔ کامل مطالبه و وصول: اعلامیه، رسید یا چک، سند حسابداری", "ج-۲", "فرم ۴"),
        ("ب", "اسناد مزایده، رسید سپرده، فهرست چک‌ها، الحاقیه، استرداد سپرده (شرطی)", "ج-۳", "فرم ۵؛ ۳-۱۱"),
        ("ب", f"پیشنهاد افزوده: پروندهٔ غرفهٔ کد {SB['V']} (پایانهٔ کاوه) که در بخش ۷-ب Word آمده است", "ج-۱ تا ج-۳ (انتخاب نمونه)", "پیوند بخش ۷-ب Word"),
        ("هر دو", "فرد توضیح‌دهندهٔ گردش کار؛ فرم‌های فعلی", "د-۱، د-۲", "نقش‌ها و چیدمان"),
        ("هر دو", "فهرست کاری فعالیت‌ها، پایانه‌ها، واحدها و مراکز", "د-۳، د-۴", "فعالیت، محل و مرکز پیش‌فرض"),
        ("هر دو", "واحد پول دفاتر", "هـ-۱", "نمایش مبالغ"),
        ("ب (و الف در پایانه‌ها)", "جدول نقش سازمان و شهرداری در اجاره، شارژ و هزینه‌ها", "هـ-۲", "۱-۱۰، ۳-۱۳، ۳-۱۴، ۳-۲۳"),
    ]:
        w("| " + " | ".join(r) + " |")
    w("")
    w("## ۸. نمونه‌های دفترمحور")
    w("")
    w(f"منبع: `data/work/taxi_terminals/tt_v21.sqlite` (فقط‌خواندنی؛ هش `{sha16(DB)}`؛ کپی در حافظه). قاعدهٔ انتخاب در `t09_build.py` (`sample_a`، `sample_b`) اجرا می‌شود و قابل بازتولید است. "
      "شرح ردیف‌ها و نام اشخاص نقل نشد؛ شمارهٔ حساب بانکی و چک پوشانده شد.")
    w("")
    per = SA["per"]
    w("**نمونهٔ الف (تاکسیرانی):**")
    w("")
    w(f"- انتخاب: سطرهای «خدمات پیمانکاری» (`612-5-14-0`) ۱۴۰۴ با واژهٔ «انتظامات» در شرح → {len(SA['series'])} سند. تنها سندی که به پرداختنی (`211-1-3-0`) رفته: {per['doc']} ({per['date']})؛ پرداخت: {SA['pay_doc']}.")
    w(f"- مشاهده: جمع دو سطر هزینه {per['gross']:,}؛ سطر `114-9-1-0` «مالیات» {per['m114']:,}؛ «سپرده بیمه» {per['ins']:,}؛ پرداختنی {per['pay']:,} = هزینه + «مالیات» − سپرده. در هر {len(SA['series'])} سند، سپردهٔ بیمه ۵٫۰٪ جمع هزینه است و سپردهٔ حسن انجام کار نیست.")
    w(f"- کد طرف‌حساب {SA['cp']} در {SA['other_docs']} سند هزینهٔ دیگر با شرح دیگر (تأمین نیرو) هم آمده است → کد طرف‌حساب هویت قرارداد نیست (DR-01).")
    w("- آنچه ثابت نمی‌شود: وجود یک قرارداد واحد، شمارهٔ صورت‌وضعیت، دورهٔ خدمت، نرخ قراردادی، ماهیت سطر «مالیات» و ذی‌نفع پرداخت.")
    w("")
    ff, clr = SB["first_full"], SB["clear_doc"]
    w("**نمونهٔ ب (پایانه‌ها):**")
    w("")
    w(f"- انتخاب: از {SB['n_full']} غرفهٔ ۱۴۰۴ با هر هفت مرحله و یک کد طرف‌حساب، کوچک‌ترین کد واحد: `{SB['V']}`. مراحل دیده‌شده: {'، '.join(STAGE_FA[x] for x in SB['stages'])}.")
    w(f"- مطالبه‌ها: {SB['n_claim_docs']} سند ماهانه؛ مبالغ: " + "؛ ".join(f"{a:,} × {len(v)}" for a, v in SB["amounts"].items()) + ".")
    w(f"- چک‌ها: {len(SB['recv'])} چک {SB['chq_amt']:,} در {SB['recv'][0][0]} ({SB['recv'][0][2]})، ارسال در {SB['send_this'][0]}. دورهٔ نمونه: مطالبهٔ {ff[0]} ({ff[2]})؛ وصول چک {SB['chq_mask']} و تسویه در {clr[0]} ({clr[2]})، پیش از ثبت مطالبه.")
    kinds = list(zip(SB["dep_in"], SB["dep_kind"]))
    w(f"- سپرده (`212-6-5-0`): " + "؛ ".join(f"{r[0]} ردیف {r[1]} ({r[2]}) {int(r[5]):,} [{ {'bank': 'هم‌مبلغ با واریز بانکی مجاور', 'auction': 'انتقال از سپردهٔ مزایده 212-6-2-0', 'other': 'بدون سطر مقابل مشخص'}[k] }]" for r, k in kinds)
      + f". استرداد {int(SB['refund'][4]):,} در {SB['refund'][0]}؛ ماندهٔ آغاز ۱۴۰۴ {SB['opening']:,} و پایان {SB['closing']:,}.")
    w(f"- سطر طرف مقابل مطالبه در سند ماهانه در سطح پایانه تجمیع شده است ({SB['claim_doc_cr'][0]} سطر `211-1-1-x` برای {SB['claim_doc_cr'][1]} واحد)، پس قلم مطالبه در سطر غرفه دیده نمی‌شود.")
    w("- آنچه ثابت نمی‌شود: طرف و امضاکنندهٔ قرارداد، مزایده، مبلغ و اقلام قراردادی، پیوند وصول با مطالبهٔ مشخص و مالک درآمد.")
    w("")
    w("## ۹. کنترل‌ها")
    w("")
    w("| کنترل | شرح | وضعیت | جزئیات |")
    w("|---|---|---|---|")
    for c in ctrl:
        w(f"| {c['id']} | {c['title']} | {c['status']} | {c['detail']} |")
    w("")
    w("## ۱۰. بررسی دیداری Word و PDF")
    w("")
    if visual:
        w(f"PDF مستقیم از همین Word با Microsoft Word (`ExportAsFixedFormat`) ساخته شد. Word {fa(n_pages_word)} صفحه و PDF {fa(n_pages_pdf)} صفحه گزارش کرد. صفحه‌ها با PyMuPDF به تصویر تبدیل و دیده شدند.")
        w("")
        for b in visual:
            w(f"- {b}")
    else:
        w("انجام نشده است.")
    w("")
    w("## ۱۱. آمادگی، بازتولید و محدودیت‌ها")
    w("")
    w("- **آمادگی برای بازبینی فرم:** بله. فرم‌ها، راهنما، نگاشت و وضعیت شواهد کامل‌اند.")
    w("- **آمادگی برای پایلوت یا ارجاع نهایی:** خیر. پیش‌شرط‌ها: پروندهٔ واقعی دو نمونه (T-D4، T-D5)، پاسخ T-D2، T-D3، T-D6، T-D7، T-D8، قاعدهٔ مرجع منتشرشده (هـ-۴) و رفع صدور غیراتمیک کد در نرم‌افزار.")
    w("- **بازتولید:** `py -3.14 scripts/passenger_transport/t09_build.py` (محتوا در `t09_content.py`). کنترل‌ها در `data/work/passenger_transport/t09/controls.json` و تصویر صفحه‌ها در `data/work/passenger_transport/t09/png/`.")
    w("- **محدودیت‌ها:** فرم‌ها با پروندهٔ واقعی آزموده نشده‌اند؛ نقش‌ها پیشنهادی‌اند (T03 د-۱)؛ فهرست فعالیت و مرکز T05 تأیید نشده است؛ دادهٔ ۱۴۰۵ ناقص و واحد پول تأییدنشده است؛ نشانه‌های شرح فقط نشانه‌اند.")
    OUT_MD.write_text("\n".join(L) + "\n", encoding="utf-8")


def main():
    md_only = "--md-only" in sys.argv
    WORK.mkdir(parents=True, exist_ok=True)
    ctrl, uncovered, omitted = run_controls()
    con = open_db()
    SA = sample_a(con)
    SB = sample_b(con)
    con.close()
    n_word = n_pdf = 0
    if not md_only:
        build_docx(SA, SB)
        n_word = export_pdf()
        n_pdf = render_png()
    else:
        import fitz
        d = fitz.open(OUT_PDF)
        n_pdf = d.page_count
        d.close()
        n_word = json.loads((WORK / "controls.json").read_text(encoding="utf-8")).get("pages_word", n_pdf)
    visual = json.loads(VISUAL.read_text(encoding="utf-8")) if VISUAL.exists() else []
    build_md(SA, SB, ctrl, uncovered, omitted, visual, n_word, n_pdf)
    ctrl_all = ctrl + scan_outputs()
    build_md(SA, SB, ctrl_all, uncovered, omitted, visual, n_word, n_pdf)
    (WORK / "controls.json").write_text(json.dumps(dict(controls=ctrl_all, pages_word=n_word, pages_pdf=n_pdf,
                                                        uncovered=[f"{a}/{b}" for (a, b), _ in uncovered]), ensure_ascii=False, indent=1),
                                        encoding="utf-8")
    for c in ctrl_all:
        print(c["id"], c["status"], c["detail"][:300])
    print("pages word/pdf:", n_word, n_pdf)


if __name__ == "__main__":
    main()
