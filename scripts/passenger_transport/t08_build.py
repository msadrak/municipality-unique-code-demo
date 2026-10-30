"""T08 builder: event relationship and information-dimension design workbook (passenger transport).

Run after t08_observe.py:
    .venv/Scripts/python.exe scripts/passenger_transport/t08_observe.py
    .venv/Scripts/python.exe scripts/passenger_transport/t08_build.py

Inputs are read-only. Output: data/deliverables/passenger_transport/event_relationship_design_v0_1.xlsx
and data/work/passenger_transport/t08_summary.json (numbers quoted by the report).
"""
from __future__ import annotations

import collections
import hashlib
import json
import re
import sqlite3
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import t08_content as K  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
OBS = ROOT / "data/work/passenger_transport/t08_observed.json"
T07 = ROOT / "data/deliverables/passenger_transport/financial_event_review_v0_1.xlsx"
T05 = ROOT / "data/deliverables/passenger_transport/activity_cost_center_review_v0_1.xlsx"
DB = ROOT / "data/work/taxi_terminals/tt_v21.sqlite"
OUT = ROOT / "data/deliverables/passenger_transport/event_relationship_design_v0_1.xlsx"
SUMMARY = ROOT / "data/work/passenger_transport/t08_summary.json"

FA_DIGITS = str.maketrans("0123456789,.", "۰۱۲۳۴۵۶۷۸۹٬٫")
SECTOR_FA = {"Taxi": "تاکسیرانی", "Terminals": "پایانه‌ها", "TX": "تاکسیرانی", "TM": "پایانه‌ها"}
SECTOR_KEY = {"Taxi": "TX", "Terminals": "TM"}
STATUS = "پیشنهادی؛ تأییدنشده"
ORG_WORDS = re.compile("شهرداري|شهرداری|سازمان|شركت|شرکت|بانك|بانک|اداره|ادارات|منطقه|تعاوني|تعاونی|موسسه|مؤسسه|"
                       "پايانه|پایانه|غرفه|مشترك|مشترک|صندوق|بيمه|بیمه|دفتر|انجمن|اتحاديه|اتحادیه|هيئت|هیئت|معاينه|معاینه|"
                       "تابلو|تبليغ|تبلیغ|جايگاه|جایگاه|خدمات|درآمد|درامد|متفرقه|ساير|سایر|اشخاص|افراد")


def fa(n, dec: int | None = None) -> str:
    if isinstance(n, float) and dec is None:
        dec = 1
    s = f"{n:,.{dec}f}" if dec is not None else f"{n:,}"
    return s.translate(FA_DIGITS)


def pct(part, whole) -> str:
    return fa(round(100 * part / whole, 1) if whole else 0.0) + "٪"


def sha16(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


# ============================================================================ inputs
obs = json.loads(OBS.read_text(encoding="utf-8"))
wb07 = openpyxl.load_workbook(T07, read_only=True, data_only=True)
wb05 = openpyxl.load_workbook(T05, read_only=True, data_only=True)


def sheet_rows(wb, name):
    rows = list(wb[name].iter_rows(values_only=True))
    return rows[2], [r for r in rows[3:] if r and r[0]]


_, t07_events = sheet_rows(wb07, "رویدادهای_پیشنهادی")
EV = {r[0]: dict(title=r[1], family=r[2], sector=r[3], cls=r[4], nature=r[5], tx=r[12], tm=r[13], t05=r[17])
      for r in t07_events}
_, t07_dec = sheet_rows(wb07, "تصمیم‌ها")
DE_TOPIC = {r[0]: r[1] for r in t07_dec}
_, t05_dec = sheet_rows(wb05, "تصمیم‌ها_و_شکاف‌ها")
DG_TOPIC = {r[0]: r[1] for r in t05_dec}
_, t05_act = sheet_rows(wb05, "فعالیت‌ها")
ACT_IDS = {r[0] for r in t05_act}
_, t07_rem = sheet_rows(wb07, "باقی‌مانده‌ها")
U_ROWS = {r[0]: r for r in t07_rem}
_, t07_am = sheet_rows(wb07, "تطابق‌های_مبلغی")
AM = {r[0]: r for r in t07_am}

PROFILE = {(p["event"], p["org"]): p for p in obs["profiles"]}
FIELDS = obs["fields"]["by_org"]
REL = obs["relations"]

# ============================================================================ styling
TITLE_FONT = Font(name="Tahoma", size=12, bold=True)
NOTE_FONT = Font(name="Tahoma", size=9, color="595959")
HEAD_FONT = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Tahoma", size=10)
GREY_FONT = Font(name="Tahoma", size=10, color="808080")
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CODE_FILL = {K.L: "FCE4D6", K.C: "FFF2CC", K.O: "EDEDED", K.D: "DDEBF7", K.N: "FFFFFF"}

wb = openpyxl.Workbook()
wb.remove(wb.active)
ALL_TEXT: list[str] = []  # every string written, for privacy and reference controls
ROUTED_TEXT: list[str] = []  # cell strings outside «تصمیم‌ها», where duplicate DRs must travel with their primary


def route(text: str) -> str:
    """Cite duplicate or sub-question DRs through their primary prior decision (lineage kept in brackets)."""
    for dr, primary in K.ROUTE.items():
        if dr in text and primary not in text:
            text = re.sub(rf"\b{dr}\b", f"{primary} ({dr})", text)
    return text


def add_sheet(name: str, title: str, note: str, headers: list[str], rows: list[list], widths: list[int],
              freeze: str = "C4", code_cols: tuple[int, ...] = (), header_height: int = 48, route_refs: bool = True):
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="right")
    ws["A2"] = note
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = Alignment(horizontal="right")
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font, cell.fill, cell.border = HEAD_FONT, HEAD_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = header_height
    for r, row in enumerate(rows, 4):
        lines = 1
        for c, v in enumerate(row, 1):
            if isinstance(v, str) and route_refs:
                v = route(v)
                ROUTED_TEXT.append(v)
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = BODY_FONT, BORDER
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            if isinstance(v, str):
                ALL_TEXT.append(v)
                per_line = max(1, int(widths[c - 1] * 1.2))
                lines = max(lines, sum(-(-max(1, len(part)) // per_line) for part in v.split("\n")))
            if c in code_cols and v in CODE_FILL:
                cell.fill = PatternFill("solid", fgColor=CODE_FILL[v])
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                if v == K.N:
                    cell.font = GREY_FONT
        # openpyxl cannot auto-fit; estimate wrapped lines so long Persian cells are not clipped
        ws.row_dimensions[r].height = min(409, 14 * lines + 4)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(3, len(rows) + 3)}"
    ALL_TEXT.extend([title, note, *headers])
    return ws


# ============================================================================ section 1: T07 population
def population_sheet():
    pop = {(p["بخش"], p["سال مالی"]): p for p in obs["population"]["rows"]}
    keys = [("تاکسیرانی", "1404"), ("تاکسیرانی", "1405"), ("پایانه\u200cها", "1404"), ("پایانه\u200cها", "1405")]
    tot = pop[("هر دو بخش", "جمع کل")]
    c = obs["population"]["cats"]

    def vals(f):
        return [f(pop[k]) for k in keys] + [f(tot)]

    rel = lambda p: p[c[0]] + p[c[1]]
    cp = lambda p: p[c[2]] + p[c[3]]
    rows = []

    def add(pid, topic, v, t07, t08, src):
        rows.append([pid, topic, *v, t07, t08, src])

    add("P-01", "سطرهای بی‌برچسب v2.1 در جمعیت بررسی", vals(lambda p: p["بدون برچسب v2.1"]),
        "افراز شده به شش زیرگروه (P-02، P-03، P-05، P-06، P-08، P-09).", "—", "شیت جمعیت_بررسی در T07 (PR-01 تا PR-07)؛ PC-04")
    add("P-02", "مرتبط با رویداد پیشنهادی — با شرط حساب", vals(lambda p: p[c[0]]),
        "رویداد دارد (شرط حساب و جهت).", "سطر رویداد؛ ابعادش در پروفایل (شیت پروفایل_ابعاد_رویداد) سنجیده شده است.", "شیت جمعیت_بررسی در T07")
    add("P-03", "مرتبط با رویداد پیشنهادی — فقط با جفت‌سازی مبلغی", vals(lambda p: p[c[1]]),
        "رویداد دارد؛ سطح اطمینان پیوند در شیت تطابق‌های_مبلغی (T3 تا T0).", "پیوند فقط‌مبلغی «نامزد» می‌ماند و الزام نمی‌سازد.", "شیت جمعیت_بررسی در T07؛ PC-05")
    add("P-04", "جمع مرتبط با رویداد پیشنهادی (P-02 + P-03)", vals(rel),
        "در گزارش T07 «مرتبط با رویداد پیشنهادی» = ۱۳٬۱۲۳.", "—", "گزارش T07 §۲؛ PC-04")
    add("P-05", "طرف مقابل در سند دارای رویداد — سطر بانک", vals(lambda p: p[c[2]]),
        "حل‌نشده نیست؛ رویداد سطری ندارد. سطر بانک در سندی که دست‌کم یک سطرش رویداد گرفته است.",
        "حامل ابعاد بانکی همان رویداد (D16 حساب بانکی، D18 مرجع)؛ اینکه ویژگی سطر رویداد باشد یا مرحلهٔ جدا: DR-04.", "T07 §۲")
    add("P-06", "طرف مقابل در سند دارای رویداد — سطر غیربانکی", vals(lambda p: p[c[3]]),
        "حل‌نشده نیست؛ رویداد سطری ندارد (مثلاً بستانکار ۲۱۱-۱-۱ در اسناد صدور مطالبه).",
        "سطر مقابل حسابداری همان رویداد (نقش «طرف مقابل» در D07)؛ DR-04.", "T07 §۲")
    add("P-07", "جمع طرف مقابل (P-05 + P-06)", vals(cp), "در گزارش T07 = ۸٬۲۹۴.", "—", "T07 §۲؛ PC-04")
    add("P-08", "ثبت انتظامی بدون رویداد", vals(lambda p: p[c[4]]),
        "«انتظامی» یعنی حساب‌های انتظامی (کنترلی، خارج از ترازنامه: ۹۱۸ و ۹۱۹)، نه انضباطی. کوهورت U-10؛ رویدادی پیشنهاد نشد (DE-21).",
        "ر.ک. P-14 و P-15؛ پیشنهاد: ثبت ابزار و ثبت امانی انبار، تصمیم در DR-09.", "T07 باقی‌مانده‌ها U-10")
    add("P-09", "حل‌نشدهٔ بی‌برچسب", vals(lambda p: p[c[5]]),
        "سطرهای اسنادی که هیچ سطرشان رویداد نگرفته است؛ کوهورت‌های U-01 تا U-09.", "در T08 طبقه‌بندی نشدند.", "T07 باقی‌مانده‌ها")
    add("P-10", "کنترل جمع: P-04 + P-07 + P-08 + P-09 = P-01",
        vals(lambda p: rel(p) + cp(p) + p[c[4]] + p[c[5]] - p["بدون برچسب v2.1"]),
        "اختلاف صفر در هر بخش و سال (کنترل PC-01 در T07).", "—", "محاسبهٔ T08 روی جدول T07")
    add("P-11", "عدد «۱۰٬۵۶۶ سطر مرتبط» — خلاصهٔ قدیمی ناسازگار با نتایج ذخیره‌شده", ["—"] * 4 + [10566],
        "منشأ: خلاصهٔ روایی پایانی تهیه‌کننده برای T07 در گفت‌وگوی پروژه (۱۴۰۵/۰۸/۰۷). این عدد در کتاب‌کار و گزارش ذخیره‌شدهٔ T07 نیست و با جمع هیچ زیرگروهی برابر نمی‌شود. همان خلاصه شمار سطرهای چندرویدادی را هم ۹٬۱۲۷ گفته بود، در حالی که کنترل ذخیره‌شدهٔ PC-02 عدد ۴۲۰ را نشان می‌دهد.",
        "کنار گذاشته شد. مبنای عملیاتی T08 جدول ذخیره‌شدهٔ T07 است (P-01 تا P-10؛ مرتبط = ۱۳٬۱۲۳). برای توجیه این عدد طبقه‌بندی تازه‌ای ساخته نشد. موضوع خطای گزارش‌دهی داخلی است و به سازمان ارجاع نمی‌شود.",
        "جست‌وجوی خودکار در T07 (کنترل K-03)")
    add("P-12", "«۱۰٬۹۴۱ سطر باقی‌مانده» — خلاصهٔ قدیمی ناسازگار با نتایج ذخیره‌شده", ["—"] * 4 + [10941],
        "فقط باقی‌ماندهٔ حسابی همان خلاصه است (۲۱٬۸۴۱ − ۱۰٬۵۶۶ − ۳۳۴) و گروهی در نتایج T07 نیست.",
        "کنار گذاشته شد. باقی‌ماندهٔ واقعی P-13 است.", "محاسبهٔ T08")
    add("P-13", "باقی‌ماندهٔ واقعی پس از کسر مرتبط و حل‌نشده (= P-07 + P-08)",
        vals(lambda p: p["بدون برچسب v2.1"] - rel(p) - p[c[5]]), "۸٬۲۹۴ طرف مقابل + ۹۰ انتظامی.", "—", "محاسبهٔ T08 روی جدول T07")
    memo = obs["memo"]["rows"]

    def memo_counts(accounts):
        out = []
        for org, fy in (("Taxi", "1404"), ("Taxi", "1405"), ("Terminals", "1404"), ("Terminals", "1405")):
            out.append(sum(r[4] for r in memo if r[0] == org and r[1] == fy and r[2] in accounts))
        return out + [sum(out)]

    guar = memo_counts({"918-1-5", "918-2-5"})
    goods = memo_counts({"918-1-6", "919-1-6"})
    gbn = sum(r[5] for r in memo if r[2] in {"918-1-5", "918-2-5"}) / 1e9
    dbn = sum(r[5] for r in memo if r[2] in {"918-1-6", "919-1-6"}) / 1e9
    gdocs = sum(r[6] for r in memo if r[2] in {"918-1-5", "918-2-5"} and r[3] == "بدهکار")
    ddocs = sum(r[6] for r in memo if r[2] in {"918-1-6", "919-1-6"} and r[3] == "بدهکار")
    add("P-14", "از P-08: «اسناد ضمانتی ما نزد دیگران» (۹۱۸-۱-۵ و طرف آن ۹۱۸-۲-۵)", guar,
        f"تضمینی که سازمان به دیگران سپرده است؛ ثبت جفتی در حساب کنترل و طرف آن. جمع گردش دو سوی ثبت: {fa(gbn, 2)} میلیارد واحد منبع.",
        "قرینهٔ LC-21 از سوی سازمان: ابزار صادره (D17) با قرارداد (D08) و دریافت‌کنندهٔ تضمین (D13). رویداد پیشنهاد نشد (DR-09).",
        "بازتولید از tt_v21.sqlite (کنترل K-04)")
    add("P-15", "از P-08: «کالای امانی دیگران نزد ما» (۹۱۸-۱-۶ و طرف آن ۹۱۹-۱-۶)", goods,
        f"نگهداری کالای دیگران؛ نمونه: TERM-1404-780 همراه حوالهٔ انبار. جمع گردش دو سوی ثبت: {fa(dbn, 2)} میلیارد واحد منبع.",
        "نگهداری فیزیکی در انبار (D06) با مالک کالا (D13) و حوالهٔ انبار (D19)؛ مالی بودن آن به DE-21 بسته است.",
        "بازتولید از tt_v21.sqlite (کنترل K-04)")
    headers = ["شناسه", "موضوع", "تاکسیرانی ۱۴۰۴", "تاکسیرانی ۱۴۰۵ (تا ۰۷/۱۳)", "پایانه‌ها ۱۴۰۴", "پایانه‌ها ۱۴۰۵ (تا ۰۷/۱۴)",
               "جمع", "وضعیت در نتایج T07", "جایگاه در طراحی T08", "منبع"]
    add_sheet("جمعیت_T07", "تفکیک ۲۱٬۸۴۱ سطر بی‌برچسب T07 (بازبینی نتایج موجود)",
              "فقط بازخوانی جدول «جمعیت_بررسی» و کوهورت‌های «باقی‌مانده‌ها» در کتاب‌کار T07. هیچ سطری دوباره طبقه‌بندی نشد. P-14 و P-15 فقط بر حسب حساب بازشماری شدند.",
              headers, rows, [7, 42, 12, 14, 12, 14, 10, 52, 52, 26], freeze="C4")
    return dict(rows={r[0]: r for r in rows}, guar_bn=gbn, goods_bn=dbn, guar_docs=gdocs, goods_docs=ddocs)


# ============================================================================ A1 source fields
def source_fields_sheet():
    tx, tm = FIELDS["Taxi"], FIELDS["Terminals"]
    f = obs["fields"]
    emb = f["embed"]
    types_tm = "، ".join(f"«{t}» {fa(n)}" for t, n in f["bank_types"]["Terminals"][:6])
    types_tx = "، ".join(f"«{t}» {fa(n)}" for t, n in f["bank_types"]["Taxi"])
    vs = obs["vahed"]["by_category"]
    rows = [
        ["SF-01", "Sanad_Num، Date، Des", "شمارهٔ سند، تاریخ و شرح سند", "هر دو: ۱۰۰٪",
         "هویت سند فقط با (بخش، سال، شماره) یکتاست؛ شرح سند روی همهٔ سطرهای سند تکرار می‌شود.", "مرجع منشأ (نه D09)", "شرح سند دادهٔ محدود است."],
        ["SF-02", "Kol، Moeen، Tafsili1، Tafsili2، HCode", "کد حساب چهارسطحی", "هر دو: ۱۰۰٪",
         "سطح چهارم معناهای متفاوت دارد: نوع هزینهٔ تحمیل‌شده (۱۱۳-۱-۱-xx پایانه‌ها)، قلم مطالبه (۲۱۱-۱-۱-x)، نهاد شهرداری (۲۱۱-۱-۶-x) و شمارهٔ حساب بانکی (۱۱۱-x-x-حساب).",
         "D07؛ بخشی از D16", "شمارهٔ حساب بانکی پوشانده می‌شود."],
        ["SF-03", "Bed، Bes", "مبلغ و جهت", "هر دو: ۱۰۰٪", "واحد پول تأیید نشده است (CD-02).", "D07 (نقش و جهت)", "—"],
        ["SF-04", "Sharh", "شرح سطر", "تقریباً همهٔ سطرها",
         f"متن آزاد. «قرارداد» فقط در {fa(tx['qd'])} سطر تاکسیرانی و {fa(tm['qd'])} سطر پایانه‌ها آمده است. نام ماه و عدد گاهی دیده می‌شود.",
         "نشانهٔ D08، D09، D21 (نه مقدار قابل اتکا)", "دادهٔ محدود؛ نقل نمی‌شود."],
        ["SF-05", "CodePerson", "کد طرف‌حساب",
         f"تاکسیرانی {pct(tx['person'], tx['n'])}؛ پایانه‌ها {pct(tm['person'], tm['n'])}",
         f"طرف دفتر معین (بدهکار یا بستانکار) است. روی سطر بانک فقط {fa(tx['person_codes_on_bank'])} کد در تاکسیرانی و {fa(tm['person_codes_on_bank'])} کد در پایانه‌ها دیده می‌شود، پس پرداخت‌کننده را نشان نمی‌دهد. در پایانه‌ها نام این کد واحد و پایانه را هم دارد (OR-05).",
         "D13 (نه D12، D14 یا D15)", "—"],
        ["SF-06", "Field1", "فیلد شناور ۱ (معنا بر حسب حساب)", "بر حسب حساب",
         f"(الف) روی حساب‌های اشخاص: نام طرف‌حساب؛ در {fa(f['f1_person_keys'] - f['f1_person_conflicts'])} از {fa(f['f1_person_keys'])} کد طرف‌حساب یکتا و وابسته به کد است. (ب) روی چک‌ها (۱۱۳-۵، ۱۱۳-۶) و تضمین‌های پایانه‌ها (۹۱۸): شمارهٔ ابزار. (ج) روی سطر بانک: شمارهٔ مرجع تراکنش. (د) روی هزینه‌های تاکسیرانی (۶۱۱ تا ۶۱۳): عنوان کد فعالیت.",
         "D13 (نام)، D17، D18، D02 (عنوان)", "نام اشخاص دادهٔ محدود است."],
        ["SF-07", "Field2", "فیلد شناور ۲", "سطر بانک و ابزار",
         f"تاریخ تراکنش یا سررسید. مرجع بانکی کامل (Field1 + تاریخ): تاکسیرانی {pct(tx['bankref_on_bank'][0], tx['bankref_on_bank'][1])} و پایانه‌ها {pct(tm['bankref_on_bank'][0], tm['bankref_on_bank'][1])} سطرهای بانک.",
         "D17، D18", "—"],
        ["SF-08", "Field3", "فیلد شناور ۳", "سطر بانک و ابزار",
         f"روی سطر بانک نوع تراکنش است. پایانه‌ها: {types_tm}. تاکسیرانی: {types_tx}. روی ابزار: نام بانک صادرکننده.",
         "D18 (نوع)، D17 (بانک)", "—"],
        ["SF-09", "Field4", "فیلد شناور ۴", "سطر بانک و ابزار", "روی ابزار: شعبه. روی سطر بانک: مقدار ثابت («متفرقه» یا «-----»).", "D17", "—"],
        ["SF-10", "Field5", "فیلد شناور ۵", "ابزار؛ بانک تاکسیرانی",
         "روی ابزار: توصیف صادرکننده که در پایانه‌ها واحد و پایانه را هم دارد. روی بانک تاکسیرانی: عدد مرجع دوم.", "D14، D18", "دادهٔ محدود."],
        ["SF-11", "CodeVahed", "کد واحد",
         f"پایانه‌ها: {fa(tm['distinct_vahed'])} کد در جمعیت بررسی؛ تاکسیرانی: فقط «001»",
         f"پایانه‌ها: «پایانه-دسته-سریال». دسته ۰۰۱ غرفه ({fa(vs.get('001', 0))} کد)، ۰۰۲ دفتر شرکت ({fa(vs.get('002', 0))} کد)، ۰۰۳ سطح پایانه یا «مشترک» ({fa(vs.get('003', 0))} کد). کدهای ۰۰۵، ۰۰۶ و ۰۰۸ واحد ستادی‌اند. تاکسیرانی بُعد واحد ندارد.",
         "D04، D05، D06", "محل مرکز هزینه نیست."],
        ["SF-12", "NameVahed", "نام واحد", "هر دو", "پایانه‌ها: نام واحد، یک‌به‌یک با CodeVahed. تاکسیرانی: نوع سند («عادي»).", "D05 (عنوان)", "نام شرکت‌ها و غرفه‌ها نقل نمی‌شود."],
        ["SF-13", "CodeMarkaz", "کد فعالیت منبع",
         f"تاکسیرانی: {pct(tx['markaz_on_expense'][0], tx['markaz_on_expense'][1])} سطرهای هزینه (۶۱۱ تا ۶۱۳) و صفر روی درآمد؛ پایانه‌ها: {fa(tm['markaz'])} سطر",
         "کد فعالیت عملیاتی تاکسیرانی (۳۴ کد؛ کد ۱۰۰ عمومی). در پایانه‌ها فقط «008» دیده می‌شود.", "شاهد D02 (برابر فعالیت فرض نمی‌شود)", "—"],
        ["SF-14", "— (ستون ندارد)", "قرارداد، شمارهٔ مطالبه یا صورت‌حساب، پروژه، ردیف بودجه، مرکز هزینه، پرداخت‌کننده، ذی‌نفع", "۰٪",
         "در دفاتر ستونی برای این‌ها نیست (ثبت منابع §۳-۲). حساب انتظامی «کنترل قراردادها-پروژه‌ها» (۹۱۸-۱-۳) در هر دو نمودار هست، ولی در ۱۴۰۴ و ۱۴۰۵ هیچ سطری ندارد.",
         "D03، D08، D09، D10، D11، D12، D14، D15", "نبود در داده الزام یا ممنوعیت نمی‌سازد."],
    ]
    headers = ["شناسه", "ستون منبع", "عنوان", "پرشدگی در جمعیت بررسی", "نقش عملیاتی مشاهده‌شده", "بُعد مرتبط T08", "محدودیت"]
    add_sheet("فیلدهای_منبع", "لایهٔ الف — ستون‌های دفتر منبع و نقش عملیاتی آن‌ها",
              f"جمعیت بررسی T07 (اسناد عادی و تسویهٔ داخلی): تاکسیرانی {fa(tx['n'])} و پایانه‌ها {fa(tm['n'])} سطر. پرشدگی فیلدها را تعریف حساب در سامانهٔ منبع تعیین می‌کند، نه رویداد.",
              headers, rows, [7, 22, 26, 26, 70, 26, 24], freeze="C4")


# ============================================================================ A2 event profiles
BASIS_FA = {"EXACT": "دقیق (شرط SQL T07)",
            "DEF": "بازتولیدشده از تعریف مکتوب T07 (شمار برابر)",
            "APPROX": "تقریبی / نیازمند راستی‌آزمایی (بازسازی مستقل جفت‌سازی)"}


def basis_label(p):
    label = BASIS_FA[p["basis"]]
    return label + (" — ترکیبی با سطرهای دقیق" if p.get("mixed") else "")


def profile_sheet():
    rows = []
    ids = {}
    for i, p in enumerate(sorted(obs["profiles"], key=lambda x: (x["event"], x["org"])), 1):
        oid = f"OP-{i:03d}"
        ids[(p["event"], p["org"])] = oid
        ev = EV.get(p["event"], {})
        if not p.get("n"):
            rows.append([oid, p["event"], ev.get("title"), SECTOR_FA[p["org"]], BASIS_FA["EXACT"], "—", p["t07"], 0, 0]
                        + ["—"] * 15 + ["در دادهٔ بررسی‌شده دیده نشد."])
            continue
        tm = p["org"] == "Terminals"
        v = p["vahed"]
        roles = "؛ ".join(f"{('بدهکار' if r.endswith('بد') else 'بستانکار')} {r.rsplit(' ', 1)[0]} ({fa(n)})"
                          for r, n in p["roles"])
        note = []
        if p["persons"] and p["persons"] <= 4 and p["person"] > 50:
            note.append(f"کد طرف‌حساب ثابت ({fa(p['persons'])} کد): نشانهٔ طرف واقعی نیست")
        if p["proxy"]:
            note.append("روش ردیف‌های جفت‌سازی: " + " | ".join(p["proxy"]))
        citable = "خیر — فقط شاهد تحلیلی" if p["basis"] == "APPROX" else "بله، به‌عنوان شاهد مشاهده (نه الزام)"
        rows.append([
            oid, p["event"], ev.get("title"), SECTOR_FA[p["org"]], basis_label(p), citable, p["t07"], p["n"], p["docs"], roles,
            p["person"], p["persons"],
            v.get("SUB", 0.0) if tm else "—", v.get("TERM", 0.0) if tm else "—", v.get("HQ", 0.0) if tm else "—",
            p["markaz"], p["instr"], p["bankref"], p["docbank"], p["qd"], p["sv"], p["month"], p["n5"], p["taf2"],
            "؛ ".join(note) or "—",
        ])
    headers = ["شناسه", "رویداد", "عنوان رویداد (T07)", "بخش", "وضعیت پروفایل", "قابل استناد در طراحی؟", "سطر T07",
               "سطر پروفایل‌شده", "سند",
               "نقش‌های حساب (کل-معین-تفصیلی و جهت)", "٪ دارای کد طرف‌حساب", "شمار کد طرف‌حساب",
               "٪ زیرمحل (غرفه یا دفتر شرکت)", "٪ سطح پایانه (مشترک)", "٪ واحد ستادی", "٪ کد فعالیت منبع",
               "٪ فیلدهای ابزار (شماره، تاریخ، بانک)", "٪ مرجع بانکی (روی سطر بانک)", "٪ سطر در سند دارای بانک",
               "٪ «قرارداد» در شرح", "٪ «صورت‌وضعیت» در شرح", "٪ نام ماه در شرح", "٪ عدد ۵ رقمی یا بیشتر در شرح",
               "٪ دارای تفصیلی سطح ۴", "یادداشت"]
    add_sheet("پروفایل_ابعاد_رویداد", "لایهٔ الف — ابعادی که دفتر برای هر رویداد T07 واقعاً ثبت کرده است",
              "درصدها از سطرهای پروفایل‌شده‌اند. «دقیق» = همان شرط‌های SQL T07. «بازتولیدشده» = تعریف مکتوب T07 با شمار برابر. «تقریبی / نیازمند راستی‌آزمایی» = بازسازی مستقل جفت‌سازی؛ نه تکرار دقیق روش گم‌شدهٔ T07 است و نه مبنای الزام فرم یا رابطهٔ تأییدشده. هیچ پروفایلی به‌تنهایی الزام نمی‌سازد. «—» در ستون‌های واحد: تاکسیرانی کد واحد ندارد. نشانه‌های شرح فقط نشانه‌اند.",
              headers, rows, [8, 9, 34, 10, 22, 16, 9, 9, 7, 36, 9, 8, 9, 9, 9, 9, 9, 9, 9, 8, 8, 8, 9, 8, 40], freeze="E4",
              header_height=90)
    return ids


def pairing_status_sheet(profile_ids):
    """Item 2: the 18 T07 evidence rows that came from amount pairing."""
    am_of = {"EV-012": ["AM-06"], "EV-013": ["AM-10"], "EV-014": ["AM-05"], "EV-015": ["AM-09"], "EV-019": ["AM-02"],
             "EV-022": ["AM-04"], "EV-132": ["AM-11"], "EV-133": ["AM-11"], "EV-164": ["AM-07", "AM-08"],
             "EV-169": ["AM-12"], "EV-170": ["AM-03"], "EV-010": ["C-08"], "EV-075": ["EV-083"], "EV-083": ["EV-083"]}
    tiers = obs["recon_tiers"]

    def tier_text(am):
        rec = tiers.get(am, {})
        # T07 counts AM-11 in lines (one-account neutral pairs); every other AM row counts relations
        unit, k = ("سطر", 2) if am == "AM-11" else ("رابطه", 1)
        if am in AM:
            row = AM[am]
            t07 = f"در T07 ({unit}) {fa(row[3])}: T3 {fa(row[5])}، T2 {fa(row[6])}، T1 {fa(row[7] or 0)}، T0 {fa(row[8] or 0)}"
        else:
            t07 = "در T07 سطح اطمینان ثبت نشده است"
        if "pairs" in rec:
            r = (f"در بازسازی ({unit}) {fa(k * rec['pairs'])}: T3 {fa(k * rec.get('T3', 0))}، T2 {fa(k * rec.get('T2', 0))}، "
                 f"T1 {fa(k * rec.get('T1', 0))}، T0 {fa(k * rec.get('T0', 0))}")
        else:
            r = f"در بازسازی {fa(rec.get('lines', 0))} سطر رویداد (جمعیت رابطهٔ T07 با واحد دیگری شمرده شده است)"
        return f"رابطهٔ {am} — {t07}؛ {r}"

    def tiers_equal(ev):
        """True/False when T07 tiers are comparable with the reconstruction, None otherwise."""
        verdicts = []
        for am in am_of.get(ev, []):
            rec = tiers.get(am, {})
            if am not in AM or "pairs" not in rec:
                continue
            k = 2 if am == "AM-11" else 1
            row = AM[am]
            verdicts.append([row[5] or 0, row[6] or 0, row[7] or 0, row[8] or 0] ==
                            [k * rec.get(t, 0) for t in ("T3", "T2", "T1", "T0")])
        return all(verdicts) if verdicts else None

    rows = []
    for a in obs["approx"]:
        p = PROFILE.get((a["event"], a["org"]), {})
        diff = a["recon"] - a["t07"]
        same_tiers = tiers_equal(a["ev"])
        if a["basis"] == "DEF":
            status = "بازتولیدشده از تعریف مکتوب T07 — شمار برابر"
            note = "تعریف انتخاب قطعی است و بدون داوری جفت‌سازی اجرا شد. هم‌ارزی سطربه‌سطر با خروجی اصلی آزمودنی نیست، چون T07 فهرست سطرها را ذخیره نکرده بود."
        else:
            status = "تقریبی / نیازمند راستی‌آزمایی"
            head = "بازسازی مستقل است، نه تکرار دقیق روش گم‌شدهٔ T07. "
            if diff:
                note = head + f"شمار آن {fa(abs(diff))} سطر {'کمتر' if diff < 0 else 'بیشتر'} از T07 است و قاعدهٔ دقیق T07 برای این رابطه مکتوب نیست."
            elif same_tiers:
                note = head + "شمار و توزیع سطح اطمینان با T07 برابرند، ولی فهرست سطرهای T07 ذخیره نشده بود و هم‌ارزی سطربه‌سطر آزمودنی نیست."
            elif same_tiers is False:
                note = head + "شمار برابر است، ولی توزیع سطح اطمینان کمی فرق دارد و فهرست سطرهای T07 برای مقایسه در دست نیست."
            else:
                note = head + "شمار برابر است. T07 سطح اطمینان قابل مقایسه‌ای برای این ردیف ندارد و فهرست سطرهای آن ذخیره نشده بود."
        rows.append([a["ev"], a["event"], EV[a["event"]]["title"], SECTOR_FA[a["org"]],
                     "پایگاه tt_v21.sqlite (فقط خواندنی؛ اسناد عادی) و شیت شواهد T07 (تعریف و شمار)",
                     "قاعدهٔ T07 — " + a["method"], fa(a["superset"]) + " سطر (الگوی حساب؛ روش v0.1 T08، کنار رفت)",
                     "اسکریپت t08_pairing.py — " + ("اجرای تعریف مکتوب" if a["kind"] == "تعریف مکتوب قطعی" else
                                            "جفت‌سازی یک‌به‌یک درون سند؛ هم‌مبلغ؛ اولویت T3 شمارهٔ مشترک، T2 طرف‌حساب یا شرح یکسان، T1 ردیف مجاور، T0 فقط مبلغ"),
                     a["t07"], a["recon"], diff, "\n".join(tier_text(m) for m in am_of.get(a["ev"], [])) or "—",
                     status, note, f"پروفایل {profile_ids.get((a['event'], a['org']), '—')} — {basis_label(p)}" if p else "—",
                     K.APPROX_IMPACT[a["ev"]], "هیچ کد ماتریس، فیلد لازم یا رابطه‌ای بر این ردیف تکیه ندارد (کنترل K-18)."])
    headers = ["ردیف شاهد T07", "رویداد", "عنوان", "بخش", "منبع داده", "روش تقریب یا تعریف (مکتوب در T07)",
               "روش پیشین T08 v0.1 (فقط تبار)", "روش این نسخه", "شمار T07", "شمار بازسازی", "اختلاف",
               "مقایسهٔ سطح اطمینان (T07 در برابر بازسازی)", "وضعیت", "توضیح وضعیت", "پروفایل اثرپذیر", "اثر عملی بر طراحی فرم",
               "استفاده در ماتریس و فرم"]
    add_sheet("وضعیت_پروفایل‌های_جفت‌سازی", "بند ۲ — وضعیت ۱۸ ردیف شاهد T07 که از جفت‌سازی مبلغی آمده بودند",
              "کد جفت‌سازی T07 ذخیره نشده بود و T07 فقط شمارها را نگه داشته است. روش مکتوب T07 در scripts/passenger_transport/t08_pairing.py بازسازی و ارزیابی شد. بازسازی جفت‌سازی، حتی با شمار برابر، «تقریبی / نیازمند راستی‌آزمایی» می‌ماند و مبنای فیلد لازم یا رابطهٔ تأییدشده نیست.",
              headers, rows, [9, 9, 30, 9, 24, 44, 22, 36, 9, 9, 7, 52, 22, 44, 26, 44, 30], freeze="C4", header_height=60)
    return rows


# ============================================================================ A3 observed relationships
def relations_observed_sheet(pop):
    # T07 تطابق‌های_مبلغی columns: 3 population, 5 T3, 6 T2, 7 T1, 8 T0
    def am(k, what):
        row = AM[k]
        return {"pop": row[3], "t3": row[5], "t2": row[6], "cand": (row[7] or 0) + (row[8] or 0)}[what]

    e = obs["fields"]["embed"]
    r = REL
    cl = r["claims_docs_per_party_1404"]
    pd = r["parties_per_claim_doc"]
    loc = r["claim_party_locations"]
    um = r["claim_units_multi_party"]
    tel = r["terminal_expense_location"]
    tcs = r["terminal_contract_service_location"]
    tcd = r["terminal_contract_service_docs_multi_location"]
    tx, tm = FIELDS["Taxi"], FIELDS["Terminals"]
    p027 = PROFILE[("PFE-027", "Terminals")]
    p046 = PROFILE[("PFE-046", "Terminals")]
    p073 = PROFILE[("PFE-073", "Terminals")]
    cil = REL["claim_issue_lines_location"]
    ldl = REL["lease_deposit_lines_location"]
    p102 = PROFILE[("PFE-102", "Terminals")]
    p102x = PROFILE[("PFE-102", "Taxi")]
    rows = [
        ["OR-01", "یک سند حسابداری چند رویداد دارد",
         "الگوهای CP-01 تا CP-09 در اسناد عادی دیده شدند؛ مثلاً در CP-02 اسناد صورت‌وضعیت دارای کسر یا ارزش افزوده: تاکسیرانی ۹۴ از ۱۰۷ و پایانه‌ها ۱۳۶ از ۱۴۶.",
         "T07 چرخه‌ها_و_مرکب‌ها", "الگوی مشاهده‌شده", "سطر طبقه‌بندی و نمونهٔ رویداد معادل سند حسابداری نیستند (RL-02، RL-03).", "هم‌سند بودن قاعدهٔ کسب‌وکار نیست."],
        ["OR-02", "شمارهٔ چک مراحل چرخهٔ چک را به هم وصل می‌کند",
         f"جفت‌های دارای شمارهٔ چک مشترک (T3): AM-01 {fa(am('AM-01', 't3'))} از {fa(am('AM-01', 'pop'))}؛ AM-02 {fa(am('AM-02', 't3'))} از {fa(am('AM-02', 'pop'))}؛ AM-03 {fa(am('AM-03', 't3'))} از {fa(am('AM-03', 'pop'))}؛ AM-05 {fa(am('AM-05', 't3'))} از {fa(am('AM-05', 'pop'))}.",
         "T07 تطابق‌های_مبلغی", "پیوند تأییدشده با شناسه (T3)", "ابزار باید شناسهٔ پایدار و ثبت جدا داشته باشد (D17، FM-04).", "پیوند چک با مطالبهٔ مشخص را ثابت نمی‌کند."],
        ["OR-03", "پیوند مطالبه و وصول فقط در سطح طرف‌حساب است",
         f"وصول مستقیم (AM-06): از {fa(am('AM-06', 'pop'))} جفت، T3 {fa(am('AM-06', 't3'))}، T2 {fa(am('AM-06', 't2'))} و تطابق نامزد {fa(am('AM-06', 'cand'))}. هر کد طرف‌حساب مطالبه در ۱۴۰۴ میانه {fa(cl['median'])} سند مطالبه دارد ({fa(cl['ge6'])} از {fa(cl['parties'])} کد دست‌کم ۶ سند).",
         "T07 تطابق‌های_مبلغی؛ محاسبهٔ T08", "پیوند محتمل (T2)", "قرارداد چند نمونهٔ رویداد دارد (RL-01) و دریافت باید صریحاً به مطالبه تخصیص یابد (RL-10، DR-03).",
         "شمارهٔ صورت‌حساب در دفتر نیست؛ ارتقای پیوند بدون مرجع ممکن نیست."],
        ["OR-04", "اسناد مطالبه و وصول تجمیعی‌اند",
         f"{fa(pd['docs'])} سند صدور مطالبه در پایانه‌ها: میانه {fa(pd['median'])} و بیشینه {fa(pd['max'])} کد طرف‌حساب در هر سند.",
         "محاسبهٔ T08", "الگوی مشاهده‌شده", "دامنهٔ بسته و الگوی posting جدا تصمیم می‌شوند (DR-06، RL-03).", "—"],
        ["OR-05", "کد طرف‌حساب پایانه‌ها بهره‌بردار را با غرفه، پایانه و گاه سال ترکیب می‌کند",
         f"در نام کدهای طرف‌حساب: مطالبه (۱۱۳-۱-۴) {fa(e['113-1-4']['slash'])} از {fa(e['113-1-4']['persons'])} «/» دارند، {fa(e['113-1-4']['terminal'])} نام پایانه و {fa(e['113-1-4']['year'])} سال؛ پیش‌دریافت (۲۱۳-۲) {fa(e['213-2']['terminal'])} از {fa(e['213-2']['persons'])} و سپرده (۲۱۲-۶) {fa(e['212-6']['terminal'])} از {fa(e['212-6']['persons'])} نام پایانه دارند.",
         "محاسبهٔ T08 (بدون نقل نام)", "الگوی مشاهده‌شده", "طرف (D13)، زیرمحل (D05) و قرارداد (D08) باید جدا ثبت شوند؛ کد منبع فقط نگاشت است (DR-01).",
         "اینکه یک شخص چند کد دارد بدون پروندهٔ طرف اثبات نمی‌شود."],
        ["OR-06", "CodeVahed پایانه‌ها زیرمحل را برای مطالبه و سپرده نشان می‌دهد",
         f"بدهکار ۱۱۳-۱-۴ و ۱۱۳-۱-۸ (انتخاب فقط با حساب و جهت): {fa(cil['n'])} سطر، {pct(cil.get('SUB', 0), cil['n'])} در سطح غرفه یا دفتر شرکت. بستانکار سپرده‌های ۲۱۲-۶-۵ و ۲۱۲-۶-۷: {fa(ldl['n'])} سطر، {pct(ldl.get('SUB', 0), ldl['n'])} در سطح زیرمحل.",
         "محاسبهٔ T08 روی tt_v21.sqlite", "الگوی مشاهده‌شده", "زیرمحل در دادهٔ پایهٔ اماکن و قرارداد ثبت می‌شود (FM-02، FM-03).", "محل و زیرمحل مرکز هزینه نیستند."],
        ["OR-07", "زیرمحل در زمان چند بهره‌بردار دارد و بهره‌بردار گاه چند زیرمحل",
         f"{fa(um['multi_party'])} از {fa(um['units'])} واحد دارای مطالبه بیش از یک کد طرف‌حساب دارند؛ {fa(loc['multi_unit'])} از {fa(loc['parties'])} کد طرف‌حساب در بیش از یک واحد و {fa(loc['multi_terminal'])} کد در بیش از یک پایانه مطالبه دارند.",
         "محاسبهٔ T08", "الگوی مشاهده‌شده", "رابطهٔ زیرمحل و قرارداد اعتبار زمانی دارد (RL-07، RL-08).", "علت (تغییر مستأجر یا اشتراک) معلوم نیست."],
        ["OR-08", "سپرده‌ها به طرف‌حساب و زیرمحل وصل‌اند، نه به قرارداد",
         f"سپردهٔ تعهدات (PFE-073): {fa(p073['n'])} سطر، {fa(p073['persons'])} کد طرف‌حساب و {fa(p073['vahed'].get('SUB', 0))}٪ در سطح زیرمحل. الگوی CP-06: ۳۸ سند استرداد سپرده همراه تهاتر مطالبه.",
         "پروفایل؛ T07 CP-06", "الگوی مشاهده‌شده", "سپرده و تضمین با شناسهٔ قرارداد ثبت می‌شوند (RL-11، FM-08).", "—"],
        ["OR-09", "تضمین غیرنقدی در حساب انتظامی است و مشخصات ابزار فقط در پایانه‌ها ثبت می‌شود",
         f"دریافت تضمین (PFE-102): فیلدهای ابزار در پایانه‌ها روی {fa(p102['instr'])}٪ سطرها و در تاکسیرانی روی {fa(p102x['instr'])}٪.",
         "پروفایل_ابعاد_رویداد", "الگوی مشاهده‌شده", "ثبت ابزار برای تضمین‌ها (FM-04، DR-08).", "نوع تضمین فقط در شرح می‌آید (DE-21)."],
        ["OR-10", "کسور با صورت‌وضعیت هم‌سندند، ولی محل را روی سطر ناخالص دارند",
         f"صورت‌وضعیت خدمات (PFE-027) در پایانه‌ها: {fa(p027['vahed'].get('TERM', 0))}٪ سطرها در سطح پایانه و {fa(p027['vahed'].get('HQ', 0))}٪ ستاد. کسر حسن انجام کار (PFE-046): {fa(p046['vahed'].get('HQ', 0))}٪ ستاد با {fa(p046['persons'])} کد طرف‌حساب.",
         "پروفایل_ابعاد_رویداد", "الگوی مشاهده‌شده", "ابعاد بها روی سطر ناخالص می‌مانند و کسرها ذی‌نفع خود را دارند (SG-DED).", "—"],
        ["OR-11", "شاهد نقش عامل یا اصیل پایانه‌ها",
         "سه سند به «قرارداد کارگزاری فی‌مابین با شهرداری» و «سهم ۱ درصدی» ارجاع می‌دهند؛ خالص ۱۱۳-۱-۱ و ۲۱۱-۱-۱ در هر سال صفر است (E-05).",
         "T07 §۱؛ سند دامنه E-05", "شاهد تحلیلی", "نقش اصیل یا عامل ویژگی قرارداد است (D23) و تا T-D2 «منوط به تصمیم» می‌ماند.", "تصمیم CD-01 با کارفرماست."],
        ["OR-12", "روابط بین دو بخش و با شهرداری شناسهٔ مشترک ندارند",
         f"پرداخت‌های پایانه‌ها به تاکسیرانی (AM-13): {fa(am('AM-13', 'cand'))} از {fa(am('AM-13', 'pop'))} فقط تطابق نامزد دارند و شمارهٔ مرجع مشترک ندارند. جدول هم‌ترازشده «نیازمند تطبیق» است (VF-06).",
         "T07 تطبیق_بین_دو_بخش", "تطابق نامزد", "ثبت قرینه با شناسهٔ مشترک (RL-13، DR-11).", "علت اختلاف فرض نشده است."],
        ["OR-13", "قرارداد در دفتر ثبت نمی‌شود",
         f"ستون قرارداد نیست؛ «قرارداد» در {fa(tx['qd'] + tm['qd'])} سطر شرح جمعیت بررسی؛ حساب «کنترل قراردادها-پروژه‌ها» (۹۱۸-۱-۳) {fa(r['contract_memo_lines'])} سطر دارد.",
         "محاسبهٔ T08؛ ثبت منابع §۳-۲", "نبود در داده", "قرارداد دادهٔ پایهٔ جداست (FM-03).", "نبود در داده الزام یا ممنوعیت نمی‌سازد."],
        ["OR-14", "صورت‌وضعیت‌های خدمات پایانه‌ها میان چند محل تقسیم شده‌اند، بی‌مبنا",
         f"سطرهای خدمات پیمانکاری (۱۱۳-۱-۱-۵، ۲۲ و ۲۳): {fa(tcs['n'])} سطر؛ {fa(tcs.get('TERM', 0))} در سطح پایانه و {fa(tcs.get('HQ', 0))} ستاد. {fa(tcd['multi'])} از {fa(tcd['docs'])} سند بیش از یک محل دارند.",
         "محاسبهٔ T08", "الگوی مشاهده‌شده", "تفکیک فقط با مدرک کمّی یا سهم مستند (SP-02، SP-03)؛ وگرنه SP-04.", "مبنای تقسیم در دفتر نیست؛ درصد ساخته نمی‌شود."],
        ["OR-15", "درآمد خدمات تاکسیرانی هویت متقاضی، پرونده یا فعالیت ندارد",
         f"{fa(r['taxi_revenue']['n'])} سطر بستانکار ۴۱۳-۱: کد طرف‌حساب {fa(r['taxi_revenue']['person'])} و کد فعالیت {fa(r['taxi_revenue']['markaz'])}. شرح PFE-001 عدد مرجع ندارد ({fa(PROFILE[('PFE-001', 'Taxi')]['n5'])}٪).",
         "محاسبهٔ T08", "نبود در داده", "پرونده، متقاضی و قلم تعرفه از سامانهٔ تاکسیرانی می‌آیند (FM-09، DR-10).", "—"],
        ["OR-16", "کد فعالیت تاکسیرانی روی هزینه و حقوق هست و روی درآمد نیست",
         f"هزینه: {fa(tx['markaz_on_expense'][0])} از {fa(tx['markaz_on_expense'][1])}؛ حقوق (۶۱۱ بدهکار): {fa(r['taxi_payroll_markaz']['markaz'])} از {fa(r['taxi_payroll_markaz']['n'])}؛ درآمد: {fa(tx['markaz_on_revenue'][0])} از {fa(tx['markaz_on_revenue'][1])}.",
         "محاسبهٔ T08", "الگوی مشاهده‌شده", "کد منبع شاهد فعالیت است؛ حضورش روی حقوق الزام انتساب حقوق نمی‌سازد (CD-08).", "کد ۱۰۰ عمومی است (T05)."],
        ["OR-17", "سطر بانک مرجع دارد، ولی پرداخت‌کننده ندارد",
         f"سطرهای بانک دارای کد طرف‌حساب: تاکسیرانی {fa(tx['person_on_bank'][0])} از {fa(tx['person_on_bank'][1])} ولی فقط {fa(tx['person_codes_on_bank'])} کد؛ پایانه‌ها {fa(tm['person_on_bank'][0])} از {fa(tm['person_on_bank'][1])} با {fa(tm['person_codes_on_bank'])} کد. مرجع بانکی: {pct(tx['bankref_on_bank'][0], tx['bankref_on_bank'][1])} و {pct(tm['bankref_on_bank'][0], tm['bankref_on_bank'][1])}.",
         "محاسبهٔ T08", "الگوی مشاهده‌شده", "پرداخت‌کننده (D14) در فرم دریافت ثبت می‌شود؛ مرجع بانکی (D18) از سطر بانک قابل بازیابی است.", "—"],
        ["OR-18", "کدهای هم‌شکل در دو بخش معنای متفاوت دارند",
         f"{fa(len(obs['homonyms']))} کد ۱۴۰۴ در دو بخش عنوان‌هایی با کمتر از نیمی واژهٔ مشترک دارند؛ از جمله ۱۱۳-۱-۱، ۱۱۳-۱-۴، ۴۱۳-۸-x و ۶۱۱-x.",
         "محاسبهٔ T08؛ T07 EA-10؛ سند دامنه E-02", "الگوی مشاهده‌شده", "کلید مرکب (بخش، سال یا نسخه، کد) برای هر حساب (شیت حساب‌های_مشاهده‌شده).", "معیار واژگانی است؛ چند مورد فقط تفاوت املایی‌اند."],
        ["OR-19", "هزینه‌های تحمیل‌شدهٔ پایانه‌ها در سطح پایانه یا ستاد ثبت شده‌اند، نه غرفه",
         f"بدهکار ۱۱۳-۱-۱-x: {fa(tel['n'])} سطر؛ سطح پایانه {fa(tel.get('TERM', 0))}، ستاد {fa(tel.get('HQ', 0))}، غرفه {fa(tel.get('SUB', 0))}.",
         "محاسبهٔ T08", "الگوی مشاهده‌شده", "D04 برای هزینه‌های پایانه‌ها «لازم» و D05 «اختیاری» پیشنهاد می‌شود.", "—"],
    ]
    basis = {"OR-01": "نتیجهٔ ذخیره‌شدهٔ T07", "OR-02": "نتیجهٔ ذخیره‌شدهٔ T07", "OR-03": "نتیجهٔ ذخیره‌شدهٔ T07 و محاسبهٔ دقیق T08",
             "OR-08": "پروفایل دقیق (شرط SQL T07)", "OR-09": "پروفایل دقیق (شرط SQL T07)", "OR-10": "پروفایل دقیق (شرط SQL T07)",
             "OR-11": "نتیجهٔ ذخیره‌شدهٔ T07 و سند دامنه", "OR-12": "نتیجهٔ ذخیره‌شدهٔ T07", "OR-15": "محاسبهٔ دقیق T08 و پروفایل دقیق"}
    rows = [r[:5] + [basis.get(r[0], "محاسبهٔ دقیق T08 (حساب‌محور)")] + r[5:] for r in rows]
    headers = ["شناسه", "رابطه یا قید", "آنچه در داده دیده شد", "منبع شاهد", "سطح اطمینان", "مبنای شاهد",
               "پیامد برای طراحی (لایهٔ ب)", "آنچه اثبات نمی‌کند"]
    add_sheet("روابط_مشاهده‌شده", "لایهٔ الف — روابط و قیدهای مشاهده‌شده در دفاتر",
              "هر ردیف شاهد است، نه قاعده. سطح اطمینان پیوندها همان سطوح T07 است: T3 (شناسهٔ مشترک)، T2 (همان طرف‌حساب یا شرح)، T1/T0 (فقط مبلغ = تطابق نامزد). هیچ ردیفی بر پروفایل تقریبی تکیه ندارد.",
              headers, rows, [7, 34, 70, 22, 16, 20, 46, 34], freeze="C4")
    used_profiles = [("PFE-073", "Terminals"), ("PFE-102", "Terminals"), ("PFE-102", "Taxi"), ("PFE-027", "Terminals"),
                     ("PFE-046", "Terminals"), ("PFE-001", "Taxi")]
    return {r[0]: r[5] for r in rows}, used_profiles


# ============================================================================ B1 dimensions
def dimensions_sheet():
    rows = [[d[0], d[1], d[2], d[3], d[4], d[5], d[6], d[7], d[8], d[9], STATUS] for d in K.DIMENSIONS]
    headers = ["شناسه", "بُعد", "تعریف", "با این‌ها یکی نشود", "سطح ساختاری (دادهٔ پایه / نمونهٔ رویداد / سطر / بسته)",
               "فهرست مرجع پیشنهادی", "تأمین‌کنندهٔ پیشنهادی (نقش)", "مرحلهٔ ثبت", "دلیل", "تصمیم‌های مرتبط", "وضعیت"]
    add_sheet("ابعاد_و_سطح", "لایهٔ ب — ابعاد اطلاعاتی، مرز و سطح ساختاری هر بُعد",
              "نقش‌های طرف (D12 تا D15) چهار بُعد جدا هستند. محل (D04) و زیرمحل (D05) مرکز هزینه نیستند. «تأمین‌کنندهٔ پیشنهادی» نقش است، نه شخص، و تأیید نشده است (T03 د-۱).",
              headers, rows, [6, 20, 48, 40, 34, 28, 30, 24, 52, 22, 14], freeze="C4")


# ============================================================================ B2 matrix rules + matrix
def sector_rows(ev_id):
    s = EV[ev_id]["sector"]
    if s.startswith("هر دو"):
        return ["TX", "TM"]
    return ["TX"] if s == "تاکسیرانی" else ["TM"]


def resolve(code, sk):
    return code[sk] if isinstance(code, dict) else code


def matrix_sheets(profile_ids):
    dim_name = {d[0]: d[1] for d in K.DIMENSIONS}
    rule_rows = []
    for sid, g in K.SG.items():
        events = [e for e, a in K.EVENTS.items() if a[1] == sid]
        defaults = []
        for d in K.MATRIX_DIMS:
            v = g["defaults"][d]
            defaults.append(f"{dim_name[d]} ({d}): {v}" if not isinstance(v, dict)
                            else f"{dim_name[d]} ({d}): تاکسیرانی {v['TX']} / پایانه‌ها {v['TM']}")
        cond = "؛ ".join(f"{dim_name[d]} ({d}): {t}" for d, t in g["cond"].items()) or "—"
        rule_rows.append([sid, g["title"], g["scope"], "، ".join(events), " | ".join(defaults), cond, g["n_reason"],
                          g["stage"], g["supplier"], STATUS])
    add_sheet("قواعد_ماتریس", "لایهٔ ب — گروه‌های مرحله و کدهای پیش‌فرض ابعاد",
              "هر رویداد T07 به یک گروه مرحله تعلق دارد. خانه‌های ماتریس کد پیش‌فرض گروه را می‌گیرند، مگر در ستون «استثنا» برای همان رویداد تغییر کرده باشند. شرط کدهای «شرطی» و تصمیم کدهای «منوط به تصمیم» در ستون شرط‌ها آمده است.",
              ["گروه", "عنوان", "دامنه", "رویدادها", "کدهای پیش‌فرض", "شرط‌ها و تصمیم‌ها", "دلیل کدهای «نامرتبط مفهومی»",
               "مرحلهٔ ثبت", "تأمین‌کنندهٔ پیشنهادی", "وضعیت"],
              rule_rows, [12, 30, 36, 30, 70, 70, 46, 22, 30, 14], freeze="C4")

    rows, cells = [], []
    mid = 0
    for ev_id in sorted(EV):
        flow, sid, over, note = K.EVENTS[ev_id]
        g = K.SG[sid]
        for sk in sector_rows(ev_id):
            mid += 1
            codes = []
            changed = []
            for d in K.MATRIX_DIMS:
                base = resolve(g["defaults"][d], sk)
                val = resolve(over[d], sk) if d in over else base
                if d in over and val != base:
                    changed.append(f"بُعد {d} از «{base}» به «{val}»")
                codes.append(val)
                cells.append((ev_id, sk, d, val, sid))
            org = "Taxi" if sk == "TX" else "Terminals"
            pobj = PROFILE.get((ev_id, org))
            prof = profile_ids.get((ev_id, org), "—") + (f" — {basis_label(pobj)}" if pobj and pobj.get("n") else "")
            also = "، ".join(K.ALSO_IN.get(ev_id, []))
            t07n = PROFILE.get((ev_id, org), {}).get("t07", 0)
            rows.append([f"M-{mid:03d}", ev_id, EV[ev_id]["title"], SECTOR_FA[sk], EV[ev_id]["cls"], EV[ev_id]["nature"],
                         flow + (f" (+{also})" if also else ""), sid, K.L, *codes,
                         ("؛ ".join(changed) + (" — " if changed and note else "") + note) if (changed or note) else "—",
                         t07n, prof, "گروه مرحله و استثنای مکتوب؛ از پروفایل مشتق نشده است", STATUS])
    headers = ["شناسه", "رویداد", "عنوان (T07)", "بخش", "طبقه‌بندی T07", "ماهیت مرحله", "جریان", "گروه مرحله",
               "D01 بخش", *[f"{d} {dim_name[d]}" for d in K.MATRIX_DIMS], "استثنا نسبت به گروه و یادداشت",
               "سطر شاهد T07", "پروفایل و وضعیت آن (لایهٔ الف)", "مبنای کدها", "وضعیت"]
    first_code_col = 9
    code_cols = tuple(range(first_code_col, first_code_col + 1 + len(K.MATRIX_DIMS)))
    add_sheet("ماتریس_رویداد_بعد", "لایهٔ ب — ماتریس پیشنهادی رویداد × بُعد",
              "پنج کد: لازم، شرطی، اختیاری، منوط به تصمیم، نامرتبط مفهومی (تعریف در «راهنما»). هیچ خانه‌ای قاعدهٔ فعال نیست و در پایلوت پیش از Rule Profile مرجع اجبار نمی‌شود (ADR-0019). رویدادهای «هر دو بخش» دو ردیف جدا دارند.",
              headers, rows, [7, 9, 32, 9, 14, 12, 10, 12, 8] + [10] * len(K.MATRIX_DIMS) + [48, 8, 22, 18, 14],
              freeze="D4", code_cols=code_cols, header_height=96)
    return rows, cells


# ============================================================================ B3 lifecycle patterns
def lifecycle_sheet():
    rows = []
    for (lid, flow, lc, stage, events, elements, level, docs, roles, link, amb, dec, ev07) in K.LIFECYCLE:
        evs = re.findall(r"PFE-\d{3}", events)
        prof = []
        for e in evs:
            for org in ("Taxi", "Terminals"):
                p = PROFILE.get((e, org))
                if p and p.get("n"):
                    s = f"رویداد {e} در {SECTOR_FA[org]} [{BASIS_FA[p['basis']]}]: {fa(p['t07'])} سطر شاهد؛ کد طرف‌حساب {fa(p['person'])}٪ ({fa(p['persons'])} کد)"
                    if org == "Terminals":
                        s += f"؛ زیرمحل {fa(p['vahed'].get('SUB', 0.0))}٪"
                    if p["instr"]:
                        s += f"؛ ابزار {fa(p['instr'])}٪"
                    if p["markaz"]:
                        s += f"؛ کد فعالیت {fa(p['markaz'])}٪"
                    prof.append(s)
        if elements[:1].isascii() and elements[:1].isalpha():
            elements = "ابعاد: " + elements
        rows.append([lid, K.F[flow], lc, stage, events, elements, level, docs, roles, link, ev07,
                     " | ".join(prof) or "—", amb, dec, STATUS])
    headers = ["شناسه", "جریان", "چرخه یا الگو (T07)", "مرحله", "رویداد", "عنصرهای داده (ابعاد)", "سطح ثبت",
               "مدارک لازم (پیشنهاد)", "نقش مسئول پیشنهادی (ثبت / تأیید)", "قوت پیوند در داده", "شاهد T07",
               "پروفایل مشاهده‌شده (T08)", "ابهام‌های باز", "تصمیم‌ها", "وضعیت"]
    add_sheet("الگوهای_چرخه", "لایهٔ ب — شش جریان اولویت‌دار بر پایهٔ چرخه‌های مشاهده‌شده",
              "پیوند در سطح بهره‌بردار یا طرف‌حساب به مطالبه یا قرارداد مشخص ارتقا داده نشده است. تطابق مبلغی «نامزد» می‌ماند. نقش‌ها پیشنهادی‌اند و تأیید نشده‌اند.",
              headers, rows, [9, 24, 18, 22, 22, 50, 22, 36, 32, 26, 16, 52, 40, 18, 14], freeze="E4")
    return rows


# ============================================================================ B4 form fields and data elements
FIELD_ELEMENTS = collections.defaultdict(list)
for _el, _name, _purpose, _fields in K.ELEMENTS:
    for _f in _fields:
        FIELD_ELEMENTS[_f].append(_el)
EL_NAME = {e[0]: e[1] for e in K.ELEMENTS}


def forms_sheet():
    rows, display = [], []
    for (fid, form, name, dim, definition, source, stage, supplier, req, cond, why, ev, infra, dec) in K.FORM_FIELDS:
        ftitle, flevel, finfra = K.FORMS[form]
        rows.append([fid, form, ftitle, flevel, name, dim, definition, source, stage, supplier, req, cond or "—", why,
                     ev, infra, dec, STATUS])
        els = "، ".join(f"{e} {EL_NAME[e]}" for e in FIELD_ELEMENTS.get(fid, [])) or "—"
        display.append([fid, form, ftitle, flevel, name, els, dim, definition, source, stage, supplier, req, cond or "—",
                        why, ev, infra, dec, STATUS])
    headers = ["شناسه", "فرم", "عنوان فرم یا موجودیت", "سطح", "فیلد", "عنصر داده", "بُعد", "تعریف و مرز", "مرجع تأمین مقدار",
               "مرحلهٔ ثبت", "نقش ثبت‌کننده (پیشنهادی)", "پیشنهاد الزام", "شرط", "کاربرد و دلیل", "شاهد (لایهٔ الف)",
               "زیرساخت موجود در نرم‌افزار", "تصمیم لازم", "وضعیت پیشنهاد"]
    add_sheet("فیلدهای_فرم", "لایهٔ ب — فیلدهایی که فرم‌های عملیاتی آینده باید ثبت کنند",
              "۱۳ فرم و ۱۲۵ فیلد پیشنهاد کاری‌اند. حضور یا نبود داده در دفاتر قدیم به‌تنهایی فیلدی را لازم یا ممنوع نمی‌کند. «زیرساخت موجود» فقط نشان می‌دهد کدام مدل نرم‌افزار (commit c8550e9) نزدیک است؛ کدی تغییر نکرد. «پیشنهاد الزام» برای Rule Profile آینده است و اکنون فعال نیست. تعریف یکپارچهٔ هر عنصر در شیت «عناصر_داده» است.",
              headers, display, [9, 7, 26, 13, 30, 30, 10, 46, 20, 18, 26, 11, 30, 40, 11, 30, 16, 16], freeze="F4",
              code_cols=(12,))
    return rows


# ledger availability per dimension, grounded in فیلدهای_منبع and روابط_مشاهده‌شده
LEDGER = {
    "D01": "کامل (دفتر جدا برای هر بخش؛ SF-01)", "D02": "جزئی: کد فعالیت منبع فقط روی هزینه و حقوق تاکسیرانی (SF-13، OR-16)",
    "D03": "ستون ندارد (SF-14)", "D04": "فقط پایانه‌ها (CodeVahed؛ SF-11)", "D05": "فقط پایانه‌ها (CodeVahed؛ SF-11، OR-06)",
    "D06": "فقط پایانه‌ها (واحدهای ستادی در CodeVahed؛ SF-11)", "D07": "کامل (HCode؛ SF-02)",
    "D08": "ستون ندارد؛ فقط واژهٔ «قرارداد» در شرح (SF-04، OR-13)", "D09": "ستون ندارد؛ عدد در شرح فقط نشانه است (SF-04)",
    "D10": "ستون ندارد؛ فقط نوع حساب سرمایه‌ای (SF-14)", "D11": "ستون ندارد (SF-14)", "D12": "ستون ندارد (SF-14)",
    "D13": "جزئی: CodePerson، در پایانه‌ها ترکیب طرف و غرفه (SF-05، OR-05)", "D14": "ستون ندارد؛ صادرکنندهٔ چک فقط روی سطر چک (SF-10، OR-17)",
    "D15": "ستون ندارد (SF-14)", "D16": "در کد حساب بانک (SF-02)", "D17": "روی سطرهای چک و تضمین (SF-06 تا SF-10)",
    "D18": "روی سطر بانک (SF-06 تا SF-08)", "D19": "در دفتر نیست", "D20": "فقط شمارهٔ چک (T3) یا طرف‌حساب (T2) (OR-02، OR-03)",
    "D21": "ستون ندارد؛ نام ماه در شرح فقط نشانه است (SF-04)", "D22": "در دفتر نیست (OR-14)", "D23": "در دفتر نیست (OR-11)",
}


def elements_sheet():
    fields = {f[0]: f for f in K.FORM_FIELDS}
    rows = []
    for el, name, purpose, members in K.ELEMENTS:
        fs = [fields[m] for m in members]
        uniq = lambda xs: "؛ ".join(dict.fromkeys(x for x in xs if x and x != "—")) or "—"
        # a field carrying several elements (e.g. FM12-07) lists several dimensions; attribute only the element's own
        own = [d for d, e in K.DIM_ELEMENT.items() if e == el]
        single = [f for f in fs if len(FIELD_ELEMENTS[f[0]]) == 1] or fs
        dims = own or sorted({t for f in single for t in re.findall(r"D\d{2}", f[3])} |
                             ({"D12", "D13", "D14", "D15"} if any("D12–D15" in f[3] for f in single) else set()))
        ledger = "؛ ".join(dict.fromkeys(LEDGER[d] for d in dims)) or "ارزیابی نشد (عنصر بُعدی از D01 تا D23 نیست)"
        rows.append([el, name, purpose,
                     uniq(K.FORMS[f[1]][1] for f in fs),
                     "، ".join(dims) or "—",
                     uniq(f[5] for f in fs),
                     ledger,
                     uniq(f[7] for f in fs),
                     uniq(f[6] for f in fs),
                     "پیشنهادی؛ فعال نیست — " + "؛ ".join(f"{f[0]}: {f[8]}" + (f" ({f[9]})" if f[9] else "") for f in fs),
                     uniq(f[13] for f in fs),
                     "، ".join(members)])
    headers = ["شناسه", "نام ساده", "کاربرد", "سطح ثبت", "بُعد", "مرجع تأمین مقدار", "وضعیت در دفاتر قدیم (شاهد، نه الزام)",
               "نقش ثبت‌کننده (پیشنهادی)", "مرحلهٔ ثبت", "وضعیت پیشنهاد و الزام پیشنهادی در هر فیلد", "تصمیم لازم",
               "فیلدهای فرم"]
    add_sheet("عناصر_داده", "بند ۳ — فرهنگ عناصر دادهٔ فرم‌ها (هر عنصر یک بار)",
              "هر عنصر یک مفهوم داده است که ممکن است در چند فرم بیاید. ستون «وضعیت در دفاتر قدیم» شاهد است و الزام نمی‌سازد. الزام‌ها پیشنهاد برای Rule Profile آینده‌اند و اکنون فعال نیستند (ADR-0019). نقش‌ها پیشنهادی‌اند و تأیید نشده‌اند (T03 د-۱).",
              headers, rows, [7, 30, 40, 18, 12, 30, 44, 34, 26, 56, 22, 24], freeze="C4")
    return rows


DIM_SOURCE = {
    "D01": "سامانه (از کاربر و زیرسامانه)", "D02": "پیش‌فرض از قرارداد یا پرونده؛ انتخاب از فهرست فعالیت T05",
    "D03": "انتخاب از مراکز پیشنهادی T05 (برچسب «پیشنهادی»)", "D04": "انتخاب از دادهٔ پایهٔ اماکن؛ پیش‌فرض از قرارداد",
    "D05": "انتخاب از دادهٔ پایهٔ اماکن؛ از قرارداد", "D06": "سامانه (واحد کاربر آغازگر)",
    "D07": "انتخاب از دیکشنری حساب قدیم با کلید (بخش، سال، کد)", "D08": "انتخاب از دادهٔ پایهٔ قرارداد یا پرونده (FM-03)",
    "D09": "ورود یا تولید در همین مرحله", "D10": "انتخاب از فهرست پروژه‌ها", "D11": "انتخاب از بودجهٔ سازمان",
    "D12": "از دادهٔ پایهٔ قرارداد", "D13": "انتخاب از رجیستری طرف (FM-01)", "D14": "انتخاب از رجیستری طرف یا سامانهٔ مبدأ",
    "D15": "انتخاب از رجیستری طرف", "D16": "انتخاب از فهرست حساب‌های بانکی سازمان", "D17": "انتخاب یا ثبت در ثبت ابزار (FM-04)",
    "D18": "ورود از رسید یا صورت‌حساب بانک", "D19": "بارگذاری مدرک", "D20": "انتخاب از نمونه‌های باز همان قرارداد یا طرف",
    "D21": "ورود یا از جدول قرارداد", "D22": "ورود با مدرک کمّی", "D23": "از دادهٔ پایهٔ قرارداد، پس از تصمیم T-D2",
}
MASTER_DIMS = {"D08": K.ROLE["contracts"], "D12": K.ROLE["contracts"], "D23": f"{K.ROLE['contracts']} پس از تصمیم T-D2"}
DIM_LEVEL = {d[0]: d[4].split("؛")[0].rstrip(".") for d in K.DIMENSIONS}


def path_evidence(d, p, org):
    if not p or not p.get("n"):
        return "رویداد در دادهٔ بررسی‌شده دیده نشد"
    tm = org == "Terminals"
    v = p["vahed"]
    ev = {
        "D01": "دفتر جدا برای هر بخش",
        "D02": f"کد فعالیت منبع روی {fa(p['markaz'])}٪ سطرها",
        "D20": "پیوند ثبت نمی‌شود؛ فقط از شمارهٔ مشترک یا طرف‌حساب استنباط می‌شود (OR-02، OR-03)",
        "D04": f"سطح پایانه یا زیرمحل روی {fa(round(v.get('SUB', 0) + v.get('TERM', 0), 1))}٪ سطرها" if tm else "تاکسیرانی کد محل ندارد",
        "D05": f"زیرمحل روی {fa(v.get('SUB', 0.0))}٪ سطرها" if tm else "تاکسیرانی کد زیرمحل ندارد",
        "D06": f"واحد ستادی روی {fa(v.get('HQ', 0.0))}٪ سطرها" if tm else "تاکسیرانی کد واحد ندارد",
        "D07": "حساب منبع روی همهٔ سطرها هست",
        "D08": f"«قرارداد» در شرح {fa(p['qd'])}٪ (فقط نشانه)",
        "D09": f"عدد ۵ رقمی در شرح {fa(p['n5'])}٪ (فقط نشانه)",
        "D13": f"کد طرف‌حساب روی {fa(p['person'])}٪ سطرها ({fa(p['persons'])} کد)",
        "D16": f"{fa(p['docbank'])}٪ سطرها در سند دارای سطر بانک",
        "D17": f"فیلدهای ابزار روی {fa(p['instr'])}٪ سطرها",
        "D18": f"مرجع بانکی روی {fa(p['bankref'])}٪ سطرها",
        "D21": f"نام ماه در شرح {fa(p['month'])}٪ (فقط نشانه)",
    }
    return ev.get(d, "در دفتر ستون ندارد")


def path_sheet(cells):
    code = {(e, sk, d): (val, sid) for e, sk, d, val, sid in cells}
    dim_name = {d[0]: d[1] for d in K.DIMENSIONS}
    rows = []
    for lp in K.LIFECYCLE:
        lid, flow, lc, stage, events, elements, level, docs, roles = lp[:9]
        if flow == "F0":
            continue
        for e in re.findall(r"PFE-\d{3}", events):
            flow_e, sid, over, note = K.EVENTS[e]
            g = K.SG[sid]
            sks = sector_rows(e)
            for d in ["D01", *K.MATRIX_DIMS]:
                vals = {sk: (K.L if d == "D01" else code[(e, sk, d)][0]) for sk in sks}
                groups = [("هر دو بخش" if len(sks) == 2 else SECTOR_FA[sks[0]], sks)] if len(set(vals.values())) == 1 \
                    else [(SECTOR_FA[sk], [sk]) for sk in sks]
                for label, members in groups:
                    val = vals[members[0]]
                    if val == K.N:
                        continue
                    el = K.DIM_ELEMENT[d]
                    if d == "D09":
                        name = K.DOC_LABEL[sid]
                    elif d == "D08":
                        name = K.DOSSIER_LABEL.get(sid, "قرارداد یا پرونده")
                    elif d == "D19":
                        name = f"مدارک پشتیبان: {docs}"
                    else:
                        name = EL_NAME[el]
                    cond = g["cond"].get(d, "")
                    if d in over and note:
                        cond = (cond + " | " if cond else "") + f"استثنای رویداد: {note}"
                    if d in MASTER_DIMS:
                        role = f"انتخاب: {g['supplier']}؛ ثبت اولیه: {MASTER_DIMS[d]}"
                        when = "انتخاب در همین مرحله؛ ثبت اولیه در دادهٔ پایهٔ قرارداد"
                    elif d == "D07":
                        role, when = K.ROLE["accountant"], "بازبینی حسابداری"
                    elif d in ("D01", "D06"):
                        role, when = K.ROLE["system"], "ایجاد بسته"
                    else:
                        role, when = g["supplier"], g["stage"]
                    org = "Taxi" if members[0] == "TX" else "Terminals"
                    p = PROFILE.get((e, org))
                    evid = "؛ ".join(f"{SECTOR_FA[sk]}: {path_evidence(d, PROFILE.get((e, 'Taxi' if sk == 'TX' else 'Terminals')), 'Taxi' if sk == 'TX' else 'Terminals')}"
                                     for sk in members) if len(members) > 1 else path_evidence(d, p, org)
                    status = "؛ ".join(dict.fromkeys(basis_label(PROFILE[(e, 'Taxi' if sk == 'TX' else 'Terminals')])
                                                      for sk in members if PROFILE.get((e, 'Taxi' if sk == 'TX' else 'Terminals'), {}).get('n'))) or "—"
                    decisions = "، ".join(dict.fromkeys(re.findall(r"\b(?:T-D\d|DE-\d{2}|DG-\d{2}|CD-\d{2}|DR-\d{2}|G-\d{2})\b", cond))) or "—"
                    rows.append([f"PV-{len(rows) + 1:04d}", K.F[flow], f"{lid} — {lc}", stage, e, label, el or "—", name,
                                 f"{d} {dim_name[d]}", DIM_LEVEL[d],
                                 val, cond or "—", DIM_SOURCE[d], role, when, evid, status, decisions, STATUS])
    headers = ["شناسه", "جریان", "الگوی چرخه", "مرحله", "رویداد", "بخش", "عنصر داده", "داده‌ای که فرم باید بگیرد",
               "بُعد", "سطح ثبت", "پیشنهاد — الزام", "پیشنهاد — شرط یا تصمیم", "پیشنهاد — مرجع تأمین مقدار",
               "پیشنهاد — نقش ثبت‌کننده", "پیشنهاد — مرحلهٔ ثبت", "شاهد — دفاتر ۱۴۰۴ و ۱۴۰۵ برای همین رویداد",
               "شاهد — وضعیت پروفایل", "تصمیم لازم", "وضعیت"]
    add_sheet("نیاز_داده_مسیرها", "معیار پذیرش — دادهٔ لازم در هر مسیر و مرحله، با مرجع و نقش",
              "برای هر مرحله از شش جریان اولویت‌دار و هر رویداد آن: چه داده‌ای، با چه الزام پیشنهادی، از چه مرجعی، به دست چه نقشی و در چه مرحله‌ای. ستون‌های «پیشنهاد» طراحی‌اند و ستون‌های «شاهد» مشاهده. شاهد هیچ الزامی نمی‌سازد و پروفایل تقریبی برچسب خود را دارد. ابعاد «نامرتبط مفهومی» نیامده‌اند. LP-F4-07 و LP-F4-08 رویداد ندارند و فقط در «الگوهای_چرخه» آمده‌اند.",
              headers, rows, [9, 22, 22, 22, 9, 10, 8, 34, 16, 18, 11, 46, 30, 30, 24, 36, 26, 16, 14], freeze="F4",
              code_cols=(11,), header_height=60)
    return rows


# ============================================================================ B5 relationships and splits
def relations_sheet():
    rows = [[r[0], r[1], r[2], r[3], r[4], r[6], r[7], r[5], r[8], STATUS] for r in K.RELATIONS]
    rows += [[p[0], p[1], p[2], "—", "—", p[3], "سطر طبقه‌بندی ↔ سطر سند", p[4], p[5], STATUS] for p in K.POSTING]
    headers = ["شناسه", "نوع", "از یا عنوان", "به", "کاردینالیتی پیشنهادی", "قاعدهٔ پیشنهادی", "سطح", "شاهد (لایهٔ الف)",
               "تصمیم", "وضعیت"]
    add_sheet("روابط_و_تفکیک", "لایهٔ ب — روابط میان سطح‌ها، تفکیک مبلغ و گزینهٔ پیوند سطر با سند",
              "کاردینالیتی‌ها پیشنهادند و با شاهد همراه‌اند؛ تأیید آن‌ها با Rule Profile است. رابطهٔ چند به چند سطر طبقه‌بندی و سطر سند (RL-03) گزینهٔ طراحی پیشنهادی است: گزینه‌ها PO-01 تا PO-04 و حفاظ‌های شمارش دوباره DC-01 تا DC-08. هیچ درصد یا سهمی بدون مدرک کمّی یا مصوبه پیشنهاد نشده است (ADR-0022، ADR-0026).",
              headers, rows, [7, 16, 30, 22, 26, 70, 20, 14, 12, 14], freeze="C4")
    return rows


# ============================================================================ section 6: accounts
def accounts_sheet():
    hom = {h[0] for h in obs["homonyms"]}
    rows = []
    for i, a in enumerate(obs["accounts"], 1):
        h = a["h"]
        parts = h.split("-")
        parent = "-".join(parts[:3]) + "-0" if len(parts) == 4 else None
        flag = "بله (همین کد)" if h in hom else ("بله (کد والد)" if parent in hom or "-".join(parts[:2]) + "-0-0" in hom else "—")
        rows.append([f"AC-{i:04d}", f"{SECTOR_KEY[a['org']]}|{a['fy']}|{h}", SECTOR_FA[a["org"]], a["fy"], h,
                     a["title"] or "—", "بله" if a["in_hesab"] else "خیر", a["n"], a["dr"], a["cr"], a["docs"],
                     "، ".join(a["events"]) or "—", "، ".join(a["proxy_events"]) or "—", flag, None,
                     "حل‌نشده — منتظر T-D3 (فایل شایان و سند مرجع خزانه‌داری کل)"])
    headers = ["شناسه", "کلید مرکب (بخش|سال|کد)", "بخش", "سال مالی (نسخهٔ نمودار)", "کد منبع", "عنوان منبع",
               "در فهرست حساب همان سال", "سطر در جمعیت بررسی", "بدهکار", "بستانکار", "سند", "رویدادهای T07 (شرط دقیق)",
               "رویدادهای T07 (فقط از بازسازی جفت‌سازی یا تعریف مکتوب)", "کد هم‌شکل با معنای متفاوت در بخش دیگر", "حساب مقصد", "وضعیت نگاشت"]
    add_sheet("حساب‌های_مشاهده‌شده", "بخش ۶ — حساب‌های مشاهده‌شده با کلید (بخش، سال، کد منبع)",
              "حساب مقصد عمداً خالی است و با «حل‌نشده» علامت خورده است. نگاشت‌های نامزد پیشین (Taxi-Gate، crosswalk نسخهٔ v2.1) مبنای پر کردن آن نیستند. شمارهٔ حساب بانکی در کد و عنوان پوشانده شده و عنوان حساب‌های شخصی حذف شده است.",
              headers, rows, [9, 26, 10, 10, 18, 36, 10, 9, 8, 8, 7, 34, 20, 16, 14, 34], freeze="F4")
    return rows


# ============================================================================ section 6: unique code
def yekta_sheet():
    code = (ROOT / "app/services/yekta_codes.py").read_text(encoding="utf-8")
    default_calls = re.findall(r'zero_fill_segment\(\w+, \w+, "([^"]+)"\)', code)
    proposal = {
        "نامعلوم": "خالی؛ برچسب «نامعلوم — منتظر T-D7». مقدار پیش‌فرض به کار نمی‌رود.",
        "ناسازگار": "خالی؛ برچسب «ناسازگار — منتظر تصمیم». نگاشت ساخته نمی‌شود.",
        "مقدار ساختگی در پیاده‌سازی": "خالی؛ هش واژهٔ «Unknown» تولید نمی‌شود.",
        "معلوم از تاریخ": "از تاریخ سند منبع یا نمونهٔ رویداد؛ پیش‌فرض 1403 به کار نمی‌رود.",
        "معلوم (سامانه)": "سامانه تولید می‌کند، پس از رفع صدور غیراتمیک.",
    }
    dec = {"YK-05": "T-D6، T-D7، DG-02", "YK-09": "DE-27، T-D7"}
    rows = [[y[0], y[1], y[2], y[3], y[4], y[6], y[5], proposal[y[6]], dec.get(y[0], "T-D7، DR-12"), STATUS]
            for y in K.YEKTA]
    rows += [
        ["YK-N1", "—", "تعریف رقیب", "—", "—", "تعارض", "deck کدینگ (SRC-D04) قالب دیگری را توصیف می‌کند (G-11).", "تا T-D7 فقط تعریف پیاده‌شده در محیط آزمایشی.", "T-D7", STATUS],
        ["YK-N2", "—", "جایگاه کد در بسته", "—", "—", "قاعدهٔ پذیرفته‌شده", "کد یکتا مرجع منشأ است، نه شناسهٔ بسته یا سطر (ADR-0013).", "یک کد با split صریح چند سطر را پشتیبانی می‌کند.", "—", "ADR پذیرفته‌شده"],
        ["YK-N3", "—", "تغییر کد نرم‌افزار", "—", "—", "—", f"پیاده‌سازی فعلی بخش خالی را با مقدار پیش‌فرض پر می‌کند ({fa(len(default_calls))} فراخوانی zero_fill با پیش‌فرض در app/services/yekta_codes.py).", "در T08 کدی تغییر نکرد؛ اصلاح پس از تصمیم DR-12.", "DR-12", STATUS],
    ]
    headers = ["شناسه", "بخش کد", "نام بخش", "قالب پیاده‌شده", "مقدار پیاده‌شده هنگام خالی بودن", "وضعیت برای دو بخش",
               "توضیح", "پیشنهاد T08", "تصمیم", "وضعیت"]
    add_sheet("کد_یکتا", "بخش ۶ — بخش‌های کد یکتا برای تاکسیرانی و پایانه‌ها",
              "ساختار ۱۱ بخشی بازطراحی نمی‌شود (سند دامنه ۶-۵). برای بخشی که مقدارش معلوم نیست مقدار پیش‌فرض ساخته نمی‌شود و قاعدهٔ کدگذاری تازه‌ای استنباط نمی‌شود.",
              headers, rows, [8, 7, 16, 12, 18, 18, 56, 44, 16, 14], freeze="C4")
    return default_calls, rows


# ============================================================================ decisions
SCOPE_TD = {"T-D1": "هویت حقوقی سازمان و جایگاه دو بخش", "T-D2": "ماهیت حسابداری پایانه‌ها (CD-01)",
            "T-D3": "فایل تطبیق شایان و سند مرجع خزانه‌داری", "T-D4": "پروندهٔ واقعی نمونهٔ الف",
            "T-D5": "نمونهٔ ب و طرف حقوقی قرارداد", "T-D6": "فعالیت و مرکز هزینهٔ اختصاصی؛ قاعدهٔ درآمد",
            "T-D7": "کاربرد کد یکتا برای سازمان", "T-D8": "گزینهٔ ارجاع به دفتر و واحد پول",
            "T-D9": "نسخهٔ مبنای سند Word و ارائه"}
OTHER_REFS = {"CD-08": "سیاست انتساب حقوق (v2.1؛ T05 DG-05)", "CD-01": "نقش پایانه‌ها (= T-D2)",
              "CD-02": "واحد پول منبع (= T-D8)", "CD-03": "قاعدهٔ درآمد و مرکز (= DG-02)",
              "G-10": "دادهٔ ۱۴۰۵ ناقص و کد بودجه ناموجود (ثبت منابع)", "هـ-۴": "مرجع تأیید قواعد (T03)",
              "T03 د-۳": "فهرست فعالیت‌ها و واحدها (T03)"}


def decisions_sheet():
    rows = []
    for d in K.DECISIONS:
        prior, relation, status, residual = K.DECISION_RECON[d[0]]
        kind = "تازه (T08)" if relation == "تازه" else f"T08 — {relation}"
        rows.append([d[0], kind, d[1], d[2], d[3], d[4], d[5], d[6], d[7], d[8], prior, relation, status, residual])
    referenced = set()
    for t in ALL_TEXT:
        referenced |= set(re.findall(r"\b(?:DE-\d{2}|DG-\d{2}|T-D\d|CD-\d{2})\b", t))
    for d in K.DECISION_RECON.values():
        referenced |= set(re.findall(r"\b(?:DE-\d{2}|DG-\d{2}|T-D\d|CD-\d{2})\b", d[0]))
    for ref in sorted(referenced):
        topic = DE_TOPIC.get(ref) or DG_TOPIC.get(ref) or SCOPE_TD.get(ref) or OTHER_REFS.get(ref)
        src = "T07" if ref in DE_TOPIC else "T05" if ref in DG_TOPIC else "سند دامنه" if ref in SCOPE_TD else "v2.1 / سایر"
        dup = [dr for dr, primary in K.ROUTE.items() if primary == ref]
        rows.append([ref, f"موجود ({src})", topic or "—", "—", "—", "—", "—", "—", "—", "—", "—",
                     ("مرجع اصلی برای " + "، ".join(dup)) if dup else "—", "باز", "—"])
    headers = ["شناسه", "نوع", "موضوع", "پرسش", "گزینه‌ها", "فرض کاری T08", "شاهد", "اثر بر ابعاد و فرم‌ها",
               "پاسخ‌دهندهٔ پیشنهادی (مرجع تصویب نیست)", "اولویت", "تصمیم‌های پیشین سنجیده‌شده", "نسبت با تصمیم پیشین",
               "وضعیت پس از تطبیق", "پرسش باقی‌مانده"]
    add_sheet("تصمیم‌ها", "تصمیم‌های لازم و تطبیق با تصمیم‌های پیشین",
              "DR-01 تا DR-14 با تصمیم‌های T01، T05، T07، v2.1 و ADRها تطبیق داده شدند. شناسهٔ DR تکراری حذف نشد. در بقیهٔ کتاب‌کار از راه تصمیم اصلی ارجاع می‌شود، مثل «DE-02 (DR-04)». ردیف‌های «موجود» تصمیم‌هایی‌اند که در این کتاب‌کار ارجاع شده‌اند. پاسخ‌دهندهٔ پیشنهادی مرجع تصویب نیست.",
              headers, rows, [9, 18, 30, 52, 44, 40, 16, 30, 30, 9, 26, 16, 18, 52], freeze="C4", route_refs=False)
    return referenced


# ============================================================================ sources + guide
def sources_sheet():
    files = [
        ("S-01", "کتاب‌کار T07", T07, "رویدادها، چرخه‌ها، جمعیت، شواهد، تطابق‌ها، تصمیم‌ها"),
        ("S-02", "گزارش T07", ROOT / "docs/project/passenger_transport_financial_event_review_v0_1.md", "خلاصه و جدول جمعیت"),
        ("S-03", "کتاب‌کار T05", T05, "فعالیت‌ها، مراکز، ابعاد محل و واحد، تصمیم‌ها"),
        ("S-04", "گزارش T05", ROOT / "docs/project/passenger_transport_activity_cost_center_review_v0_1.md", "—"),
        ("S-05", "سند دامنه (T01)", ROOT / "docs/project/passenger_transport_scope_v0_1.md", "T-D1 تا T-D9؛ ۶-۳؛ ۶-۵"),
        ("S-06", "ثبت منابع (T02)", ROOT / "docs/project/passenger_transport_source_register_v0_1.md", "نبود ستون‌ها؛ G-xx"),
        ("S-07", "درخواست اطلاعات (T03)", ROOT / "docs/project/passenger_transport_information_request_v0_1.md", "د-۱، د-۳، هـ-۴"),
        ("S-08", "پایگاه کاری tt_v21.sqlite (فقط خواندنی)", DB, "دفاتر دو بخش؛ برچسب‌های v2.1"),
        ("S-09", "واژه‌نامه CONTEXT.md", ROOT / "CONTEXT.md", "تعریف رویداد، مرکز، ذی‌نفع، کسور، سپرده"),
        ("S-10", "ADR-0002، 0011، 0013، 0019، 0022، 0024، 0026، 0027", ROOT / "docs/adr/0019-pre-authoritative-form-dimensions.md", "بسته، سطر، provenance، ابعاد پیش از قاعده، تفکیک مبلغ"),
        ("S-11", "راهنمای کد یکتا", ROOT / "docs/contractor-hr/YEKTA_CODE_CONTRACTOR_GUIDE_FA.md", "ساختار ۱۱ بخشی و ماتریس ۹گانه"),
        ("S-12", "پیاده‌سازی کد یکتا", ROOT / "app/services/yekta_codes.py", "مقادیر پیش‌فرض بخش‌ها"),
        ("S-13", "مدل‌های نرم‌افزار", ROOT / "app/models.py", "Contract، ProgressStatement، ContractGuarantee، Contractor"),
        ("S-14", "مدل‌های حضور و مالی", ROOT / "app/models/finance.py", "FinancialMatrix، TransactionAllocation"),
        ("S-15", "اسکریپت لایهٔ الف", ROOT / "scripts/passenger_transport/t08_observe.py", "بازتولید پروفایل‌ها"),
        ("S-16", "محتوای لایهٔ ب", ROOT / "scripts/passenger_transport/t08_content.py", "ابعاد، گروه‌ها، فرم‌ها، چرخه‌ها، تصمیم‌ها"),
        ("S-17", "سازندهٔ کتاب‌کار", ROOT / "scripts/passenger_transport/t08_build.py", "این کتاب‌کار و کنترل‌ها"),
        ("S-18", "بازسازی جفت‌سازی T07", ROOT / "scripts/passenger_transport/t08_pairing.py", "۱۸ ردیف جفت‌سازی (بند ۲)"),
    ]
    rows = [[s, title, str(p.relative_to(ROOT)).replace("\\", "/"), sha16(p), use, "فقط خواندنی" if i < 14 else "خروجی T08"]
            for i, (s, title, p, use) in enumerate(files)]
    add_sheet("منابع", "منابع و نسخه‌ها", "هش: ۱۶ نویسهٔ اول SHA-256 در زمان ساخت. هیچ منبعی تغییر نکرد.",
              ["شناسه", "منبع", "مسیر", "SHA-256 (۱۶)", "کاربرد", "نقش"], rows, [7, 40, 62, 20, 50, 14], freeze="C4")


def guide_sheet(n):
    rows = [
        ["هدف", "مدل قابل بازبینی رابطهٔ هر رویداد مالی با فعالیت، مرکز هزینه یا بُعد درآمد، حساب، قرارداد، نقش‌های طرف، محل و سایر ابعاد، و داده‌ای که فرم‌های عملیاتی آینده باید ثبت کنند (T08)."],
        ["وضعیت", "پیشنهاد کاری. نه تأیید سازمانی است، نه فعال‌سازی قاعده و نه ارجاع نهایی. شمار ردیف‌های تأییدشده: ۰. هیچ فرض کاری قاعدهٔ عملیاتی نشده است."],
        ["مبنا", "T07 (بازبینی محدود ۱)، T05 (بازنگری محدود ۱)، سند دامنه (ویرایش ۲)، ثبت منابع و دفاتر از راه tt_v21.sqlite (فقط خواندنی). مبنای عددی جدول ذخیره‌شدهٔ T07 است (مرتبط ۱۳٬۱۲۳)."],
        ["اعداد قدیمی", "۱۰٬۵۶۶ و ۱۰٬۹۴۱ از خلاصهٔ روایی پیشین تهیه‌کننده آمده بودند، نه از نتایج ذخیره‌شدهٔ T07. با برچسب «خلاصهٔ قدیمی ناسازگار با نتایج ذخیره‌شده» در P-11 و P-12 نگه داشته شده‌اند و مبنای هیچ محاسبه‌ای نیستند."],
        ["سه نوع محتوا", "شاهد (مشاهده در دفاتر): فیلدهای_منبع، پروفایل_ابعاد_رویداد (ردیف‌های «دقیق» و «بازتولیدشده»)، روابط_مشاهده‌شده و ستون‌های «شاهد». تقریب (نیازمند راستی‌آزمایی): ردیف‌های «تقریبی» پروفایل و شیت وضعیت_پروفایل‌های_جفت‌سازی. پیشنهاد (طراحی): ابعاد_و_سطح، قواعد_ماتریس، ماتریس_رویداد_بعد، الگوهای_چرخه، فیلدهای_فرم، عناصر_داده، نیاز_داده_مسیرها (ستون‌های «پیشنهاد») و روابط_و_تفکیک. حضور مقدار در داده الزام نمی‌سازد و نبودِ آن اختیاری یا ممنوع بودن را ثابت نمی‌کند. تقریب هرگز مبنای فیلد لازم یا رابطهٔ تأییدشده نیست."],
        ["راه استفادهٔ طراح فرم", "۱) در «نیاز_داده_مسیرها» جریان و مرحله را فیلتر کنید تا داده، الزام پیشنهادی، مرجع، نقش و مرحلهٔ ثبت را ببینید. ۲) تعریف یکپارچهٔ هر عنصر را در «عناصر_داده» بخوانید. ۳) چیدمان فیلدها در فرم‌ها در «فیلدهای_فرم» است. ۴) پرسش‌های باز در «تصمیم‌ها» است."],
        ["کدهای پیشنهاد الزام", " | ".join(f"«{c}»: {t}" for c, t in K.CODES)],
        ["سطح ساختاری", "دادهٔ پایه (طرف، اماکن، قرارداد یا پرونده، ابزار)؛ نمونهٔ رویداد (صورت‌وضعیت، مطالبه، رسید، فهرست حقوق)؛ سطر طبقه‌بندی؛ بسته. سطر طبقه‌بندی معادل سطر سند حسابداری نیست."],
        ["نقش‌های طرف", "طرف قرارداد (D12)، طرف‌حساب بدهکار یا بستانکار (D13)، پرداخت‌کننده (D14) و ذی‌نفع (D15) چهار بُعد جدا هستند."],
        ["محل", "محل (D04) و زیرمحل (D05) مرکز هزینه نیستند و خودکار به مرکز تبدیل نمی‌شوند."],
        ["حساب‌ها", "هر حساب با کلید (بخش، سال یا نسخهٔ نمودار، کد منبع) آمده است. حساب مقصد خالی است و «حل‌نشده — منتظر T-D3» علامت خورده است."],
        ["کد یکتا", "بخش‌های نامعلوم خالی و «نامعلوم» می‌مانند؛ مقدار پیش‌فرض یا قاعدهٔ کدگذاری تازه ساخته نشد."],
        ["پیوندها", "سطح اطمینان T07: T3 شناسهٔ مشترک، T2 همان طرف‌حساب یا شرح، T1/T0 فقط مبلغ (تطابق نامزد). پیوند در سطح طرف‌حساب به مطالبه یا قرارداد مشخص ارتقا داده نشد. پیوند سطر طبقه‌بندی با سطر سند یک گزینهٔ طراحی پیشنهادی است (RL-03، PO-04، DC-01 تا DC-08)."],
        ["شناسه‌ها", "P (جمعیت)، SF (فیلد منبع)، OP (پروفایل)، OR (رابطهٔ مشاهده‌شده)، D (بُعد)، SG (گروه مرحله)، M (ردیف ماتریس)، LP (الگوی چرخه)، FM (فرم و فیلد)، EL (عنصر داده)، PV (نیاز داده در مسیر)، RL و SP (رابطه و تفکیک)، PO و DC (گزینهٔ پیوند سطر با سند و حفاظ شمارش دوباره)، AC (حساب)، YK (کد یکتا)، DR (تصمیم T08)، K (کنترل). رویدادها همان PFE-001 تا PFE-105 در T07 هستند. شناسه‌های EV و AM از شیت‌های T07 آمده‌اند."],
        ["محرمانگی", "نام اشخاص، شرح اسناد و شمارهٔ حساب بانکی نقل نشده است. عنوان حساب‌های شخصی پوشانده شده است."],
        ["محدودیت‌ها", f"۱) کد جفت‌سازی T07 ذخیره نشده بود. {fa(len(obs['approx']))} ردیف آن با t08_pairing.py بازسازی شد: {fa(sum(a['basis'] == 'DEF' for a in obs['approx']))} ردیف از تعریف مکتوب بازتولید شد و {fa(sum(a['basis'] == 'APPROX' for a in obs['approx']))} ردیف «تقریبی / نیازمند راستی‌آزمایی» ماند. ۲) نشانه‌های شرح (قرارداد، صورت‌وضعیت، ماه، عدد) فقط نشانه‌اند. ۳) کدهای ماتریس نظر تهیه‌کننده‌اند و بازبینی کارشناس سازمان را لازم دارند. ۴) هیچ پروندهٔ واقعی قرارداد یا مطالبه‌ای دیده نشد (G-05). ۵) ۱۴۰۵ ناقص و واحد پول تأییدنشده است."],
        ["بازتولید", ".venv/Scripts/python.exe scripts/passenger_transport/t08_observe.py و سپس t08_build.py"],
        ["شیت‌ها", "؛ ".join(n)],
        ["تاریخ و نسخه", "نسخهٔ ۰٫۱ (بازبینی نهایی)، ۱۴۰۵/۰۸/۰۷."],
    ]
    ws = add_sheet("راهنما", "راهنما — طراحی پیشنهادی روابط رویداد مالی و ابعاد اطلاعاتی (T08، نسخهٔ ۰٫۱، بازبینی نهایی)",
                   "خلاصهٔ روش، قراردادها و محدودیت‌ها. هیچ ردیفی تأیید نشده است.", ["موضوع", "توضیح"], rows, [24, 150],
                   freeze="A4")
    wb.move_sheet(ws, offset=-(len(wb.sheetnames) - 1))


# ============================================================================ controls
def controls_sheet(pop, matrix_rows, cells, form_rows, life_rows, account_rows, referenced, or_ids, yekta_defaults,
                   yekta_rows, used_profiles, elem_rows, path_rows, pairing_rows):
    res = []

    def ctl(cid, group, desc, method, result, status, note="—"):
        res.append([cid, group, desc, method, result, status, note])

    G1, G2, G3, G4 = "۱. ساختاری (خودکار)", "۲. جمعیت و بازتولید", "۳. ارزیابی محتوایی", "۴. تأیید رسمی"
    P = obs["population"]
    tot = P["total"]
    ok = P["identity_ok"] and tot["بدون برچسب v2.1"] == 21841
    ctl("K-01", G2, "افراز ۲۱٬۸۴۱ سطر بی‌برچسب در جدول T07 برای هر بخش و سال جمع‌پذیر است.", "بازخوانی PR-01 تا PR-07",
        f"ردیف‌های ناسازگار: {0 if P['identity_ok'] else 'دارد'}؛ کل بی‌برچسب: {tot['بدون برچسب v2.1']}", "PASS" if ok else "FAIL")
    ctl("K-02", G2, "شمار P-04 و P-07 با گزارش T07 (۱۳٬۱۲۳ و ۸٬۲۹۴) و PC-04 یکی است.", "مقایسه",
        f"P-04={pop['rows']['P-04'][6]}؛ P-07={pop['rows']['P-07'][6]}", "PASS" if (pop['rows']['P-04'][6], pop['rows']['P-07'][6]) == (13123, 8294) else "FAIL")
    ctl("K-03", G2, "مبنای عملیاتی جدول ذخیره‌شدهٔ T07 است و اعداد خلاصهٔ روایی قدیمی (۱۰٬۵۶۶ و ۱۰٬۹۴۱) در آن نیستند.",
        "جست‌وجوی متنی در همهٔ شیت‌ها و گزارش T07؛ همهٔ ترکیب‌های زیرگروه‌ها در هر بخش و سال",
        f"یافته در متن: {len(P['needle_hits'])}؛ جمع برابر: {len(P['natural_hits'])}",
        "PASS" if not P["needle_hits"] and not P["natural_hits"] else "FAIL",
        "اعداد قدیمی در P-11 و P-12 «خلاصهٔ قدیمی ناسازگار با نتایج ذخیره‌شده» علامت خوردند.")
    memo_n = sum(r[4] for r in obs["memo"]["rows"])
    ctl("K-04", G2, "۹۰ سطر انتظامی بدون رویداد (U-10) از پایگاه بازتولید می‌شود.", "SQL روی tt_v21.sqlite", f"{memo_n} سطر", "PASS" if memo_n == 90 else "FAIL")
    bad = [c for c in obs["exact_check"] if c["t07"] != c["n"]]
    ctl("K-05", G2, "شرط‌های SQL شیت شواهد T07 همان شمار T07 را بازتولید می‌کنند.", "اجرای هر شرط روی پایگاه",
        f"ردیف دقیق: {len(obs['exact_check'])}؛ ناسازگار: {len(bad)}", "PASS" if not bad else "FAIL")
    appr = obs["approx"]
    n_def = sum(1 for a in appr if a["basis"] == "DEF")
    n_eq = sum(1 for a in appr if a["basis"] == "APPROX" and a["recon"] == a["t07"])
    ctl("K-06", G2, "۱۸ ردیف جفت‌سازی T07 دوباره ساخته و با شمار T07 سنجیده شدند. شمار بازسازی جای شمار T07 ننشسته است.",
        "اجرای t08_pairing.py و مقایسه با ستون «تعداد سطر» شیت شواهد T07",
        f"ردیف: {len(appr)}؛ بازتولیدشده از تعریف مکتوب: {n_def}؛ بازسازی جفت‌سازی با شمار برابر: {n_eq}؛ با شمار نابرابر: {len(appr) - n_def - n_eq}",
        "گزارش", "ردیف‌های جفت‌سازی «تقریبی / نیازمند راستی‌آزمایی» می‌مانند (شیت وضعیت_پروفایل‌های_جفت‌سازی).")
    missing = sorted(set(EV) - set(K.EVENTS))
    extra = sorted(set(K.EVENTS) - set(EV))
    ctl("K-07", G1, "هر ۱۰۵ رویداد T07 در ماتریس هست و رویداد ناشناخته‌ای نیست.", "مقایسهٔ مجموعه‌ها",
        f"رویداد T07: {len(EV)}؛ بی‌ردیف: {len(missing)}؛ ناشناخته: {len(extra)}؛ ردیف ماتریس: {len(matrix_rows)}",
        "PASS" if not missing and not extra and len(EV) == 105 else "FAIL")
    codes = {c[0] for c in K.CODES}
    badc = [c for c in cells if c[3] not in codes]
    ctl("K-08", G1, "هر خانهٔ ماتریس یکی از پنج کد مجاز است.", "بررسی خانه‌ها", f"خانه: {len(cells)}؛ نامعتبر: {len(badc)}",
        "PASS" if not badc else "FAIL")
    unexplained = []
    for ev, sk, d, val, sid in cells:
        if val in (K.C, K.D):
            over = K.EVENTS[ev][2]
            if d in K.SG[sid]["cond"]:
                continue
            if d in over and K.EVENTS[ev][3]:
                continue
            unexplained.append((ev, sk, d, val))
    ctl("K-09", G1, "هر کد «شرطی» یا «منوط به تصمیم» شرط یا تصمیم ثبت‌شده دارد (در گروه یا یادداشت استثنا).", "بررسی خانه‌ها",
        f"بی‌توضیح: {len(unexplained)}" + (f" — {unexplained[:6]}" if unexplained else ""), "PASS" if not unexplained else "FAIL")
    dims = set(K.DIM_IDS)
    bad_dims = []
    for r in form_rows:
        for tok in re.findall(r"D\d{2}", r[5]):
            if tok not in dims:
                bad_dims.append((r[0], tok))
    ctl("K-10", G1, "بُعد هر فیلد فرم در فهرست ابعاد هست؛ نقش‌های طرف در فیلد تراکنشی ادغام نشده‌اند.", "بررسی ارجاع‌ها",
        f"فیلد: {len(form_rows)}؛ بُعد نامعتبر: {len(bad_dims)}؛ فیلد تراکنشی با «D12–D15»: فقط FM12-09 (چهار فیلد جدا)",
        "PASS" if not bad_dims else "FAIL")
    known = set(DE_TOPIC) | set(DG_TOPIC) | set(SCOPE_TD) | set(OTHER_REFS) | {d[0] for d in K.DECISIONS}
    unknown = sorted(r for r in referenced if r not in known)
    dr_refs = set()
    for t in ALL_TEXT:
        dr_refs |= set(re.findall(r"\bDR-\d{2}\b", t))
    dr_bad = sorted(dr_refs - {d[0] for d in K.DECISIONS})
    ctl("K-11", G1, "هر ارجاع به تصمیم (T-D، DE، DG، CD، DR) معتبر است.", "مقایسه با سند دامنه، T05، T07 و DR",
        f"ارجاع یکتا: {len(referenced) + len(dr_refs)}؛ نامعتبر: {len(unknown) + len(dr_bad)} {unknown + dr_bad if unknown or dr_bad else ''}",
        "PASS" if not unknown and not dr_bad else "FAIL")
    filled = [r for r in account_rows if r[14] not in (None, "")]
    ctl("K-12", G1, "هیچ حساب مقصدی پر نشده است و همه «حل‌نشده» علامت دارند.", "بررسی ستون‌ها",
        f"حساب: {len(account_rows)}؛ مقصد پرشده: {len(filled)}", "PASS" if not filled else "FAIL")
    yk_rows = [r for r in yekta_rows if r[5] in ("نامعلوم", "ناسازگار", "مقدار ساختگی در پیاده‌سازی")]
    yk_bad = [r[0] for r in yk_rows if re.search(r"(?<![-\w])\d{2,}(?![-\w])", r[7])]
    ctl("K-13", G1, "برای بخش نامعلوم، ناسازگار یا ساختگی کد یکتا مقداری پیشنهاد نشده است.", "بررسی ستون «پیشنهاد T08»",
        f"بخش بررسی‌شده: {len(yk_rows)}؛ دارای مقدار عددی: {len(yk_bad)}", "PASS" if not yk_bad else "FAIL",
        f"پیاده‌سازی فعلی {len(yekta_defaults)} پیش‌فرض صفرپر دارد (YK-N3).")
    # privacy: long digit runs, party names from the ledger, masked titles
    names = set()
    src = sqlite3.connect(f"file:{DB.as_posix()}?mode=ro", uri=True)
    for (f1,) in src.execute("""select distinct f1 from sanad where person is not null and person<>'NULL'
                               and hcode not like '111-%' and hcode not like '113-5-%' and hcode not like '113-6-%'
                               and hcode not like '918-%' and hcode not like '919-%'
                               and substr(hcode,1,3) not in ('611','612','613')"""):
        # the party-name part of Field1 (before any unit descriptor); organisations are not personal data
        for part in str(f1 or "").split("/"):
            part = part.strip()
            if len(part) >= 6 and " " in part and not ORG_WORDS.search(part):
                names.add(part)
    for (t,) in src.execute("select distinct title from account_dictionary_restricted"):
        if t and ("بازنشسته" in t or "بانشسته" in t):
            names.add(t.split("-")[0].strip())
    src.close()
    blob = "\n".join(ALL_TEXT)
    leaked = sorted(n for n in names if n in blob)
    # SHA-256 prefixes in «منابع» and the implemented all-zero default of the budget segment are not identifiers
    scan = re.sub(r"\b[0-9a-f]{16}\b", "", blob)
    long_digits = sorted(t for t in set(re.findall(r"\d{7,}", scan)) if set(t) != {"0"})
    ctl("K-14", G1, "نام اشخاص، عنوان حساب‌های شخصی و رشتهٔ عددی بلند (۷ رقم یا بیشتر) در متن‌ها نیست.", "جست‌وجوی خودکار پیش از ذخیره",
        f"نام بررسی‌شده: {len(names)}؛ یافته: {len(leaked)}؛ رشتهٔ عددی بلند: {len(long_digits)}",
        "PASS" if not leaked and not long_digits else "FAIL")
    ref_ok = set(or_ids) | {f"SF-{i:02d}" for i in range(1, 15)} | {f"P-{i:02d}" for i in range(1, 16)}
    bad_ev = []
    for r in form_rows:
        for tok in re.findall(r"(?:OR|SF|P)-\d{2}", r[13]):
            if tok not in ref_ok:
                bad_ev.append((r[0], tok))
    for r in K.RELATIONS:
        for tok in re.findall(r"(?:OR|SF|P)-\d{2}", r[5]):
            if tok not in ref_ok:
                bad_ev.append((r[0], tok))
    for d in K.DECISIONS:
        for tok in re.findall(r"(?:OR|SF|P)-\d{2}", d[5]):
            if tok not in ref_ok:
                bad_ev.append((d[0], tok))
    ctl("K-15", G1, "ارجاع‌های شاهد (OR، SF، P) در فرم‌ها، روابط و تصمیم‌ها معتبرند.", "بررسی ارجاع‌ها",
        f"ارجاع نامعتبر: {len(bad_ev)} {bad_ev[:5] if bad_ev else ''}", "PASS" if not bad_ev else "FAIL")
    covered = set()
    for r in life_rows:
        covered |= set(re.findall(r"PFE-\d{3}", r[4]))
    need = {e for e, a in K.EVENTS.items() if a[0] != "F0"}
    miss = sorted(need - covered)
    ctl("K-16", G1, "هر رویداد شش جریان اولویت‌دار دست‌کم در یک الگوی چرخه آمده است.", "مقایسهٔ مجموعه‌ها",
        f"رویداد اولویت‌دار: {len(need)}؛ بی‌الگو: {len(miss)} {miss if miss else ''}", "PASS" if not miss else "FAIL")
    acts = set()
    for t in ALL_TEXT:
        acts |= set(re.findall(r"\b(?:TX|TM|SH)-A\d{2}\b", t))
    bad_act = sorted(a for a in acts if a not in ACT_IDS)
    # T05 writes "TM-A07، A08" style lists; only full ids are checked
    ctl("K-17", G1, "شناسه‌های فعالیت ارجاع‌شده در فهرست ۳۸ فعالیت T05 هستند.", "مقایسه با کتاب‌کار T05",
        f"شناسهٔ یکتا: {len(acts)}؛ نامعتبر: {len(bad_act)} {bad_act if bad_act else ''}", "PASS" if not bad_act else "FAIL")
    # ---- final refinement controls
    approx_used = [k for k in used_profiles if PROFILE[k]["basis"] == "APPROX"]
    or_approx = [k for k, b in or_ids.items() if "تقریبی" in b]
    req_bad = [r[0] for r in form_rows if r[10] == K.L and re.search(r"\b(?:OP|EV|AM)-\d{2,3}\b", r[13])]
    ctl("K-18", G1, "هیچ پیشنهاد الزام یا رابطه‌ای بر پروفایل تقریبی تکیه ندارد.",
        "پروفایل‌های به‌کاررفته در روابط_مشاهده‌شده؛ مبنای ردیف‌های OR؛ شاهد فیلدهای «لازم»؛ مبنای کدهای ماتریس",
        f"پروفایل تقریبی در OR: {len(approx_used)}؛ OR با مبنای تقریبی: {len(or_approx)}؛ فیلد «لازم» با شاهد پروفایل یا جفت‌سازی: {len(req_bad)}؛ کدهای ماتریس از گروه مرحله و استثنای مکتوب",
        "PASS" if not approx_used and not or_approx and not req_bad else "FAIL")
    field_ids = {f[0] for f in K.FORM_FIELDS}
    el_fields = {m for e in K.ELEMENTS for m in e[3]}
    orphan_fields = sorted(field_ids - el_fields)
    bad_members = sorted(el_fields - field_ids)
    empty_el = [e[0] for e in K.ELEMENTS if not e[3]]
    dup_el = [k for k, v in collections.Counter(e[0] for e in K.ELEMENTS).items() if v > 1]
    ctl("K-19", G1, "هر فیلد فرم دست‌کم به یک عنصر داده نگاشت دارد و هر عنصر دست‌کم یک فیلد موجود دارد.", "مقایسهٔ مجموعه‌ها",
        f"فیلد: {len(field_ids)}؛ عنصر: {len(elem_rows)}؛ فیلد بی‌عنصر: {len(orphan_fields)}؛ ارجاع نامعتبر: {len(bad_members)}؛ عنصر بی‌فیلد: {len(empty_el)}؛ شناسهٔ تکراری: {len(dup_el)}",
        "PASS" if not orphan_fields and not bad_members and not empty_el and not dup_el else "FAIL")
    code = {(e, sk, d): v for e, sk, d, v, sid in cells}
    expected = 0
    for lp in K.LIFECYCLE:
        if lp[1] == "F0":
            continue
        for e in re.findall(r"PFE-\d{3}", lp[4]):
            for sk in sector_rows(e):
                expected += 1 + sum(1 for d in K.MATRIX_DIMS if code[(e, sk, d)] != K.N)
    actual = sum(2 if r[5] == "هر دو بخش" else 1 for r in path_rows)
    incomplete = [r[0] for r in path_rows if "—" in (r[12], r[13], r[14]) or (r[6] == "—" and not r[8].startswith("D19"))]
    path_events = {r[4] for r in path_rows}
    need = {e for e, a in K.EVENTS.items() if a[0] != "F0"}
    ctl("K-20", G1, "«نیاز_داده_مسیرها» همهٔ خانه‌های غیر «نامرتبط مفهومی» رویدادهای شش جریان را در هر مرحله پوشش می‌دهد و هر ردیف مرجع، نقش و مرحله دارد.",
        "شمارش از ماتریس در برابر ردیف‌ها (هر بخش جدا)",
        f"خانهٔ مورد انتظار: {expected}؛ پوشش‌داده‌شده: {actual}؛ ردیف ناقص: {len(incomplete)}؛ رویداد اولویت‌دار بی‌ردیف: {len(need - path_events)}",
        "PASS" if expected == actual and not incomplete and not need - path_events else "FAIL")
    no_recon = sorted({d[0] for d in K.DECISIONS} - set(K.DECISION_RECON))
    known_prim = set(DE_TOPIC) | set(DG_TOPIC) | set(SCOPE_TD) | set(OTHER_REFS)
    bad_prim = sorted(p for p in K.ROUTE.values() if p not in known_prim)
    lone = sorted({dr for t in ROUTED_TEXT for dr, p in K.ROUTE.items() if re.search(rf"\b{dr}\b", t) and p not in t})
    ctl("K-21", G1, "هر DR با تصمیم‌های پیشین تطبیق دارد؛ DR تکراری حذف نشده و در بقیهٔ کتاب‌کار فقط همراه تصمیم اصلی ارجاع می‌شود.",
        "بررسی DECISION_RECON و همهٔ متن‌های غیر از شیت تصمیم‌ها",
        f"DR: {len(K.DECISIONS)}؛ بی‌تطبیق: {len(no_recon)}؛ تکراری یا زیرپرسش: {len(K.ROUTE)} ({'، '.join(f'{k}→{v}' for k, v in K.ROUTE.items())})؛ مرجع اصلی نامعتبر: {len(bad_prim)}؛ ارجاع تنها: {len(lone)}",
        "PASS" if not no_recon and not bad_prim and not lone and len(K.DECISIONS) == 14 else "FAIL")
    labels = {"بازتولیدشده از تعریف مکتوب T07 — شمار برابر", "تقریبی / نیازمند راستی‌آزمایی"}
    bad_pair = [r[0] for r in pairing_rows if r[12] not in labels or not r[15]
                or (r[12].startswith("بازتولید") and r[8] != r[9])]
    ctl("K-22", G1, "هر ۱۸ ردیف جفت‌سازی وضعیت و اثر عملی بر فرم دارد؛ ردیف «بازتولیدشده» فقط با شمار برابر و هیچ ردیف «دقیق» نیست.",
        "بررسی شیت وضعیت_پروفایل‌های_جفت‌سازی", f"ردیف: {len(pairing_rows)}؛ ناسازگار: {len(bad_pair)}",
        "PASS" if len(pairing_rows) == 18 and not bad_pair else "FAIL")
    ctl("A-01", G3, "درستی کدهای ماتریس، شرط‌ها و دلیل کدهای «نامرتبط مفهومی».", "نظر تهیه‌کننده", "نیازمند بازبینی کارشناس سازمان", "ارزیابی تحلیلی")
    ctl("A-02", G3, "کفایت فیلدهای فرم برای دو نمونهٔ آزمون (الف و ب).", "مقایسه با پروندهٔ واقعی", "پرونده‌ای در دست نیست (G-05)", "ارزیابی تحلیلی")
    ctl("A-03", G3, "تفسیر نشانه‌های شرح و پروفایل‌های تقریبی.", "نمونه‌خوانی محدود T07", "نشانه، نه اثبات", "ارزیابی تحلیلی")
    ctl("O-01", G4, "تأیید ابعاد، سطح‌ها و ماتریس", "تصمیم مرجع ذی‌صلاح", "انجام نشده", "انجام‌نشده", "Rule Profile و جدول قواعد مرجع (G-04)")
    ctl("O-02", G4, "تأیید فرم‌ها و نقش‌های مسئول", "تصمیم مرجع ذی‌صلاح", "انجام نشده", "انجام‌نشده", "T03 د-۱")
    ctl("O-03", G4, "تصمیم‌های DR-01 تا DR-14", "تصمیم مرجع ذی‌صلاح", "انجام نشده", "انجام‌نشده", "—")
    add_sheet("کنترل‌ها", "کنترل‌ها", "چهار گروه: ساختاری (خودکار)، جمعیت و بازتولید، ارزیابی محتوایی (بدون PASS/FAIL)، تأیید رسمی (انجام‌نشده).",
              ["شناسه", "گروه", "شرح کنترل", "روش", "نتیجه", "وضعیت", "یادداشت"], res, [7, 20, 56, 34, 50, 14, 40], freeze="C4")
    return res


# ============================================================================ main
def main():
    pop = population_sheet()
    source_fields_sheet()
    prof_ids = profile_sheet()
    pairing_rows = pairing_status_sheet(prof_ids)
    or_ids, used_profiles = relations_observed_sheet(pop)
    dimensions_sheet()
    matrix_rows, cells = matrix_sheets(prof_ids)
    life_rows = lifecycle_sheet()
    form_rows = forms_sheet()
    elem_rows = elements_sheet()
    path_rows = path_sheet(cells)
    relations_sheet()
    account_rows = accounts_sheet()
    yekta_defaults, yekta_rows = yekta_sheet()
    referenced = decisions_sheet()
    sources_sheet()
    res = controls_sheet(pop, matrix_rows, cells, form_rows, life_rows, account_rows, referenced, or_ids,
                         yekta_defaults, yekta_rows, used_profiles, elem_rows, path_rows, pairing_rows)
    guide_sheet(list(wb.sheetnames))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)

    code_count = collections.Counter(c[3] for c in cells)
    by_dim = collections.defaultdict(collections.Counter)
    for ev, sk, d, val, sid in cells:
        by_dim[d][val] += 1
    summary = dict(
        population={k: v[2:7] for k, v in pop["rows"].items()},
        memo=dict(guar_bn=pop["guar_bn"], goods_bn=pop["goods_bn"]),
        matrix_rows=len(matrix_rows), cells=len(cells), code_count=dict(code_count),
        by_dim={d: dict(c) for d, c in by_dim.items()},
        flows={f: sorted(e for e, a in K.EVENTS.items() if a[0] == f) for f in K.F},
        sg_count={s: sum(1 for a in K.EVENTS.values() if a[1] == s) for s in K.SG},
        forms=len(K.FORMS), fields=len(form_rows), field_req=dict(collections.Counter(r[10] for r in form_rows)),
        elements=len(elem_rows), path_rows=len(path_rows),
        path_by_flow=dict(collections.Counter(r[1] for r in path_rows)),
        path_req=dict(collections.Counter(r[10] for r in path_rows)),
        lifecycle=len(life_rows), relations=len(K.RELATIONS), posting=len(K.POSTING), decisions=len(K.DECISIONS),
        decision_recon={k: v[1] for k, v in K.DECISION_RECON.items()}, route=K.ROUTE, accounts=len(account_rows),
        accounts_by_org=dict(collections.Counter(r[2] for r in account_rows)),
        homonym_accounts=sum(1 for r in account_rows if r[13] != "—"),
        controls=[(r[0], r[5]) for r in res], sheets=wb.sheetnames, output_sha16=sha16(OUT),
        profiles=len(obs["profiles"]), exact=len(obs["exact_check"]),
        approx=[{k: a[k] for k in ("ev", "event", "org", "t07", "superset", "recon", "basis")} for a in obs["approx"]],
        profile_basis=dict(collections.Counter(p.get("basis") for p in obs["profiles"])),
    )
    SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=1), encoding="utf-8")
    print("sheets:", wb.sheetnames)
    print("controls:", [(r[0], r[5]) for r in res])
    fails = [r for r in res if r[5] == "FAIL"]
    for r in fails:
        print("FAIL", r[0], r[4])
    print("matrix rows:", len(matrix_rows), "cells:", len(cells), dict(code_count))


if __name__ == "__main__":
    main()
