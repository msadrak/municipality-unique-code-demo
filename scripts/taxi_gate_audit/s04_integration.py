"""Step 4 — connect the reassessed crosswalk to the prior Unique Code artefacts and the 1404–1405 Taxi/Terminals ledger.

* line-level lineage (111,600 source lines; narration excluded — RESTRICTED) and monetary reconciliation
* sub-account split evidence for 1→N accounts
* Account ↔ canonical financial event (observed in the prior v2.1 provisional classification)
* Activity ↔ cost center (observed in the prior v2.1 provisional classification)
* dimensional (observed / proposed / approved) cardinality per event
* taxi / terminal domain evidence (colleague files vs ledger)
* CD-01 terminal agency evidence recomputed from the ledger

Outputs: data/work/taxi_gate_audit/integration.json and data/work/taxi_gate_audit/lineage_lines.csv
"""
import csv
import json
import os
import re
from collections import Counter, OrderedDict, defaultdict

import openpyxl

from common import (TT_DB, FE_V71, CC_MASTER, V21_EVENTS, V21_DIR, V21_DECISIONS, V21_AGENCY, WORK, ROOT, dump, load, norm_fa,
                    ro_sqlite, mask_title, sha256)
from expert_crosswalk import ENTRIES, SUB_SPLIT, SPLIT_ACCOUNTS, DOMAIN_TERMS, REL_EN

MIG_SINGLE = ('EQ', 'RN', 'RC', 'CR')


def mask_hcode(hc):
    p = hc.strip().split('-')
    if len(p) == 4 and len(p[3]) >= 5:
        p[3] = '•' * (len(p[3]) - 3) + p[3][-3:]
    return '-'.join(p)


def main():
    coa = load('coa.json')
    mt = load('matching.json')
    old = {x['code']: x for x in coa['old_rows']}
    new_title = {}
    for x in coa['new_rows']:
        if x['code'] is not None:
            new_title.setdefault(x['code'], str(x['title']).strip())
    audit = {a['old_code']: a for a in mt['audit']}
    ent = defaultdict(dict)
    for e in ENTRIES:
        ent[e['code']][e['org']] = e

    def entry_for(code, org):
        d = ent[code]
        return d.get(org) or d.get('BOTH') or next(iter(d.values()))

    tc = ro_sqlite(TT_DB)
    # ---------------------------------------------------------------- lineage
    l3 = {(x['kol'], x['moeen'], x['taf1']) for x in coa['old_rows'] if x['level'] == 3}
    l2 = {(x['kol'], x['moeen']) for x in coa['old_rows'] if x['level'] == 2}
    q = """SELECT s.sid,s.org,s.fy,s.src_file,s.src_sheet,s.src_row,s.doc_key,s.sanad_num,s.date,s.kol,s.moeen,s.taf1,s.taf2,s.bed,s.bes,s.hcode,
                  e.event_id,e.role,e.outcome,e.conf,c.eligible,c.cc_status,c.cc_id,c.activity_id,c.conf,d.kind,s.person,s.vahed,s.markaz
           FROM sanad s LEFT JOIN line_event e ON e.sid=s.sid LEFT JOIN line_cc c ON c.sid=s.sid LEFT JOIN doc_class d ON d.doc_key=s.doc_key
           ORDER BY s.org,s.fy,s.src_row"""
    lineage_path = os.path.join(WORK, 'lineage_lines.csv')
    fh = open(lineage_path, 'w', encoding='utf-8-sig', newline='')
    w = csv.writer(fh)
    cols = ['Lineage_ID', 'Source_Record_ID', 'Source_File', 'Source_Sheet', 'Source_Row', 'Organization', 'Fiscal_Year', 'Doc_Key',
            'Doc_No', 'Date', 'Line_Ref', 'Original_Account_HCode', 'Old_List_Code', 'Colleague_Best_New', 'Colleague_Verdict',
            'Proposed_New_Code', 'Proposed_Relation', 'Match_Status', 'Matching_Key', 'Matching_Method', 'Debit', 'Credit',
            'Budget_Code', 'Prior_Event', 'Prior_Event_Role', 'Prior_CC_Status', 'Prior_CC', 'Prior_Activity', 'Doc_Kind', 'Description']
    w.writerow(cols)
    n = 0
    tot = Counter()
    by_bucket = defaultdict(lambda: [0, 0, 0])
    by_verdict = defaultdict(lambda: [0, 0, 0])
    sid_seen = set()
    doc_dest = defaultdict(set)
    doc_bucket = defaultdict(set)
    acc_event = defaultdict(lambda: defaultdict(lambda: [0, 0, 0, Counter()]))
    acc_ops = defaultdict(int)
    acc_cc = defaultdict(lambda: defaultdict(lambda: [0, 0, 0]))
    act_cc = defaultdict(lambda: [0, 0, 0, Counter()])
    split_sub = defaultdict(lambda: defaultdict(lambda: [0, 0, 0]))
    ev_dims = defaultdict(lambda: Counter())
    ev_docs = defaultdict(lambda: defaultdict(set))
    per_of = defaultdict(lambda: [0, 0, 0])
    for row in tc.execute(q):
        (sid, org, fy, sfile, ssheet, srow, doc_key, sanad_num, date, kol, mo, t1, t2, bed, bes, hc, ev, role, outcome, econf,
         elig, ccst, ccid, act, cconf, kind, person, vahed, markaz) = row
        k, m, t = int(kol), int(mo), int(t1)
        bed, bes = bed or 0, bes or 0
        if sid in sid_seen:
            raise SystemExit(f'duplicate join for {sid}')
        sid_seen.add(sid)
        n += 1
        tot['lines'] += 1
        tot['debit'] += bed
        tot['credit'] += bes
        p = per_of[(org, fy)]
        p[0] += 1; p[1] += bed; p[2] += bes
        if (k, m, t) in l3:
            code = f'{k}-{m}-{t}-0'
        elif (k, m) in l2:
            code = f'{k}-{m}-0-0'
        else:
            code = None
        if code is None:
            bucket, dest, rel, verdict, best = 'NOT_IN_COLLEAGUE_LIST', '', '', 'NOT_IN_COLLEAGUE_LIST', ''
        else:
            a = audit[code]
            verdict, best = a['verdict'], a['best_code']
            e = entry_for(code, org)
            if code in SPLIT_ACCOUNTS and e['status'] in ('SPLIT_1_TO_N', 'EXPERT_REVIEW_REQUIRED'):
                sub = SUB_SPLIT.get((org, f'{k}-{m}-{t}'))
                if sub is None:
                    bucket, dest, rel = 'SPLIT_SUBACCOUNT_UNRESOLVED', '', 'NO_SUB_RULE'
                elif sub[0] is None:
                    bucket, dest, rel = 'SPLIT_NO_EQUIVALENT', '', REL_EN[sub[1]]
                else:
                    bucket = 'SPLIT_ALLOCATED' if sub[1] == 'SP' else 'SPLIT_ALLOCATED_REVIEW'
                    dest, rel = sub[0], REL_EN[sub[1]]
                split_sub[(code, org)][(f'{k}-{m}-{t}', sub[0] if sub else 'NO_RULE', REL_EN[sub[1]] if sub else '')][0] += 1
                ss = split_sub[(code, org)][(f'{k}-{m}-{t}', sub[0] if sub else 'NO_RULE', REL_EN[sub[1]] if sub else '')]
                ss[1] += bed; ss[2] += bes
            else:
                tg = [x for x in e['targets'] if x[1] in MIG_SINGLE and x[0] is not None]
                if e['status'] == 'SPLIT_1_TO_N':
                    # split basis is not a sub-account (e.g. statutory vs voluntary payroll deductions): needs narration or a decision
                    bucket, dest, rel = 'SPLIT_BASIS_NOT_IN_ACCOUNT_CODE', '', 'SPLIT_COMPONENT'
                elif e['status'] in ('NO_SUITABLE_CANDIDATE',) or not tg:
                    bucket, dest, rel = 'NO_SUITABLE_CANDIDATE' if e['status'] == 'NO_SUITABLE_CANDIDATE' else 'EXPERT_REVIEW_NO_TARGET', '', ''
                else:
                    dest, rel = tg[0][0], REL_EN[tg[0][1]]
                    bucket = 'SINGLE_TARGET_STRONG' if tg[0][1] in ('EQ', 'RN') else 'SINGLE_TARGET_REVIEW'
        b = by_bucket[bucket]
        b[0] += 1; b[1] += bed; b[2] += bes
        v = by_verdict[verdict]
        v[0] += 1; v[1] += bed; v[2] += bes
        doc_dest[doc_key].add(dest if dest != '' else bucket)
        doc_bucket[doc_key].add(bucket)
        if code is not None:
            if role in ('CLOSING',):
                pass
            else:
                key = (code, org)
                acc_ops[key] += 1
                ae = acc_event[key][ev or 'UNASSIGNED']
                ae[0] += 1; ae[1] += bed; ae[2] += bes; ae[3][role] += 1
            if elig == 1:
                cc = ccid or ccst
                x = acc_cc[(code, org)][cc]
                x[0] += 1; x[1] += bed; x[2] += bes
        if elig == 1:
            ac = act_cc[(org, act or f'NO_ACTIVITY (CodeMarkaz={markaz})' if org == 'Taxi' else act or 'NO_ACTIVITY', ccid or ccst)]
            ac[0] += 1; ac[1] += bed; ac[2] += bes; ac[3][cconf] += 1
        if ev and role not in ('CLOSING',):
            d = ev_dims[ev]
            d['lines'] += 1
            d['person'] += person not in (None, '', 'NULL', '0')
            d['unit'] += vahed not in (None, '', 'NULL')
            d['activity'] += markaz not in (None, '', 'NULL') and org == 'Taxi'
            d['taxi_lines'] += org == 'Taxi'
            d['cc_eligible'] += elig == 1
            d['cc_assigned'] += (elig == 1 and ccst == 'PROVISIONAL_EXISTING_CC')
            d['cc_multi'] += (elig == 1 and ccst == 'MULTIPLE_CC_CANDIDATES')
            ev_docs[ev][doc_key].add(f'{k}-{m}')
        w.writerow([f'LN-{n:06d}', sid, os.path.basename(sfile or ''), ssheet, srow, org, fy, doc_key, sanad_num, date, srow, mask_hcode(hc),
                    code or '', best, verdict, dest, rel, bucket, f'{org}|{k}-{m}-{t}', 'CURATED_PROPOSAL_v1+SUB_SPLIT' if code else '',
                    bed, bes, 'NOT_AVAILABLE_IN_SOURCE', ev or '', role or '', ccst or '', ccid or '', act or '', kind or '',
                    'RESTRICTED — narration not exported'])
    fh.close()
    shared_docs = sum(1 for d, s in doc_dest.items() if len(s) > 1)
    docs_multi_bucket = sum(1 for d, s in doc_bucket.items() if len(s) > 1)
    src_tot = tc.execute('SELECT count(*),sum(bed),sum(bes) FROM sanad').fetchone()
    reconciliation = OrderedDict(
        source=dict(lines=src_tot[0], debit=src_tot[1], credit=src_tot[2]),
        lineage=dict(lines=tot['lines'], debit=tot['debit'], credit=tot['credit']),
        per_org_fy={f'{o}-{f}': v for (o, f), v in sorted(per_of.items())},
        proposal_buckets={k: v for k, v in sorted(by_bucket.items())},
        colleague_verdict_buckets={k: v for k, v in sorted(by_verdict.items())},
        documents=len(doc_dest), shared_documents_multi_destination=shared_docs, documents_multi_bucket=docs_multi_bucket,
        unique_sid=len(sid_seen))

    # ---------------------------------------------------------------- split evidence
    l3names = defaultdict(dict)
    for org, hc, name, lvl in tc.execute('SELECT org,hcode,name,level FROM hesab WHERE level=3'):
        p = hc.strip().split('-')
        l3names[(org, f'{p[0]}-{p[1]}-{p[2]}')] = mask_title(name)
    split_rows = []
    split_recon = []
    for (code, org), subs in sorted(split_sub.items()):
        tl = sum(v[0] for v in subs.values())
        td = sum(v[1] for v in subs.values())
        acc = old[code]['usage']
        al = sum(v[0] for kk, v in acc.items() if kk.startswith(org))
        ad = sum(v[1] for kk, v in acc.items() if kk.startswith(org))
        split_recon.append(OrderedDict(old_code=code, org=org, sub_lines=tl, sub_debit=td, account_lines=al, account_debit=ad,
                                       equal=(tl == al and td == ad)))
        for (sub, dest, rel), v in sorted(subs.items(), key=lambda kv: -kv[1][1]):
            split_rows.append(OrderedDict(old_code=code, org=org, sub_account=sub, sub_title=l3names.get((org, sub), '— (بدون تفصیلی/سطح ۲)'),
                                          proposed_new=dest if dest not in (None, 'NO_RULE') else '', new_title=new_title.get(dest, ''),
                                          relation=rel or ('NO_SUB_RULE' if dest == 'NO_RULE' else ''), lines=v[0], debit=v[1], credit=v[2],
                                          share_of_account_debit=round(v[1] / td, 4) if td else None))

    # ---------------------------------------------------------------- canonical events
    wbe = openpyxl.load_workbook(FE_V71, read_only=True, data_only=True)
    FE = OrderedDict()
    for r in wbe['فرهنگ_رویداد_مالی_کاندید'].iter_rows(min_row=2, values_only=True):
        if r[0]:
            FE[r[0]] = dict(title=r[1], family=r[4], nature=r[5], origin='FE v7.1 (candidate dictionary, not approved)')
    v21ev = list(csv.DictReader(open(V21_EVENTS, encoding='utf-8-sig')))
    for r in v21ev:
        if r['event_id'] not in FE:
            FE[r['event_id']] = dict(title=r['title_fa'], family=r['family'], nature=r['nature'], origin=r['origin'] + ' / ' + r['governance_status'][:40])
    v21_shares = {r['event_id']: r for r in v21ev}

    def sem(acc_nat, ev_id, role):
        if ev_id in ('UNASSIGNED',) or ev_id not in FE:
            return 'NOT_TESTABLE'
        nat = str(FE[ev_id]['nature'] or '')
        if acc_nat in ('EXPENSE', 'NONCASH_EXPENSE', 'COST_OF_SALES'):
            return 'COMPATIBLE' if 'هزینه' in nat else 'REVIEW'
        if acc_nat == 'REVENUE':
            return 'COMPATIBLE' if 'درآمد' in nat else 'REVIEW'
        if acc_nat in ('MEMO', 'MEMO_CONTRA'):
            return 'COMPATIBLE' if 'انتظامی' in nat else 'REVIEW'
        if acc_nat in ('NONOP_MIXED', 'EQUITY', 'CLOSING', 'TEST'):
            return 'REVIEW'          # mixed revenue/expense or equity accounts: event relation always needs expert review
        return 'COMPATIBLE_AS_BALANCE_SHEET_LEG'

    acc_ev_rows = []
    for (code, org), evs in sorted(acc_event.items()):
        tot_ops = acc_ops[(code, org)]
        e = entry_for(code, org)
        strong = e['status'] in ('EXACT_EQUIVALENCE_PROPOSED', 'RENAMED_EQUIVALENT', 'RECLASSIFIED')
        new_codes = '; '.join(f"{t[0]} ({REL_EN[t[1]]})" for t in e['targets'] if t[0] is not None and t[1] != 'RL') or '—'
        for ev_id, (nl, deb, cre, roles) in sorted(evs.items(), key=lambda kv: -kv[1][0]):
            share = nl / tot_ops if tot_ops else 0
            support = ('STRONG' if share >= 0.2 and nl >= 10 else 'MODERATE' if (share >= 0.05 and nl >= 3) or nl >= 10 else 'WEAK')
            role = roles.most_common(1)[0][0]
            s = sem(old[code]['nature'], ev_id, role)
            analytically = (support in ('STRONG', 'MODERATE') and s in ('COMPATIBLE', 'COMPATIBLE_AS_BALANCE_SHEET_LEG') and strong
                            and ev_id not in ('UNASSIGNED',) and not str(ev_id).startswith('EXC'))
            acc_ev_rows.append(OrderedDict(
                rel_id=f'AEX-{len(acc_ev_rows) + 1:04d}', old_code=code, org=org, old_title=old[code]['title_masked'], account_nature=old[code]['nature'],
                account_coding_version='Taxi/Terminals affiliated-organisation chart (hesab 1404–1405)', proposed_new_accounts=new_codes,
                crosswalk_status=e['status'], event_id=ev_id, event_title=FE.get(ev_id, {}).get('title', '—'), event_origin=FE.get(ev_id, {}).get('origin', ''),
                relation_type={'DRIVER': 'DRIVER_ACCOUNT', 'COMPONENT': 'COMPONENT_ACCOUNT', 'SETTLEMENT': 'SETTLEMENT_ACCOUNT', 'MEMO': 'MEMO_ACCOUNT'}.get(role, role or ''),
                lines=nl, share_of_account_operational_lines=round(share, 4), debit=deb, credit=cre, roles=dict(roles),
                evidence_source='tt_v21.sqlite line_event (prior v2.1 provisional classification of 1404–1405 Sanad lines)',
                observed_or_inferred='OBSERVED (co-occurrence in provisional classification)', matching_method='ledger co-occurrence by account × event',
                support=support, semantic_compatibility=s,
                analytical_status='ANALYTICALLY_SUPPORTED' if analytically else 'EXPERT_REVIEW_REQUIRED',
                expert_review_status='REQUIRED', approval_status='NOT_APPROVED'))
    # ---------------------------------------------------------------- cost centers
    CC = json.load(open(CC_MASTER, encoding='utf8'))
    ncc = json.load(open(os.path.join(V21_DIR, 'machine_readable', 'new_cost_center_candidates.json'), encoding='utf8'))
    w21 = openpyxl.load_workbook(os.path.join(V21_DIR, 'taxi_terminals_financial_mapping_1404_1405_v2.1.xlsx'), read_only=True, data_only=True)
    act_title = {}
    for r in w21['Activity_Service_Map'].iter_rows(min_row=5, values_only=True):
        if r[0] and str(r[0]).startswith('ACT'):
            act_title.setdefault(r[0], (r[1], r[2], r[3]))

    def cc_title(c):
        if c in CC:
            return CC[c]['title'], CC[c]['status']
        if c in ncc:
            return ncc[c].get('title_v2') or ncc[c].get('title', ''), 'PROVISIONAL TTCC (v2.1) — ' + ncc[c].get('recommendation', '')
        return '', ''
    act_rows = []
    for (org, act, cc), (nl, deb, cre, confs) in sorted(act_cc.items(), key=lambda kv: (kv[0][0], str(kv[0][1]), -kv[1][0])):
        is_cc = str(cc).startswith(('CC-', 'TTCC'))
        t, st = cc_title(cc) if is_cc else ('', '')
        at = act_title.get(act, ('', '', ''))
        act_rows.append(OrderedDict(
            activity_id=act, activity_title=at[0] or ('INSUFFICIENT SOURCE EVIDENCE' if str(act).startswith('NO_ACTIVITY') else ''),
            organization=org, source_code=at[2] or '', cost_center_id=cc if is_cc else '', cost_center_title=t,
            cc_domain=CC[cc]['domain'] if cc in CC else ('حمل‌ونقل (پیشنهاد TTCC)' if str(cc).startswith('TTCC') else ''),
            cc_master_status=st, prior_cc_status=cc if not is_cc else 'PROVISIONAL_EXISTING_CC' if cc.startswith('CC-') else 'NEW_CC_CANDIDATE',
            mapping_method='prior v2.1 rule (CodeMarkaz / CodeVahed / event default / v4 activity rules)', lines=nl, net_debit=deb - cre,
            confidence=dict(confs), historical_source='tt_v21.sqlite line_cc', status='PROPOSED — NOT APPROVED',
            evidence_in_taxi_gate_files='NONE — the colleague workbooks contain no activity, unit, budget or document fields'))
    acc_cc_rows = []
    for (code, org), ccs in sorted(acc_cc.items()):
        tl = sum(v[0] for v in ccs.values())
        named = {c: v for c, v in ccs.items() if str(c).startswith(('CC-', 'TTCC'))}
        acc_cc_rows.append(OrderedDict(old_code=code, org=org, old_title=old[code]['title_masked'], eligible_lines=tl,
                                       distinct_cost_centers=len(named),
                                       cc_distribution='; '.join(f'{c}: {v[0]}' for c, v in sorted(ccs.items(), key=lambda kv: -kv[1][0])[:8]),
                                       derivable_from_account_alone='NO' if len(named) != 1 else 'NOT_PROVEN (single CC observed — still activity-driven)'))

    # ---------------------------------------------------------------- dimensional cardinality per event
    dims = []
    for ev_id, d in sorted(ev_dims.items(), key=lambda kv: -kv[1]['lines']):
        nl = d['lines']
        if nl < 5:
            continue
        per_doc = [len(s) for s in ev_docs[ev_id].values()]
        sh = v21_shares.get(ev_id, {})

        def obs(x):
            if x is None:
                return 'NOT_AVAILABLE'
            return '0 (not observed)' if x == 0 else '1..1 (always present)' if x >= 0.999 else f'0..1 (present on {x:.0%})'
        fam = str(FE.get(ev_id, {}).get('family') or '')
        nat = str(FE.get(ev_id, {}).get('nature') or '')
        revenue = 'درآمد' in nat
        contract_prop = '1..1' if 'پیمان' in fam else 'Needs Decision'
        cc_prop = 'Needs Decision (CD-03: درآمد ≠ مرکز هزینه)' if revenue else ('1..N (تخصیص)' if 'هزینه' in nat else '0..0 (ترازنامه‌ای)')
        dims.append(OrderedDict(
            event_id=ev_id, event_title=FE.get(ev_id, {}).get('title', '—'), lines=nl, documents=len(per_doc),
            accounts_per_doc_observed=f'{min(per_doc)}..{max(per_doc)}' if per_doc else '',
            account_observed='1..N', account_proposed='1..N', account_approved='Needs Decision',
            cc_observed=obs(d['cc_assigned'] / d['cc_eligible']) if d['cc_eligible'] else '0 (not eligible)',
            cc_proposed=cc_prop, cc_approved='Needs Decision',
            contract_observed=obs(float(sh['share_contract_mention'])) if sh.get('share_contract_mention') not in (None, '') else 'NOT_AVAILABLE',
            contract_proposed=contract_prop, contract_approved='Needs Decision',
            project_observed='NOT_AVAILABLE (no project field in Sanad)', project_proposed='Needs Decision', project_approved='Needs Decision',
            person_observed=obs(d['person'] / nl), person_proposed='1..1' if not revenue else '0..1', person_approved='Needs Decision',
            org_unit_observed=obs(d['unit'] / nl), org_unit_proposed='1..1', org_unit_approved='Needs Decision',
            source_activity_observed=obs(d['activity'] / d['taxi_lines']) if d['taxi_lines'] else 'NOT_AVAILABLE (Terminals has no activity code)'))

    # ---------------------------------------------------------------- domain evidence
    def classify(title):
        t = norm_fa(title).lower()
        for cls, terms in DOMAIN_TERMS:
            for term in terms:
                if norm_fa(term).lower() in t:
                    return cls, term
        return 'GENERAL', ''
    dom = []
    for x in coa['new_rows']:
        if x['title']:
            c, term = classify(x['title'])
            dom.append(OrderedDict(source='colleague workbook — new chart', ref=x['src'], code=x['code'], title=str(x['title']).strip(),
                                   classification=c, matched_term=term, usage='', note='حساب عمومی شهرداری' if c == 'GENERAL' else
                                   'مرتبط با حمل‌ونقل ولی نه اختصاصی تاکسیرانی/پایانه' if c == 'RELATED_INSUFFICIENT' else ''))
    for x in coa['old_rows']:
        c, term = classify(x['title'])          # what the colleague file itself shows
        cur = sorted({e['domain'] for e in ENTRIES if e['code'] == x['code'] and e['domain'] != 'GENERAL'})
        orgs = x['orgs']
        used = sorted({kk.split('-')[0] for kk, v in x['usage'].items() if v[0]})
        dom.append(OrderedDict(source='colleague workbook — old chart (moen)', ref=x['src'], code=x['code'], title=x['title_masked'], classification=c,
                               matched_term=term, usage=f"{x['lines_total']} ردیف؛ استفاده‌کننده: {','.join(used) or '—'}",
                               curated_domain=', '.join(cur) or 'GENERAL',
                               note=f"حساب در منشور {','.join(orgs)}" + ('؛ فقط پایانه‌ها' if orgs == ['Terminals'] else '؛ فقط تاکسیرانی' if orgs == ['Taxi'] else '')))
    for r, code, title in coa['old88']:
        c, term = classify(title)
        if c != 'GENERAL':
            dom.append(OrderedDict(source="colleague workbook — '88888888'!A:B", ref=f"'88888888'!A{r}", code=code, title=mask_title(title),
                                   classification=c, matched_term=term, usage='', note=''))
    # ledger sub-accounts (prior source) — domain evidence the colleague files do not contain
    led = []
    for org, hc, name, lvl in tc.execute('SELECT org,hcode,name,level FROM hesab WHERE level>=3'):
        c, term = classify(name)
        if c in ('DIRECT_TAXI', 'DIRECT_TERMINAL'):
            p = hc.strip().split('-')
            u = tc.execute('SELECT count(*),coalesce(sum(bed),0),coalesce(sum(bes),0) FROM sanad WHERE org=? AND kol=? AND moeen=? AND taf1=?' +
                           (' AND taf2=?' if lvl == 4 else ''), (org, p[0], p[1], p[2]) + ((p[3],) if lvl == 4 else ())).fetchone()
            led.append(OrderedDict(source='tt_v21 hesab (Taxi/Terminals account master; prior project)', ref=f'{org}:{mask_hcode(hc)}', code=mask_hcode(hc),
                                   title=mask_title(name), classification=c, matched_term=term, usage=f'{u[0]} ردیف؛ بدهکار {u[1]:,}؛ بستانکار {u[2]:,}',
                                   note=f'سطح {lvl}؛ {org}'))
    # terminal facilities (unit dimension) — only in the ledger
    fac = []
    for seg, nm in (('001', 'پایانه کاوه'), ('002', 'پایانه صفه'), ('003', 'پایانه زاینده‌رود'), ('004', 'پایانه جی'), ('007', 'پایانه صمدیه'),
                    ('005', 'دفتر مرکزی'), ('008', 'دفتر کار طالقانی'), ('006', 'انبار سازمان')):
        u = tc.execute("SELECT count(*),coalesce(sum(bed),0) FROM sanad WHERE org='Terminals' AND substr(vahed,1,3)=?", (seg,)).fetchone()
        fac.append(OrderedDict(unit_segment=seg, name=nm, lines=u[0], debit=u[1], source='Sanad.CodeVahed (prior project interpretation, v1 BASELINE)'))

    # ---------------------------------------------------------------- CD-01 recomputation
    cd01 = []
    for fy in ('1404', '1405'):
        for acc in ('113-1-1', '211-1-1', '113-1-4', '113-1-8', '212-6-5'):
            k, m, t = acc.split('-')
            for kind in ('OPENING', 'REGULAR', 'INTERNAL_SETTLEMENT', 'CLOSING_TEMPORARY', 'CLOSING_PERMANENT'):
                u = tc.execute("""SELECT count(*),coalesce(sum(s.bed),0),coalesce(sum(s.bes),0) FROM sanad s JOIN doc_class d ON d.doc_key=s.doc_key
                                  WHERE s.org='Terminals' AND s.fy=? AND s.kol=? AND s.moeen=? AND s.taf1=? AND d.kind=?""", (fy, k, m, t, kind)).fetchone()
                if u[0]:
                    cd01.append(OrderedDict(fy=fy, account=acc, doc_kind=kind, lines=u[0], debit=u[1], credit=u[2], net=u[1] - u[2]))
    cd01_bal = []
    for fy in ('1404', '1405'):
        for acc in ('113-1-1', '211-1-1'):
            k, m, t = acc.split('-')
            u = tc.execute("""SELECT coalesce(sum(bed),0)-coalesce(sum(bes),0) FROM sanad WHERE org='Terminals' AND fy=? AND kol=? AND moeen=? AND taf1=?""",
                           (fy, k, m, t)).fetchone()
            o = tc.execute("""SELECT coalesce(sum(s.bed),0)-coalesce(sum(s.bes),0) FROM sanad s JOIN doc_class d ON d.doc_key=s.doc_key WHERE s.org='Terminals'
                              AND s.fy=? AND s.kol=? AND s.moeen=? AND s.taf1=? AND d.kind='OPENING'""", (fy, k, m, t)).fetchone()
            cd01_bal.append(OrderedDict(fy=fy, account=acc, opening_net=o[0], all_documents_net=u[0]))
    agency_prior = list(csv.DictReader(open(V21_AGENCY, encoding='utf-8-sig')))

    # ---------------------------------------------------------------- prior decisions (v2.1)
    wcd = openpyxl.load_workbook(V21_DECISIONS, read_only=True, data_only=True)
    prior_dec = []
    for r in wcd['Decisions'].iter_rows(min_row=4, values_only=True):
        if r[0] and str(r[0]).startswith('CD'):
            prior_dec.append(OrderedDict(id=r[0], org=r[1], subject=r[2], status=r[11] if len(r) > 11 else ''))

    # ---------------------------------------------------------------- prior artefact inventory
    def art(name, rel, function, status, link_fields, inspected):
        p = os.path.join(ROOT, rel)
        exists = os.path.exists(p)
        return OrderedDict(candidate=name, path=rel if exists else '', found='YES' if exists else 'NOT FOUND', function=function, status=status,
                           link_fields=link_fields, inspected=inspected if exists else 'NOT AVAILABLE — limitation documented',
                           sha256=sha256(p) if exists and os.path.isfile(p) and os.path.getsize(p) < 300_000_000 else '')
    fe_dir = 'data/deliverables/financial_event_discovery/revisions'
    artefacts = [
        art('Financial event discovery v4', f'{fe_dir}/1404-r20260922-canonical-events-v4', 'event discovery revision', 'superseded candidate', 'FE ids', 'folder listed only'),
        art('Financial event discovery v5', f'{fe_dir}/1404-r20260922-canonical-events-v5', 'event discovery revision', 'superseded candidate', 'FE ids', 'folder listed only'),
        art('Financial event discovery v6', f'{fe_dir}/1404-r20260922-canonical-events-v6', 'event discovery revision', 'superseded candidate', 'FE ids', 'folder listed only'),
        art('Financial event discovery v7', f'{fe_dir}/1404-r20260922-canonical-events-v7', 'canonical events v7', 'candidate', 'FE ids', 'folder listed only'),
        art('Canonical events v7.1 workbook', 'data/deliverables/financial_event_discovery/revisions/1404-r20260922-canonical-events-v7.1/financial_event_discovery_1404-r20260922_canonical-events_v7.1.xlsx',
            'FE-01…FE-70 candidate dictionary (+ Doc_No_Mapping_Lineage sheet)', 'CANDIDATE — not approved; on-disk SHA differs from its manifest (prior finding)', 'event_id ↔ Account_Event_Crosswalk',
            'DIRECTLY INSPECTED (sheet فرهنگ_رویداد_مالی_کاندید)'),
        art('Cost center master (cc_master.json)', 'cc_master.json', '75 cost centers (v3 baseline + CC-109)', 'PROPOSED baseline, not ERP-active', 'cc_id ↔ Activity_CC_Crosswalk', 'DIRECTLY INSPECTED'),
        art('Service-to-CC mapping workbook', 'data/deliverables/service_activity_coverage/revisions/1405-r20260922-service-coverage/service_activity_coverage_1405-r20260922_QA-remapped_v2.xlsx',
            'service ↔ cost center coverage 1405', 'review-ready', 'cc_id', 'NOT OPENED in this audit (described in prior BASELINE_DISCOVERY)'),
        art('Doc_No_Mapping (standalone file)', 'data/deliverables/Doc_No_Mapping.xlsx', 'document lineage', '', 'Doc No', ''),
        art('Master_Data_ERP', 'data/deliverables/Master_Data_ERP.xlsx', 'ERP master data', '', 'account/cc codes', ''),
        art('committee_decisions_v2.1.xlsx', 'data/deliverables/taxi_terminals_1404_1405/revisions/v2.1/committee_decisions_v2.1.xlsx', 'CD-01…CD-LOWCC decision matrix',
            'READINESS ONLY — no decision taken', 'Decision IDs', 'DIRECTLY INSPECTED'),
        art('taxi_terminals_financial_mapping_1404_1405_v2.1.xlsx', 'data/deliverables/taxi_terminals_1404_1405/revisions/v2.1/taxi_terminals_financial_mapping_1404_1405_v2.1.xlsx',
            'Taxi/Terminals event, CC, KPI mapping v2.1', 'PROPOSED/OBSERVED — not approved', 'HCode, Source_Record_ID, FE/TT-NE, CC', 'DIRECTLY INSPECTED (Activity_Service_Map) + its SQLite tt_v21'),
        art('tt_v21.sqlite (scratch DB of v2.1)', 'data/work/taxi_terminals/tt_v21.sqlite', 'source ledger (111,600 lines), hesab (2,178), line_event, line_cc',
            'working copy of verified source extracts', 'HCode (kol-moeen-taf1) = old code of Taxi-Gate', 'DIRECTLY QUERIED (read-only)'),
        art('account_crosswalk.csv v2.1', 'data/deliverables/taxi_terminals_1404_1405/revisions/v2.1/machine_readable/account_crosswalk.csv', 'org chart → municipality chart (L4/L5/NONE)',
            'PROVISIONAL', 'Org_HCode = old code; Mun_Moin = new code', 'DIRECTLY INSPECTED'),
        art('CD-01 memo (redacted)', 'data/deliverables/taxi_terminals_1404_1405/revisions/v2.1/CD-01_terminal_agency_expert_memo_REDACTED.md', 'terminal agency question',
            'REQUIRES_AUTHORITY_DECISION', '113-1-1 / 211-1-1', 'DIRECTLY INSPECTED'),
        art('QA registry v2.1', 'data/deliverables/taxi_terminals_1404_1405/revisions/v2.1/QA_Test_Registry_v2.1.csv', 'prior QA tests', 'prior QA', 'test ids', 'listed only'),
        art('1404 municipality journal records', 'data/deliverables/journal_1404_enrichment_v1_independent_audit/journal_records.sqlite', '553,667-line 1404 journal (X:AC new coding)',
            'source extract (review-ready enrichment)', 'class/kol/moin (X,Z,AB)', 'DIRECTLY QUERIED (provenance test of Sheet1)'),
        art('ADR 0001–0027 / financial-event-control-layer.md', 'docs/adr', 'accepted governance decisions', 'ACCEPTED (documented)', 'governance layers', 'listed; decisions recalled from ADR titles'),
    ]
    dump('integration.json', OrderedDict(
        reconciliation=reconciliation, split_rows=split_rows, split_recon=split_recon, account_event=acc_ev_rows, activity_cc=act_rows,
        account_cc=acc_cc_rows, dims=dims, domain=dom, domain_ledger=led, facilities=fac, cd01=cd01, cd01_balances=cd01_bal,
        agency_prior=agency_prior[:40], prior_decisions=prior_dec, artefacts=artefacts,
        fe_dictionary=[OrderedDict(event_id=k, **v) for k, v in FE.items()],
        cc_master_size=len(CC), lineage_file=lineage_path))
    print('lines', n, 'buckets', {k: v[0] for k, v in by_bucket.items()}, 'acc_ev', len(acc_ev_rows), 'act_cc', len(act_rows), 'dims', len(dims),
          'domain', Counter(d['classification'] for d in dom), 'ledger_dom', Counter(d['classification'] for d in led), 'split_recon_bad',
          [r for r in split_recon if not r['equal']])


if __name__ == '__main__':
    main()
