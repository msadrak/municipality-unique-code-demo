"""QA of Taxi_Terminals_Final_Mapping_v2.xlsx, run against the delivered workbook and the ledger (tt_v21.sqlite).

Writes data/work/taxi_gate_audit/final_v2_QA_log.md, final_v2_QA_log.csv and final_v2_summary.json (internal; the deliverable
folder holds only the workbook).
"""
import csv
import json
import os
import re
from collections import Counter, defaultdict

import openpyxl

from common import WORK, TT_DB, FE_V71, CC_MASTER, V21_DIR, V21_EVENTS, ro_sqlite, mask_title
from m03_build_final_v2 import XLSX, MAIN, DETAIL, LINEAGE, MAIN_COLS, DETAIL_COLS

LOG_MD = os.path.join(WORK, 'final_v2_QA_log.md')
LOG_CSV = os.path.join(WORK, 'final_v2_QA_log.csv')
SUMMARY = os.path.join(WORK, 'final_v2_summary.json')
ORG = {'تاکسیرانی': 'Taxi', 'پایانه‌ها': 'Terminals'}
STATUSES = {'مشاهده‌شده', 'پیشنهادی', 'حل‌نشده', 'تأییدشده'}
RELS = {'اصلی', 'ثانوی', 'نامشخص', 'غیررویدادی'}
CONFS = {'بالا', 'متوسط', 'پایین'}
NA_CC = 'بدون مرکز هزینهٔ عملیاتی / موضوعیت ندارد'
AMOUNT_KIND = 'دارای مبلغ (هر سطر دفتر یک بار)'


def main():
    results = []

    def check(tid, desc, ok, detail='', review=False):
        results.append((tid, desc, 'PASS' if ok else ('REVIEW' if review else 'FAIL'), str(detail)[:900]))

    wb = openpyxl.load_workbook(XLSX)          # full mode: views, filters and comments are needed
    check('Q01', 'exactly two sheets with the required names and order', wb.sheetnames == [MAIN, DETAIL], wb.sheetnames)
    wm, wd = wb[MAIN], wb[DETAIL]
    mh = [c.value for c in wm[1]]
    dh = [c.value for c in wd[1]]
    check('Q02', 'main sheet headers: row id + the 15 requested columns in order', mh == [h for _, h, _ in MAIN_COLS], mh)
    check('Q03', 'details sheet headers as specified and keyed by the same row id', dh == [h for _, h, _ in DETAIL_COLS] and dh[0] == mh[0], dh[:3])
    M = [dict(zip(mh, r)) for r in wm.iter_rows(min_row=2, values_only=True)]
    D = [dict(zip(dh, r)) for r in wd.iter_rows(min_row=2, values_only=True)]
    ids_m = [r['شناسه ردیف'] for r in M]
    ids_d = [r['شناسه ردیف'] for r in D]
    check('Q04', 'row ids unique, no blank rows, main ↔ details one-to-one in the same order',
          len(set(ids_m)) == len(ids_m) and ids_m == ids_d and all(ids_m), f'main {len(ids_m)}, details {len(ids_d)}')
    det = {r['شناسه ردیف']: r for r in D}

    # ---------------------------------------------------------------- presentation
    views = all(ws.sheet_view.rightToLeft for ws in (wm, wd))
    frozen = all(ws.freeze_panes for ws in (wm, wd))
    filt = all(ws.auto_filter.ref for ws in (wm, wd))
    comments = sum(1 for c in wm[1] if c.comment)
    check('Q05', 'right-to-left, frozen header, active AutoFilter on both sheets; header help comments on the main sheet',
          views and frozen and filt and comments >= 4, f'rtl={views} freeze={wm.freeze_panes},{wd.freeze_panes} filter={wm.auto_filter.ref} comments={comments}')
    code_cells = [wm.cell(row=i, column=1).data_type for i in range(2, wm.max_row + 1)]
    code_cells += [wd.cell(row=i, column=c).data_type for i in range(2, wd.max_row + 1) for c in (4, 5, 6)
                   if wd.cell(row=i, column=c).value not in (None, '')]
    check('Q06', 'ids and account codes stored as text', set(code_cells) == {'s'}, Counter(code_cells))

    # ---------------------------------------------------------------- vocabularies and identifiers
    check('Q07', 'status / relationship / confidence values only from the allowed vocabularies',
          {r['وضعیت رابطه'] for r in M} <= STATUSES and {r['نوع رابطه حساب و رویداد'] for r in M} <= RELS and {r['درجه اطمینان'] for r in M} <= CONFS,
          (Counter(r['وضعیت رابطه'] for r in M), Counter(r['نوع رابطه حساب و رویداد'] for r in M)))
    approved = sum(r['وضعیت رابطه'] == 'تأییدشده' for r in M)
    check('Q08', 'no relationship is APPROVED (no documented approval exists)', approved == 0, approved)
    wbe = openpyxl.load_workbook(FE_V71, read_only=True, data_only=True)
    EVN = {r[0]: str(r[5] or '') for r in wbe['فرهنگ_رویداد_مالی_کاندید'].iter_rows(min_row=2, values_only=True) if r[0]}
    for r in csv.DictReader(open(V21_EVENTS, encoding='utf-8-sig')):
        EVN.setdefault(r['event_id'], r['nature'])
    CC = json.load(open(CC_MASTER, encoding='utf8'))
    NCC = json.load(open(os.path.join(V21_DIR, 'machine_readable', 'new_cost_center_candidates.json'), encoding='utf8'))
    bad_ev = {r['کد رویداد مالی'] for r in M if r['کد رویداد مالی'] != '—' and r['کد رویداد مالی'] not in EVN}
    bad_cc = {r['کد مرکز هزینه'] for r in M if r['کد مرکز هزینه'] != '—' and r['کد مرکز هزینه'] not in CC and r['کد مرکز هزینه'] not in NCC}
    check('Q09', 'every event / cost-center id exists in FE v7.1 + v2.1 events / cc_master + TTCC candidates (no invented ids)',
          not bad_ev and not bad_cc, (bad_ev, bad_cc))
    proposed_ids = [r['شناسه ردیف'] for r in M if (r['کد رویداد مالی'].startswith(('TT-NE', 'PEO', 'EXC')) or r['کد مرکز هزینه'].startswith('TTCC'))
                    and r['وضعیت رابطه'] == 'مشاهده‌شده']
    check('Q10', 'candidate events (TT-NE/PEO/EXC) and candidate centers (TTCC) are never shown as observed', not proposed_ids, proposed_ids[:10])
    key = lambda r: tuple(r[h] for h in mh[1:13])
    dup = [k for k, n in Counter(key(r) for r in M).items() if n > 1]
    check('Q11', 'no duplicate relationship rows (area, accounts, event, activity, center, type, status)', not dup, dup[:3])

    # ---------------------------------------------------------------- coverage
    tc = ro_sqlite(TT_DB)
    universe = defaultdict(set)
    level = {}
    for org, hc, lv in tc.execute('SELECT DISTINCT org, hcode, level FROM hesab'):
        level[(org, hc.strip())] = lv
        if lv in (2, 3):
            universe[org].add(hc.strip())
    ledger_heads = defaultdict(set)
    for org, k, m, t in tc.execute('SELECT DISTINCT org, kol, moeen, taf1 FROM sanad'):
        ledger_heads[org].add(f'{int(k)}-{int(m)}-{int(t)}-0')
    present = defaultdict(set)
    for r in D:
        present[ORG[r['حوزه']]].add(r['کد حساب تفصیلی'] or r['کد حساب معین'])
    missing = []
    for org in ('Taxi', 'Terminals'):
        for hc in universe[org] | ledger_heads[org]:
            if hc in present[org]:
                continue
            k, m, t = hc.split('-')[:3]
            has_children = any(x.startswith(f'{k}-{m}-') and x != hc and level.get((org, x)) == 3 for x in universe[org])
            if t == '0' and has_children and hc not in ledger_heads[org]:
                continue
            missing.append((org, hc))
    check('Q12', 'every Taxi and Terminals account (Hesab L2/L3 + every ledger head) appears in the table', not missing,
          f"missing {missing[:10]}; covered heads {sum(len(v) for v in present.values())}")

    # ---------------------------------------------------------------- amounts
    amt = [r for r in D if r['نوع ردیف'] == AMOUNT_KIND]
    led = tc.execute('SELECT COUNT(*), SUM(bed), SUM(bes) FROM sanad').fetchone()
    s_lines = sum(r['تعداد سطرهای دفتر'] for r in amt)
    s_dr = sum(r['جمع بدهکار'] for r in amt)
    s_cr = sum(r['جمع بستانکار'] for r in amt)
    check('Q13', 'Σ ledger lines / debit / credit of amount rows equals the ledger', (s_lines, s_dr, s_cr) == tuple(led),
          f'table {s_lines:,} / {s_dr:,} / {s_cr:,} — ledger {led[0]:,} / {led[1]:,} / {led[2]:,}')
    per_head = defaultdict(lambda: [0, 0, 0])
    for r in amt:
        h = per_head[(ORG[r['حوزه']], r['کد حساب تفصیلی'] or r['کد حساب معین'])]
        h[0] += r['تعداد سطرهای دفتر']
        h[1] += r['جمع بدهکار']
        h[2] += r['جمع بستانکار']
    led_head = {(o, f'{int(k)}-{int(m)}-{int(t)}-0'): [n, dr, cr] for o, k, m, t, n, dr, cr in
                tc.execute('SELECT org, kol, moeen, taf1, COUNT(*), SUM(bed), SUM(bes) FROM sanad GROUP BY 1,2,3,4')}
    diff = [h for h in set(per_head) | set(led_head) if per_head.get(h) != led_head.get(h)]
    check('Q14', 'per organisation × account: lines, debit and credit equal the ledger', not diff, diff[:5])
    free = [r for r in D if r['نوع ردیف'] != AMOUNT_KIND]
    free_amt = [r['شناسه ردیف'] for r in free if r['جمع بدهکار'] not in (None, '') or r['جمع بستانکار'] not in (None, '')]
    amt_ids = {r['شناسه ردیف'] for r in amt}
    refs_bad = []
    for r in free:
        if 'ردیف‌های:' in (r['وضعیت شمارش مبلغ'] or ''):
            refs = re.findall(r'T[XM]-\d{4}', r['وضعیت شمارش مبلغ'])
            if not refs or not set(refs) <= amt_ids:
                refs_bad.append(r['شناسه ردیف'])
    check('Q15', 'amount-free rows carry no amounts and point to existing amount rows (no double counting)', not free_amt and not refs_bad,
          f'with amounts {free_amt[:5]}; bad refs {refs_bad[:5]}; amount-free rows {len(free)}')

    # ---------------------------------------------------------------- traceability (sid → row id)
    lin = list(csv.DictReader(open(LINEAGE, encoding='utf8')))
    sids = [x['sid'] for x in lin]
    agg = defaultdict(lambda: [0, 0, 0])
    sanad = {sid: (bed or 0, bes or 0) for sid, bed, bes in tc.execute('SELECT sid, bed, bes FROM sanad')}
    for x in lin:
        a = agg[x['row_id']]
        bed, bes = sanad[x['sid']]
        a[0] += 1
        a[1] += bed
        a[2] += bes
    trace_bad = [r['شناسه ردیف'] for r in amt if agg.get(r['شناسه ردیف']) != [r['تعداد سطرهای دفتر'], r['جمع بدهکار'], r['جمع بستانکار']]]
    check('Q16', 'every ledger line maps to exactly one amount row; lineage reproduces every amount row',
          len(sids) == len(set(sids)) == led[0] and set(sids) <= set(sanad) and set(agg) <= amt_ids and not trace_bad,
          f'lineage {len(sids):,} unique {len(set(sids)):,}; mismatched rows {trace_bad[:5]}')
    docs_bad = [r['شناسه ردیف'] for r in M if r['تعداد اسناد مرتبط'] != det[r['شناسه ردیف']]['تعداد اسناد']]
    check('Q17', 'document counts identical in main and details', not docs_bad, docs_bad[:5])

    # ---------------------------------------------------------------- semantics
    def evclass(e):
        if e.startswith('PEO'):
            return 'PEO'
        n = EVN.get(e, '')
        for pre, c in (('هزینه', 'EXPENSE'), ('سرمایه', 'CAPITAL'), ('درآمد', 'REVENUE'), ('ترازنامه', 'BALANCE'), ('انتظامی', 'MEMO'), ('واسط', 'TRANSFER')):
            if n.startswith(pre):
                return c
        return 'OTHER'

    viol, cash_primary = [], []
    for r in M:
        if r['نوع رابطه حساب و رویداد'] != 'اصلی':
            continue
        d = det[r['شناسه ردیف']]
        k, m = int(d['کد حساب کل'].split('-')[0]), int(d['کد حساب معین'].split('-')[1])
        c, e = evclass(r['کد رویداد مالی']), r['کد رویداد مالی']
        cls = k // 100
        ok = {'EXPENSE': cls == 6 or (k in (115, 116, 121, 123, 124, 125) and 'دارایی' in EVN.get(e, '')),
              'CAPITAL': cls == 6 or k in (115, 116, 121, 123, 124, 125),
              'REVENUE': cls == 4 or k == 812,
              'MEMO': k in (918, 919),
              'BALANCE': cls in (1, 2, 3) or (cls == 6 and e == 'FE-48'),
              'TRANSFER': cls in (1, 2, 3) or k == 812,
              'PEO': cls == 6 or k in (122, 214, 217, 223, 224) or (k, m) == (212, 2)}.get(c, False)
        if not ok:
            viol.append((r['شناسه ردیف'], d['کد حساب تفصیلی'] or d['کد حساب معین'], e))
        if k == 111 and c != 'TRANSFER':
            cash_primary.append(r['شناسه ردیف'])
    check('Q18', 'primary rows: account nature consistent with the event nature', not viol, viol[:8])
    ar_ap_primary = [r['شناسه ردیف'] for r in M if r['نوع رابطه حساب و رویداد'] == 'اصلی'
                     and int(det[r['شناسه ردیف']]['کد حساب کل'].split('-')[0]) in (113, 114, 211, 212, 213)
                     and evclass(r['کد رویداد مالی']) in ('EXPENSE', 'REVENUE')]
    check('Q19', 'counterpart accounts never primary: bank/cash only for cash transfers; receivables/payables never for expense or revenue events',
          not cash_primary and not ar_ap_primary, (cash_primary[:5], ar_ap_primary[:5]))
    rev_cc = [r['شناسه ردیف'] for r in M if det[r['شناسه ردیف']]['کد حساب کل'].startswith('4') and r['کد مرکز هزینه'] != '—']
    na_ok = all(r['عنوان مرکز هزینه'] == NA_CC for r in M if det[r['شناسه ردیف']]['کد حساب کل'].startswith('4') and r['کد مرکز هزینه'] == '—'
                and r['کد رویداد مالی'] != '—')
    check('Q20', 'revenue accounts carry no expense cost center and are labelled «بدون مرکز هزینهٔ عملیاتی / موضوعیت ندارد»', not rev_cc and na_ok, rev_cc[:5])
    weak = []
    for r in M:
        if r['وضعیت رابطه'] == 'مشاهده‌شده' and r['کد مرکز هزینه'] in CC:
            d = det[r['شناسه ردیف']]
            basis, chk = d['مبنای تعیین مرکز هزینه'] or '', d['بازبینی مرکز با فعالیت'] or ''
            if any(x in basis for x in ('پیش‌فرض', 'واژهٔ کلیدی', 'محل (پایانه)', 'تکمیل v2')) or 'ناسازگار' in chk or not basis:
                weak.append(r['شناسه ردیف'])
    check('Q21', 'observed rows with a cost center are grounded in the activity (activity code or narration), never in event/account defaults, keywords, location or conflicts',
          not weak, weak[:8])
    cd01 = [r for r in D if r['حوزه'] == 'پایانه‌ها' and ((r['کد حساب تفصیلی'] or '') in ('113-1-1-0', '211-1-1-0') or r['کد حساب معین'] == '413-8-0-0')]
    note_of = {r['شناسه ردیف']: r['توضیحات و موارد نیازمند بررسی'] or '' for r in M}
    cd01_bad = [r['شناسه ردیف'] for r in cd01 if r['وضعیت رابطه'] != 'حل‌نشده' or 'CD-01' not in note_of[r['شناسه ردیف']]]
    check('Q22', 'CD-01 accounts of Terminals (113-1-1, 211-1-1, 413-8) are all UNRESOLVED and say so in the notes', cd01 and not cd01_bad,
          f'{len(cd01)} rows; bad {cd01_bad[:5]}')

    # ---------------------------------------------------------------- privacy
    raw_names = set()
    for (title,) in tc.execute('SELECT DISTINCT name FROM hesab'):
        mt = mask_title(title)
        if mt != title.strip() and '‹نام شخص حذف شد›' in mt:
            raw_names.add(title.strip().split('-')[0].split('=')[0].strip())
    cells = [str(v) for r in M + D for v in r.values() if isinstance(v, str)]
    leaked = sorted({n for n in raw_names if len(n) > 4 and any(n in c for c in cells)})
    long_digits = sorted({c for r in M for c in (r['حساب کل'], r['حساب معین'], r['حساب تفصیلی'])
                          if re.search(r'\d{6,}', re.sub(r'^[\d\-]+ — ', '', c or ''))})
    narr = {s.strip() for (s,) in tc.execute('SELECT DISTINCT sharh FROM sanad WHERE length(sharh) > 25')}
    narr_hits = [c for c in cells if c in narr]
    check('Q23', 'privacy: no personal names, no unmasked bank-account digits in titles, no ledger narration text', not leaked and not long_digits and not narr_hits,
          f'names {leaked[:3]}; digits {long_digits[:3]}; narration cells {len(narr_hits)}')

    # ---------------------------------------------------------------- figures (baselines from the previous audit are reference, not targets)
    obs_cc = [r for r in M if r['وضعیت رابطه'] == 'مشاهده‌شده' and r['کد مرکز هزینه'] in CC]
    obs_cc_act = [r for r in obs_cc if r['فعالیت مرتبط'].startswith('ACT-')]
    exp_amt = [r for r in M if det[r['شناسه ردیف']]['کد حساب کل'].startswith('6') and det[r['شناسه ردیف']]['نوع ردیف'] == AMOUNT_KIND]
    exp_nocc = Counter(r['عنوان مرکز هزینه'] for r in exp_amt if r['کد مرکز هزینه'] == '—')
    heads_nocc = {(r['حوزه'], det[r['شناسه ردیف']]['کد حساب تفصیلی'] or det[r['شناسه ردیف']]['کد حساب معین']) for r in exp_amt
                  if r['کد مرکز هزینه'] == '—' and r['عنوان مرکز هزینه'].startswith('نیازمند بررسی')}
    heads_with_cc = {(r['حوزه'], det[r['شناسه ردیف']]['کد حساب تفصیلی'] or det[r['شناسه ردیف']]['کد حساب معین']) for r in exp_amt if r['کد مرکز هزینه'] != '—'}
    resolved = Counter()
    for r in exp_amt:
        for part in (det[r['شناسه ردیف']]['تکمیل مرکز برای سطر هزینهٔ بدون مرکز'] or '').split('؛ '):
            if part:
                name, n = part.rsplit(' ', 1)
                resolved[name] += int(n.replace(',', ''))
    check('Q24', 'baseline comparison: observed relationships with a cost center (audit baseline 165) and activity-linked ones (137)', True,
          f'now {len(obs_cc)} observed with center, {len(obs_cc_act)} of them with a specific activity code')
    check('Q25', 'expense rows without a center (audit baseline 162 rows) — completed / suggestions / needs review / not relevant', True,
          f'expense amount rows without center by label {dict(exp_nocc)}; lines completed or classified {dict(resolved)}; '
          f'expense accounts with lines still lacking a center {len(heads_nocc)} (of which with no center at all {len(heads_nocc - heads_with_cc)})')

    rel = Counter(r['نوع رابطه حساب و رویداد'] for r in M)
    kinds = Counter(r['نوع ردیف'] for r in D)
    st = Counter(r['وضعیت رابطه'] for r in M)
    evs = {r['کد رویداد مالی'] for r in M if r['کد رویداد مالی'] != '—'}
    ccs = {r['کد مرکز هزینه'] for r in M if r['کد مرکز هزینه'] != '—'}
    by_org = Counter(o for o, _ in {(r['حوزه'], det[r['شناسه ردیف']]['کد حساب تفصیلی'] or det[r['شناسه ردیف']]['کد حساب معین']) for r in M})
    gl = Counter(o for o, _ in {(r['حوزه'], det[r['شناسه ردیف']]['کد حساب کل']) for r in M})
    khoda = [r for r in M if 'خداپرست' in (r['توضیحات و موارد نیازمند بررسی'] or '')]
    summary = dict(
        rows=len(M), row_kinds=dict(kinds), heads=sum(by_org.values()), heads_by_org=dict(by_org), gl_accounts_by_org=dict(gl),
        events_fe=len([e for e in evs if e.startswith('FE')]), events_ttne=len([e for e in evs if e.startswith('TT-NE')]),
        events_other=sorted(e for e in evs if e.startswith(('PEO', 'EXC'))), cost_centers=len([c for c in ccs if c in CC]),
        cost_center_candidates=sorted(c for c in ccs if c in NCC), status=dict(st), relationship=dict(rel),
        amount_rows_relationship=dict(Counter(r['نوع رابطه حساب و رویداد'] for r in M if det[r['شناسه ردیف']]['نوع ردیف'] == AMOUNT_KIND)),
        observed_with_cc=len(obs_cc), observed_with_cc_activity=len(obs_cc_act), expense_amount_rows_without_cc=dict(exp_nocc),
        expense_heads_needing_cc=len(heads_nocc), expense_heads_without_any_cc=len(heads_nocc - heads_with_cc), c_resolution_lines=dict(resolved),
        khodaparast_rows=len(khoda), khodaparast_cd01_rows=sum('CD-01' in r['توضیحات و موارد نیازمند بررسی'] for r in khoda),
        khodaparast_peo_rows=sum('CD-PEO' in r['توضیحات و موارد نیازمند بررسی'] for r in khoda),
        ledger_lines=led[0], ledger_debit=led[1], ledger_credit=led[2])
    json.dump(summary, open(SUMMARY, 'w', encoding='utf8'), ensure_ascii=False, indent=1)
    with open(LOG_CSV, 'w', newline='', encoding='utf-8-sig') as fh:
        w = csv.writer(fh)
        w.writerow(['test_id', 'description', 'status', 'detail'])
        w.writerows(results)
    with open(LOG_MD, 'w', encoding='utf8') as fh:
        fh.write('# QA — Taxi_Terminals_Final_Mapping_v2.xlsx\n\n| Test | Description | Status | Detail |\n|---|---|---|---|\n')
        for tid, desc, s, d in results:
            fh.write(f"| {tid} | {desc} | {s} | {d.replace('|', '/')} |\n")
    print(Counter(r[2] for r in results), [r for r in results if r[2] != 'PASS'])
    print(json.dumps(summary, ensure_ascii=False))


if __name__ == '__main__':
    main()
