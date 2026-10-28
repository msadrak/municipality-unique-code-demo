"""Step 3 — reverse-engineer the colleague's V2 matching, audit every suggested pair, and build the controlled
(constraint-first) crosswalk proposal with explicit cardinality.

Output: data/work/taxi_gate_audit/matching.json
"""
import csv
import re
from collections import Counter, OrderedDict, defaultdict

import openpyxl

from common import (F_V2, V21_CROSSWALK, S_MATCH, S_ALLM, S_SUMMARY, dump, load, norm_fa, norm_key, colleague_score,
                    indel_similarity, structural_compat, old_nature)
from expert_crosswalk import ENTRIES, REL_EN, MEANING_COLLISIONS

NOT_FOUND = 'پیدا نشد'

# Stage A: kol families (old kol → admissible new kols).  Derived from account nature, not from titles.
KOL_FAMILY = {
    111: [110], 112: [125, 120], 113: [115, 180], 114: [120, 115], 115: [130], 116: [140, 120], 120: [150], 121: [150, 160],
    122: [], 123: [160], 124: [170], 125: [180, 140, 160, 150], 211: [210], 212: [220, 225, 210], 213: [230, 225],
    214: [220], 215: [220, 210], 216: [235, 220], 221: [250, 210], 222: [235, 220], 223: [], 224: [], 311: [310], 313: [310],
    314: [310], 315: [310], 316: [310], 411: [412, 413, 417], 412: [], 413: [412, 413, 411, 415, 414], 414: [], 611: [610, 614],
    612: [611, 614], 613: [613, 614, 630, 620, 650, 651, 652], 812: [413, 415, 416, 417, 651, 310], 917: [], 918: [810, 820],
    919: [820], 920: []}
# Stage C synonym normalisation (domain vocabulary)
SYN = [('اثاثیه', 'اثاثه'), ('وسایط', 'وسایل'), ('دارائی', 'دارایی'), ('بستانکاران', 'بدهی به'), ('بستانکارن', 'بدهی به'),
       ('بدهکاران', 'مطالبات از'), ('بدهکارن', 'مطالبات از'), ('کارکنان', 'کارکنان'), ('حقوق و دستمزد', 'حقوق ثابت'),
       ('ماموریت', 'ماموریت'), ('اسناد دریافتنی جاری', 'اسناد دریافتنی عملیاتی سال جاری'), ('سپرده دیگران نزد ما', 'سپرده های پرداختنی'),
       ('سپرده ما نزد دیگران', 'سپرده ها و ودایع دریافتنی'), ('دارایی های ثابت در جریان ساخت', 'دارایی در جریان تکمیل'),
       ('اجاره بها', 'اجاره'), ('کرایه', 'کرایه'), ('ساختمان', 'ساختمان و مستحدثات')]
STOP = {'و', 'از', 'به', 'در', 'با', 'ها', 'های', 'حساب', 'حسابهای', 'سایر'}


def syn_norm(s):
    s = norm_fa(s)
    for a, b in SYN:
        s = s.replace(a, b)
    return s


def tokens(s):
    return {t for t in syn_norm(s).split() if t not in STOP and len(t) > 1}


def nums(s):
    return set(re.findall(r'\d+', norm_fa(s)))


def controlled_score(a, b):
    ta, tb = tokens(a), tokens(b)
    jac = 100.0 * len(ta & tb) / len(ta | tb) if ta | tb else 0.0
    ind = indel_similarity(syn_norm(a), syn_norm(b))
    score = 0.5 * jac + 0.5 * ind
    na, nb = nums(a), nums(b)
    entity_conflict = bool(na and nb and na != nb)
    if entity_conflict:
        score -= 40
    return round(score, 1), entity_conflict


def main():
    coa = load('coa.json')
    new_rows = coa['new_rows']
    parents = defaultdict(set)
    new_title = {}
    new_nat = {}
    for x in new_rows:
        if x['code'] is not None:
            parents[x['code']].add((x['cls'], x['kol']))
            new_title.setdefault(x['code'], str(x['title']).strip())
            new_nat.setdefault(x['code'], set()).add(x['nature'])
    kol_title = {x['kol']: str(x['kol_title']).strip() for x in new_rows}
    kol_nature = {x['kol']: x['nature'] for x in new_rows}
    old = {x['code']: x for x in coa['old_rows']}

    wb = openpyxl.load_workbook(F_V2, data_only=True)
    ws = wb[S_MATCH]
    match = []
    for r in range(2, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 16)]
        match.append(OrderedDict(src=f"'{S_MATCH}'!A{r}:O{r}", row=r, moen_row=v[0], old_code=str(v[1]).strip(), old_title=str(v[2]).strip(),
                                 best_code=v[3], best_title=str(v[4]).strip(), best_score=v[5], status=v[6], grp_code=v[7], grp_title=v[8],
                                 kol_code=v[9], kol_title=v[10], alt2_title=str(v[11]).strip(), alt2_score=v[12], alt3_title=str(v[13]).strip(),
                                 alt3_score=v[14]))
    wa = wb[S_ALLM]
    allm = []
    for r in range(2, wa.max_row + 1):
        v = [wa.cell(r, c).value for c in range(1, 10)]
        allm.append(OrderedDict(src=f"'{S_ALLM}'!A{r}:I{r}", row=r, old_code=str(v[0]).strip(), old_title=str(v[1]).strip(), new_code=v[2],
                                new_title=str(v[3]).strip(), score=v[4], grp_code=v[5], grp_title=v[6], kol_code=v[7], kol_title=v[8]))
    wsum = wb[S_SUMMARY]
    summary = [(wsum.cell(r, 1).value, wsum.cell(r, 2).value) for r in range(2, 10)]

    # ---------------------------------------------------------------- reproduction
    # candidate universe = the 234 code-1 rows of moen, in moen order (ties are resolved by that order);
    # code-less chart rows appear there as 0 / 0 and score against an empty title
    wm0 = wb['moen']
    cand_universe = []
    for r in range(2, 393):
        if wm0.cell(r, 3).value == 1:
            a, b = wm0.cell(r, 1).value, wm0.cell(r, 2).value
            cand_universe.append(dict(code=a if a not in (None, '') else 0, title='' if b in (None, 0) else str(b)))
    rep_score_ok = sum(1 for p in allm if abs(colleague_score(p['old_title'], p['new_title']) - p['score']) < 0.051)
    rank_ok, rank_tie_only, rank_bad = 0, 0, []
    for m in match:
        sc = []
        for x in cand_universe:
            t = x['title']
            sc.append((colleague_score(m['old_title'], t), x['code'], t.strip()))
        ranked = sorted(sc, key=lambda z: -z[0])          # stable: ties keep chart order
        top = ranked[:3]
        m['_rep_top3'] = top
        best_ok = top[0][1] == m['best_code'] and abs(top[0][0] - m['best_score']) < 0.051
        alts_ok = (norm_key(top[1][2]) == norm_key(m['alt2_title']) and abs(top[1][0] - m['alt2_score']) < 0.051 and
                   norm_key(top[2][2]) == norm_key(m['alt3_title']) and abs(top[2][0] - m['alt3_score']) < 0.051)
        if best_ok and alts_ok:
            rank_ok += 1
        elif best_ok and {round(top[1][0], 1), round(top[2][0], 1)} == {round(m['alt2_score'], 1), round(m['alt3_score'], 1)} and top[1][0] == top[2][0]:
            rank_tie_only += 1
        else:
            rank_bad.append(m['old_code'])
    bands = Counter()
    for m in match:
        s = m['best_score']
        bands['exact' if m['status'] == 'تطبیق دقیق نام' else '>=90' if s >= 90 else '80-90' if s >= 80 else '70-80' if s >= 70 else '<70'] += 1
    exact_after_norm = sum(1 for m in match if norm_fa(m['old_title']) == norm_fa(m['best_title']))
    rep_summary = OrderedDict([('تعداد حساب‌های جدید (کد ۱)', len(new_rows)), ('تعداد حساب‌های قدیمی (کد ۲)', len(match)),
                               ('نام‌های دقیقاً یکسان پس از یکسان‌سازی حروف و فاصله', exact_after_norm), ('>=90', bands['>=90']),
                               ('80-90', bands['80-90']), ('70-80', bands['70-80']), ('<70', bands['<70'])])
    status_by_band = Counter((m['status'], 'exact' if m['best_score'] == 100 else '>=90' if m['best_score'] >= 90 else
                              '80-90' if m['best_score'] >= 80 else '70-80' if m['best_score'] >= 70 else '<70') for m in match)

    # parent columns: hypothesis "taken from moen!D:G of the row where the new code sits in column A"
    moen_dg = {}
    import openpyxl as _o
    wbv1 = _o.load_workbook(F_V2, data_only=True)
    wm = wbv1['moen']
    for r in range(4, 393):
        a = wm.cell(r, 1).value
        moen_dg.setdefault(a, []).append((wm.cell(r, 4).value, wm.cell(r, 6).value, r))
    hyp_ok, parent_true, parent_wrong, parent_nf = 0, 0, 0, 0
    for p in allm:
        cands = [(d, f) for d, f, r in moen_dg.get(p['new_code'], [])]
        hyp_ok += (p['grp_code'], p['kol_code']) in cands
        if p['grp_code'] == NOT_FOUND:
            parent_nf += 1
        elif (p['grp_code'], p['kol_code']) in parents.get(p['new_code'], set()):
            parent_true += 1
        else:
            parent_wrong += 1
        p['parent_shown'] = f"{p['grp_code']}/{p['kol_code']}"
        p['parent_true'] = '; '.join(f'{c}/{k}' for c, k in sorted(parents.get(p['new_code'], set())))
        p['parent_status'] = 'NOT_FOUND_SHOWN' if p['grp_code'] == NOT_FOUND else ('CORRECT' if (p['grp_code'], p['kol_code']) in parents.get(p['new_code'], set()) else 'WRONG')

    # ---------------------------------------------------------------- expert proposal index
    entries = defaultdict(list)
    for e in ENTRIES:
        entries[e['code']].append(e)

    def pick_entry(code):
        lst = entries[code]
        src = old[code]['title_source_org']
        for e in lst:
            if e['org'] == 'BOTH':
                return e
        for e in lst:
            if e['org'] in src:
                return e
        return lst[0]

    def struct_for(code_old, new_code):
        o = old[code_old]
        if new_code in (None, 0) or new_code not in parents:
            return 'NOT_TESTABLE', ''
        nats = new_nat[new_code]
        res = sorted({structural_compat(o['nature'], n) for n in nats})
        best = 'COMPATIBLE' if 'COMPATIBLE' in res else 'CONDITIONAL' if 'CONDITIONAL' in res else res[0]
        return best, '/'.join(sorted(nats))

    # ---------------------------------------------------------------- audit of the 157 best matches
    audit = []
    for m in match:
        e = pick_entry(m['old_code'])
        o = old[m['old_code']]
        st, nn = struct_for(m['old_code'], m['best_code'])
        targets = [t[0] for t in e['targets'] if t[1] != 'RL']
        agrees = m['best_code'] in targets
        other = [x for x in entries[m['old_code']] if x is not e]
        audit.append(OrderedDict(
            audit_id=f'OMA-{len(audit) + 1:03d}', src=m['src'], moen_row=m['moen_row'], old_code=m['old_code'], old_title=o['title_masked'],
            old_nature=o['nature'], orgs=','.join(o['orgs']), title_source_org=','.join(o['title_source_org']),
            meaning_collision=m['old_code'] in MEANING_COLLISIONS,
            lines=o['lines_total'], debit=o['debit_total'], credit=o['credit_total'],
            best_code=m['best_code'], best_title=m['best_title'], best_score=m['best_score'], colleague_status=m['status'],
            parent_shown=f"{m['grp_code']}/{m['kol_code']}", parent_true='; '.join(f'{c}/{k}' for c, k in sorted(parents.get(m['best_code'], set()))),
            new_nature=nn, structural=st, verdict=e['verdict'], agrees_with_proposal=agrees,
            proposal_status=e['status'], proposal_targets='; '.join(f"{t[0]} ({REL_EN[t[1]]})" for t in e['targets']) or '—',
            reason=e['reason'], other_org_rows='; '.join(f"{x['org']}: {x['status']} / {x['verdict']}" for x in other),
            decision=e['decision'] or ''))
    # ---------------------------------------------------------------- audit of every candidate (471)
    cand_audit = []
    rank_of = {}
    for p in allm:
        rank_of.setdefault(p['old_code'], 0)
        rank_of[p['old_code']] += 1
        e = pick_entry(p['old_code'])
        st, nn = struct_for(p['old_code'], p['new_code'])
        tset = [t[0] for t in e['targets'] if t[1] != 'RL']
        rl = [t[0] for t in e['targets'] if t[1] == 'RL']
        if p['new_code'] in tset:
            v = 'AGREES_WITH_PROPOSAL'
        elif p['new_code'] in rl:
            v = 'RELATED_NOT_EQUIVALENT'
        elif st == 'INCOMPATIBLE':
            v = 'INVALID_STRUCTURAL'
        elif st == 'NOT_TESTABLE':
            v = 'NOT_TESTABLE'
        else:
            v = 'COMPATIBLE_CLASS_NOT_EQUIVALENT'
        nc = controlled_score(p['old_title'], p['new_title'])
        cand_audit.append(OrderedDict(src=p['src'], old_code=p['old_code'], rank=rank_of[p['old_code']], new_code=p['new_code'], new_title=p['new_title'],
                                      score=p['score'], reproduced_score=colleague_score(p['old_title'], p['new_title']),
                                      parent_shown=p['parent_shown'], parent_true=p['parent_true'], parent_status=p['parent_status'],
                                      structural=st, new_nature=nn, verdict=v, entity_number_conflict=nc[1]))

    # ---------------------------------------------------------------- controlled matcher (Stage A→B→C) for every old account
    v21 = list(csv.DictReader(open(V21_CROSSWALK, encoding='utf-8-sig')))
    prior = defaultdict(set)
    prior_lvl = defaultdict(set)
    for r in v21:
        prior[r['Org_HCode'].strip()].add(r['Mun_Moin'] or '')
        prior_lvl[r['Org_HCode'].strip()].add(r['Match_Level'])
    alternatives = []
    controlled_top = {}
    for code, o in old.items():
        fam = KOL_FAMILY.get(o['kol'], [])
        pool = []
        for x in new_rows:
            if x['code'] is None or x['kol'] not in fam:
                continue
            stc = structural_compat(o['nature'], x['nature'])
            if stc == 'INCOMPATIBLE':
                continue
            sc, ent = controlled_score(o['title'], x['title'])
            pool.append((sc, x['code'], str(x['title']).strip(), x['kol'], stc, ent))
        pool = sorted({p[1]: p for p in pool}.values(), key=lambda z: -z[0])[:3]
        controlled_top[code] = [p[1] for p in pool]
        e_all = entries[code]
        seen = set()
        rank = 0

        def add(new_code, source, relation, evidence, struct=None, sc=None):
            nonlocal rank
            key = (new_code, source)
            if key in seen:
                return
            seen.add(key)
            rank += 1
            stc, nn = struct_for(code, new_code) if struct is None else (struct, '')
            alternatives.append(OrderedDict(old_code=code, old_title=o['title_masked'], rank=rank, new_code=new_code,
                                            new_title=new_title.get(new_code, kol_title.get(new_code, '—')) if new_code else '—',
                                            new_parent='; '.join(f'{c}/{k}' for c, k in sorted(parents.get(new_code, set()))) if new_code in parents else (str(new_code) if new_code else ''),
                                            source=source, relation=relation, structural=stc, evidence=evidence,
                                            similarity=sc if sc is not None else (colleague_score(o['title'], new_title.get(new_code, '')) if new_code in new_title else '')))
        for e in e_all:
            for t in e['targets']:
                add(t[0], f"INDEPENDENT_PROPOSAL ({e['org']})", REL_EN[t[1]], t[2] or e['reason'][:120])
        for p in [p for p in allm if p['old_code'] == code]:
            add(p['new_code'], 'COLLEAGUE_V2_TEXT_MATCH', 'TEXT_SIMILARITY_ONLY', f"شباهت {p['score']}٪؛ والد نمایش‌داده {p['parent_shown']}", sc=p['score'])
        for mm in sorted(prior.get(code, set())):
            for c in re.split(r'[/]', mm) if mm else []:
                if c.isdigit():
                    add(int(c), 'PRIOR_V2.1_CROSSWALK', '/'.join(sorted(prior_lvl[code])), 'account_crosswalk.csv v2.1 (Match_Level ' + '/'.join(sorted(prior_lvl[code])) + ')')
                elif '-' in c:
                    for cc in range(int(c.split('-')[0]), int(c.split('-')[1]) + 1):
                        add(cc, 'PRIOR_V2.1_CROSSWALK', '/'.join(sorted(prior_lvl[code])), 'account_crosswalk.csv v2.1 (range)')
        for p in pool:
            add(p[1], 'CONTROLLED_MATCHER_STAGE_ABC', 'CANDIDATE_GENERATOR', f"کل مجاز {fam}؛ سازگاری {p[4]}؛ تعارض عدد نهاد: {p[5]}", struct=p[4], sc=p[0])

    # ---------------------------------------------------------------- proposed crosswalk rows + cardinality
    rows = []
    for e in ENTRIES:
        o = old[e['code']]
        orgs = o['orgs'] if e['org'] == 'BOTH' else [e['org']]
        usage = {k: v for k, v in o['usage'].items() if k.split('-')[0] in orgs}
        lines = sum(v[0] for v in usage.values())
        deb = sum(v[1] for v in usage.values())
        cre = sum(v[2] for v in usage.values())
        tlist = e['targets'] or [(None, None, '')]
        for t in tlist:
            rows.append(OrderedDict(old_code=e['code'], org_scope=e['org'], old_title=o['title_masked'], old_nature=o['nature'],
                                    account_status=e['status'], new_code=t[0],
                                    new_title=new_title.get(t[0], kol_title.get(t[0], '—')) if t[0] else '—',
                                    new_parent='; '.join(f'{c}/{k}' for c, k in sorted(parents.get(t[0], set()))) if t[0] in parents else (str(t[0]) if t[0] else ''),
                                    relation=REL_EN[t[1]] if t[1] else 'NONE', relation_note=t[2], structural=struct_for(e['code'], t[0])[0] if t[0] else 'NOT_APPLICABLE',
                                    lines=lines, debit=deb, credit=cre, reason=e['reason'], decision=e['decision'] or '',
                                    domain=e['domain'], review_status='EXPERT_REVIEW_REQUIRED', approval_status='NOT_APPROVED'))
    # cardinality (only migration-type relations count)
    MIG = {'EQUIVALENT', 'RENAMED', 'SPLIT_COMPONENT', 'RECLASSIFIED', 'CANDIDATE_REVIEW'}
    out_deg = defaultdict(set)
    in_deg = defaultdict(set)
    for r in rows:
        if r['relation'] in MIG and r['new_code'] is not None:
            out_deg[(r['old_code'], r['org_scope'])].add(r['new_code'])
            in_deg[r['new_code']].add((r['old_code'], r['org_scope']))
    for r in rows:
        k = (r['old_code'], r['org_scope'])
        n_out = len(out_deg.get(k, ()))
        n_in = len(in_deg.get(r['new_code'], ())) if r['new_code'] is not None else 0
        r['old_to_new_degree'] = n_out
        r['new_from_old_degree'] = n_in
        if n_out == 0:
            r['cardinality'] = 'UNMAPPED'
        elif n_out == 1 and n_in <= 1:
            r['cardinality'] = '1→1'
        elif n_out == 1 and n_in > 1:
            r['cardinality'] = 'N→1'
        elif n_out > 1 and all(len(in_deg[c]) <= 1 for c in out_deg[k]):
            r['cardinality'] = '1→N'
        else:
            r['cardinality'] = 'N→N'
        r['row_id'] = f"XW-{len([x for x in rows if 'row_id' in x]) + 1:04d}"

    # kol-level pairs of the colleague (حساب کل کد 1)
    kol_pairs = []
    judge = {'220/212': ('VALID', 'سایر پرداختنی‌ها ↔ سایر حساب‌ها و اسناد پرداختنی'),
             '230/213': ('PARTIALLY_VALID', 'پیش‌دریافت‌ها؛ ولی ۲۳۰ جدید شامل «اسناد دریافتنی...» (۲۳۰۰۴–۶) است که دارایی‌اند'),
             '250/221': ('PARTIALLY_VALID', 'پرداختنی بلندمدت؛ ۲۵۰ جدید فقط اوراق مشارکت دارد'),
             '820/919': ('VALID', 'طرف حساب انتظامی پایانه‌ها'),
             '810/918': ('PARTIALLY_VALID', '۹۱۸ تاکسیرانی هم حساب انتظامی (۹۱۸-۱) و هم طرف آن (۹۱۸-۲) را دارد؛ ۹۱۸-۲ باید به ۸۲۰ برود')}
    for p in coa['kol1_pairs']:
        v, why = judge.get(p['code'], ('NOT_TESTABLE', ''))
        kol_pairs.append(OrderedDict(src=p['src'], pair=p['code'], title=p['title'], new_kol=p['code'].split('/')[0], old_kol=p['code'].split('/')[1],
                                     verdict=v, reason=why))

    # candidate generators vs the independent proposal: top-1 precision and top-3 recall over the comparable accounts
    agree_ctrl = recall_ctrl = agree_col = recall_col = comparable = 0
    col3 = defaultdict(list)
    for p in allm:
        col3[p['old_code']].append(p['new_code'])
    for code in old:
        e = pick_entry(code)
        tg = [t[0] for t in e['targets'] if t[1] in ('EQ', 'RN', 'RC', 'CR', 'SP') and t[0] is not None]
        if tg:
            comparable += 1
            ct = controlled_top.get(code) or []
            agree_ctrl += bool(ct) and ct[0] in tg
            recall_ctrl += any(c in tg for c in ct)
            agree_col += col3[code][0] in tg
            recall_col += any(c in tg for c in col3[code])
    colleague_agree = sum(1 for a in audit if a['agrees_with_proposal'])
    # textual similarity vs validity cross-tab
    xtab = Counter()
    for a in audit:
        s = a['best_score']
        band = '100' if s == 100 else '90–99.9' if s >= 90 else '80–89.9' if s >= 80 else '70–79.9' if s >= 70 else '<70'
        xtab[(band, a['verdict'])] += 1
    # prior v2.1 vs proposal vs colleague
    prior_cmp = []
    for a in audit:
        pm = sorted(prior.get(a['old_code'], set()))
        e = pick_entry(a['old_code'])
        tg = {t[0] for t in e['targets'] if t[1] != 'RL' and t[0] is not None}
        pcodes = set()
        for mm in pm:
            for c in re.split(r'[/]', mm) if mm else []:
                if c.isdigit():
                    pcodes.add(int(c))
                elif '-' in c:
                    pcodes.update(range(int(c.split('-')[0]), int(c.split('-')[1]) + 1))
        if not pm:
            rel = 'NOT_IN_PRIOR (level-3 account)'
        elif not pcodes and not tg:
            rel = 'BOTH_NO_EQUIVALENT'
        elif pcodes and tg and pcodes & tg:
            rel = 'PRIOR_AGREES' if pcodes <= tg or tg <= pcodes else 'PRIOR_OVERLAPS'
        elif not pcodes and tg:
            rel = 'PRIOR_NONE_NOW_PROPOSED'
        elif pcodes and not tg:
            rel = 'PRIOR_MAPPED_NOW_NONE'
        else:
            rel = 'PRIOR_CONTRADICTS'
        prior_cmp.append(OrderedDict(old_code=a['old_code'], old_title=a['old_title'], prior_mun_moin=' | '.join(m or '∅' for m in pm),
                                     prior_level='/'.join(sorted(prior_lvl.get(a['old_code'], set()))), colleague_best=a['best_code'],
                                     colleague_verdict=a['verdict'], proposal='; '.join(str(t) for t in sorted(tg)) or '—', relation=rel))

    dump('matching.json', OrderedDict(
        colleague_summary=summary, reproduced_summary=list(rep_summary.items()), status_by_band=[[k[0], k[1], v] for k, v in status_by_band.items()],
        reproduction=dict(score_pairs=len(allm), score_reproduced=rep_score_ok, accounts=len(match), ranking_reproduced=rank_ok,
                          ranking_tie_order_only=rank_tie_only, ranking_failed=rank_bad, exact_after_norm=exact_after_norm,
                          candidate_universe=len(cand_universe),
                          parent_hypothesis_reproduced=hyp_ok, parent_correct=parent_true, parent_wrong=parent_wrong, parent_not_found=parent_nf),
        audit=audit, candidates=cand_audit, alternatives=alternatives, crosswalk=rows, kol_pairs=kol_pairs,
        controlled=dict(comparable=comparable, controlled_top1=agree_ctrl, controlled_recall3=recall_ctrl, colleague_top1=agree_col,
                        colleague_recall3=recall_col, colleague_agrees_all_accounts=colleague_agree),
        xtab=[[k[0], k[1], v] for k, v in sorted(xtab.items())], prior_cmp=prior_cmp))
    print('score', rep_score_ok, '/', len(allm), 'rank', rank_ok, '+tie', rank_tie_only, 'bad', rank_bad, '| parents ok/wrong/nf', parent_true, parent_wrong, parent_nf,
          '| hyp', hyp_ok, '| verdicts', Counter(a['verdict'] for a in audit), '| crosswalk rows', len(rows),
          '| controlled top1/recall3', agree_ctrl, recall_ctrl, '| colleague top1/recall3', agree_col, recall_col, '/', comparable)


if __name__ == '__main__':
    main()
