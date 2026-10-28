"""Consolidated mapping: Organization × Accounting Head × Financial Event × Activity × Cost Center (Taxi & Terminals).

Builds ONE review worksheet from the already-validated project data — no new forensic analysis:
  * ledger lines 1404–1405 and their prior provisional classification   tt_v21.sqlite (sanad, line_event, line_cc, doc_class, hesab)
  * reassessed old→new account crosswalk (Taxi-Gate audit v1)             expert_crosswalk.py + data/work/taxi_gate_audit/matching.json
  * canonical financial events FE-01…FE-70 (v7.1) + TT-NE/PEO/EXC ids     FE v7.1 workbook, v2.1 canonical_events.csv
  * cost centers CC-xxx (cc_master.json) + TTCC-P-01…03 proposals         cc_master.json, v2.1 new_cost_center_candidates.json
  * activity map (CodeMarkaz → ACT-TX-xx; ACT-TM-01; terminal units)       v2.1 workbook sheet Activity_Service_Map

Every ledger line lands in exactly one evidence row (amounts never duplicated).  Alternative cost centers and the companion events
of composite documents are added as amount-free PROPOSED rows.  Accounts without transactions still get a row (proposal or
NO RELATION ESTABLISHED).  Nothing is APPROVED.
"""
import csv
import json
import os
import re
from collections import Counter, OrderedDict, defaultdict

import openpyxl
import xlsxwriter

from common import (ROOT, SRC_DIR, TT_DB, FE_V71, CC_MASTER, V21_DIR, V21_EVENTS, NATURE_FA, load, ro_sqlite, mask_title, norm_key, old_nature,
                    sha256)
from expert_crosswalk import ENTRIES, SUB_SPLIT, SPLIT_ACCOUNTS, MEANING_COLLISIONS

OUT_DIR = os.path.join(SRC_DIR, 'account_event_cc_mapping_v1')
os.makedirs(OUT_DIR, exist_ok=True)
XLSX = os.path.join(OUT_DIR, 'Taxi_Terminals_Account_Event_CostCenter_Mapping.xlsx')
SHEET = 'نگاشت_جامع_سرفصل_رویداد_مرکز'
ROWS_JSON = os.path.join(ROOT, 'data', 'work', 'taxi_gate_audit', 'mapping_rows.json')

ORG_FA = {'Taxi': 'تاکسیرانی', 'Terminals': 'پایانه‌ها'}
ORG_NAME = {'Taxi': 'سازمان تاکسیرانی', 'Terminals': 'سازمان پایانه‌های مسافربری'}
CLASS_FA = {'1': 'دارایی', '2': 'بدهی', '3': 'حقوق صاحبان سرمایه', '4': 'درآمد', '6': 'هزینه', '8': 'غیرعملیاتی',
            '9': 'انتظامی و تراز'}
ROLE_FA = {'DRIVER': 'محرک — حساب اصلی رویداد', 'COMPONENT': 'مؤلفه — طرف مقابل/پرداخت/کسور در همان سند', 'SETTLEMENT': 'تسویه — پرداخت/تهاتر بدهی قبلی',
           'MEMO': 'انتظامی — ثبت تضمین/کنترل', 'CLOSING': 'سند افتتاحیه/اختتامیه'}
VERDICT_FA = {'VALID': 'معتبر', 'PARTIALLY_VALID': 'معتبر فقط برای یک جزء', 'PLAUSIBLE_REVIEW': 'قابل‌دفاع — بازبینی', 'RELATED_NOT_EQUIVALENT': 'مرتبط غیرهم‌ارز',
              'INVALID_STRUCTURAL': 'نامعتبر (ماهیت/طبقه)', 'INVALID_SEMANTIC': 'نامعتبر (معنا)', 'INVALID_ENTITY': 'نامعتبر (نهاد)', 'INVALID_PERSPECTIVE': 'نامعتبر (آینهٔ طرف مقابل)'}
STATUS_FA = {'EXACT_EQUIVALENCE_PROPOSED': 'هم‌ارز (پیشنهادی)', 'RENAMED_EQUIVALENT': 'هم‌ارز با تغییر نام (پیشنهادی)',
             'POSSIBLE_EQUIVALENT_INCOMPLETE_EVIDENCE': 'هم‌ارز احتمالی — شواهد ناقص', 'SPLIT_1_TO_N': 'تفکیک ۱←N', 'RECLASSIFIED': 'طبقه‌بندی مجدد (پیشنهادی)',
             'RELATED_ONLY': 'فقط مرتبط', 'NO_SUITABLE_CANDIDATE': 'بدون معادل در کدینگ جدید', 'EXPERT_REVIEW_REQUIRED': 'نیازمند تصمیم کارشناس',
             'MODEL_AS_DIMENSION': 'به‌صورت بُعد (طرف حساب/دارایی) نه معین'}
REL_FA = {'EQ': 'هم‌ارز', 'RN': 'هم‌ارز با تغییر نام', 'SP': 'جزء تفکیک بر اساس زیرحساب', 'RC': 'طبقه‌بندی مجدد', 'CR': 'کاندید نیازمند بازبینی', 'RL': 'مرتبط غیرهم‌ارز'}

PSEUDO_EVENT = {
    'NON_EVENT_CLOSING': ('', 'غیررویدادی — سند افتتاحیه/اختتامیه/بستن حساب‌ها', '—', 'ردیف در سند افتتاحیه یا اختتامیه؛ رویداد مالی عملیاتی نیست (NOT_AN_OPERATIONAL_EVENT در طبقه‌بندی v2.1).'),
    'UNRESOLVED_AMBIGUOUS': ('', 'UNRESOLVED — رویداد سند مبهم', '—', 'هیچ سطر محرکی برای تعیین رویداد سند یافت نشد (AMBIGUOUS_EVENT در v2.1).'),
    'NO_RELATION': ('', 'NO RELATION ESTABLISHED', '—', 'حساب در دفتر ۱۴۰۴–۱۴۰۵ این سازمان گردش ندارد و شاهد یا پیشنهاد معتبری برای رویداد وجود ندارد.'),
}
PSEUDO_CC = {
    'NA_NONCOST': ('', 'NOT APPLICABLE — NON-COST TRANSACTION', 'غیرهزینه‌ای (ترازنامه‌ای)', 'غیرهزینه‌ای — طرف ترازنامه‌ای سند؛ مرکز هزینه روی سطر هزینهٔ همان سند است'),
    'NA_REVENUE': ('', 'NOT APPLICABLE — REVENUE (بُعد منبع درآمد، CD-03)', 'درآمد — خارج از مرکز هزینه', 'درآمد — به مرکز هزینهٔ مخارج تحمیل نمی‌شود'),
    'NA_NONOP': ('', 'NOT APPLICABLE — NON-OPERATING (درآمد/هزینهٔ غیرعملیاتی)', 'غیرعملیاتی', 'غیرعملیاتی — مرکز هزینه ندارد'),
    'NA_MEMO': ('', 'NOT APPLICABLE — MEMORANDUM', 'انتظامی', 'حساب انتظامی — بدون اثر هزینه‌ای'),
    'NA_CLOSING': ('', 'NOT APPLICABLE — OPENING/CLOSING', 'بستن/افتتاح حساب‌ها', 'سند افتتاحیه/اختتامیه — مرکز هزینه ندارد'),
    'UNRESOLVED_MULTIPLE': ('', 'UNRESOLVED — چند مرکز هزینهٔ کاندید', 'تعیین نشده', 'چند مرکز کاندید؛ گزینه‌ها در ردیف‌های جداگانهٔ PROPOSED بدون مبلغ'),
    'UNRESOLVED_INSUFFICIENT': ('', 'UNRESOLVED — اطلاعات فعالیت کافی نیست', 'تعیین نشده', 'کد فعالیت/واحد برای تعیین مرکز کافی نیست'),
    'UNRESOLVED_NO_CC': ('', 'UNRESOLVED — سطر هزینه بدون مرکز در v2.1', 'تعیین نشده', 'سطر هزینه‌ای (مؤلفه/عملیات پایان دوره) که v2.1 مرکز مستقل به آن نداده است'),
    'UNRESOLVED_NO_EVIDENCE': ('', 'UNRESOLVED — بدون شاهد فعالیت', 'تعیین نشده', 'مرکز هزینه بدون شاهد فعالیت قابل تعیین نیست (از عنوان حساب استخراج نشد)'),
    'NONE': ('', '', '', ''),
}


def hkey(k, m, t):
    return f'{k}-{m}-{t}-0'


def main():
    tc = ro_sqlite(TT_DB)
    mt = load('matching.json')
    coa = load('coa.json')
    audit = {a['old_code']: a for a in mt['audit']}
    prior_rel = {p['old_code']: p for p in mt['prior_cmp']}
    new_title = {}
    for x in coa['new_rows']:
        if x['code'] is not None:
            new_title.setdefault(x['code'], str(x['title']).strip())
    for x in coa['new_rows']:
        new_title.setdefault(x['kol'], str(x['kol_title']).strip())
    ent = defaultdict(dict)
    for e in ENTRIES:
        ent[e['code']][e['org']] = e

    def entry_for(code, org):
        d = ent.get(code)
        if not d:
            return None
        return d.get(org) or d.get('BOTH') or (next(iter(d.values())) if org not in ('Taxi', 'Terminals') else None)

    # ------------------------------------------------------------------ account master (hesab), names: latest FY
    acct = defaultdict(dict)        # org -> hcode -> dict
    for org, fy, hc, name, lvl, kol, mo, t1 in tc.execute('SELECT org,fy,hcode,name,level,kol,moeen,taf1 FROM hesab ORDER BY fy'):
        hc = hc.strip()
        d = acct[org].setdefault(hc, dict(level=lvl, kol=int(kol), moeen=int(mo), taf1=int(t1), names=OrderedDict(), fys=[]))
        d['names'][fy] = name.strip()
        d['fys'].append(fy)

    def aname(org, hc):
        d = acct[org].get(hc)
        if not d:
            return ''
        return list(d['names'].values())[-1]

    # ------------------------------------------------------------------ dictionaries
    wbe = openpyxl.load_workbook(FE_V71, read_only=True, data_only=True)
    EV = OrderedDict()
    for r in wbe['فرهنگ_رویداد_مالی_کاندید'].iter_rows(min_row=2, values_only=True):
        if r[0]:
            EV[r[0]] = dict(title=r[1], family=r[4], desc=f'{r[3]} | ماهیت: {r[5]}', kind='FE')
    for r in csv.DictReader(open(V21_EVENTS, encoding='utf-8-sig')):
        i = r['event_id']
        if i in EV:
            continue
        if i.startswith('TT-NE'):
            EV[i] = dict(title=r['title_fa'], family=r['family'],
                         desc=f"[رویداد کاندید TT-NE — پیشنهادی، خارج از فرهنگ FE v7.1] {r['definition']} | ماهیت: {r['nature']}", kind='TTNE')
        elif i.startswith('PEO'):
            EV[i] = dict(title=r['title_fa'], family='عملیات حسابداری پایان دوره (خارج از FE؛ CD-PEO)',
                         desc=f"[شناسهٔ PEO پروژهٔ قبلی — عملیات پایان دوره، خارج از FE؛ پیشنهادی] {r['definition'] or r['title_fa']}", kind='PEO')
        elif i.startswith('EXC'):
            EV[i] = dict(title=r['title_fa'], family='صف استثنا (خارج از FE)', desc=f"[استثنا — در انتظار شناسایی] {r['definition'] or r['title_fa']}", kind='EXC')
    CC = json.load(open(CC_MASTER, encoding='utf8'))
    NCC = json.load(open(os.path.join(V21_DIR, 'machine_readable', 'new_cost_center_candidates.json'), encoding='utf8'))
    w21 = openpyxl.load_workbook(os.path.join(V21_DIR, 'taxi_terminals_financial_mapping_1404_1405_v2.1.xlsx'), read_only=True, data_only=True)
    act_by_code, act_title, unit_name = {}, {}, {}
    for r in w21['Activity_Service_Map'].iter_rows(min_row=5, values_only=True):
        if not r[0]:
            continue
        rid = str(r[0])
        if rid.startswith('ACT-TX') and r[2] == 'Taxi':
            act_by_code[str(r[3])] = rid
            act_title.setdefault(rid, r[1])
        elif rid.startswith('ACT-TM'):
            act_title.setdefault(rid, r[1])
        elif rid.startswith('ORGUNIT-TM-'):
            unit_name[rid[-3:]] = r[1]
    wcd = openpyxl.load_workbook(os.path.join(V21_DIR, 'committee_decisions_v2.1.xlsx'), read_only=True, data_only=True)
    prior_dec = {str(r[0]) for r in wcd['Decisions'].iter_rows(min_row=4, values_only=True) if r[0]}

    def cc_info(c):
        if c in CC:
            return CC[c]['title'], f"{CC[c]['domain']} | {CC[c]['status']}", CC[c]['domain']
        if c in NCC:
            n = NCC[c]
            return (n.get('title_v2') or n.get('title', ''), f"کاندید TTCC — پیشنهادی v2.1 (خارج از cc_master) | توصیه: {n.get('recommendation', '')}",
                    'حمل‌ونقل، ترافیک و مهندسی عبور و مرور (کاندید)')
        return '', '', ''

    # ------------------------------------------------------------------ old → new coding (reassessed crosswalk; V2 suggestions never reused)
    def new_map(org, k, m, t):
        c2, c3 = hkey(k, m, 0), hkey(k, m, t)
        code_l = c3 if (t and c3 in ent) else c2
        e = entry_for(code_l, org)
        notes = []
        if e is None:
            if code_l not in ent:
                return '', '', 'خارج از فهرست Taxi-Gate — نگاشت تعریف نشده (TG-D-22)', None
            return '', '', 'معنای این کد در این سازمان در ممیزی بررسی نشده', None
        new_code, ntitle, st = '', '', STATUS_FA[e['status']]
        if code_l == c2 and c2 in SPLIT_ACCOUNTS and e['status'] in ('SPLIT_1_TO_N', 'EXPERT_REVIEW_REQUIRED'):
            sub = SUB_SPLIT.get((org, f'{k}-{m}-{t}'))
            if t == 0:
                st += ' — سطر بدون تفصیلی در حساب تفکیکی؛ مقصد نامعین'
            elif sub is None:
                st += ' — زیرحساب بدون قاعدهٔ تفکیک'
            elif sub[0] is None:
                st += f' — زیرحساب بدون معادل در کدینگ جدید ({REL_FA[sub[1]]}؛ {sub[2]})'
            else:
                new_code, st = sub[0], st + f' — {REL_FA[sub[1]]}' + (f' ({sub[2]})' if sub[2] else '')
        else:
            tg = [x for x in e['targets'] if x[1] != 'RL' and x[0] is not None]
            if len(tg) == 1:
                new_code = tg[0][0]
                st += f' — {REL_FA[tg[0][1]]}'
            elif len(tg) > 1:
                st += ' — چند مقصد کاندید: ' + '؛ '.join(f'{x[0]} ({REL_FA[x[1]]})' for x in tg)
            rl = [x for x in e['targets'] if x[1] == 'RL' and x[0] is not None]
            if rl:
                notes.append('مرتبط غیرهم‌ارز (مقصد مهاجرت نیست): ' + '؛ '.join(str(x[0]) for x in rl))
        if new_code != '':
            ntitle = new_title.get(new_code, '')
        a = audit.get(code_l)
        if a:
            same = (str(a['best_code']) == str(new_code))
            v2 = f"پیشنهاد متنی Taxi-Gate V2: {a['best_code']} ({VERDICT_FA[a['verdict']]})" + ('' if same else ' — بازاستفاده نشد')
            notes.append(v2)
        p = prior_rel.get(code_l)
        if p and p['relation'] in ('PRIOR_CONTRADICTS', 'PRIOR_MAPPED_NOW_NONE'):
            notes.append(f"متعارض با crosswalk v2.1 ({p['prior_mun_moin']}) — بازبینی")
        if code_l in MEANING_COLLISIONS:
            notes.append('کد با معنای متفاوت در دو سازمان (TG-D-04)')
        return (str(new_code) if new_code != '' else ''), ntitle, st + ('؛ ' + '؛ '.join(notes) if notes else ''), e

    # ------------------------------------------------------------------ ledger lines → evidence groups
    q = """SELECT s.sid,s.org,s.fy,s.doc_key,s.kol,s.moeen,s.taf1,s.bed,s.bes,s.sharh,s.des,s.markaz,
                  e.event_id,e.role,e.outcome,e.conf,e.rule_id,
                  c.eligible,c.cc_status,c.cc_id,c.alt_ccs,c.activity_id,c.source_markaz,c.source_unit,c.conf,c.rule_id,
                  d.kind,d.composite,d.driver_events
           FROM sanad s JOIN line_event e ON e.sid=s.sid JOIN line_cc c ON c.sid=s.sid JOIN doc_class d ON d.doc_key=s.doc_key"""
    G = OrderedDict()
    cand = defaultdict(lambda: dict(lines=0, docs=set(), sources=Counter()))
    n_lines = 0
    for (sid, org, fy, doc, kol, mo, t1, bed, bes, sharh, des, markaz, ev, role, outcome, econf, erule, elig, ccst, ccid, alt, act, smz, sunit, cconf, crule,
         kind, composite, drivers) in tc.execute(q):
        n_lines += 1
        k, m, t = int(kol), int(mo), int(t1)
        bed, bes = bed or 0, bes or 0
        nat = old_nature(k, m)
        unit = ''
        if org == 'Terminals' and elig == 1:
            unit = (sunit or '')[:3]
        # event
        if ev:
            ekey = ev
        elif role == 'CLOSING':
            ekey = 'NON_EVENT_CLOSING'
        else:
            ekey = 'UNRESOLVED_AMBIGUOUS'
        # activity
        akey = ''
        if elig == 1:
            if org == 'Taxi':
                mz = str(smz or markaz or '').strip()
                akey = act or act_by_code.get(mz) or ('NO_ACTIVITY_CODE' if mz in ('', 'NULL', 'None') else f'MARKAZ:{mz}')
            else:
                akey = 'ACT-TM-01' if str(smz or markaz) == '008' else 'NO_ACTIVITY_CODE'
        elif nat == 'REVENUE' and role != 'CLOSING':
            akey = 'REVENUE_SOURCE'
        # cost center
        if elig == 1:
            ckey = {'PROVISIONAL_EXISTING_CC': ccid, 'NEW_CC_CANDIDATE': ccid, 'MULTIPLE_CC_CANDIDATES': 'UNRESOLVED_MULTIPLE',
                    'INSUFFICIENT_INFORMATION': 'UNRESOLVED_INSUFFICIENT'}[ccst]
        elif role == 'CLOSING':
            ckey = 'NA_CLOSING'
        elif nat == 'REVENUE':
            ckey = 'NA_REVENUE'
        elif nat == 'NONOP_MIXED':
            ckey = 'NA_NONOP'
        elif nat in ('MEMO', 'MEMO_CONTRA'):
            ckey = 'NA_MEMO'
        elif nat in ('EXPENSE', 'NONCASH_EXPENSE', 'COST_OF_SALES'):
            ckey = 'UNRESOLVED_NO_CC'
        else:
            ckey = 'NA_NONCOST'
        key = (org, unit, k, m, t, ekey, akey, ckey)
        g = G.get(key)
        if g is None:
            g = G[key] = dict(org=org, unit=unit, k=k, m=m, t=t, ekey=ekey, akey=akey, ckey=ckey, lines=0, docs=set(), debit=0, credit=0, fys=set(),
                              roles=Counter(), outcomes=Counter(), econf=Counter(), cconf=Counter(), erule=Counter(), crule=Counter(), narr=Counter(),
                              narr_doc={}, big=(0, ''), alt_for=0, comp_for=0)
        g['lines'] += 1
        g['docs'].add(doc)
        g['debit'] += bed
        g['credit'] += bes
        g['fys'].add(fy)
        g['roles'][role] += 1
        g['outcomes'][outcome] += 1
        g['econf'][econf or ''] += 1
        if elig == 1:
            g['cconf'][cconf or ''] += 1
            g['crule'][crule or ''] += 1
        g['erule'][erule or ''] += 1
        txt = (sharh or des or '').strip()
        if txt:
            nk = re.sub(r'\d+', '#', re.sub(r'\s+', ' ', txt))[:120]
            g['narr'][nk] += 1
            g['narr_doc'].setdefault(nk, (doc, txt))
        if bed + bes > g['big'][0]:
            g['big'] = (bed + bes, doc)
        # alternative cost centers → amount-free candidate rows
        if elig == 1 and alt:
            for a in [x.strip() for x in alt.split(';') if x.strip()]:
                if a != ccid:
                    ck = (org, unit, k, m, t, ekey, akey, a)
                    cand[ck]['lines'] += 1
                    cand[ck]['docs'].add(doc)
                    cand[ck]['sources'][f'گزینهٔ جایگزین برای سطرهای «{ckey}»'] += 1
        # companion events of composite documents → amount-free candidate rows
        if composite and outcome == 'COMPONENT_OF_COMPOSITE_TRANSACTION' and drivers:
            for de in json.loads(drivers):
                if de and de != ev:
                    ck = (org, unit, k, m, t, de, akey, ckey)
                    cand[ck]['lines'] += 1
                    cand[ck]['docs'].add(doc)
                    cand[ck]['sources'][f'رویداد همراه در سند ترکیبی (مبلغ در ردیف رویداد اصلی {ev})'] += 1

    # merge candidates that coincide with evidence rows
    cand_rows = OrderedDict()
    for ck, c in cand.items():
        if ck in G:
            if ck[7] == G[ck]['ckey'] and ck[5] == G[ck]['ekey']:
                G[ck]['also_candidate'] = f"همچنین گزینه/رویداد همراه برای {c['lines']:,} ردیف دیگر ({len(c['docs'])} سند)"
        else:
            cand_rows[ck] = c

    # ------------------------------------------------------------------ account universe and proposals for accounts without an established event
    established = defaultdict(set)
    heads_with_rows = defaultdict(set)
    ev_by_head = defaultdict(Counter)
    for g in G.values():
        heads_with_rows[g['org']].add(hkey(g['k'], g['m'], g['t']))
        if g['ekey'] in EV and EV[g['ekey']]['kind'] in ('FE', 'TTNE', 'PEO'):
            established[g['org']].add(hkey(g['k'], g['m'], g['t']))
            ev_by_head[(g['org'], hkey(g['k'], g['m'], g['t']))][g['ekey']] += g['lines']
    prop_rows = OrderedDict()

    def add_prop(org, k, m, t, ekey, ckey, method, why):
        key = (org, '', k, m, t, ekey, '', ckey)
        if key in G or key in cand_rows or key in prop_rows:
            return
        prop_rows[key] = dict(method=method, why=why)

    universe = defaultdict(set)
    for org, accs in acct.items():
        for hc, d in accs.items():
            if d['level'] in (2, 3):
                universe[org].add(hc)
    for org in heads_with_rows:
        universe[org] |= heads_with_rows[org]
    for org in ('Taxi', 'Terminals'):
        other = 'Terminals' if org == 'Taxi' else 'Taxi'
        for hc in sorted(universe[org]):
            if hc in established[org]:
                continue
            k, m, t = [int(x) for x in hc.split('-')[:3]]
            d = acct[org].get(hc, {})
            # leaf rule: an L2 with children is covered through its children
            if t == 0 and hc not in heads_with_rows[org] and any(x.startswith(f'{k}-{m}-') and x != hc and acct[org][x]['level'] == 3 for x in acct[org]):
                continue
            nat = old_nature(k, m)
            # (a) period-end operations defined by the prior governance (CD-PEO): accumulated depreciation / provisions / depreciation expense
            peo = None
            if k == 122 or (k == 613 and m == 20 and t not in (9,)):
                peo = 'PEO-01'
            elif (k, m) in ((112, 9), (124, 4), (115, 9), (113, 9), (212, 2)) or k in (214, 217, 223, 224) or (k == 613 and m == 20 and t == 9):
                peo = 'PEO-02'
            if peo:
                add_prop(org, k, m, t, peo, 'UNRESOLVED_NO_CC' if nat in ('EXPENSE', 'NONCASH_EXPENSE') else 'NA_NONCOST',
                         'پیشنهاد معنایی کنترل‌شده بر اساس تعریف عملیات پایان دوره (CD-PEO) و ماهیت حساب (کاهنده/ذخیره/هزینهٔ استهلاک)',
                         'ماهیت حساب (نه واژهٔ عنوان) با تعریف PEO هم‌خوان است؛ گردش مستقیم در ۱۴۰۴–۱۴۰۵ برای این سازمان ثبت نشده')
            # (b) the same account (same code AND same title) is used by the other organisation
            same = hc in acct[other] and norm_key(aname(other, hc)) == norm_key(aname(org, hc)) and hkey(k, m, 0) not in MEANING_COLLISIONS
            if same and ev_by_head.get((other, hc)):
                tot = sum(ev_by_head[(other, hc)].values())
                for e2, n2 in ev_by_head[(other, hc)].most_common(3):
                    if n2 >= 3 and n2 / tot >= 0.10:
                        add_prop(org, k, m, t, e2, 'UNRESOLVED_NO_EVIDENCE' if nat in ('EXPENSE', 'NONCASH_EXPENSE', 'COST_OF_SALES') else
                                 ('NA_REVENUE' if nat == 'REVENUE' else 'NA_MEMO' if nat in ('MEMO', 'MEMO_CONTRA') else 'NA_NONOP' if nat == 'NONOP_MIXED' else 'NA_NONCOST'),
                                 f'پیشنهاد از شواهد همان حساب (همان کد و همان عنوان) در سازمان {ORG_FA[other]}',
                                 f'در {ORG_FA[other]} {n2:,} از {tot:,} ردیف عملیاتی این حساب به این رویداد طبقه‌بندی شده‌اند؛ در این سازمان گردشی ثبت نشده')
            if hc not in heads_with_rows[org] and not any(kk[0] == org and hkey(kk[2], kk[3], kk[4]) == hc for kk in prop_rows):
                prop_rows[(org, '', k, m, t, 'NO_RELATION', '', 'NONE')] = dict(method='—', why='')

    # ------------------------------------------------------------------ assemble rows
    def head_cols(org, k, m, t):
        l1, l2, l3 = hkey(k, 0, 0), hkey(k, m, 0), hkey(k, m, t)
        nat = old_nature(k, m)
        nm3 = aname(org, l3) if t else ''
        masked3 = mask_title(nm3) if nm3 else ''
        return OrderedDict([
            ('c5', l1), ('c6', aname(org, l1)), ('c7', l2), ('c8', mask_title(aname(org, l2))), ('c9', l3 if t else ''), ('c10', masked3),
            ('c11', f"{CLASS_FA.get(str(k)[0], '')} / {NATURE_FA.get(nat, nat)}"),
        ]), nat, (nm3 and masked3 != nm3)

    def ev_cols(ekey):
        if ekey in PSEUDO_EVENT:
            c, t_, f, dsc = PSEUDO_EVENT[ekey]
            return c, t_, f, dsc
        e = EV.get(ekey)
        if not e:
            return ekey, '(شناسه در فرهنگ رویداد یافت نشد)', '', ''
        return ekey, e['title'], e['family'], e['desc']

    def act_cols(org, akey, k, m, t):
        if akey == '':
            return ''
        if akey == 'REVENUE_SOURCE':
            l3 = hkey(k, m, t)
            return f"منبع درآمد: {mask_title(aname(org, l3) if t else aname(org, hkey(k, m, 0)))}"
        if akey.startswith('ACT-'):
            return f"{akey} — {act_title.get(akey, '')}"
        if akey.startswith('MARKAZ:'):
            return f"کد فعالیت منبع {akey[7:]} — در نقشهٔ فعالیت v2.1 نیست"
        if akey == 'NO_ACTIVITY_CODE':
            return 'کد فعالیت منبع ثبت نشده (فقط واحد سازمانی معلوم است)'
        return akey

    def decision_refs(org, k, m, t, ekey, ckey, entry):
        refs = []
        if org == 'Terminals' and (k, m, t) in ((113, 1, 1), (211, 1, 1)):
            refs.append('CD-01 / TG-D-01')
        if org == 'Terminals' and (k, m) == (413, 8):
            refs.append('CD-01 / TG-D-01')
        if ckey == 'NA_REVENUE':
            refs.append('CD-03 / TG-D-14')
        if ckey in ('TTCC-P-01',):
            refs.append('CD-04')
        if ckey == 'TTCC-P-02':
            refs.append('CD-05')
        if ckey == 'TTCC-P-03':
            refs.append('CD-06')
        if ckey == 'CC-1301':
            refs.append('CD-08')
        if ckey.startswith('UNRESOLVED_MULTIPLE') or ckey.startswith('UNRESOLVED_INSUFF'):
            refs.append('CD-LOWCC / CD-07')
        if str(ekey).startswith('PEO'):
            refs.append('CD-PEO / TG-D-03')
        if ekey == 'EXC-01':
            refs.append('CD-EXC01')
        if ekey == 'FE-60':
            refs.append('CD-FE60')
        if str(ekey).startswith('TT-NE-'):
            cd = 'CD-EV-' + ekey[-2:]
            if cd in prior_dec:
                refs.append(cd)
        if ckey == 'NA_MEMO':
            refs.append('CD-GUAR / TG-D-15')
        if entry and entry.get('decision'):
            refs.append(entry['decision'])
        if ekey not in ('NON_EVENT_CLOSING', 'NO_RELATION'):
            refs.append('CD-CARD / TG-D-20')
        return ' ; '.join(dict.fromkeys(refs))

    def conf_label(level, strength):
        return f"{ {'HIGH': 'بالا', 'MEDIUM': 'متوسط', 'LOW': 'پایین'}.get(level, 'نامشخص') } — { {'STRONG': 'شواهد قوی', 'PARTIAL': 'شواهد جزئی', 'INCOMPLETE': 'زمینهٔ ناقص', 'CONFLICT': 'شواهد متعارض'}[strength] }"

    def sem_conflict(nat, ekey, role):
        e = EV.get(ekey)
        if not e or role != 'DRIVER':
            return False
        n = str(e['desc'])
        if nat in ('EXPENSE',) and 'هزینه' not in n:
            return True
        if nat == 'REVENUE' and 'درآمد' not in n:
            return True
        return False

    out = []
    rank = {'HIGH': 3, 'MEDIUM': 2, 'LOW': 1, '': 0}
    inv = {3: 'HIGH', 2: 'MEDIUM', 1: 'LOW', 0: 'LOW'}
    for key, g in G.items():
        org, unit, k, m, t, ekey, akey, ckey = key
        hc, nat, masked = head_cols(org, k, m, t)
        nc, nt, nst, entry = new_map(org, k, m, t)
        ecode, etitle, efam, edesc = ev_cols(ekey)
        ccode, ctitle, cclass, crel = (ckey, *cc_info(ckey)[:2], '') if ckey in CC or ckey in NCC else PSEUDO_CC[ckey]
        if ckey in CC or ckey in NCC:
            crel = 'مرکز کاندید TTCC (پیشنهادی)' if ckey in NCC else 'مستقیم — طبقه‌بندی موقت v2.1'
            top_rule = g['crule'].most_common(1)[0][0] if g['crule'] else ''
            if top_rule:
                crel += f' (قاعده {top_rule})'
        role, _ = g['roles'].most_common(1)[0]
        rel = ROLE_FA.get(role, role)
        if len(g['roles']) > 1:
            rel += ' (ترکیبی: ' + '؛ '.join(f"{ROLE_FA.get(r, r).split(' —')[0]} {n:,}" for r, n in g['roles'].most_common()) + ')'
        econf = g['econf'].most_common(1)[0][0]
        lvl = rank.get(econf, 0)
        if g['cconf']:
            lvl = min(lvl, rank.get(g['cconf'].most_common(1)[0][0], 0))
        ev_kind = EV.get(ekey, {}).get('kind')
        if ev_kind in ('TTNE', 'PEO', 'EXC') or ckey in NCC:
            lvl = min(lvl, 2)
        conflict = sem_conflict(nat, ekey, role)
        unresolved = ekey.startswith('UNRESOLVED') or ekey == 'EXC-01' or ckey.startswith('UNRESOLVED')
        cd01 = org == 'Terminals' and ((k, m, t) in ((113, 1, 1), (211, 1, 1)) or (k, m) == (413, 8))
        if conflict:
            strength, lvl = 'CONFLICT', 1
        elif unresolved:
            strength, lvl = 'INCOMPLETE', min(lvl, 1)
        elif lvl == 3 and g['lines'] >= 10:
            strength = 'STRONG'
        else:
            strength = 'PARTIAL'
        if unresolved or cd01:
            status = 'UNRESOLVED'
        elif ev_kind in ('TTNE', 'PEO') or ckey in NCC:
            status = 'PROPOSED'
        else:
            status = 'OBSERVED'
        notes = []
        if cd01:
            notes.append('ماهیت عاملیتی/اصیل پایانه‌ها حل نشده (CD-01)؛ رابطه تا تصمیم UNRESOLVED است')
        if ev_kind == 'TTNE':
            notes.append('رویداد کاندید TT-NE (پیشنهادی)')
        if ckey in NCC:
            notes.append('مرکز هزینهٔ کاندید TTCC (پیشنهادی)')
        if conflict:
            notes.append('ماهیت حساب با ماهیت رویداد ناسازگار به نظر می‌رسد — بازبینی')
        if g.get('also_candidate'):
            notes.append(g['also_candidate'])
        if ckey == 'UNRESOLVED_NO_CC':
            notes.append('مرکز باید از سطر محرک همان سند یا قاعدهٔ تخصیص تعیین شود' if role != 'DRIVER' else 'v2.1 مرکز تخصیص نداد')
        if masked:
            notes.append('عنوان تفصیلی برای حریم خصوصی ماسک شد')
        if '1405' in g['fys']:
            notes.append('۱۴۰۵ ناقص است (تا مهر ۱۴۰۵)')
        if ckey.startswith('NA_') and akey == 'REVENUE_SOURCE':
            notes.append('«فعالیت» برای درآمد = منبع درآمد (همان حساب درآمد)؛ مرکز هزینه تخصیص نمی‌یابد')
        mode_n = g['narr'].most_common(1)[0][0] if g['narr'] else None
        sample_doc, sample_txt = g['narr_doc'][mode_n] if mode_n else (g['big'][1], '')
        docs_s = [sample_doc] + ([g['big'][1]] if g['big'][1] and g['big'][1] != sample_doc else [])
        row = OrderedDict()
        row['c1'] = ORG_FA[org]
        row['c2'] = ORG_NAME[org] + (f" — {unit_name.get(unit, 'واحد ' + unit)}" if unit else '')
        row['c3'] = (cc_info(ckey)[2] if ckey in CC or ckey in NCC else
                     {'NA_REVENUE': 'درآمد عملیاتی سازمان', 'NA_NONCOST': 'غیرهزینه‌ای (ترازنامه‌ای/تسویه)', 'NA_MEMO': 'حساب‌های انتظامی',
                      'NA_CLOSING': 'بستن/افتتاح حساب‌ها', 'NA_NONOP': 'درآمد/هزینهٔ غیرعملیاتی'}.get(ckey, 'هزینه — مرکز تعیین نشده'))
        row['c4'] = f"دفتر اسناد {ORG_NAME[org]} {'، '.join(sorted(g['fys']))} (Sanad) + طبقه‌بندی v2.1 (line_event/line_cc)"
        row.update(hc)
        row['c12'], row['c13'], row['c14'] = nc, nt, nst
        row['c15'], row['c16'], row['c17'], row['c18'] = ecode, etitle, efam, edesc
        row['c19'] = rel
        row['c20'] = act_cols(org, akey, k, m, t)
        row['c21'], row['c22'], row['c23'] = ccode, ctitle, cclass
        row['c24'] = unit_name.get(unit, '') if unit else ''
        row['c25'] = crel
        row['c26'], row['c27'], row['c28'], row['c29'] = g['lines'], len(g['docs']), g['debit'], g['credit']
        row['c30'] = '؛ '.join(sorted(g['fys']))
        row['c31'] = '؛ '.join(docs_s[:2])
        row['c32'] = mask_title(sample_txt)[:220] if sample_txt else ''
        row['c33'] = 'هم‌رخدادی سطر واقعی دفتر × رویداد (line_event) × مرکز هزینه (line_cc) — طبقه‌بندی موقت v2.1'
        row['c34'] = 'tt_v21.sqlite (Sanad/Hesab ۱۴۰۴–۱۴۰۵)؛ crosswalk ممیزی Taxi-Gate v1؛ FE v7.1؛ cc_master.json؛ Activity_Service_Map v2.1'
        row['c35'] = conf_label(inv[lvl], strength)
        row['c36'] = status
        refs = decision_refs(org, k, m, t, ekey, ckey, entry)
        row['c37'] = 'نیازمند تأیید کارشناسی' + (f' — {refs}' if refs else '')
        row['c38'] = '؛ '.join(notes)
        row['_key'] = list(key)
        row['_kind'] = 'EVIDENCE'
        out.append(row)
    for key, c in cand_rows.items():
        org, unit, k, m, t, ekey, akey, ckey = key
        hc, nat, masked = head_cols(org, k, m, t)
        nc, nt, nst, entry = new_map(org, k, m, t)
        ecode, etitle, efam, edesc = ev_cols(ekey)
        if ckey in CC or ckey in NCC:
            ccode, (ctitle, cclass, _) = ckey, cc_info(ckey)
        else:
            ccode, ctitle, cclass, _ = PSEUDO_CC[ckey]
        src = c['sources'].most_common(1)[0][0]
        is_alt = src.startswith('گزینهٔ جایگزین')
        row = OrderedDict()
        row['c1'] = ORG_FA[org]
        row['c2'] = ORG_NAME[org] + (f" — {unit_name.get(unit, 'واحد ' + unit)}" if unit else '')
        row['c3'] = cc_info(ckey)[2] if (ckey in CC or ckey in NCC) else PSEUDO_CC[ckey][2]
        row['c4'] = f'دفتر اسناد {ORG_NAME[org]} (Sanad) + طبقه‌بندی v2.1 — ردیف کاندید بدون مبلغ'
        row.update(hc)
        row['c12'], row['c13'], row['c14'] = nc, nt, nst
        row['c15'], row['c16'], row['c17'], row['c18'] = ecode, etitle, efam, edesc
        row['c19'] = 'مؤلفهٔ سند ترکیبی — رویداد همراه (بدون مبلغ)' if not is_alt else 'همان رابطهٔ حساب و رویداد ردیف اصلی'
        row['c20'] = act_cols(org, akey, k, m, t)
        row['c21'], row['c22'], row['c23'] = ccode, ctitle, cclass
        row['c24'] = unit_name.get(unit, '') if unit else ''
        row['c25'] = 'گزینهٔ جایگزین مرکز هزینه (v2.1 alt_ccs) — تعیین نشده' if is_alt else PSEUDO_CC.get(ckey, ('', '', '', 'طبقه‌بندی موقت v2.1'))[3] if ckey in PSEUDO_CC else 'طبقه‌بندی موقت v2.1'
        for col in ('c26', 'c27', 'c28', 'c29', 'c30', 'c31', 'c32'):
            row[col] = ''
        row['c33'] = ('گزینهٔ جایگزین مرکز هزینه در طبقه‌بندی v2.1 برای همان سطرها' if is_alt else
                      'رویداد محرک دیگر در همان سند ترکیبی (doc_class.driver_events)')
        row['c34'] = 'tt_v21.sqlite (line_cc.alt_ccs / doc_class.driver_events)'
        row['c35'] = conf_label('LOW', 'INCOMPLETE')
        cd01c = org == 'Terminals' and ((k, m, t) in ((113, 1, 1), (211, 1, 1)) or (k, m) == (413, 8))
        unres_c = ckey.startswith('UNRESOLVED') or ekey.startswith('UNRESOLVED') or ekey == 'EXC-01'
        row['c36'] = 'UNRESOLVED' if (cd01c or unres_c) else 'PROPOSED'
        refs = decision_refs(org, k, m, t, ekey, ckey, entry)
        row['c37'] = 'نیازمند تأیید کارشناسی' + (f' — {refs}' if refs else '')
        row['c38'] = (('ماهیت عاملیتی/اصیل پایانه‌ها حل نشده (CD-01)؛ ' if cd01c else '') +
                      f"مربوط به {c['lines']:,} ردیف / {len(c['docs']):,} سند که مبلغشان فقط در ردیف اصلی (همان حساب و رویداد) آمده است؛ جمع نزنید" +
                      ('؛ رویداد کاندید TT-NE (پیشنهادی)' if EV.get(ekey, {}).get('kind') == 'TTNE' else ''))
        row['_key'] = list(key)
        row['_kind'] = 'CANDIDATE'
        out.append(row)
    for key, p in prop_rows.items():
        org, unit, k, m, t, ekey, akey, ckey = key
        hc, nat, masked = head_cols(org, k, m, t)
        nc, nt, nst, entry = new_map(org, k, m, t)
        ecode, etitle, efam, edesc = ev_cols(ekey)
        ccode, ctitle, cclass, crel = PSEUDO_CC[ckey]
        no_rel = ekey == 'NO_RELATION'
        row = OrderedDict()
        row['c1'] = ORG_FA[org]
        row['c2'] = ORG_NAME[org]
        row['c3'] = 'بدون گردش ۱۴۰۴–۱۴۰۵' if no_rel else PSEUDO_CC[ckey][2]
        row['c4'] = f"سرفصل حساب‌های {ORG_NAME[org]} (Hesab ۱۴۰۴–۱۴۰۵)" + ('' if no_rel else ' + ' + p['method'])
        row.update(hc)
        row['c12'], row['c13'], row['c14'] = nc, nt, nst
        row['c15'], row['c16'], row['c17'], row['c18'] = ecode, etitle, efam, edesc
        row['c19'] = '—' if no_rel else 'پیشنهادی (بدون ردیف دفتر در این سازمان)'
        row['c20'] = ''
        row['c21'], row['c22'], row['c23'] = ccode, ctitle, cclass
        row['c24'] = ''
        row['c25'] = crel
        for col in ('c26', 'c27', 'c28', 'c29', 'c30', 'c31', 'c32'):
            row[col] = ''
        row['c33'] = '—' if no_rel else p['method']
        row['c34'] = 'tt_v21.sqlite (Hesab)' + ('' if no_rel else '؛ ' + ('CD-PEO (committee_decisions_v2.1)' if str(ekey).startswith('PEO') else 'line_event سازمان دیگر'))
        row['c35'] = 'نامشخص — بدون شاهد' if no_rel else conf_label('MEDIUM' if str(ekey).startswith('PEO') else 'LOW', 'INCOMPLETE' if not str(ekey).startswith('PEO') else 'PARTIAL')
        cd01p = org == 'Terminals' and ((k, m, t) in ((113, 1, 1), (211, 1, 1)) or (k, m) == (413, 8))
        row['c36'] = 'UNRESOLVED' if (no_rel or cd01p or ckey.startswith('UNRESOLVED')) else 'PROPOSED'
        refs = decision_refs(org, k, m, t, ekey, ckey, entry)
        row['c37'] = 'نیازمند تأیید کارشناسی' + (f' — {refs}' if refs else '')
        row['c38'] = ('حساب در دفتر ۱۴۰۴–۱۴۰۵ این سازمان گردش ندارد؛ رویدادی از شواهد قابل استنتاج نیست' if no_rel else p['why'])
        if acct[org].get(hkey(k, m, t), {}).get('level') is None:
            row['c38'] += '؛ سرفصل در Hesab این سازمان یافت نشد'
        if hkey(k, m, 0).startswith('920-'):
            row['c38'] += '؛ کل ۹۲۰ «تست» — احتمالاً باقی‌ماندهٔ آزمایشی'
        row['_key'] = list(key)
        row['_kind'] = 'PROPOSAL' if not no_rel else 'NO_RELATION'
        out.append(row)

    # ------------------------------------------------------------------ order
    def sk(r):
        org = 0 if r['c1'] == ORG_FA['Taxi'] else 1
        k, m, t = r['_key'][2], r['_key'][3], r['_key'][4]
        kind = {'EVIDENCE': 0, 'PROPOSAL': 1, 'CANDIDATE': 2, 'NO_RELATION': 3}[r['_kind']]
        return (org, str(k)[0], k, m, t, r['c15'] == '', r['c15'], r['c16'], kind, r['c21'] == '', r['c21'], r['c22'], r['c2'], r['c20'])
    out.sort(key=sk)
    json.dump(dict(rows=out, n_lines=n_lines, universe={o: sorted(v) for o, v in universe.items()}), open(ROWS_JSON, 'w', encoding='utf8'), ensure_ascii=False)
    write_xlsx(out)
    print('rows', len(out), Counter(r['_kind'] for r in out), Counter(r['c36'] for r in out), 'lines', n_lines)


HEADERS = [
    ('c1', 'حوزه سازمانی', 10), ('c2', 'سازمان / واحد مرتبط', 22), ('c3', 'حوزه فعالیت', 22), ('c4', 'منبع اطلاعات', 26),
    ('c5', 'کد حساب کل', 11), ('c6', 'عنوان حساب کل', 20), ('c7', 'کد حساب معین', 11), ('c8', 'عنوان حساب معین', 24), ('c9', 'کد حساب تفصیلی', 11),
    ('c10', 'عنوان حساب تفصیلی', 24), ('c11', 'طبقه / ماهیت حساب', 22), ('c12', 'کد حساب در کدینگ جدید', 11), ('c13', 'عنوان حساب در کدینگ جدید', 24),
    ('c14', 'وضعیت تطبیق حساب قدیم و جدید', 40), ('c15', 'کد رویداد مالی', 10), ('c16', 'عنوان رویداد مالی', 32), ('c17', 'خانواده رویداد مالی', 18),
    ('c18', 'شرح و ماهیت رویداد', 45), ('c19', 'نوع رابطه حساب و رویداد مالی', 28), ('c20', 'عنوان فعالیت / خدمت', 32), ('c21', 'کد مرکز هزینه', 11),
    ('c22', 'عنوان مرکز هزینه', 32), ('c23', 'طبقه‌بندی مرکز هزینه', 30), ('c24', 'واحد متولی فعالیت', 16), ('c25', 'نوع ارتباط رویداد و مرکز هزینه', 30),
    ('c26', 'تعداد ردیف‌های حسابداری مرتبط', 11), ('c27', 'تعداد اسناد حسابداری یکتا', 11), ('c28', 'مبلغ بدهکار', 18), ('c29', 'مبلغ بستانکار', 18),
    ('c30', 'سال مالی', 10), ('c31', 'نمونه شماره سند حسابداری', 22), ('c32', 'شرح نمونه سند واقعی', 45), ('c33', 'روش استخراج رابطه', 32),
    ('c34', 'منبع شواهد تطبیق', 30), ('c35', 'درجه اطمینان', 18), ('c36', 'وضعیت رابطه', 13), ('c37', 'وضعیت نیاز به تأیید کارشناسی', 30),
    ('c38', 'توضیحات و ابهامات', 50)]
WRAP = {'c14', 'c18', 'c19', 'c20', 'c22', 'c23', 'c25', 'c32', 'c33', 'c34', 'c37', 'c38', 'c4', 'c3'}
NUM = {'c26', 'c27', 'c28', 'c29'}


def write_xlsx(rows):
    wb = xlsxwriter.Workbook(XLSX, {'strings_to_numbers': False, 'strings_to_formulas': False, 'strings_to_urls': False})
    head = wb.add_format({'bold': True, 'font_name': 'Tahoma', 'font_size': 9, 'bg_color': '#1F4E78', 'font_color': 'white', 'text_wrap': True,
                          'valign': 'vcenter', 'align': 'center', 'border': 1, 'reading_order': 2})
    base = dict(font_name='Tahoma', font_size=9, valign='top', reading_order=2, border=1, border_color='#D9D9D9')
    f_txt = wb.add_format(dict(base))
    f_wrap = wb.add_format(dict(base, text_wrap=True))
    f_code = wb.add_format(dict(base, num_format='@', align='center'))
    f_num = wb.add_format(dict(base, num_format='#,##0'))
    ws = wb.add_worksheet(SHEET)
    ws.right_to_left()
    for i, (k, h, w) in enumerate(HEADERS):
        ws.set_column(i, i, w)
        ws.write_string(0, i, h, head)
    ws.set_row(0, 48)
    code_cols = {'c5', 'c7', 'c9', 'c12', 'c15', 'c21', 'c31'}
    for r, row in enumerate(rows, start=1):
        for i, (k, h, w) in enumerate(HEADERS):
            v = row.get(k, '')
            if k in NUM:
                if v == '' or v is None:
                    ws.write_blank(r, i, None, f_num)
                else:
                    ws.write_number(r, i, v, f_num)
            elif k in code_cols:
                ws.write_string(r, i, str(v), f_code) if v != '' else ws.write_blank(r, i, None, f_code)
            else:
                ws.write_string(r, i, str(v), f_wrap if k in WRAP else f_txt) if v != '' else ws.write_blank(r, i, None, f_wrap if k in WRAP else f_txt)
    n = len(rows)
    ws.freeze_panes(1, 10)
    ws.autofilter(0, 0, n, len(HEADERS) - 1)
    st_col = [h[0] for h in HEADERS].index('c36')
    cf_col = [h[0] for h in HEADERS].index('c35')
    for val, color in (('OBSERVED', '#C6EFCE'), ('PROPOSED', '#FFEB9C'), ('UNRESOLVED', '#F8CBAD'), ('APPROVED', '#BDD7EE')):
        ws.conditional_format(1, st_col, n, st_col, {'type': 'cell', 'criteria': '==', 'value': f'"{val}"',
                                                     'format': wb.add_format({'bg_color': color, 'bold': True})})
    ws.conditional_format(1, cf_col, n, cf_col, {'type': 'text', 'criteria': 'begins with', 'value': 'پایین',
                                                 'format': wb.add_format({'bg_color': '#FCE4D6', 'font_color': '#9C0006'})})
    ws.conditional_format(1, cf_col, n, cf_col, {'type': 'text', 'criteria': 'begins with', 'value': 'نامشخص',
                                                 'format': wb.add_format({'bg_color': '#EDEDED', 'font_color': '#7F7F7F'})})
    # short guide (legend) — not a QA sheet
    g = wb.add_worksheet('راهنما')
    g.right_to_left()
    g.set_column(0, 0, 28)
    g.set_column(1, 1, 140)
    gw = wb.add_format(dict(base, text_wrap=True))
    gh = wb.add_format({'bold': True, 'font_name': 'Tahoma', 'font_size': 10, 'bg_color': '#1F4E78', 'font_color': 'white', 'reading_order': 2})
    items = [
        ('موضوع', 'نگاشت جامع سرفصل حسابداری ← رویداد مالی ← فعالیت ← مرکز هزینه برای سازمان تاکسیرانی و سازمان پایانه‌های مسافربری (۱۴۰۴–۱۴۰۵). همهٔ روابط در شیت «' + SHEET + '» است.'),
        ('کلید یکتای هر ردیف', 'حوزه سازمانی + سازمان/واحد مرتبط + کد معین/تفصیلی + رویداد مالی (کد یا برچسب) + فعالیت/خدمت + مرکز هزینه (کد یا برچسب). برای پایانه‌ها واحد سازمانی (پایانه) بخشی از «سازمان/واحد» است.'),
        ('ردیف شاهد (با مبلغ)', 'هر سطر واقعی دفتر دقیقاً در یک ردیف شمرده شده است؛ جمع ستون‌های تعداد ردیف، بدهکار و بستانکار با کل دفتر برابر است.'),
        ('ردیف بدون مبلغ', 'گزینه‌های جایگزین مرکز هزینه، رویدادهای همراه اسناد ترکیبی، پیشنهادهای تحلیلی و حساب‌های بدون گردش؛ مبلغ آن‌ها عمداً خالی است تا جمع تکراری نشود.'),
        ('OBSERVED', 'رابطه با سطرهای واقعی دفتر و شناسهٔ رویداد/مرکز موجود (FE v7.1 / cc_master) پشتیبانی می‌شود؛ طبقه‌بندی سطرها موقت (v2.1) است و تأیید نشده.'),
        ('PROPOSED', 'شاهد دارد ولی شناسهٔ رویداد یا مرکز «کاندید» است (TT-NE، PEO، TTCC-P) یا رابطه پیشنهاد تحلیلی/گزینهٔ جایگزین است.'),
        ('UNRESOLVED', 'رویداد یا مرکز تعیین نشده، یا رابطه به تصمیم باز CD-01 (پایانه‌ها) وابسته است، یا حساب گردش و شاهد ندارد (NO RELATION ESTABLISHED).'),
        ('APPROVED', 'هیچ رابطه‌ای تأییدشده نیست؛ هیچ تصمیم کمیته‌ای ثبت نشده است.'),
        ('NOT APPLICABLE', 'برای طرف‌های ترازنامه‌ای، انتظامی، افتتاحیه/اختتامیه و درآمد مرکز هزینه تعریف نمی‌شود؛ مرکز هزینه روی سطر هزینهٔ همان سند است. درآمد در بُعد منبع درآمد (CD-03) دنبال می‌شود.'),
        ('کدینگ جدید', 'از crosswalk بازارزیابی‌شدهٔ ممیزی Taxi-Gate v1 (سازمان‌محور و در سطح زیرحساب برای حساب‌های تفکیکی). پیشنهادهای متنی Taxi-Gate V2 بازاستفاده نشده‌اند.'),
        ('حریم خصوصی', 'نام اشخاص و ارقام شمارهٔ حساب در عناوین و شرح‌ها ماسک شده‌اند؛ سطح چهارم حساب (حساب‌های بانکی/زیرنوع‌ها) در سطح تفصیلی تجمیع شده است. فایل برای استفادهٔ داخلی است.'),
        ('مبالغ', 'واحد پول در منبع تصریح نشده (CD-02)؛ ۱۴۰۵ ناقص است (تاکسیرانی تا ۱۴۰۵/۰۷/۱۳، پایانه‌ها تا ۱۴۰۵/۰۷/۱۴).'),
        ('منابع', 'tt_v21.sqlite (Sanad ۱۱۱٬۶۰۰ سطر، Hesab ۲٬۱۷۸ حساب)؛ expert_crosswalk.py و matching.json (ممیزی Taxi-Gate v1)؛ FE v7.1؛ v2.1 canonical_events.csv؛ cc_master.json؛ new_cost_center_candidates.json؛ Activity_Service_Map و committee_decisions_v2.1.xlsx.'),
    ]
    g.write_string(0, 0, 'مورد', gh)
    g.write_string(0, 1, 'توضیح', gh)
    for i, (a, b) in enumerate(items, start=1):
        g.write_string(i, 0, a, gw)
        g.write_string(i, 1, b, gw)
    wb.close()


if __name__ == '__main__':
    main()
