"""Build the client decision workbook from immutable independent-audit evidence.

Run locally: python scripts/build_handover_1404_expert_workbook.py
No journal mutation, enrichment rerun, or authoritative decision is performed.
"""
from __future__ import annotations

import collections
import csv
import hashlib
import json
import sqlite3
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parents[1]
AUDIT = BASE / 'data/deliverables/journal_1404_enrichment_v1_independent_audit'
FINAL = BASE / 'data/deliverables/FINAL_CLIENT_HANDOVER_1404'
BUILD = BASE / 'data/deliverables/handover_1404_build'
OUTPUT = FINAL / '03_Expert_Review_and_Decisions.xlsx'
SHEETS = ['راهنما_و_خلاصه', 'تصمیمات_گروه‌بندی‌شده', 'بازبینی_مراکز_هزینه',
          'نگاشت‌های_پیشنهادی', 'تعارض‌های_مراکز_هزینه', 'بازبینی_بانک',
          'چندتطبیقی‌ها', 'اصلاحات_پیشنهادی', 'مراجع_مفقود', 'ثبت_تصمیمات']
AUTH = ['تصمیم بازبین', 'نگاشت یا قاعده مصوب', 'نام بازبین', 'تاریخ تصمیم', 'مرجع تصمیم', 'یادداشت']
GROUP_HEADERS = ['شناسه گروه تصمیم', 'دسته مسئله', 'حوزه مالی', 'زمینه حسابداری',
                 'شرح مسئله', 'نگاشت موجود', 'گزینه‌های پیشنهادی', 'تعداد ردیف روزنامه متأثر',
                 'تعداد اسناد یکتای متأثر', 'نمونه‌های واقعی روزنامه', 'شواهد منبع',
                 'توصیه', 'اطمینان', 'نیازمند تصمیم کارشناس'] + AUTH
RAW_COLS = ['B', 'C', 'E', 'H', 'N', 'O', 'P', 'T', 'W']
RAW_TITLES = ['شناسه سند (خام B)', 'منطقه (خام C)', 'شماره سند (خام E)',
              'کد کامل حساب (خام H)', 'عنوان معین (خام N)', 'عنوان تفصیلی (خام O)',
              'عنوان جزء (خام P)', 'شرح / نوع سند (خام T)', 'کد بودجه (خام W)']
DETAIL_HEADERS = ['شناسه گروه‌های تصمیم', 'ردیف منبع روزنامه'] + RAW_TITLES + [
    'مرکز هزینه موجود (خام)', 'مرکز هزینه نسخه یک / نام بانک', 'نگاشت پیشنهادی (صرفاً پیشنهاد)',
    'نوع مسئله / وضعیت ممیزی', 'ردیف‌های مرجع', 'شواهد منبع (خام)', 'داده اصلی کامل (JSON)',
    'کد ماهیت حساب (خام X)', 'ماهیت حساب (خام Y)',
    'کد مرکز هزینه سطح یک نسخه یک (AF)', 'عنوان مرکز هزینه سطح یک نسخه یک (AG)',
    'کد مرکز هزینه سطح دو نسخه یک (AH)', 'عنوان مرکز هزینه سطح دو نسخه یک (AI)'] + AUTH
RISK_FA = {'administration': 'اداری و پشتیبانی', 'construction_infrastructure': 'عمران و زیرساخت',
           'overlapping_domains': 'همپوشانی حوزه‌ها', 'culture': 'فرهنگی و اجتماعی',
           'urban_services': 'خدمات شهری', 'transportation': 'حمل‌ونقل',
           'missing_authority': 'نبود مرجع مصوب'}
RISK_FA.update({
    'EQUIPMENT_TO_BUILDING_CENTER': 'آیا هزینه ماشین‌آلات و تجهیزات واقعاً به مرکز نگهداری ساختمان تعلق دارد؟',
    'NATIVE_OPERATIONAL_BUDGET_TO_EXECUTION_CENTER': 'آیا هزینه عملیاتی بودجه باید به مرکز اجرای عمرانی تخصیص یابد؟',
    'PURPOSE_BUDGET_DOMAIN_DIVERGENCE': 'آیا هدف واقعی ردیف با حوزه بودجه و مرکز پیشنهادی سازگار است؟',
    'NATIVE_CULTURAL_BUDGET_TO_SPORTS': 'آیا خدمت فرهنگی، تخصیص به مرکز ورزشی را توجیه می‌کند؟',
    'CONSTRUCTION_TO_BUILDING_MAINTENANCE': 'آیا هزینه ساخت به مرکز نگهداری ساختمان تعلق دارد؟',
    'OPERATIONAL_SPENDING_TO_CAPITAL_BUILDING': 'آیا هزینه جاری به مرکز سرمایه‌ای ساختمان تعلق دارد؟',
    'AFFILIATED_SUPPORT_NOT_DIRECT_EXECUTION': 'آیا حمایت از نهاد وابسته معادل اجرای مستقیم خدمت است؟',
    'BENEFITS_ACCOUNT_UNRELATED_REFERENCE': 'آیا مرجع بودجه با هزینه بیمه و مزایای این حساب مرتبط است؟',
    'CAPITAL_ACCOUNT_OPERATING_REFERENCE': 'آیا حساب سرمایه‌ای با مرجع هزینه جاری سازگار است؟',
    'CONTRACT_TRANSFER_NOT_DIRECT_PAYROLL': 'آیا پرداخت پیمانکاری یا انتقالی معادل حقوق مستقیم کارکنان است؟',
    'CULTURAL_FACILITY_TO_SPORT_REFERENCE': 'آیا مرکز فرهنگی با مرجع ورزشی هم‌دامنه است؟',
    'GENERIC_FALLBACK_DOMAIN': 'آیا مرکز عمومی پیشنهادی هدف تخصصی هزینه را پوشش می‌دهد؟',
    'NONFINANCE_TO_TREASURY_REFERENCE': 'آیا هزینه غیرمالی به مرجع خزانه و حسابداری تعلق دارد؟',
    'REFERENCE_REQUIRES_EXPERT_DECISION': 'کدام رابطه مرجع بومی با هدف ردیف قابل تأیید است؟',
    'RENTAL_RELIEF_EQUIPMENT_NOT_CONSTRUCTION': 'آیا اجاره یا کمک یا تجهیزات واقعاً هزینه ساخت است؟',
    'SPECIALIZED_DOMAIN_TO_GENERIC_PROCUREMENT': 'آیا خدمت تخصصی به مرکز عمومی تدارکات تعلق دارد؟',
    'TRANSFER_ACCOUNT_INCOMPATIBLE_REFERENCE': 'آیا حساب انتقالی با مرکز مرجع سازگار است؟',
    'AMBIGUOUS': 'از میان مراکز متفاوت مرجع، کدام با هدف واقعی ردیف سازگار است؟',
    'CONFLICT': 'آیا مرکز موجود با مرکز مرجع و هدف هزینه سازگار است؟',
    'SOURCE_CORRECTION_REQUIRES_AUTHORITY': 'آیا اصلاح مرکز موجود با تأیید مالک منبع مجاز است؟',
    'REVIEW_ACCOUNT_TITLE_NO_MATCHING_BANK_NAME_TOKEN': 'آیا نام بانک با عنوان حساب و مرجع رسمی بانک قابل تأیید است؟'})
CASE_FA = {'cc_review': 'بازبینی تخصیص جدید مرکز هزینه', 'proposed': 'نگاشت پیشنهادی هزینه',
           'conflict': 'تعارض مرکز هزینه موجود', 'bank': 'بازبینی نام بانک',
           'ambiguous': 'چندتطبیقی مرجع', 'correction': 'اصلاح پیشنهادی مرکز هزینه'}
EXPECTED = {'cc_review': 884, 'proposed': 6977, 'conflict': 8042, 'bank': 245, 'ambiguous': 31, 'correction': 51}


def csv_rows(name):
    with (AUDIT / name).open(encoding='utf-8-sig', newline='') as handle:
        return list(csv.DictReader(handle))


def parsed(value, default=None):
    if not value:
        return default
    try:
        return json.loads(value)
    except (TypeError, json.JSONDecodeError):
        return value


def compact(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'), sort_keys=True)


def text_cell(ws, row):
    # Explicit strings prevent formula injection and preserve zero-prefixed codes.
    for col, value in enumerate(row, 1):
        if isinstance(value, str) and len(value) > 32767:
            raise ValueError(f'Excel cell overflow: {ws.title}, row {ws._current_row + 1}, column {col}')
    ws.append(row)
    for col, value in enumerate(row, 1):
        cell = ws.cell(ws._current_row, col)
        if isinstance(value, str):
            cell.data_type = 's'
            cell.number_format = '@'


def decorate(ws, table_index, auth_start=None):
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 42
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(name='Tahoma', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True, readingOrder=2)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='Tahoma', size=10)
            cell.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True, readingOrder=2)
        ws.row_dimensions[row[0].row].height = 80 if ws.title == SHEETS[0] else 100
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 24 if col < 5 else 38
    if ws.title in SHEETS[2:8]:
        ws.column_dimensions.group('Q', 'R', hidden=True)
    ws.column_dimensions['A'].width = 31
    if ws.max_row > 1:
        table = Table(displayName=f'Handover{table_index}', ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
    if auth_start:
        letter = get_column_letter(auth_start)
        dv = DataValidation(type='list', formula1='"تأیید,رد,نیازمند مدرک,ارجاع به مرجع صلاحیت‌دار"', allow_blank=True)
        dv.errorTitle = 'انتخاب تصمیم'; dv.error = 'یکی از گزینه‌ها را انتخاب کنید.'; dv.showErrorMessage = True
        ws.add_data_validation(dv); dv.add(f'{letter}2:{letter}{ws.max_row}')
        ws.conditional_formatting.add(f'{letter}2:{letter}{ws.max_row}', FormulaRule(
            formula=[f'{letter}2="نیازمند مدرک"'], fill=PatternFill('solid', fgColor='FFF2CC')))
        ws.conditional_formatting.add(f'{letter}2:{letter}{ws.max_row}', FormulaRule(
            formula=[f'{letter}2="رد"'], fill=PatternFill('solid', fgColor='FCE4D6')))


def build():
    FINAL.mkdir(parents=True, exist_ok=True); BUILD.mkdir(parents=True, exist_ok=True)
    population = csv_rows('full_population_findings.csv')
    cc = [r for r in population if r['assignment_origin'].startswith('NEW')]
    coverage = csv_rows('Cost_Center_Coverage_Reconciliation.csv')
    datasets = {'cc_review': cc,
                'proposed': [r for r in coverage if r['final_status'] == 'PROPOSED_MATCH'],
                'conflict': [r for r in coverage if r['final_status'] == 'CONFLICT'],
                'bank': csv_rows('Bank_Assignment_Exceptions.csv'),
                'ambiguous': [r for r in coverage if r['final_status'] == 'AMBIGUOUS'],
                'correction': csv_rows('proposed_cc_corrections.csv')}
    counts = {key: len(rows) for key, rows in datasets.items()}
    assert counts == EXPECTED, counts
    source_ids = sorted({int(r['source_row']) for rows in datasets.values() for r in rows})
    originals = {}; enriched = {}
    with sqlite3.connect(f'file:{AUDIT / "journal_records.sqlite"}?mode=ro', uri=True) as db:
        for start in range(0, len(source_ids), 500):
            chunk = source_ids[start:start + 500]
            for source_row, original, final in db.execute(
                    f'SELECT source_row, original, enriched FROM journal WHERE source_row IN ({",".join("?" for _ in chunk)})', chunk):
                originals[source_row] = json.loads(original); enriched[source_row] = json.loads(final)
    assert len(originals) == len(source_ids)
    records = []; groups = {}
    for case, rows in datasets.items():
        for finding in rows:
            rid = int(finding['source_row']); raw = originals[rid]
            current = finding.get('original_cc', finding.get('current_cc', compact({c: raw.get(c) for c in ['AF','AG','AH','AI']})))
            proposal = finding.get('candidate_cc', '')
            final_cc = finding.get('enriched_cc', finding.get('new_value', compact({c: enriched[rid].get(c) for c in ['AF','AG','AH','AI']})))
            refs = finding.get('candidate_rows', finding.get('native_reference_rows', finding.get('account_reference_rows', '')))
            native = parsed(finding.get('native_budget_references'), [])
            if native:
                refs = compact(sorted({str(n.get('reference_row', n.get('row', ''))) for n in native}))
                proposal = compact(sorted({compact({k: n.get(k) for k in ['l1_code','l1_name','l2_code','l2_name','domain']}) for n in native}))
            evidence = compact(finding)
            if case == 'cc_review':
                categories = parsed(finding.get('applicable_rules'), []) or ['missing_authority']
            elif case in ['proposed','conflict','ambiguous']:
                gates = parsed(finding.get('candidate_gate_decisions'), [])
                categories = sorted({g.get('reason') or finding['final_status'] for g in gates
                                     if g.get('status') != 'V1_GATE_ELIGIBLE'}) or [finding['final_status']]
            elif case == 'bank':
                categories = [finding['account_detail_consistency']]
            else:
                categories = ['SOURCE_CORRECTION_REQUIRES_AUTHORITY']
            member = {'case': case, 'id': rid, 'raw': raw, 'current': current, 'final': final_cc,
                      'proposal': proposal, 'refs': refs, 'evidence': evidence, 'groups': []}
            for category in categories:
                # Reuse a decision across budget codes and documents sharing semantic risk/context.
                # Exact budget/reference codes remain on every detail record and in group evidence.
                key = compact([case, category, raw.get('I'), raw.get('N'), current, proposal, final_cc])
                gid = 'DG-' + hashlib.sha256(key.encode('utf8')).hexdigest()[:16].upper()
                member['groups'].append(gid)
                if gid not in groups:
                    groups[gid] = {'case': case, 'category': category, 'members': [], 'context': key}
                groups[gid]['members'].append(member)
            records.append(member)
    wb = Workbook(); wb.remove(wb.active)
    for name in SHEETS: wb.create_sheet(name)
    detail_sheets = dict(zip(EXPECTED, SHEETS[2:8]))
    detail_locations = {}
    for case, sheet in detail_sheets.items():
        ws = wb[sheet]; text_cell(ws, DETAIL_HEADERS)
        for rec in (r for r in records if r['case'] == case):
            row = ['; '.join(rec['groups']), rec['id']] + [rec['raw'].get(c) for c in RAW_COLS] + [
                rec['current'], rec['final'], rec['proposal'], CASE_FA[case], rec['refs'],
                rec['evidence'], compact(rec['raw']), rec['raw'].get('X'), rec['raw'].get('Y')
                ] + [enriched[rec['id']].get(c) for c in ['AF','AG','AH','AI']] + [None] * 6
            text_cell(ws, row); detail_locations[(case, rec['id'])] = (sheet, ws._current_row)
    group_ws = wb[SHEETS[1]]; text_cell(group_ws, GROUP_HEADERS)
    group_locations = {}
    for gid, group in sorted(groups.items()):
        members = group['members']; first = members[0]; raw = first['raw']
        ids = sorted({m['id'] for m in members})
        docs = {(str(m['raw'].get('C')), str(m['raw'].get('E'))) for m in members}
        references = sorted({m['refs'] for m in members if m['refs']})
        budget_codes = sorted({str(m['raw'].get('W')) for m in members})
        examples = '; '.join(f'ردیف {m["id"]} | C={m["raw"].get("C")} | E={m["raw"].get("E")}' for m in members[:3])
        question = f'{CASE_FA[group["case"]]}؛ علت ممیزی: {RISK_FA.get(group["category"], group["category"])}. تعیین رابطه مجاز فقط پس از بررسی هدف واقعی ردیف و مرجع مصوب.'
        recommendation = ('مالک حسابداری و بودجه و متولی مرکز هزینه، دامنه و شروط و محدودیت‌های قاعده را با توجه به هدف واقعی هر عضو تأیید کنند؛ تأیید گروه مجوز تخصیص یکسان و بی‌قید به همه ردیف‌ها نیست؛ نسخه یک تغییر نکند.'
                          if group['case'] != 'bank' else 'مالک خزانه، بانک و حساب را با مرجع رسمی بانک و شناسه حساب تأیید کند؛ نام موجود معادل شناسه مصوب بانک نیست.')
        text_cell(group_ws, [gid, CASE_FA[group['case']] + ' / ' + RISK_FA.get(group['category'], group['category']),
                            str(raw.get('F', '')) + ' | ' + str(raw.get('Y', '')),
                            compact({'I': raw.get('I'), 'N': raw.get('N'), 'کدهای بودجه اعضا': budget_codes}), question,
                            first['current'], first['proposal'] or 'گزینه‌ها از شواهد ردیف‌های تفصیلی بررسی شود؛ مصوب نیست',
                            len(ids), len(docs), examples,
                            f'ممیزی مستقل؛ مراجع بومی: {compact(references)}؛ کدهای W: {compact(budget_codes)}؛ شواهد خام کامل در تفصیل',
                            recommendation, 'شاهد ممیزی؛ اعتبار کسب‌وکار تصویب نشده', 'بله'] + [None] * 6)
        group_locations[gid] = group_ws._current_row
        sheet, rownum = detail_locations[(first['case'], first['id'])]
        group_ws.cell(group_ws._current_row, 10).hyperlink = f"#'{sheet}'!A{rownum}"
    # Non-row governance questions are distinct from financial mapping rules.
    governance = json.loads((AUDIT/'governance_review.json').read_text(encoding='utf8'))
    missing_ws = wb[SHEETS[8]]
    text_cell(missing_ws, ['مرجع یا قاعده مفقود', 'هدف استفاده', 'حداقل فیلدهای لازم', 'مالک مرجع (نیازمند تأیید)',
                           'مسئله', 'اولویت', 'پیامد نبود مرجع'])
    for item in governance['authority_evidence']['missing_registered_sources']:
        text_cell(missing_ws, [item['filename'], item['purpose'],
            'شناسه پایدار، عنوان، نسخه، وضعیت اعتبار، تاریخ اثر، ارتباط با ردیف / سند و مرجع تصویب',
            'مالک داده و مرجع صلاحیت‌دار باید تعیین و تأیید شوند',
            'مرجع ثبت‌شده در ممیزی در دسترس / مصوب نیست', 'بالا', 'تخصیص خودکار یا تأیید قطعی مجاز نیست'])
    for title, purpose, minimum, consequence in [
        ('قواعد مصوب رویداد مالی و ابعاد', 'الزام / منع / امکان ابعاد برای هر زمینه', 'شناسه قاعده، دامنه، نسخه، تاریخ اثر، سند تصویب', 'نبود جدول، منع صریح نیست؛ تحویل قطعی بدون قاعده مجاز نیست'),
        ('قواعد کاربرد قرارداد و پروژه', 'تشخیص کاربرد ابعاد پیش از تخصیص', 'زمینه، الزام، شناسه مرجع، نگاشت ردیفی تأییدشده', 'خالی بودن، عدم کاربرد را اثبات نمی‌کند'),
        ('مجوز محرمانگی و توزیع', 'کنترل گیرندگان و نگهداری اطلاعات شخصی و مالی', 'گیرندگان، سطح دسترسی، حفاظت، مدت نگهداری، مرجع مجوز', 'توزیع نامحدود تا مجوز مستند ممنوع است'),
        ('فرایند تصویب اصلاح و استثنا', 'تفکیک اصلاح منبع، استثنا و پیشنهاد', 'نوع اختیار، مالک، دامنه، نسخه، شناسه ردیف و مستند', 'نظر کارشناس به‌تنهایی مرجع قطعی حسابداری نیست')]:
        text_cell(missing_ws, [title, purpose, minimum, 'مالک داده و مقام تصویب باید تأیید شوند', 'نیازمند تعیین و تصویب', 'بالا', consequence])
    for offset, title in enumerate(['تأمین مراجع پایه و نگاشت‌های ردیفی', 'تصویب قواعد کاربرد ابعاد و روابط مالی', 'تصویب حفاظت و مجوز توزیع'], 1):
        gid = f'GOV-{offset:03d}'
        text_cell(group_ws, [gid, 'حاکمیت / مرجع مفقود', 'سراسری', 'مراجع مفقود؛ فاقد انتساب ردیفی تأییدشده', title,
                           'مرجع مصوب در ممیزی احراز نشده', 'تأمین مرجع رسمی و تعیین دامنه', None, None,
                           'دامنه ردیفی / اسناد از فقدان مرجع قابل استنتاج نیست', 'governance_review.json؛ شیت مراجع مفقود',
                           'مرجع صلاحیت‌دار مستندات نسخه‌دار را بررسی و تصویب کند', 'نبود مرجع؛ نه تأیید رابطه', 'بله'] + [None]*6)
        group_ws.cell(group_ws._current_row, 11).hyperlink = f"#'{SHEETS[8]}'!A2"
    for rec in records:
        sheet, rownum = detail_locations[(rec['case'], rec['id'])]
        wb[sheet].cell(rownum, 1).hyperlink = f"#'{SHEETS[1]}'!A{group_locations[rec['groups'][0]]}"
    register = wb[SHEETS[9]]
    text_cell(register, ['شناسه تصمیم', 'شناسه گروه‌های مرتبط', 'بازبین', 'تصمیم', 'کد یا قاعده مصوب',
                         'تاریخ اثر', 'مرجع تصویب', 'وضعیت اجرا', 'یادداشت'])
    for _ in range(100): text_cell(register, [None]*9)
    dv = DataValidation(type='list', formula1='"اجرا نشده,نیازمند مجوز,اجرا شده با مستندات"', allow_blank=True)
    register.add_data_validation(dv); dv.add('H2:H101')
    memberships = sum(len(rec['groups']) for rec in records)
    summary = wb[SHEETS[0]]; text_cell(summary, ['موضوع', 'مقدار / راهنمای استفاده'])
    summary_rows = [
        ('ماهیت فایل', 'بسته بازبینی کارشناسی بر پایه ممیزی مستقل نسخه یک؛ هیچ پیشنهاد یا شاهد، تصویب مالی نیست.'),
        ('برچسب دفتر روزنامه', 'نسخه تکمیل‌شده پیشنهادی — نیازمند بازبینی و تأیید کارشناسی'),
        ('حفظ داده اصلی', 'نسخه یک و دفتر اصلی تغییر نمی‌کنند؛ B/C/E/H/N/O/P/T/W از JSON اصلی بدون نرمال‌سازی آمده‌اند.'),
        ('شناسه سند', 'سند یکتا برابر زوج دقیق C و E است؛ صفرهای ابتدایی حفظ شده‌اند.'),
        ('گروه‌بندی', 'گروه بر اساس ریسک، زمینه معنایی حساب و گزینه مرکز هزینه ساخته می‌شود؛ کد بودجه و ردیف مرجع، شواهد اعضا هستند و به‌تنهایی گروه جدید نمی‌سازند.'),
        ('چندعضویتی', 'هر ردیف می‌تواند چند ریسک داشته باشد؛ همه گروه‌ها در ستون نخست تفصیل با ; آمده‌اند. لینک، نخستین گروه را باز می‌کند.'),
        ('شواهد', 'شواهد خام کامل هر مورد و JSON اصلی در تفصیل موجود است؛ نمونه‌های گروه از ردیف واقعی آمده‌اند.'),
        ('تصمیم و تصویب', 'ستون‌های تصمیم، قاعده مصوب، بازبین، تاریخ، مرجع و یادداشت عمداً خالی‌اند؛ فهرست کشویی تصمیم پیش‌فرض ندارد.'),
        ('ثبت تصمیمات', '۱۰۰ ردیف خالی آماده؛ تصمیم می‌تواند به چند گروه ارجاع دهد. از تکرار تصمیم مالی در پرونده‌های مشابه خودداری کنید.'),
        ('ایمنی', 'نظر بازبین باید با مرجع صلاحیت‌دار و نسخه قاعده ثبت شود؛ فایل، دستور اعمال اصلاح یا تحویل قطعی حسابداری نیست.'),
        ('محرمانگی', 'فایل شامل شواهد مالی و داده خام است؛ فقط گیرندگان مجاز و ذخیره‌سازی حفاظت‌شده.'),
        ('تعداد گروه‌های مالی', len(groups)), ('تعداد پرسش‌های حاکمیتی', 3),
        ('کل گروه‌های تصمیم', len(groups)+3), ('ردیف‌های منبع یکتای مرتبط', len(source_ids)),
        ('تعداد موارد تفصیلی (با تکرار بین شیت‌ها)', len(records)), ('تعداد عضویت مورد در گروه', memberships),
        ('عضویت یکتای ردیف منبع در گروه', len({(r['id'], g) for r in records for g in r['groups']}))]
    summary_rows.extend((CASE_FA[key], value) for key,value in counts.items())
    for row in summary_rows: text_cell(summary, row)
    summary.column_dimensions['B'].width = 110
    for index, ws in enumerate(wb.worksheets, 1):
        auth_start = 15 if ws.title == SHEETS[1] else (len(DETAIL_HEADERS)-5 if ws.title in detail_sheets.values() else (4 if ws.title == SHEETS[9] else None))
        decorate(ws, index, auth_start)
    summary.column_dimensions['B'].width = 110
    wb.save(OUTPUT)
    return verify(OUTPUT, counts, records, groups, originals)


def verify(path, counts, records, groups, originals):
    """Independently reopen the delivered artifact and reconcile actual cells."""
    wb = load_workbook(path, read_only=False, data_only=False)
    assert wb.sheetnames == SHEETS
    by_group = collections.defaultdict(set); docs_by_group = collections.defaultdict(set)
    physical = {}; actual = {}; blank_checked = 0
    case_sheets = dict(zip(EXPECTED, SHEETS[2:8]))
    for ws in wb.worksheets:
        physical[ws.title] = {'physical_rows': ws.max_row, 'data_rows': ws.max_row - 1, 'columns': ws.max_column}
        assert ws.sheet_view.rightToLeft and ws.freeze_panes == 'C2'
        assert ws.auto_filter.ref == ws.dimensions
        assert len(ws.tables) == 1 and next(iter(ws.tables.values())).ref == ws.dimensions
    for case, sheet in case_sheets.items():
        ws = wb[sheet]; rows = list(ws.iter_rows(min_row=2, values_only=True)); actual[case] = len(rows)
        physical[sheet]['unique_source_rows'] = len({row[1] for row in rows})
        physical[sheet]['unique_documents_C_E'] = len({(row[3], row[4]) for row in rows})
        assert len(ws.data_validations.dataValidation) == 1
        for rownum, row in enumerate(rows, 2):
            source_id = row[1]; raw = originals[source_id]
            assert list(row[2:11]) == [raw.get(c) for c in RAW_COLS]
            assert parsed(row[17]) == raw
            assert list(row[18:20]) == [raw.get('X'), raw.get('Y')]
            assert ws.cell(rownum, 1).hyperlink is not None
            for idx, col in enumerate(RAW_COLS, 2):
                if isinstance(raw.get(col), str): assert ws.cell(rownum, idx+1).data_type == 's'
            assert all(value is None for value in row[-6:]); blank_checked += 6
            for gid in row[0].split('; '):
                by_group[gid].add(source_id)
                docs_by_group[gid].add((str(raw.get('C')), str(raw.get('E'))))
    assert actual == EXPECTED
    for row in wb[SHEETS[1]].iter_rows(min_row=2, values_only=True):
        assert all(value is None for value in row[14:20]); blank_checked += 6
        if row[0].startswith('DG-'):
            assert row[7] == len(by_group[row[0]])
            assert row[8] == len(docs_by_group[row[0]])
            assert by_group[row[0]] == {m['id'] for m in groups[row[0]]['members']}
    assert set(by_group) == set(groups)
    register = wb[SHEETS[9]]
    assert all(cell.value is None for row in register.iter_rows(min_row=2) for cell in row)
    assert len(register.data_validations.dataValidation) == 2
    assert len(wb[SHEETS[1]].data_validations.dataValidation) == 1
    metrics = {'artifact': str(path.relative_to(BASE)), 'counts': actual,
               'financial_group_count': len(groups), 'governance_group_count': 3,
               'grouped_decision_count': len(groups)+3,
               'unique_linked_records': len({rid for ids in by_group.values() for rid in ids}),
               'detail_records': sum(actual.values()),
               'memberships': sum(len(r['groups']) for r in records),
               'unique_source_group_memberships': sum(len(ids) for ids in by_group.values()),
               'sheets': physical, 'verification': {'reopened_artifact': True, 'exact_10_sheets': True,
                   'authoritative_cells_checked_blank': blank_checked, 'decision_register_blank_cells': 900,
                   'exact_original_raw_fields_and_json': True, 'leading_zero_string_type_preserved': True,
                   'exact_original_account_nature_X_Y': True,
                   'group_member_and_document_counts_reconciled': True,
                   'actual_counts_independently_reconciled': actual == EXPECTED,
                   'rtl_freeze_tables_filters_validations': True}}
    (BUILD/'workbook_build_metrics.json').write_text(json.dumps(metrics,ensure_ascii=False,indent=2),encoding='utf8')
    print(json.dumps(metrics,ensure_ascii=False,indent=2))
    return metrics


if __name__ == '__main__':
    build()
