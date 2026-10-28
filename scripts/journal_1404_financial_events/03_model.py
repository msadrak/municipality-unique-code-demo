"""Deliverable A: Unique_Code_Financial_Event_Relationship_Model_DRAFT.xlsx
Event × account × dimension relationship model for client approval, built from the canonical v7.1 dictionary,
the activity-based CC model (service coverage v2 + v4 activity crosswalk) and the provisional 1404 journal application."""
import pickle, json, re, ast, sys, statistics
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
sys.path.insert(0, 'H:/Freelancing_Project/KalaniProject/municipality_demo/scripts/journal_1404_financial_events')
import fe_rules as F
g, s, num = F.g, F.s, F.num

ROOT = 'H:/Freelancing_Project/KalaniProject/municipality_demo/'
W = ROOT + 'data/work/fe_model/'
OUT = 'C:/Users/sadra/Desktop/Unique_Code_Financial_Event_Relationship_Model_DRAFT.xlsx'
ROWS = pickle.load(open(W + 'journal_v4.pkl', 'rb'))['rows']
K = pickle.load(open(W + 'classified.pkl', 'rb'))
REC, REL, EV, DOCINFO = K['REC'], K['REL'], K['EV'], K['DOCINFO']
SRC = json.load(open(W + 'model_sources.json', encoding='utf-8'))
DICT = {r[0]: r for r in SRC['dictionary'][1:] if r[0]}
DH = SRC['dictionary'][0]
COMM = {r[0]: r for r in SRC['committee'][1:] if r[0]}
MATRIX = defaultdict(list)
for r in SRC['matrix'][1:]:
    if r[0]: MATRIX[r[0]].append(f'{r[2]} ({r[4]} سطر؛ {r[15]})')
XW = ast.literal_eval(re.search(r'XW = (\{.*?\n\})', open(ROOT + 'scripts/journal_1404_activity_cc_v4/04_classify.py', encoding='utf-8').read(), re.S).group(1))
CCM = SRC['cc_master']
N = len(ROWS)

APP_FA = {F.R_: 'REQUIRED_CANDIDATE', F.C_: 'CONDITIONAL_CANDIDATE', F.O_: 'OPTIONAL_CANDIDATE', F.N_: 'NOT_APPLICABLE_CANDIDATE', F.U_: 'UNKNOWN'}
CANON = sorted([k for k in EV if EV[k]['canonical']], key=lambda k: int(k[3:]))
GAPS = sorted([k for k in EV if not EV[k]['canonical']])
ALL = CANON + GAPS
def definition(e):
    return DICT[e][3] if e in DICT else F.GAPS[e][4]

# ------------------------------------------------------------ dimension logic text (proposals)
def cc_logic(e, dm):
    a = dm[0]
    if a == F.N_: return 'مرکز هزینه هزینه‌ای به این رویداد نسبت داده نمی‌شود (تأمین مالی/انتقال/درآمد/انتظامی). مقدار موجود در AF:AI سطرهای درآمدی فقط علامت‌گذاری می‌شود.'
    base = {F.R_: 'سطر محرک (هزینه/سرمایه‌ای) باید مرکز هزینه داشته باشد؛ مأخذ: AF:AI بومی سطر یا تعریف فعالیت ردیف بودجه/عنوان جزء. ',
            F.C_: 'مرکز هزینه فقط وقتی فعالیت زیربنایی قابل تشخیص است (عنوان جزء/تفضیلی، قرارداد، پروژه) در سطح رویداد ثبت می‌شود؛ ثبت آن به معنی شناسایی هزینه نیست. '}.get(a, '')
    if e in ('FE-01', 'FE-02', 'GAP-07', 'FE-37'): base += 'پیش‌پرداخت/علی‌الحساب مرحله مقدماتی هزینه آینده است؛ ارتباط با مرکز هزینه در سطح رویداد مجاز است ولی سطر همچنان دارایی (پیش‌پرداخت) می‌ماند.'
    if e in ('FE-07', 'FE-08', 'FE-09', 'FE-10', 'FE-11', 'FE-48', 'FE-49'): base += 'کسور تابع رویداد اصلی سند (صورت‌وضعیت/هزینه/حقوق) است و مرکز هزینه را فقط در سطح رویداد از ردیف محرک همان سند می‌گیرد.'
    if e in ('FE-03', 'GAP-10', 'FE-08', 'FE-10'): base += 'در سند پرداخت/استرداد، مرکز هزینه از ردیف هزینه همان شماره قرارداد (همان منطقه) در سند ثبت صورت‌وضعیت پیوند می‌شود (سطح رویداد، بین‌سندی). '
    if e == 'FE-67': base += 'طبق نظر کارفرما تعهدات سنواتی وقتی عنوان جزء فعالیت مشخصی دارد به مرکز هزینه فعالیت مرتبط می‌شود.'
    if e == 'FE-70': base += 'اسناد اصلاحی/افتتاحیه/اختتامیه مرکز هزینه جدید ایجاد نمی‌کنند؛ ابعاد از تراکنش اصلی به ارث می‌رسد.'
    return base
def con_logic(e, dm):
    a = dm[1]
    if a == F.N_: return 'قرارداد کاربرد ندارد.'
    t = 'شماره قرارداد فقط از متن واقعی منبع (عنوان تفضیلی/جزء/شرح سند) استخراج می‌شود؛ شماره ساخته نمی‌شود. در بودجه مستمر و غیرمستمر هر دو کاربرد دارد. '
    if e in ('GAP-01', 'FE-01', 'FE-02', 'FE-03', 'FE-07', 'FE-08', 'FE-09', 'FE-10'): t += 'حساب پیمانکار/سپرده بدون شماره → «قرارداد اشاره‌شده، شماره نامعلوم».'
    if a == F.C_: t += 'فقط در صورت وجود قرارداد واقعی (مثلاً قرارداد اجاره، تأمین نیرو، وام) پر می‌شود.'
    return t
def bank_logic(e, dm):
    a = dm[2]
    if e == 'FE-68': return 'حداقل دو حساب بانکی (مبدأ و مقصد) یا حساب رابط تمرکز وجوه؛ حساب بانکی از عنوان جزء سطر بانک (نوع حساب + شماره + شعبه) خوانده می‌شود.'
    if a == F.N_: return 'رویداد غیرنقدی است (شناسایی/کسور/بستن حساب)؛ بانک به آن نسبت داده نمی‌شود.'
    if a == F.R_: return 'وقتی وجه در همان سند جابه‌جا می‌شود حساب بانکی سطر بانک سند به رویداد نسبت داده می‌شود؛ اگر سند بانک ندارد → «بانک در سند نیست» (کمبود اطلاعات، نه عدم کاربرد).'
    return 'فقط وقتی پرداخت/دریافت در همان سند رخ داده باشد؛ در غیر این صورت ثبت نمی‌شود.'
def proj_logic(e, dm):
    a = dm[3]
    if a == F.N_: return 'پروژه کاربرد ندارد.'
    return ('قاعده کارفرما: پروژه برای بودجه‌های غیرمستمر. ردیف بودجه با نوع «سرمایه‌ای غیرمستمر» در مرجع بودجه منبع = شناسه پروژه (کد ردیف بودجه). '
            '«سرمایه‌ای مستمر» و «جاری» پروژه نیستند. «سرمایه‌ای» تفکیک‌نشده و کد 31 (تعهد سنواتی) → کاندید پروژه نیازمند اطلاعات منبع. '
            'پیش‌پرداخت پروژه‌ها (14001) → پروژه اشاره‌شده بدون شناسه. نام محل یا خدمت عمومی شناسه پروژه نیست.')

PROP_CARD = {  # analytical proposals (not approved): account, cc, contract, bank, project, person
    'default': ('1..N', '0..N', '0..1', '0..N', '0..1', '1..N'),
    'FE-68': ('2..N', '0..0', '0..0', '1..N', '0..0', '0..0'),
    'FE-70': ('1..N', '0..N', '0..N', '0..N', '0..N', '0..N'),
    'FE-44': ('1..N', '1..N', '0..1', '0..N', '0..0', '1..N'),
    'GAP-01': ('1..N', '1..N', '1..1', '0..0', '0..1', '1..1'),
    'FE-03': ('1..N', '1..N', '1..1', '1..N', '0..1', '1..1'),
    'FE-01': ('1..N', '0..N', '1..1', '1..N', '0..1', '1..1'),
}
def prop_card(e):
    if e in PROP_CARD: return PROP_CARD[e]
    dm = F.dims(e)
    c = list(PROP_CARD['default'])
    for k, a in zip(range(1, 5), dm[:4]):
        if a == F.N_: c[k] = '0..0'
        elif a == F.R_ and k in (1,): c[k] = '1..N'
        elif a == F.R_ and k == 2: c[k] = '1..1'
        elif a == F.R_ and k == 3: c[k] = '1..N'
        elif a == F.R_ and k == 4: c[k] = '1..1'
    if dm[4] == F.N_: c[5] = '0..0'
    return tuple(c)

# ------------------------------------------------------------ journal evidence per event and per pattern
inst_lines = defaultdict(list)
for i, rec in enumerate(REC):
    if rec['ev']: inst_lines[(rec['doc'], rec['ev'])].append(i)
ev_rows = Counter(r['ev'] for r in REC if r['ev'])
ev_docs = Counter(e for (d, e) in inst_lines)
def card_obs(e):
    acc, cc, con, bank, proj, per = [], [], [], [], [], []
    for (d, ee), ix in inst_lines.items():
        if ee != e: continue
        acc.append(len({s(g(ROWS[i], F.C_MOIN)) or s(g(ROWS[i], F.C_KOL)) for i in ix}))
        cc.append(len({REC[i]['cc_direct'] or REC[i]['cc_event'] for i in ix if (REC[i]['cc_direct'] or REC[i]['cc_event'])}))
        con.append(len({REC[i]['con_value'] for i in ix if REC[i]['con_status'].startswith('CONTRACT_IDENTIFIED')}))
        bank.append(len({REC[i]['bank_value'] for i in ix if REC[i]['bank_status'] == 'BANK_DIRECT_LINE'}))
        proj.append(len({REC[i]['proj_value'] for i in ix if REC[i]['proj_status'].startswith('PROJECT_IDENTIFIED')}))
        per.append(len({s(g(ROWS[i], F.C_PER)) for i in ix if s(g(ROWS[i], F.C_PER)) not in ('', 'NULL')}))
    def f(v):
        if not v: return 'بدون مشاهده'
        return f'{min(v)}..{max(v)} (میانه {statistics.median(v):g}؛ {len(v)} نمونه رویداد)'
    return [f(v) for v in (acc, cc, con, bank, proj, per)]
CARD_OBS = {e: card_obs(e) for e in ALL}

pat = defaultdict(lambda: dict(rows=0, docs=set(), D=0.0, C=0.0, st=Counter(), cc=Counter(), bt=Counter(), con=Counter(), proj=Counter(), bank=Counter(), role=Counter(), method=Counter(), ex=[]))
def raw_key(i, rec):
    r = ROWS[i]
    cls = s(g(r, F.C_CLS)); kol = s(g(r, F.C_KOL)); moin = s(g(r, F.C_MOIN))
    if moin in ('', '#N/A'): moin = f"قدیم {s(g(r, F.C_OKOL))}-{s(g(r, F.C_OMOIN))}"
    defining = rec['kind'] in ('DRIVER', 'WITHHOLD')
    kind = 'محرک' if rec['kind'] == 'DRIVER' else 'کسور/تابع' if rec['kind'] == 'WITHHOLD' else 'جزء/طرف تسویه'
    if rec['ev'] == 'FE-70': return (rec['ev'], cls, '*', '*', rec['side'], kind), defining
    return (rec['ev'], cls, kol, moin if defining else '*', rec['side'], kind), defining
KEYN = Counter(raw_key(i, rec)[0] for i, rec in enumerate(REC) if rec['ev'])
for i, rec in enumerate(REC):
    e = rec['ev']
    if not e: continue
    r = ROWS[i]
    key, defining = raw_key(i, rec)
    if KEYN[key] < 10 and key[3] != '*': key = key[:3] + ('سایر',) + key[4:]   # rare account patterns merged per general account
    p = pat[key]
    p['rows'] += 1; p['docs'].add(rec['doc']); p['D'] += num(g(r, F.C_DR)); p['C'] += num(g(r, F.C_CR))
    p['st'][rec['status']] += 1; p['role'][rec['role']] += 1; p['method'][rec['basis'].split(' | ')[0][:90] if defining else rec['method']] += 1
    if rec['cc_direct'] or rec['cc_event']:
        p['cc'][(rec['cc_direct'] or rec['cc_event']).split(' || ')[0][:70]] += 1
    if rec['btype']: p['bt'][rec['btype']] += 1
    p['con'][rec['con_status']] += 1; p['proj'][rec['proj_status']] += 1; p['bank'][rec['bank_status']] += 1
    if len(p['ex']) < 3: p['ex'].append(f"سطر {r[0]} / سند {rec['doc']}")
KOLT = {}; MOINT = {'*': '', 'سایر': '(سایر معین‌های کم‌تکرار این کل)'}; CLST = {}
for r in ROWS:
    KOLT.setdefault(s(g(r, F.C_KOL)), s(g(r, F.C_KOLN))); CLST.setdefault(s(g(r, F.C_CLS)), s(g(r, F.C_CLST)))
    m = s(g(r, F.C_MOIN))
    if m in ('', '#N/A'): MOINT.setdefault(f"قدیم {s(g(r, F.C_OKOL))}-{s(g(r, F.C_OMOIN))}", s(g(r, F.C_MOINT)))
    else: MOINT.setdefault(m, s(g(r, F.C_MOINN)))
KOLT['*'] = '(همه حساب‌ها)'

def top(c, n=3, total=None):
    tot = total or sum(c.values())
    return '؛ '.join(f'{k} ({v:,}{"، " + str(round(100 * v / tot)) + "%" if tot else ""})' for k, v in c.most_common(n)) if c else ''

# ------------------------------------------------------------ workbook helpers
HFILL = PatternFill('solid', fgColor='1F4E79'); HFONT = Font(bold=True, color='FFFFFF', name='B Nazanin', size=11)
BFONT = Font(name='B Nazanin', size=11); BOLD = Font(name='B Nazanin', size=11, bold=True)
SEC = PatternFill('solid', fgColor='DDEBF7'); YEL = PatternFill('solid', fgColor='FFF2CC'); GAPF = PatternFill('solid', fgColor='FCE4D6'); GRN = PatternFill('solid', fgColor='E2EFDA')
WRAP = Alignment(wrap_text=True, vertical='top', horizontal='right', readingOrder=2)
thin = Side(style='thin', color='BFBFBF'); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = openpyxl.Workbook()
def sheet(title, head, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.sheet_view.rightToLeft = True
    ws.append(head)
    for k, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(k)].width = w
        c = ws.cell(1, k); c.fill = HFILL; c.font = HFONT; c.alignment = WRAP; c.border = BORDER
    ws.row_dimensions[1].height = 48
    ws.freeze_panes = 'C2'
    return ws
def style_rows(ws, fill_fn=None):
    for row in ws.iter_rows(min_row=2):
        f = fill_fn(row) if fill_fn else None
        for c in row:
            c.font = BFONT; c.alignment = WRAP; c.border = BORDER
            if f: c.fill = f

# ------------------------------------------------------------ 1. guide
ws = sheet('راهنما', ['موضوع', 'توضیح'], [34, 150], first=True)
ws.freeze_panes = 'A2'
GUIDE = [
 ('عنوان', 'مدل یکپارچه رویداد مالی و ابعاد حسابداری (حساب، مرکز هزینه، قرارداد، بانک، پروژه، شخص، واحد سازمانی) — پروژه کد یکتا / پیش‌حسابداری شهرداری اصفهان'),
 ('وضعیت سند', 'پیش‌نویس برای بررسی و تأیید آقای کلانی / آقای خداپرست. هیچ رابطه‌ای مصوب نیست و این فایل پیکربندی ERP نیست.'),
 ('ترتیب کار کارفرما', '۱) مدل‌سازی روابط (این فایل) ← ۲) تأیید کارشناسی در شیت «جدول_تایید_کارفرما» ← ۳) اعمال قواعد مصوب به دفتر روزنامه و فرم‌های پیش‌حسابداری. کاربرد آزمایشی مدل روی کل دفتر ۱۴۰۴ در فایل جداگانه «دفتر_روزنامه_1404_با_رویدادهای_مالی_و_شناورها_DRAFT.xlsx» آمده است.'),
 ('منابع', 'فرهنگ رویداد کانونی v7.1 (۷۰ رویداد FE-01..FE-70، شناسه‌ها حفظ شد)؛ مدل مراکز هزینه فعالیت‌محور (service_activity_coverage v2، ۷۵ مرکز پیشنهادی) و crosswalk فعالیت v4؛ دفتر روزنامه ۱۴۰۴ (۵۵۳٬۶۶۷ سطر، ۷۵٬۲۷۹ سند)؛ مرجع «بودجه 1404 + مراکز هزینه» منبع (نوع بودجه: جاری / سرمایه‌ای مستمر / سرمایه‌ای غیرمستمر).'),
 ('شکاف‌های فرهنگ', f'{len(GAPS)} شکاف (GAP-01..GAP-12) جایی که رویداد واقعی دفتر در فرهنگ v7.1 نیست. این‌ها رویداد مصوب نیستند و فقط برای تصمیم کارفرما ثبت شده‌اند.'),
 ('قاعده ۱ کارفرما', 'رویداد مالی مخصوص هزینه نیست؛ همه طبقات حساب (دارایی، بدهی، ارزش خالص، درآمد، هزینه، رابط، انتظامی) بررسی شده‌اند. وقوع رویداد (نمونه رویداد = سند + رویداد) از ردیف‌های بدهکار/بستانکار آن جداست.'),
 ('قاعده ۲ کارفرما', 'پیش‌پرداخت و علی‌الحساب می‌توانند در سطح رویداد به مرکز هزینه مرتبط شوند؛ این ارتباط شناسایی هزینه نیست.'),
 ('قاعده ۳ کارفرما', 'بُعد پروژه برای بودجه‌های غیرمستمر؛ شناسه پروژه فقط از ردیف بودجه واقعی ساخته می‌شود و هیچ شناسه‌ای جعل نمی‌شود.'),
 ('قاعده ۴ کارفرما', 'بُعد قرارداد برای هر دو نوع بودجه؛ شماره قرارداد فقط از متن منبع.'),
 ('قاعده ۵ کارفرما', 'مدل روابط پیش از طبقه‌بندی دفتر؛ ستون‌های تأیید همه «Needs Decision» هستند.'),
 ('سطوح کاربرد ابعاد', 'REQUIRED_CANDIDATE / CONDITIONAL_CANDIDATE / OPTIONAL_CANDIDATE / NOT_APPLICABLE_CANDIDATE / UNKNOWN = پیشنهاد تحلیلی. APPROVED_REQUIRED / APPROVED_OPTIONAL / APPROVED_NOT_APPLICABLE فقط پس از تصمیم کارفرما؛ در این نسخه هیچ‌کدام استفاده نشده است.'),
 ('کاردینالیتی', '0..1، 1..1، 0..N، 1..N یا Needs Decision. ستون «مشاهده‌شده» از نمونه‌های رویداد دفتر ۱۴۰۴ محاسبه شده (شواهد، نه قاعده)؛ ستون «پیشنهاد تحلیلی» نظر کارشناسی است؛ ستون «مصوب» خالی/Needs Decision است.'),
 ('وضعیت‌های تشخیص در دفتر', 'MAPPED_DIRECT، MAPPED_EVENT_COMPONENT، MULTI_EVENT، AMBIGUOUS، INSUFFICIENT_EVIDENCE، REQUIRES_NEW_EVENT_REVIEW — هدف ۱۰۰٪ پردازش با وضعیت شفاف است، نه ۱۰۰٪ انتساب اجباری.'),
 ('هشدارهای مدل', 'هم‌وقوعی تاریخی رویداد و مرکز هزینه (ماتریس v7.1) فقط شاهد است؛ عنوان حساب، فعالیت شهری نیست؛ تملک/احداث با نگهداری، انتقال به سازمان وابسته با حقوق مستقیم، وصول درآمد با توزیع درآمد و خرید تجهیزات با تعمیر یکی نیست؛ مرکز هزینه سند به همه سطرهای سند کپی نمی‌شود.'),
 ('شیت‌ها', 'مدل_رویداد_و_شناورها: الگوهای رویداد × حساب × ابعاد | سناریوهای_تأیید: ۱۵ سناریو با سند واقعی | جدول_تایید_کارفرما: تصمیم‌گیری هر رویداد | فرهنگ_رویداد_و_شکاف‌ها | قواعد_تشخیص_و_کاربرد | مراکز_هزینه_فعالیت‌محور'),
]
for row in GUIDE: ws.append(list(row))
style_rows(ws)
for r in range(2, ws.max_row + 1): ws.cell(r, 1).font = BOLD

# ------------------------------------------------------------ 2. relationship matrix
MH = ['کد رویداد مالی', 'عنوان رویداد مالی', 'تعریف رویداد', 'خانواده رویداد', 'نوع رویداد در فرهنگ', 'طبقه حساب', 'حساب کل (کد/عنوان)', 'حساب معین (کد/عنوان)',
      'جهت ردیف', 'نقش ردیف در رویداد', 'کاربرد بودجه مستمر/غیرمستمر (مشاهده و پیشنهاد)',
      'مرکز هزینه / خانواده مرکز هزینه مرتبط (مشاهده‌شده)', 'منطق ارتباط مرکز هزینه', 'کاربرد مرکز هزینه (پیشنهاد)',
      'کاربرد قرارداد (پیشنهاد)', 'منطق ارتباط قرارداد', 'کاربرد بانک (پیشنهاد)', 'منطق ارتباط بانک', 'کاربرد پروژه (پیشنهاد)', 'منطق ارتباط پروژه',
      'کاربرد شخص / طرف حساب (پیشنهاد)', 'کاربرد واحد سازمانی (پیشنهاد)',
      'کاردینالیتی رویداد↔حساب (مشاهده‌شده)', 'کاردینالیتی رویداد↔مرکز هزینه (مشاهده‌شده)', 'کاردینالیتی رویداد↔قرارداد (مشاهده‌شده)', 'کاردینالیتی رویداد↔بانک (مشاهده‌شده)',
      'کاردینالیتی رویداد↔پروژه (مشاهده‌شده)', 'کاردینالیتی رویداد↔شخص (مشاهده‌شده)',
      'کاردینالیتی پیشنهادی (حساب | مرکز هزینه | قرارداد | بانک | پروژه | شخص)', 'کاردینالیتی مصوب', 'کاربرد ابعاد مصوب',
      'فیلدهای منبع لازم', 'شرط تحریک (قاعده تشخیص)', 'شواهد تاریخی (دفتر ۱۴۰۴ و v7.1)', 'وضعیت تشخیص در دفتر', 'قاعده پیشنهادی', 'اطمینان / وضعیت شواهد',
      'وضعیت تأیید کارفرما', 'یادداشت کارفرما']
MW = [10, 34, 50, 16, 20, 14, 26, 34, 9, 28, 34, 46, 60, 20, 20, 50, 20, 50, 20, 60, 20, 18, 22, 22, 22, 22, 22, 22, 30, 14, 14, 40, 48, 60, 44, 60, 30, 16, 30]
ws = sheet('مدل_رویداد_و_شناورها', MH, MW)
FIELDS = 'سطر منبع، شناسه سند (منطقه+شماره سند)، طبقه/کل/معین/تفضیلی/جزء، مبلغ بدهکار/بستانکار، نوع سند، شرح سند، کد بودجه'
def conf(p, e):
    if not EV[e]['canonical']: return 'شکاف فرهنگ — رویداد جدید پیشنهادی؛ تأیید نشده'
    st = p['st']; tot = p['rows']
    amb = st['AMBIGUOUS'] + st['MULTI_EVENT']
    if tot >= 100 and st['MAPPED_DIRECT'] / tot > 0.9: return 'قوی — قاعده حساب و جهت با شواهد فراوان'
    if tot >= 20 and amb / tot < 0.3: return 'متوسط — قاعده حساب/متن؛ نمونه‌برداری کارشناسی لازم'
    return 'ضعیف — شواهد کم یا چندرویدادی؛ تصمیم کارشناسی لازم'
def dim_row(e):
    dm = F.dims(e)
    return dm, [APP_FA[a] for a in dm]
mrows = []
for key in sorted(pat, key=lambda k: (ALL.index(k[0]), k[1], k[2], str(k[3]), k[4])):
    e, cls, kol, moin, side, kind = key
    p = pat[key]
    dm, ap = dim_row(e)
    obs = CARD_OBS[e]
    btxt = top(p['bt'], 4) or ('—' if cls != '6' else 'بدون کد بودجه')
    bprop = 'پروژه فقط برای ردیف‌های «سرمایه‌ای غیرمستمر»' if dm[3] != F.N_ else 'بودجه/پروژه کاربرد ندارد'
    trig = f"طبقه {cls}، کل {kol}، معین {moin if moin != '*' else '(هر معین)'}، جهت {'بدهکار' if side == 'D' else 'بستانکار'}؛ {top(p['method'], 2)}"
    if kind == 'جزء/طرف تسویه': trig = f"طرف تسویه/بدهی/بانک (کل {kol}) در سندی که رویداد محرک آن {e} است؛ روش: {top(p['method'], 2)}"
    ev_hist = f"دفتر ۱۴۰۴: {p['rows']:,} سطر در {len(p['docs']):,} سند؛ بدهکار {p['D']:,.0f} / بستانکار {p['C']:,.0f}؛ نمونه: {'، '.join(p['ex'])}"
    if e in DICT: ev_hist += f"؛ v7.1: {DICT[e][9]} سطر منبع v7.1" + (f"؛ ماتریس تاریخی (غیرمصوب): {'، '.join(MATRIX[e][:3])}" if MATRIX.get(e) else '')
    rule = f"اگر {trig.split('؛')[0]} ← {e} ({EV[e]['title']}) با نقش «{p['role'].most_common(1)[0][0]}»"
    mrows.append([e, EV[e]['title'], definition(e), EV[e]['family'], 'کانونی v7.1' if EV[e]['canonical'] else 'شکاف فرهنگ (پیشنهادی)',
                  f"{cls} {CLST.get(cls, '')}", f"{kol} {KOLT.get(kol, '')}", ('(هر معین؛ طرف تسویه)' if kind == 'جزء/طرف تسویه' else '(همه معین‌ها)') if moin == '*' else f"{moin} {MOINT.get(moin, '')}",
                  'بدهکار' if side == 'D' else 'بستانکار', f"{kind}: {top(p['role'], 2)}", f'مشاهده: {btxt} | پیشنهاد: {bprop}',
                  top(p['cc'], 3) or 'بدون مرکز هزینه در سطرهای این الگو', cc_logic(e, dm), ap[0], ap[1], con_logic(e, dm), ap[2], bank_logic(e, dm), ap[3], proj_logic(e, dm), ap[4], ap[5],
                  *obs, ' | '.join(prop_card(e)), 'Needs Decision', 'Needs Decision', FIELDS + ('، شناسه شخص (AJ)' if dm[4] != F.N_ else ''), trig, ev_hist,
                  top(p['st'], 3), rule, conf(p, e), 'Needs Decision', ''])
for r in mrows: ws.append(r)
style_rows(ws, lambda row: GAPF if row[4].value and row[4].value.startswith('شکاف') else None)
ws.auto_filter.ref = f'A1:{get_column_letter(len(MH))}{ws.max_row}'
N_PATTERNS = len(mrows)
# events in the dictionary without 1404 evidence
for e in ALL:
    if ev_rows.get(e): continue
    dm, ap = dim_row(e)
    ws.append([e, EV[e]['title'], definition(e), EV[e]['family'], 'کانونی v7.1' if EV[e]['canonical'] else 'شکاف فرهنگ (پیشنهادی)', '—', '—', '—', '—',
               'الگوی حساب در دفتر ۱۴۰۴ مشاهده نشد', '—', '—', cc_logic(e, dm), ap[0], ap[1], con_logic(e, dm), ap[2], bank_logic(e, dm), ap[3], proj_logic(e, dm), ap[4], ap[5],
               *CARD_OBS[e], ' | '.join(prop_card(e)), 'Needs Decision', 'Needs Decision', FIELDS, f"ماهیت کلی v7.1: {DICT[e][5] if e in DICT else ''}",
               f"دفتر ۱۴۰۴: بدون سطر قابل انتساب؛ v7.1: {DICT[e][9] if e in DICT else '—'} سطر منبع v7.1", 'NO_1404_JOURNAL_EVIDENCE',
               'الگوی حساب پس از ارائه نمونه واقعی تعریف شود', 'بدون شواهد ۱۴۰۴ — الگوی حساب تعریف نشده', 'Needs Decision', ''])
    for c in ws[ws.max_row]: c.font = BFONT; c.alignment = WRAP; c.border = BORDER; c.fill = YEL
ws.auto_filter.ref = f'A1:{get_column_letter(len(MH))}{ws.max_row}'

# ------------------------------------------------------------ 3. scenarios with real documents
SH = ['سناریو', 'زنجیره پیشنهادی: حساب ← رویداد ← مرکز هزینه ← قرارداد ← بانک ← پروژه', 'سطر منبع', 'شناسه سند', 'حساب (کل/معین/تفضیلی)', 'عنوان جزء',
      'بدهکار', 'بستانکار', 'رویداد / وضعیت', 'نقش ردیف', 'مرکز هزینه (وضعیت: مقدار)', 'قرارداد', 'بانک', 'پروژه', 'شخص', 'واحد سازمانی', 'شرط/یادداشت کارشناسی']
ws = sheet('سناریوهای_تأیید', SH, [30, 60, 9, 16, 34, 40, 15, 15, 30, 28, 40, 26, 34, 34, 26, 18, 60])
def pick_doc(pred_doc, prefer=None):
    best = None
    for dk, info in DOCINFO.items():
        if info['n'] > 14 or info['n'] < 2: continue
        ix = DOCIDX[dk]
        if pred_doc(ix):
            score = prefer(ix) if prefer else 0
            if best is None or score > best[0]: best = (score, dk)
            if score >= 10: break
    return best[1] if best else None
DOCIDX = defaultdict(list)
for i, rec in enumerate(REC): DOCIDX[rec['doc']].append(i)
def has(ix, ev=None, cat=None, side=None, st=None):
    return any((ev is None or REC[i]['ev'] == ev) and (cat is None or REC[i]['cat'] == cat) and (side is None or REC[i]['side'] == side) and (st is None or REC[i]['status'] == st) for i in ix)
def quality(ix):
    return (3 * any(REC[i]['con_status'] == 'CONTRACT_IDENTIFIED' for i in ix) + 3 * any(REC[i]['cc_status'].startswith(('DIRECT', 'EVENT_LEVEL')) for i in ix)
            + 2 * any(REC[i]['bank_status'] == 'BANK_DIRECT_LINE' for i in ix) + 2 * any(REC[i]['proj_status'].startswith('PROJECT_') for i in ix))
SCEN = [
 ('1. پیش‌پرداخت پیمانکار پروژه عمرانی', 'پیش‌پرداخت پروژه‌ها (14001، دارایی) ← FE-01 ← مرکز هزینه فعالیت پروژه (سطح رویداد؛ شرطی) ← قرارداد پیمان (الزامی) ← بانک پرداخت (الزامی) ← پروژه (شرطی: بودجه غیرمستمر)',
  lambda ix: has(ix, 'FE-01', 'PREPAY_PROJ', 'D') and has(ix, cat='BANK', side='C'), 'پیش‌پرداخت دارایی می‌ماند؛ مرکز هزینه فقط ارتباط سطح رویداد است. پروژه وقتی شناسه دارد که ردیف بودجه غیرمستمر واقعی موجود باشد.'),
 ('2. علی‌الحساب پیمانکار خدمات شهری', 'بدهکار شدن حساب پیمانکار در سند «پیش‌پرداخت و علی‌الحساب» ← FE-01 (علی‌الحساب) ← مرکز هزینه خدمت شهری (سطح رویداد) ← قرارداد خدمات ← بانک',
  lambda ix: has(ix, 'FE-01', 'PAY_CONTRACTOR', 'D'), 'نیازمند تصمیم: آیا علی‌الحساب پیمانکار خدمات (بدهکار کردن حساب بدهی پیمانکار) همان FE-01 است یا رویداد جداگانه؟'),
 ('3. ثبت صورت‌وضعیت پیمانکار', 'هزینه/سرمایه‌ای (615/652/611) ← GAP-01 (شکاف: ثبت صورت‌وضعیت) + کسور FE-07/FE-09/FE-11/FE-02 ← مرکز هزینه خط (الزامی) ← قرارداد (الزامی) ← بانک (کاربرد ندارد؛ غیرنقدی) ← پروژه (شرطی)',
  lambda ix: has(ix, 'GAP-01', 'EXP', 'D') and has(ix, 'FE-07'), 'فرهنگ v7.1 فقط «پرداخت» صورت‌وضعیت را دارد؛ پیشنهاد تفکیک شناسایی (GAP-01) از پرداخت (FE-03).'),
 ('4. پرداخت صورت‌وضعیت', 'بدهی به پیمانکاران (21014، بدهکار) ← FE-03 ← مرکز هزینه (از ثبت صورت‌وضعیت؛ سطح رویداد) ← قرارداد (الزامی) ← بانک (الزامی) ← پروژه (شرطی)',
  lambda ix: has(ix, 'FE-03', 'PAY_CONTRACTOR', 'D') and has(ix, cat='BANK', side='C'), 'مرکز هزینه در سند پرداخت معمولاً حاضر نیست و باید از سند ثبت صورت‌وضعیت همان قرارداد پیوند شود (کلید قرارداد/شخص).'),
 ('5. پرداخت حقوق و دستمزد', 'جبران خدمات کارکنان (610، بدهکار) ← FE-44 + کسور FE-48/FE-49 + GAP-05 ← مرکز هزینه سازمانی (الزامی) ← قرارداد (ندارد) ← بانک حقوق ← پروژه (ندارد)',
  lambda ix: has(ix, 'FE-44', 'EXP', 'D') and has(ix, cat='BANK', side='C'), 'انتقال اعتبار به سازمان‌های وابسته بابت حقوق آن‌ها حقوق مستقیم نیست (FE-69).'),
 ('6. خرید کالا و تجهیزات', 'موجودی/هزینه کالا یا دارایی ثابت (130/61109/616) ← FE-22..FE-28 ← مرکز هزینه مصرف‌کننده ← قرارداد (شرطی) ← بانک (شرطی) ← پروژه (شرطی)',
  lambda ix: any(REC[i]['ev'] in ('FE-22', 'FE-23', 'FE-24', 'FE-27', 'FE-28') and REC[i]['status'] == 'MAPPED_DIRECT' and REC[i]['cat'] in ('INVENTORY', 'EXP', 'FIXED_ASSET') for i in ix), 'خرید تجهیزات با تعمیر تجهیزات (FE-39..43) یکی نیست.'),
 ('7. خرید خدمات', 'حق‌الزحمه خدمات قراردادی (61102) ← FE-29/FE-31/FE-45 یا GAP-11 ← مرکز هزینه خدمت ← قرارداد (شرطی) ← بانک (شرطی)',
  lambda ix: any(REC[i]['ev'] in ('FE-29', 'FE-31', 'GAP-11') and REC[i]['cat'] == 'EXP' for i in ix), 'فرهنگ v7.1 رویداد عمومی «خرید خدمات» ندارد (GAP-11).'),
 ('8. ثبت و تسویه حساب‌های پرداختنی', 'ثبت: هزینه ← بدهی (210/220 بستانکار، جزء رویداد) | تسویه: بدهی (بدهکار) ← GAP-10 ← بانک (الزامی)',
  lambda ix: has(ix, 'GAP-10', 'PAY_OTHER', 'D') and has(ix, cat='BANK', side='C'), 'مرکز هزینه تسویه بدهی از سند ثبت بدهی به ارث می‌رسد و بدون پیوند سند ثبت تعیین نمی‌شود.'),
 ('9. کسر و استرداد سپرده', 'سپرده حسن انجام کار/بیمه (225) ← FE-07/FE-09 (کسر) و FE-08/FE-10 (استرداد) ← قرارداد (الزامی) ← بانک (فقط در استرداد)',
  lambda ix: has(ix, 'FE-08', 'DEP_PERF', 'D') and has(ix, cat='BANK', side='C'), 'کسر سپرده رویداد تابع صورت‌وضعیت است؛ استرداد رویداد نقدی مستقل است.'),
 ('10. شناسایی درآمد', 'اسناد دریافتنی (11513) بدهکار + درآمد (4xx) بستانکار ← FE-57..FE-61 (شناسایی هم‌زمان با دریافت چک) ← مرکز هزینه (کاربرد ندارد) ← بانک (ندارد در این مرحله)',
  lambda ix: has(ix, cat='NOTES_REC', side='D') and has(ix, cat='REV', side='C', st='MAPPED_DIRECT'), 'GAP-02: شناسایی تعهدی جدا از وصول در فرهنگ نیست و در دفتر مبنای نقدی هم مشاهده نشد.'),
 ('11. وصول درآمد و دریافت وجه', 'بانک (بدهکار) + درآمد (بستانکار) ← FE-57/FE-58/FE-60/FE-61 یا GAP-04 ← بانک (الزامی) ← مرکز هزینه (کاربرد ندارد)',
  lambda ix: has(ix, cat='BANK', side='D') and has(ix, 'FE-61', 'REV', 'C') and has(ix, cat='BANK', st='MAPPED_EVENT_COMPONENT'), 'مقادیر AF:AI موجود روی سطرهای درآمد (برخورد کد بودجه درآمد/هزینه) فقط علامت‌گذاری شده‌اند.'),
 ('12. انتقال وجه بین حساب‌های بانکی', 'بانک ← بانک/حساب رابط (710) ← FE-68 ← بانک (۲..N حساب) ← مرکز هزینه/قرارداد/پروژه (کاربرد ندارد)',
  lambda ix: has(ix, 'FE-68', 'BANK', 'C') and has(ix, 'FE-68', 'BANK', 'D'), 'توزیع درآمدهای تمرکزیافته بین مناطق با وصول درآمد یکی نیست.'),
 ('13. بازپرداخت تسهیلات', 'بازپرداخت اصل و سود (620 یا 235 بدهکار) ← FE-64/FE-65 ← قرارداد وام (الزامی) ← بانک (الزامی) ← مرکز هزینه (کاربرد ندارد)',
  lambda ix: any(REC[i]['ev'] in ('FE-64', 'FE-65') or (REC[i]['status'] == 'AMBIGUOUS' and 'FE-64' in REC[i]['cands']) for i in ix) and has(ix, cat='BANK', side='C'), 'کد بودجه 32 = بازپرداخت؛ تفکیک اصل از سود در عنوان حساب نیست.'),
 ('14. پرداخت تعهدات سنوات قبل', 'تعهدات انتقالی سنواتی (652) ← FE-67 (+GAP-01 در سند صورت‌وضعیت) ← مرکز هزینه فعالیت عنوان جزء (شرطی) ← قرارداد ← پروژه (کاندید؛ کد 31)',
  lambda ix: has(ix, 'FE-67', 'EXP', 'D'), 'تعهد سنواتی پیش‌پرداخت یا هزینه سال جاری نیست؛ ارتباط مرکز هزینه فقط از فعالیت عنوان جزء.'),
 ('15. ثبت یا اصلاح اسناد حسابداری', 'هر حساب در سند اصلاحی/افتتاحیه/اختتامیه ← FE-70 ← ابعاد از تراکنش اصلی (بدون انتساب جدید)',
  lambda ix: has(ix, 'FE-70', st='MULTI_EVENT') and 'اصلاح' in DOCINFO[REC[ix[0]]['doc']]['dtype'] and len(ix) <= 8, 'سند افتتاحیه/اختتامیه با هزاران سطر برای نمونه انتخاب نشد؛ نمونه سند اصلاحی کوچک است.'),
]
SC_EXAMPLES = {}
for title, chain, pred, note in SCEN:
    dk = pick_doc(pred, quality)
    ws.append([title, chain, '', dk or 'نمونه‌ای یافت نشد', '', '', '', '', '', '', '', '', '', '', '', '', note])
    for c in ws[ws.max_row]: c.font = BOLD; c.alignment = WRAP; c.border = BORDER; c.fill = SEC
    if not dk: continue
    SC_EXAMPLES[title] = dk
    for i in DOCIDX[dk]:
        r, rec = ROWS[i], REC[i]
        evt = f"{rec['ev'] or '—'} {EV[rec['ev']]['title'] if rec['ev'] else 'مبهم: ' + '/'.join(rec['cands'])} | {rec['status']}"
        ws.append(['', '', r[0], dk, f"{s(g(r, F.C_KOL))}/{s(g(r, F.C_MOIN))} {s(g(r, F.C_MOINN))} / {s(g(r, F.C_TAFT))}", s(g(r, F.C_JOZT)),
                   num(g(r, F.C_DR)), num(g(r, F.C_CR)), evt, rec['role'], f"{rec['cc_status']}: {rec['cc_related']}", f"{rec['con_status']}: {rec['con_value']}",
                   f"{rec['bank_status']}: {rec['bank_value']}", f"{rec['proj_status']}: {rec['proj_value']}", s(g(r, F.C_PERT)), s(g(r, F.C_ORGT)), rec['review_why']])
        for c in ws[ws.max_row]: c.font = BFONT; c.alignment = WRAP; c.border = BORDER
        ws.cell(ws.max_row, 7).number_format = '#,##0'; ws.cell(ws.max_row, 8).number_format = '#,##0'

# ------------------------------------------------------------ 4. client approval table
AH = ['کد رویداد', 'عنوان رویداد', 'تعریف پیشنهادی', 'نوع (کانونی / شکاف)', 'شواهد دفتر ۱۴۰۴ (سطر / نمونه رویداد)', 'الگوهای اصلی حساب (پیشنهادی)',
      'مرکز هزینه — پیشنهاد', 'قرارداد — پیشنهاد', 'بانک — پیشنهاد', 'پروژه — پیشنهاد', 'شخص / واحد سازمانی — پیشنهاد', 'کاردینالیتی پیشنهادی (حساب|CC|قرارداد|بانک|پروژه|شخص)',
      'تأیید رویداد', 'اصلاح تعریف', 'تأیید/رد الگوهای حساب', 'تصمیم کاربرد مرکز هزینه', 'تصمیم کاربرد قرارداد', 'تصمیم کاربرد بانک', 'تصمیم کاربرد پروژه',
      'تصمیم کاربرد شخص و واحد سازمانی', 'رابطه یک‌به‌یک / یک‌به‌چند', 'استثنا یا قاعده شرطی', 'وضعیت تأیید', 'یادداشت کارفرما']
ws = sheet('جدول_تایید_کارفرما', AH, [10, 34, 50, 16, 22, 48, 34, 30, 30, 34, 26, 30, 16, 30, 18, 20, 20, 20, 20, 20, 20, 34, 16, 34])
pat_by_ev = defaultdict(list)
for key, p in pat.items():
    if key[5] != 'جزء/طرف تسویه': pat_by_ev[key[0]].append((p['rows'], key))
SHORT = {F.R_: 'الزامی (پیشنهاد)', F.C_: 'شرطی (پیشنهاد)', F.O_: 'اختیاری (پیشنهاد)', F.N_: 'کاربرد ندارد (پیشنهاد)', F.U_: 'نامعلوم'}
for e in ALL:
    dm = F.dims(e)
    pats = sorted(pat_by_ev.get(e, []), reverse=True)[:3]
    ptxt = '؛ '.join(f"{'بدهکار' if k[4] == 'D' else 'بستانکار'} {k[3]} {MOINT.get(k[3], '')[:40]} ({n:,})" for n, k in pats) or 'بدون الگوی ۱۴۰۴'
    ws.append([e, EV[e]['title'], definition(e), 'کانونی v7.1' if EV[e]['canonical'] else 'شکاف فرهنگ — رویداد جدید پیشنهادی', f"{ev_rows.get(e, 0):,} / {ev_docs.get(e, 0):,}", ptxt,
               SHORT[dm[0]] + ' — ' + cc_logic(e, dm)[:160], SHORT[dm[1]], SHORT[dm[2]], SHORT[dm[3]] + ('؛ فقط بودجه غیرمستمر' if dm[3] != F.N_ else ''),
               f'شخص: {SHORT[dm[4]]}؛ واحد سازمانی: {SHORT[dm[5]]}', ' | '.join(prop_card(e)),
               'Needs Decision', '', 'Needs Decision', 'Needs Decision', 'Needs Decision', 'Needs Decision', 'Needs Decision', 'Needs Decision', 'Needs Decision', '', 'Needs Decision', ''])
style_rows(ws, lambda row: GAPF if row[3].value and row[3].value.startswith('شکاف') else None)
for col in range(13, 24):
    for r in range(2, ws.max_row + 1): ws.cell(r, col).fill = YEL
dv1 = DataValidation(type='list', formula1='"Needs Decision,تأیید,تأیید با اصلاح,رد,ادغام با رویداد دیگر"', allow_blank=True)
dv2 = DataValidation(type='list', formula1='"Needs Decision,APPROVED_REQUIRED,APPROVED_OPTIONAL,APPROVED_NOT_APPLICABLE,شرطی (توضیح در ستون استثنا)"', allow_blank=True)
dv3 = DataValidation(type='list', formula1='"Needs Decision,1..1,0..1,1..N,0..N"', allow_blank=True)
for dv in (dv1, dv2, dv3): ws.add_data_validation(dv)
last = ws.max_row
dv1.add(f'M2:M{last}'); dv1.add(f'O2:O{last}'); dv1.add(f'W2:W{last}')
for col in 'PQRST': dv2.add(f'{col}2:{col}{last}')
dv3.add(f'U2:U{last}')
ws.auto_filter.ref = f'A1:X{last}'

# ------------------------------------------------------------ 5. dictionary + gaps
DHX = ['کد', 'عنوان', 'عنوان کوتاه / مرحله', 'تعریف', 'خانواده', 'ماهیت کلی', 'نوع', 'شواهد v7.1 (سطر منبع v7.1)', 'مراکز هزینه مشاهده‌شده در v7.1 (غیرمصوب)',
       'سطرهای ۱۴۰۴: نگاشت مستقیم', 'سطرهای ۱۴۰۴: جزء رویداد', 'سطرهای ۱۴۰۴: چندرویدادی (نمایش اصلی)', 'سطرهای ۱۴۰۴: شکاف (برای GAP)', 'نمونه‌های رویداد ۱۴۰۴', 'علت شکاف / یادداشت']
ws = sheet('فرهنگ_رویداد_و_شکاف‌ها', DHX, [10, 36, 22, 60, 16, 22, 18, 16, 50, 14, 14, 14, 14, 14, 60])
stc = defaultdict(Counter)
for r in REC:
    if r['ev']: stc[r['ev']][r['status']] += 1
for e in ALL:
    d = DICT.get(e)
    ws.append([e, EV[e]['title'], (d[2] if d else '') + ' / ' + EV[e]['stage'], definition(e), EV[e]['family'], EV[e]['nature'],
               'کانونی v7.1' if EV[e]['canonical'] else 'شکاف فرهنگ (پیشنهادی)', d[9] if d else '', (d[14] if d else '') or '',
               stc[e]['MAPPED_DIRECT'], stc[e]['MAPPED_EVENT_COMPONENT'], stc[e]['MULTI_EVENT'], stc[e]['REQUIRES_NEW_EVENT_REVIEW'], ev_docs.get(e, 0),
               '' if EV[e]['canonical'] else F.GAPS[e][4]])
style_rows(ws, lambda row: GAPF if row[6].value and row[6].value.startswith('شکاف') else None)
ws.auto_filter.ref = f'A1:O{ws.max_row}'

# ------------------------------------------------------------ 6. rules and usage
RH = ['دسته حساب (طبقه‌بندی قاعده)', 'جهت', 'رویداد', 'عنوان رویداد', 'نقش ردیف', 'روش تشخیص', 'وضعیت', 'تعداد سطر', 'تعداد سند']
ws = sheet('قواعد_تشخیص_و_کاربرد', RH, [34, 9, 10, 36, 40, 44, 28, 12, 12])
ru = defaultdict(lambda: [0, set()])
for r in REC:
    k = (F.CAT_FA.get(r['cat'], r['cat']), 'بدهکار' if r['side'] == 'D' else 'بستانکار', r['ev'] or ('/'.join(r['cands']) if r['cands'] else '—'), r['role'], r['method'], r['status'])
    ru[k][0] += 1; ru[k][1].add(r['doc'])
for k, (n, docs) in sorted(ru.items(), key=lambda x: -x[1][0]):
    ev = k[2]
    ws.append([k[0], k[1], ev, EV[ev]['title'] if ev in EV else '', k[3], k[4], k[5], n, len(docs)])
style_rows(ws)
ws.auto_filter.ref = f'A1:I{ws.max_row}'

# ------------------------------------------------------------ 7. activity-based CC candidates
CH = ['کد مرکز هزینه پیشنهادی (CC-only)', 'عنوان', 'حوزه خدمت', 'فعالیت اصلی', 'وضعیت در مدل خدمات', 'سطح۲ بومی (AF:AI) متناظر', 'نوع رابطه با سلسله‌مراتب بومی',
      'توضیح crosswalk', 'سطرهای ۱۴۰۴ با این مرکز در سطح رویداد', 'وضعیت تأیید']
ws = sheet('مراکز_هزینه_فعالیت‌محور', CH, [14, 44, 30, 36, 24, 14, 18, 60, 14, 16])
usecc = Counter()
for r in REC:
    m = re.match(r'(CC-\d+)', r['cc_event'] or r['cc_prop'] or '')
    if m: usecc[m.group(1)] += 1
for r in CCM[1:]:
    if not r[0]: continue
    nat, rel, note = XW.get(r[0], (None, 'NOT_EVALUATED', ''))
    ws.append([r[0], r[1], r[2], r[3], r[5], nat or '—', rel, note, usecc.get(r[0], 0), 'پیشنهادی — Needs Decision'])
for r in SRC['cc_new'][1:]:
    if r[0] and r[0] not in {x[0] for x in CCM[1:]}:
        nat, rel, note = XW.get(r[0], (None, 'NOT_EVALUATED', ''))
        ws.append([r[0], r[1], r[2], '', r[8], nat or '—', rel, note, usecc.get(r[0], 0), 'پیشنهادی — Needs Decision'])
style_rows(ws)
ws.auto_filter.ref = f'A1:J{ws.max_row}'

wb.save(OUT)
json.dump(dict(patterns=N_PATTERNS, events_canonical=len(CANON), gaps=len(GAPS), scenarios=SC_EXAMPLES, events_used=sorted(ev_rows)), open(W + 'model_stats.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print('patterns', N_PATTERNS, 'canonical', len(CANON), 'gaps', len(GAPS), 'scenarios', len(SC_EXAMPLES))
