"""Dump the canonical v7.1 event dictionary, evidence matrix, committee table and the activity-based CC master (service coverage v2) to JSON."""
import openpyxl, json
ROOT = 'H:/Freelancing_Project/KalaniProject/municipality_demo/'
FE = ROOT + 'data/deliverables/financial_event_discovery/revisions/1404-r20260922-canonical-events-v7.1/financial_event_discovery_1404-r20260922_canonical-events_v7.1.xlsx'
SC = ROOT + 'data/deliverables/service_activity_coverage/revisions/1405-r20260922-service-coverage/service_activity_coverage_1405-r20260922_QA-remapped_v2.xlsx'
OUT = ROOT + 'data/work/fe_model/model_sources.json'

def sheet(wb, name):
    return [[None if v is None else (v if isinstance(v, (int, float)) else str(v)) for v in r] for r in wb[name].iter_rows(values_only=True)]

wb = openpyxl.load_workbook(FE, read_only=True)
out = {k: sheet(wb, n) for k, n in [('dictionary', 'فرهنگ_رویداد_مالی_کاندید'), ('matrix', 'ماتریس_شواهد_رویداد_مرکز_v2'),
                                    ('committee', 'جدول_تصمیم_کمیته_رویداد'), ('multi', 'QA_Multi_Event')]}
wb.close()
wb = openpyxl.load_workbook(SC, read_only=True)
out['cc_master'] = sheet(wb, 'مرجع_مراکز_هزینه_v3_با_افزودنی')
out['cc_new'] = sheet(wb, 'مراکز_هزینه_جدید_پیشنهادی')
out['service_map'] = sheet(wb, 'تطبیق_اصلاح‌شده_خدمات')
wb.close()
json.dump(out, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False)
print({k: len(v) for k, v in out.items()})
