"""Independent verification of Deliverable B against the ORIGINAL source journal and the v4 workbook values.
Opens the output with openpyxl (full XML parse of every sheet) and compares A:AQ cell-by-cell."""
import openpyxl, zipfile, pickle, time, json
from collections import Counter
OUT = 'C:/Users/sadra/Desktop/دفتر_روزنامه_1404_با_رویدادهای_مالی_و_شناورها_DRAFT.xlsx'
SRC = 'C:/Users/sadra/Desktop/دفتر روزنامه 1404 - با اشخاص جدید.xlsx'
W = 'H:/Freelancing_Project/KalaniProject/municipality_demo/data/work/fe_model/'
t0 = time.time()
assert zipfile.ZipFile(OUT).testzip() is None
V4 = pickle.load(open(W + 'journal_v4.pkl', 'rb'))['rows']
wo = openpyxl.load_workbook(OUT, read_only=True)
ws_ = openpyxl.load_workbook(SRC, read_only=True)
res = {'sheets': wo.sheetnames}
so = wo['دفتر_روزنامه_کامل'].iter_rows(min_row=1, max_col=71, values_only=True)
ss = ws_['کدینگ جامع 1404 شهرداری'].iter_rows(min_row=1, max_col=43, values_only=True)
mism = Counter(); afai_v4_mism = 0; afai_src_diff = 0; n = 0; D = C = 0.0; Ds = Cs = 0.0; formulas_out = formulas_src = formula_mism = 0
status = Counter(); blank_status = 0; header = None
for k, (ro, rs) in enumerate(zip(so, ss), start=1):
    ro = tuple(ro) + (None,) * (71 - len(ro)); rs = tuple(rs) + (None,) * (43 - len(rs))
    if k == 5: header = ro[43:71]
    if k < 6:
        for c in range(43):
            if ro[c] != rs[c]: mism[('header', c)] += 1
        continue
    n += 1
    v4 = V4[k - 6]
    assert v4[0] == k
    for c in range(43):
        if c in (31, 32, 33, 34):
            if ro[c] != v4[1 + c]: afai_v4_mism += 1
            continue
        if ro[c] != rs[c]: mism[c] += 1
    if any(ro[c] != rs[c] for c in (31, 32, 33, 34)): afai_src_diff += 1
    if isinstance(ro[29], str) and ro[29].startswith('='): formulas_out += 1
    if isinstance(rs[29], str) and rs[29].startswith('='):
        formulas_src += 1
        if ro[29] != rs[29]: formula_mism += 1
    for v, acc in ((ro[16], 'D'), (ro[17], 'C')):
        x = float(v or 0)
        if acc == 'D': D += x
        else: C += x
    Ds += float(rs[16] or 0); Cs += float(rs[17] or 0)
    st = ro[51]
    if not st: blank_status += 1
    status[str(st).split(' — ')[0]] += 1
extra_out = sum(1 for _ in so); extra_src = sum(1 for _ in ss)
res.update(rows=n, extra_rows_out=extra_out, extra_rows_src=extra_src, mismatches_vs_source=dict((str(k), v) for k, v in mism.items()),
           afai_mismatch_vs_v4=afai_v4_mism, rows_where_afai_differs_from_original_source=afai_src_diff,
           formulas_src=formulas_src, formulas_out=formulas_out, formula_text_mismatch=formula_mism,
           debit_out=D, credit_out=C, debit_src=Ds, credit_src=Cs, blank_status=blank_status, status=dict(status), new_header=header)
for name in ('مراکز_هزینه', 'هزینه_ها', 'درآمدها', 'ارتباط_سطر_رویداد', 'QA_تطبیق_و_راهنما', 'حسابهای بانکی', 'فیلتر اسناد دارای سرفصل هزینه', 'واحد سازمانی', 'بودجه 1404 + مراکز هزینه'):
    res['rows_' + name] = sum(1 for _ in wo[name].iter_rows(values_only=True)) - 1
wo.close(); ws_.close()
res['secs'] = round(time.time() - t0)
json.dump(res, open(W + 'verify_B.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)
print(json.dumps(res, ensure_ascii=False, indent=1, default=str))
