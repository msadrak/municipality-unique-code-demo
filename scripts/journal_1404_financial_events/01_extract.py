"""Extract main sheet of v4 journal (all 43 cols) + bank sheet into pickles for analysis."""
import openpyxl, pickle, time, sys
from collections import Counter
V4 = 'C:/Users/sadra/Desktop/دفتر_روزنامه_1404_مراکز_هزینه_فعالیت_محور_v4.xlsx'
OUT = 'H:/Freelancing_Project/KalaniProject/municipality_demo/data/work/fe_model/'
t0 = time.time()
wb = openpyxl.load_workbook(V4, read_only=True, data_only=False)
ws = wb['کدینگ جامع 1404 شهرداری']
rows, formula_cols, excel_row = [], Counter(), 5
for r in ws.iter_rows(min_row=6, max_col=43, values_only=True):
    excel_row += 1
    for i, v in enumerate(r):
        if isinstance(v, str) and v.startswith('='):
            formula_cols[i] += 1
    rows.append((excel_row,) + tuple(r) + (None,) * (43 - len(r)))
pickle.dump({'rows': rows, 'formula_cols': dict(formula_cols)}, open(OUT + 'journal_v4.pkl', 'wb'))
print('main rows', len(rows), dict(formula_cols), round(time.time() - t0), flush=True)
others = {}
for name in ('حسابهای بانکی', 'فیلتر اسناد دارای سرفصل هزینه', 'واحد سازمانی', 'بودجه 1404 + مراکز هزینه'):
    others[name] = [tuple(r) for r in wb[name].iter_rows(values_only=True)]
    print(name, len(others[name]), flush=True)
pickle.dump(others, open(OUT + 'aux_v4.pkl', 'wb'))
wb.close()
print('secs', round(time.time() - t0))
