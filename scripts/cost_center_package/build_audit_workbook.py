import sys
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np

print("=== STARTING AUDIT WORKBOOK GENERATION ===")

output_path = r'C:\Users\sadra\Downloads\ممیزی_نسخه_اول_مدل_مراکز_هزینه.xlsx'
wb_prev = openpyxl.load_workbook(r'C:\Users\sadra\Downloads\تحلیل_جامع_مراکز_هزینه_و_فعالیت_های_شهرداری_اصفهان.xlsx', data_only=True)
wb_org = openpyxl.load_workbook(r'C:\Users\sadra\Downloads\آمار پست های معاونتها-1405.xlsx', data_only=True)

# Load dataframes from previous output
df_mapped = pd.read_excel(r'C:\Users\sadra\Downloads\تحلیل_جامع_مراکز_هزینه_و_فعالیت_های_شهرداری_اصفهان.xlsx', sheet_name='نگاشت_ردیف_های_مالی')
df_dict = pd.read_excel(r'C:\Users\sadra\Downloads\تحلیل_جامع_مراکز_هزینه_و_فعالیت_های_شهرداری_اصفهان.xlsx', sheet_name='شناسنامه_مراکز_هزینه')
df_tax = pd.read_excel(r'C:\Users\sadra\Downloads\تحلیل_جامع_مراکز_هزینه_و_فعالیت_های_شهرداری_اصفهان.xlsx', sheet_name='درخت_فعالیت_ها')
df_org = pd.read_excel(r'C:\Users\sadra\Downloads\تحلیل_جامع_مراکز_هزینه_و_فعالیت_های_شهرداری_اصفهان.xlsx', sheet_name='ماتریس_مسئولیت_سازمانی')
df_exc = pd.read_excel(r'C:\Users\sadra\Downloads\تحلیل_جامع_مراکز_هزینه_و_فعالیت_های_شهرداری_اصفهان.xlsx', sheet_name='صف_بررسی_انسانی')

print("Loaded all previous dataframes.")
