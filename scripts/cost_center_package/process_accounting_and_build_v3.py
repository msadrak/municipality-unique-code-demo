import sys
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np
import time

print("=== STARTING ACCOUNTING DOC NO ENRICHMENT PIPELINE (v3) ===")
start_total = time.time()

# Paths
v2_path = r'C:\Users\sadra\Downloads\بسته_نهایی_بازبینی_شده_مراکز_هزینه_شهرداری_اصفهان_v2.xlsx'
v3_path = r'C:\Users\sadra\Downloads\بسته_نهایی_مراکز_هزینه_با_شماره_سند_حسابداری_v3.xlsx'
excel_mapped = r'C:\Users\sadra\Downloads\تحلیل_جامع_مراکز_هزینه_و_فعالیت_های_شهرداری_اصفهان.xlsx'
p_acc = r'H:\Freelancing_Project\KalaniProject\municipality_demo\data\reports\اطلاعات حسابداری1404.xlsx'

# 1. Load df_mapped
df_mapped = pd.read_excel(excel_mapped, sheet_name='نگاشت_ردیف_های_مالی')
code_to_meta = {}
for idx, r in df_mapped.iterrows():
    c_str = str(r['کد بودجه']).strip().replace('.0', '')
    if c_str not in code_to_meta:
        code_to_meta[c_str] = []
    code_to_meta[c_str].append({
        'cc_id': r['کد مرکز هزینه پیشنهادی'],
        'cc_name': r['مرکز هزینه پیشنهادی'],
        'domain': r['حوزه خدمت پیشنهادی'],
        'desc': r['شرح ردیف']
    })

print(f"Loaded {len(code_to_meta)} budget codes from df_mapped.")

# 2. Open accounting file and stream rows
print("Opening accounting workbook in read-only mode...")
wb_acc = openpyxl.load_workbook(p_acc, read_only=True)

doc_mapping_rows = []
cc_docs = {} # cc_id -> set of doc_nos
doc_to_ccs = {} # doc_no -> set of cc_ids
doc_to_codes = {} # doc_no -> set of budget codes
doc_to_lines = {} # doc_no -> count of lines

acc_total_lines = 0
acc_lines_with_code = 0
acc_lines_null_code = 0
acc_lines_matched = 0
acc_lines_unmatched_revenue = 0
acc_unique_all_docs = set()
acc_null_doc_count = 0

for sname in wb_acc.sheetnames:
    ws = wb_acc[sname]
    t0 = time.time()
    print(f"Streaming sheet [{sname}]...")
    
    header = [cell for cell in next(ws.iter_rows(max_row=1, values_only=True))]
    idx_area_code = header.index('کد منطقه')
    idx_area_name = header.index('منطقه')
    idx_date = header.index('تاریخ')
    idx_doc = header.index('شماره سند')
    idx_k_no = header.index('TitkNo')
    idx_m_no = header.index('TitMNo')
    idx_t_no = header.index('TitTNo')
    idx_desc = header.index('شرح سرفصل حساب جزء')
    idx_deb = header.index('مبلغ بدهکار')
    idx_cred = header.index('مبلغ بستانکار')
    idx_code = header.index('کد بودجه')
    
    sheet_lines = 0
    for r in ws.iter_rows(min_row=2, values_only=True):
        sheet_lines += 1
        acc_total_lines += 1
        doc = r[idx_doc]
        c = r[idx_code]
        
        if doc is None or str(doc).strip() in ['', 'NULL', 'None']:
            acc_null_doc_count += 1
            doc_str = 'فاقد_شماره_سند'
        else:
            doc_str = str(doc).strip().replace('.0', '')
            acc_unique_all_docs.add(doc_str)
            
        if c is None or str(c).strip() in ['', 'NULL', 'None']:
            acc_lines_null_code += 1
        else:
            c_clean = str(c).strip().replace('.0', '')
            acc_lines_with_code += 1
            
            if c_clean in code_to_meta:
                acc_lines_matched += 1
                
                # Retrieve mapping metadata
                # Primary meta is the first mapped item for this code
                m_list = code_to_meta[c_clean]
                primary_m = m_list[0]
                cid = primary_m['cc_id']
                cname = primary_m['cc_name']
                domain = primary_m['domain']
                b_desc = primary_m['desc']
                
                # Accounting fields
                deb_val = float(r[idx_deb]) if r[idx_deb] is not None and str(r[idx_deb]).strip() != 'NULL' else 0.0
                cred_val = float(r[idx_cred]) if r[idx_cred] is not None and str(r[idx_cred]).strip() != 'NULL' else 0.0
                net_val = deb_val - cred_val
                
                area_name = str(r[idx_area_name]).strip() if r[idx_area_name] else ''
                date_str = str(r[idx_date]).strip() if r[idx_date] else '1404'
                acc_desc = str(r[idx_desc]).strip() if r[idx_desc] else ''
                account_code = f"{r[idx_k_no]}-{r[idx_m_no]}-{r[idx_t_no]}"
                
                # Attach to Doc_No_Mapping
                doc_mapping_rows.append({
                    'شماره سند حسابداری': doc_str,
                    'سال مالی': 1404,
                    'ردیف سند حسابداری': sheet_lines,
                    'شرح سند حسابداری': acc_desc,
                    'کد حساب': account_code,
                    'کد بودجه / ردیف مالی': c_clean,
                    'شرح ردیف مالی': b_desc,
                    'کد مرکز هزینه پیشنهادی': cid,
                    'عنوان مرکز هزینه پیشنهادی': cname,
                    'حوزه خدمت': domain,
                    'مبلغ بدهکار': deb_val,
                    'مبلغ بستانکار': cred_val,
                    'مبلغ خالص': net_val,
                    'روش تطبیق': 'تطبیق قطعی با کد بودجه',
                    'کلید تطبیق': c_clean,
                    'وضعیت تطبیق': 'تطبیق قطعی',
                    'درجه اطمینان': 'بالا',
                    'توضیح': f"{area_name} - تاریخ: {date_str}"
                })
                
                # Update CC docs
                if cid not in cc_docs: cc_docs[cid] = set()
                cc_docs[cid].add(doc_str)
                
                # Track multi-CC docs
                if doc_str not in doc_to_ccs: doc_to_ccs[doc_str] = set()
                doc_to_ccs[doc_str].add(cid)
                
                if doc_str not in doc_to_codes: doc_to_codes[doc_str] = set()
                doc_to_codes[doc_str].add(c_clean)
                
                doc_to_lines[doc_str] = doc_to_lines.get(doc_str, 0) + 1
            else:
                acc_lines_unmatched_revenue += 1
                
    print(f"Completed sheet [{sname}] ({sheet_lines} lines) in {time.time()-t0:.1f}s.")

wb_acc.close()
print(f"Streaming finished in {time.time()-start_total:.1f}s.")
print(f"Total matched lines for Doc_No_Mapping: {len(doc_mapping_rows)}")

# Save checkpoint
