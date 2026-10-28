import sys
import os
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np

print("=== STARTING FULL PIPELINE EXECUTION ===")

# Paths
path_b = r'C:\Users\sadra\Downloads\مراکز هزینه طبق بودجه.xlsx'
path_cc = r'C:\Users\sadra\Downloads\cost center .xlsx'
path_act = r'C:\Users\sadra\Downloads\activity.xlsx'
path_org = r'C:\Users\sadra\Downloads\آمار پست های معاونتها-1405.xlsx'
output_excel = r'C:\Users\sadra\Downloads\تحلیل_جامع_مراکز_هزینه_و_فعالیت_های_شهرداری_اصفهان.xlsx'

# 1. Load Data
df_b = pd.read_excel(path_b, sheet_name='انتقال کل مراکز هزینه 1404')
df_cc = pd.read_excel(path_cc, sheet_name='اولیه', header=1)
df_cc.columns = ['کد بودجه', 'شرح ردیف', 'نوع هزینه', 'کد سطح 1', 'عنوان سطح 1', 'کد سطح 2', 'عنوان سطح 2', 'متولی']
df_act = pd.read_excel(path_act, sheet_name='نحلیل فعالیت برای زیرموضوع')

for c in df_b.columns: df_b[c] = df_b[c].fillna('').astype(str).str.strip()
for c in df_cc.columns: df_cc[c] = df_cc[c].fillna('').astype(str).str.strip()

# Topic lookups
b_lookup = {}
for _, r in df_b.iterrows():
    c = r['کد بودجه']
    d = r['شرح ردیف'].replace('\u200c', ' ').strip()
    b_lookup[(c, d)] = (r['موضوع'], r['زیرموضوع'])

code_to_topic = df_b.drop_duplicates(subset=['کد بودجه'])[['کد بودجه', 'موضوع', 'زیرموضوع']].set_index('کد بودجه')

# Master Catalogue
import generate_all_deliverables
cc_catalog = generate_all_deliverables.cc_catalog

# Helper function to classify a single row
def classify_row(code, desc, cost_type, old_l1, old_l2, old_owner, topic, subtopic):
    desc_clean = desc.replace('\u200c', ' ').strip()
    d_lower = desc_clean.lower()
    
    # Defaults
    cc_id = None
    confidence = 'بالا'
    status = 'تأیید نگاشت موجود'
    reason = ''
    
    # 1. SPECIAL CASE: Public Restrooms (سرویس‌های بهداشتی)
    if 'سرویس بهداشتی' in d_lower or 'سرویس های بهداشتی' in d_lower:
        if 'ساختمان معاونت' in d_lower or 'ساختمان منطقه' in d_lower or 'شهرداری مرکزی' in d_lower:
            cc_id = 'CC-1401'
            reason = 'احداث یا بهسازی سرویس بهداشتی اختصاصی در ساختمان اداری شهرداری.'
        elif 'گلزار شهدا' in d_lower or 'مسجد' in d_lower or 'امامزاده' in d_lower:
            cc_id = 'CC-906'
            reason = 'سرویس بهداشتی و وضوخانه اختصاصی مساجد، گلزار شهدا و اماکن متبرکه.'
        elif 'آرامستان' in d_lower:
            cc_id = 'CC-907'
            reason = 'سرویس بهداشتی اختصاصی آرامستان.'
        else:
            cc_id = 'CC-106'
            reason = 'احداث، خرید سازه پیش‌ساخته بتنی یا بهسازی سرویس بهداشتی عمومی شهر و پارک‌ها؛ تجمیع پراکندگی قبلی در مرکز بهداشت و تسهیلات عمومی.'

    # 2. CLEANLINESS & WASTE (پاکیزگی و مدیریت پسماند)
    elif cc_id is None and ('ماشین آلات جاروب' in d_lower or 'جاروب مکانیزه' in d_lower):
        cc_id = 'CC-101'
        reason = 'خرید ماشین‌آلات و تجهیزات سرمایه‌ای جهت رفت‌وروب و نظافت مکانیزه معابر شهری.'
    elif cc_id is None and ('نظافت شهری' in d_lower or (topic == 'نظافت شهری' and 'ساختمان' not in d_lower)):
        cc_id = 'CC-101'
        reason = 'قراردادهای حجمی و رفت‌وروب و نظافت روزمره معابر و انهار سطح شهر.'
    elif cc_id is None and ('نخاله' in d_lower and 'فضای سبز' not in d_lower) or 'سکوی زباله' in d_lower or 'ناوگان پسماند' in d_lower or 'کرایه و حمل نخاله' in d_lower:
        cc_id = 'CC-102'
        reason = 'جمع‌آوری و حمل نخاله‌ها و پسماندهای حجیم یا نوسازی ناوگان جمع‌آوری پسماند شهری.'
    elif cc_id is None and ('ایستگاه های سیار بازیافت' in d_lower or 'ایستگاه سیار بازیافت' in d_lower or 'ایستگاههای سیار بازیافت' in d_lower or 'شهرک بازیافت' in d_lower or 'کارخانه بازیافت' in d_lower or 'همزن' in d_lower or 'ترنر' in d_lower):
        cc_id = 'CC-103'
        reason = 'توسعه ایستگاه‌های تفکیک از مبدأ پسماند خشک یا تجهیزات پردازش و کمپوست در کارخانه پسماند؛ اصلاح خطای قبلی که به دلیل نام محل به فضای سبز یا آب اختصاص یافته بود.'
    elif cc_id is None and ('لندفیل' in d_lower or 'هاضم' in d_lower or 'زباله سوز' in d_lower or 'زباله‌سوز' in d_lower):
        cc_id = 'CC-104'
        reason = 'احداث لندفیل‌های بهداشتی مهندسی، واحدهای هاضم بیوهوازی یا زباله‌سوز نوین.'
    elif cc_id is None and ('حیوانات ناقل بیماری' in d_lower or 'جانوران موذی' in d_lower or 'سگ های بلاصاحب' in d_lower):
        cc_id = 'CC-105'
        reason = 'کنترل بیولوژیک، طعمه‌گذاری و جمع‌آوری حیوانات ناقل بیماری و بهداشت عمومی.'
    elif cc_id is None and 'افزایش سرمایه سازمان مدیریت پسماند' in d_lower:
        cc_id = 'CC-1204' # Capital investment / equity financing
        reason = 'افزایش سرمایه شرکت‌ها و سازمان‌های وابسته؛ یک سرمایه‌گذاری مالی و تأمین سرمایه است نه عملیات رفت‌وروب.'

    # 3. INTERNAL CLEANING & GUESTHOUSES (نظافت داخلی و مهمانسراها)
    elif cc_id is None and ('نظافت ساختمان' in d_lower or 'نظافت و خدمات اداری' in d_lower):
        cc_id = 'CC-1404'
        reason = 'برون‌سپاری خدمات نظافت ساختمان‌های اداری و ستادی؛ تفکیک از رفت‌وروب خیابان‌ها و انتقال به خدمات عمومی پشتیبانی.'
    elif cc_id is None and ('نظافت و خانه داری' in d_lower or 'زائر سرا' in d_lower or 'مهمانسرای' in d_lower or 'مهمانسرا' in d_lower):
        cc_id = 'CC-1302'
        reason = 'خدمات نظافت و هتل‌داری مهمانسراهای اقامتی و تفریحی پرسنل شهرداری؛ ماهیت خدمات رفاهی کارکنان دارد.'

    # 4. WATER & WASTEWATER (آب، پساب، مادی‌ها و دفع آب‌های سطحی)
    elif cc_id is None and ('پساب' in d_lower or 'تصفیه پیشرفته پساب' in d_lower or 'خط انتقال پساب' in d_lower or 'بازچرخانی' in d_lower):
        cc_id = 'CC-302'
        reason = 'تصفیه تکمیلی و بازچرخانی پساب شهری و خطوط انتقال آن؛ اصلاح انتساب اشتباه قبلی به رفت‌وروب یا فضای سبز.'
    elif cc_id is None and ('مادی' in d_lower or 'مادیها' in d_lower or 'انهار' in d_lower or topic == 'مادی و انهار'):
        cc_id = 'CC-303'
        reason = 'ساماندهی، لایروبی، دیواره‌سازی و احیای جریان آب در شبکه مادی‌های تاریخی شهر.'
    elif cc_id is None and ('چاه جذبی' in d_lower or 'دفع آبهای سطحی' in d_lower or 'دفع آب های سطحی' in d_lower or topic == 'هدایت و دفع آب های سطحی' or 'سیلاب و آبگرفتگی' in d_lower):
        cc_id = 'CC-304'
        reason = 'احداث کانال‌های جمع‌آوری روان‌آب‌ها، چاه‌های جذبی و پیشگیری از آب‌گرفتگی معابر شهری.'
    elif cc_id is None and (topic == 'شبکه آبیاری' or 'شبکه توزیع' in d_lower or 'مخزن' in d_lower or 'مخازن' in d_lower or 'آبرسانی با تانکر' in d_lower or 'بابلر' in d_lower or 'کلکتور' in d_lower or 'لوله گذاری آب' in d_lower):
        cc_id = 'CC-301'
        reason = 'توسعه، احداث و نگهداری خطوط انتقال آب خام، لوله‌گذاری قطره‌ای، مخازن ژئوممبران و ایستگاه‌های پمپاژ آب فضای سبز.'

    # 5. GREEN SPACE (فضای سبز)
    elif cc_id is None and ('پارک صفه' in d_lower or 'تفرجگاه' in d_lower or 'ناژوان' in d_lower or 'باغ گلها' in d_lower or 'باغ پرندگان' in d_lower or 'تله کابین' in d_lower or 'تله سیژ' in d_lower):
        cc_id = 'CC-204'
        reason = 'نگهداری و ساماندهی مجموعه‌ها و تفرجگاه‌های ویژه طبیعی-گردشگری شهر (صفه، ناژوان و باغ‌ها).'
    elif cc_id is None and ('گلخانه' in d_lower or 'تولیدات گیاهی' in d_lower or 'تولید گل' in d_lower or 'بذر' in d_lower):
        cc_id = 'CC-203'
        reason = 'تولید نهال و گل در گلخانه‌های تخصصی شهرداری جهت بازکاشت در معابر.'
    elif cc_id is None and ('احداث بوستان' in d_lower or 'احداث فضای سبز' in d_lower or 'توسعه فضای سبز' in d_lower or 'جنگل کاری' in d_lower or 'کمربند سبز' in d_lower):
        cc_id = 'CC-202'
        reason = 'سرمایه‌گذاری عمرانی در تملک و احداث بوستان‌ها و پارک‌های محله‌ای جدید.'
    elif cc_id is None and (topic == 'فضای سبز' or 'بهزراعی' in d_lower or 'هرس' in d_lower or 'چمن' in d_lower or 'سموم دفع آفات' in d_lower or 'ضایعات فضای سبز' in d_lower or 'نهاده های بهزراعی' in d_lower):
        cc_id = 'CC-201'
        reason = 'نگهداری روزمره، بهزراعی، چمن‌زنی و مبارزه با آفات نباتی در پارک‌ها و رفیوژهای معابر.'

    # 6. TRAFFIC & TRANSIT (ترافیک و حمل‌ونقل)
    elif cc_id is None and ('دوربین' in d_lower or 'ثبت تخلف' in d_lower or 'کنترل سرعت' in d_lower or 'سیستم حمل و نقل هوشمند' in d_lower or 'هوشمندسازی تقاطع' in d_lower or 'چراغ های راهنمایی' in d_lower or topic == 'مدیریت هوشمند ترافیک'):
        cc_id = 'CC-401'
        reason = 'سامانه‌های هوشمند مدیریت ترافیک، دوربین‌های پلاک‌خوان و نظارتی، فیبر نوری و چراغ‌های هوشمند تقاطع‌ها.'
    elif cc_id is None and ('خط کشی' in d_lower or 'تابلوهای ترافیکی' in d_lower or 'جداکننده' in d_lower or 'موانع کنترل سرعت' in d_lower or 'سرعت گیر' in d_lower or topic == 'ایمنی ترافیکی' or 'ستاد مهر' in d_lower):
        cc_id = 'CC-402'
        reason = 'تجهیزات فیزیکی ایمنی عبور و مرور، خط‌کشی طولی و عرضی معابر، تابلوهای هدایت مسیر و آرام‌سازی ترافیک.'
    elif cc_id is None and ('ایستگاه های بی آر تی' in d_lower or 'بی آر تی' in d_lower or 'سرپناه اتوبوس' in d_lower or 'پایانه' in d_lower or topic == 'پایانه و پارک سوار'):
        cc_id = 'CC-403'
        reason = 'احداث، نگهداری و تجهیز پایانه‌های اتوبوسرانی و ایستگاه‌های سامانه‌های اتوبوس تندرو (BRT).'
    elif cc_id is None and ('پل عابر' in d_lower or 'پل های عابر پیاده' in d_lower or 'پله برقی' in d_lower or topic == 'پل عابر پیاده'):
        cc_id = 'CC-404'
        reason = 'احداث سازه، مکانیزه‌سازی با پله‌برقی و آسانسور و نگهداری دوره‌ای پل‌های هوایی عابر پیاده.'
    elif cc_id is None and ('پارکینگ' in d_lower or topic == 'پارکینگ' or 'پارک حاشیه ای' in d_lower):
        cc_id = 'CC-405'
        reason = 'احداث پارکینگ‌های طبقاتی عمومی، بهسازی سیستم‌های ورودی و ساماندهی پارک خودرو در حاشیه خیابان‌ها.'
    elif cc_id is None and (topic == 'اتوبوس' or 'اتوبوسرانی' in d_lower or 'بازسازی ناوگان' in d_lower or 'شرکت اتوبوسرانی' in d_lower or topic == 'سامان دهی جابجایی بار در شهر' or topic == 'نظارت بر جابجایی مسافر با تاکسی'):
        cc_id = 'CC-406'
        reason = 'یارانه بلیت، خرید و بازسازی اتوبوس‌های مسافربری و نظارت بر ناوگان حمل‌ونقل همگانی و بار.'
    elif cc_id is None and (topic == 'قطار شهری' or 'مترو' in d_lower or 'خط ۲ قطار شهری' in d_lower or 'ایستگاه های مترو' in d_lower):
        cc_id = 'CC-407'
        reason = 'طرح‌های سرمایه‌ای حفر تونل، احداث ایستگاه‌ها، ریل‌گذاری و ناوگان قطار شهری اصفهان.'
    elif cc_id is None and ('شارژ خودرو برقی' in d_lower or 'دوچرخه' in d_lower or topic == 'دوچرخه' or 'حمل و نقل پاک' in d_lower):
        cc_id = 'CC-408'
        reason = 'توسعه مسیرهای اختصاصی دوچرخه‌سواری، ایستگاه‌های دوچرخه و جایگاه‌های شارژ خودروها و موتورسیکلت‌های برقی.'

    # 7. CIVIL WORKS (عمران و معابر)
    elif cc_id is None and (topic == 'آسفالت معابر شهری' or 'آسفالت' in d_lower or 'روکش' in d_lower or 'لکه گیری' in d_lower or 'درزگیری' in d_lower or 'قیر' in d_lower or 'ترانشه' in d_lower):
        cc_id = 'CC-502'
        reason = 'روکش مکانیزه آسفالت بزرگراه‌ها و خیابان‌ها، لکه‌گیری معابر فرعی و ترمیم ترانشه‌های آسفالتی.'
    elif cc_id is None and (topic == 'پیاده روسازی' or 'پیاده رو' in d_lower or 'سنگ فرش' in d_lower or 'موزاییک فرش' in d_lower):
        cc_id = 'CC-503'
        reason = 'احداث، بازسازی و سنگ‌فرش کالبدی پیاده‌روها و مسیرهای عابر پیاده در محلات و مراکز تجاری.'
    elif cc_id is None and (topic == 'مناسب سازی معابر جهت تردد معلولین و جانبازان' or 'معلولین' in d_lower or 'بساوایی' in d_lower or 'جانبازان' in d_lower):
        cc_id = 'CC-504'
        reason = 'مناسب‌سازی فضاهای شهری، نصب موزاییک بساوایی نابینایان و رمپ‌های ویژه تردد معلولین و جانبازان.'
    elif cc_id is None and ('تقاطع غیرهمسطح' in d_lower or 'زیرگذر' in d_lower or 'روگذر' in d_lower or 'پل سواره' in d_lower or 'عرشه پل' in d_lower or 'درز انبساط پل' in d_lower or 'دیوار حائل' in d_lower):
        cc_id = 'CC-505'
        reason = 'احداث و نگهداری تخصصی سازه پل‌های بزرگراهی، تقاطع‌های غیرهمسطح، زیرگذرها و ابنیه فنی معابر.'
    elif cc_id is None and ('حفاری' in d_lower or 'جابجایی تاسیسات' in d_lower or 'آزمایشگاه خاک' in d_lower):
        cc_id = 'CC-506'
        reason = 'مدیریت و کنترل آزمایشگاهی حفاری‌های شرکت‌های خدمات‌رسان و هماهنگی تأسیسات شهری.'
    elif cc_id is None and (topic == 'شبکه معابر' or 'احداث خیابان' in d_lower or 'جدول گذاری' in d_lower or 'اصلاح هندسی' in d_lower or 'زیرسازی' in d_lower or 'کانیو' in d_lower or 'ساماندهی ورودی' in d_lower):
        cc_id = 'CC-501'
        reason = 'احداث، بازگشایی و ساماندهی فیزیکی خیابان‌ها، جدول‌گذاری معابر و اصلاح هندسی میادین و تقاطع‌ها.'

    # 8. BEAUTIFICATION & LIGHTING (زیباسازی و روشنایی)
    elif cc_id is None and ('روشنایی' in d_lower or 'نورپردازی' in d_lower or 'لوکس نوری' in d_lower or 'پایه چراغ' in d_lower):
        cc_id = 'CC-601'
        reason = 'توسعه، احداث و نگهداری شبکه روشنایی معابر و بوستان‌ها و نورپردازی ابنیه تاریخی و پل‌ها.'
    elif cc_id is None and ('آبنما' in d_lower or 'فواره' in d_lower):
        cc_id = 'CC-603'
        reason = 'احداث، بازسازی و نگهداری تأسیسات مکانیکی و الکتریکی آبنماهای شهری و میادین.'
    elif cc_id is None and ('رنگ آمیزی' in d_lower or 'دیوارنگاره' in d_lower or 'نقاشی دیواری' in d_lower or topic == 'پیرایش شهری'):
        cc_id = 'CC-604'
        reason = 'دیوارنگاری هنری، رنگ‌آمیزی جداول و نرده‌ها و حذف زوائد و آلودگی‌های بصری نما.'
    elif cc_id is None and ('آذین بندی' in d_lower or 'فضاسازی مناسبتی' in d_lower or 'طاق نصرت' in d_lower or 'ریسه' in d_lower or 'پرچم' in d_lower):
        cc_id = 'CC-605'
        reason = 'آذین‌بندی، ریسه‌بندی نوری و فضاسازی محیطی اعیاد و مناسبت‌های مذهبی و ملی در سطح شهر.'
    elif cc_id is None and (topic == 'زیباسازی' or 'مبلمان' in d_lower or 'المان' in d_lower or 'نماد' in d_lower or 'نیمکت' in d_lower or 'پلاک های شهرداری' in d_lower or 'تابلوهای معابر' in d_lower):
        cc_id = 'CC-602'
        reason = 'تأمین و بهسازی عناصر مبلمان شهری، نیمکت‌ها، المان‌های محیطی و تابلوهای نام‌گذاری معابر.'

    # 9. URBAN PLANNING & HISTORIC (شهرسازی و بافت تاریخی)
    elif cc_id is None and (topic == 'بافت تاریخی' or 'دولتخانه صفوی' in d_lower or 'ناحیه ویژه تاریخی' in d_lower or 'بازارها و سراها' in d_lower or 'بازار خورشید' in d_lower or 'اماکن تاریخی' in d_lower or 'بافت فرسوده' in d_lower):
        cc_id = 'CC-704'
        reason = 'بازآفرینی، نوسازی و احیای هویت کالبدی بافت‌های فرسوده و ابنیه و بازارهای تاریخی شهر.'
    elif cc_id is None and (topic == 'نقشه برداری و برداشت و بروزرسانی' or 'نقشه برداری' in d_lower or 'gis' in d_lower or 'کاداستر' in d_lower or 'فتوگرامتری' in d_lower or 'اطلاعات مکانی' in d_lower):
        cc_id = 'CC-705'
        reason = 'عملیات نقشه‌برداری زمینی، عکس‌برداری هوایی، نگهداری ژئودیتابیس و اطلاعات مکانی (GIS).'
    elif cc_id is None and ('کمیسیون های ماده صد' in d_lower or 'ماده صد' in d_lower or 'گشت فنی ساختمان' in d_lower or topic == 'پیشگیری و رفع تخلفات شهری'):
        cc_id = 'CC-703'
        reason = 'امور گشت‌های پیشگیری از تخلفات ساختمانی و دبیرخانه اجرایی کمیسیون‌های ماده صد.'
    elif cc_id is None and ('مهندسین ناظر' in d_lower or 'کنترل ضوابط' in d_lower or 'ایستایی' in d_lower or 'شهرسازی الکترونیک' in d_lower or 'دفاتر پیشخوان' in d_lower):
        cc_id = 'CC-702'
        reason = 'نظارت بر نقشه‌ها، کنترل ضوابط ایستایی و ایمنی ساختمان‌ها و ارزیابی دفاتر پیشخوان شهرسازی.'
    elif cc_id is None and (topic == 'طرح های ویژه شهرسازی' or 'طرح تفصیلی' in d_lower or 'مطالعات شهرسازی' in d_lower or 'طرح جامع' in d_lower or 'حریم' in d_lower or 'خط بدنه' in d_lower):
        cc_id = 'CC-701'
        reason = 'تهیه طرح‌های توسعه شهری، تفصیلی، موضعی و موضوعی و ضوابط کالبدی شهرسازی.'

    # 10. SAFETY & DISASTER (ایمنی، بحران و آتش‌نشانی)
    elif cc_id is None and ('ایستگاه آتش نشانی' in d_lower and ('احداث' in d_lower or 'تجهیز' in d_lower or 'تکمیل' in d_lower)):
        cc_id = 'CC-802'
        reason = 'احداث و توسعه کالبدی ایستگاه‌های جدید آتش‌نشانی و خدمات ایمنی.'
    elif cc_id is None and (topic == 'ایمنی و مدیریت بحران' or 'آتش نشانی' in d_lower or 'اطفای حریق' in d_lower or 'امداد و نجات' in d_lower or 'البسه حفاظت فردی' in d_lower):
        cc_id = 'CC-801'
        reason = 'عملیات اطفای حریق، خرید و نگهداری تجهیزات تخصصی نجات و بازرسی ایمنی ساختمان‌ها.'
    elif cc_id is None and ('فرونشست' in d_lower or 'مدیریت بحران' in d_lower or 'پدافند غیرعامل' in d_lower or 'مانور' in d_lower):
        cc_id = 'CC-803'
        reason = 'مطالعات و برنامه‌ریزی مقابله با بحران‌های شهری، فرونشست دشت اصفهان و پدافند غیرعامل.'
    elif cc_id is None and ('رفع خطر' in d_lower or 'اماکن ناایمن' in d_lower or 'بند ۱۴' in d_lower or 'حفاظت ایمنی شهری' in d_lower):
        cc_id = 'CC-804'
        reason = 'رفع خطر کالبدی از دیوارهای شکسته، بناهای مخروبه و گودبرداری‌های خطرناک طبق بند ۱۴ ماده ۵۵.'

    # 11. CULTURE, SPORTS, RELIGION, CEMETERY (فرهنگ، ورزش، گردشگری و آرامستان)
    elif cc_id is None and (topic == 'خدمات امور متوفیان' or 'آرامستان' in d_lower or 'باغ رضوان' in d_lower or 'متوفیان' in d_lower):
        cc_id = 'CC-907'
        reason = 'مدیریت آرامستان‌های شهری، حمل متوفیات، تغسیل، تدفین و نگهداری قبور.'
    elif cc_id is None and (topic == 'اماکن مذهبی و تاریخی' or 'مسجد' in d_lower or 'مساجد' in d_lower or 'گلستان شهدا' in d_lower or 'امامزاده' in d_lower or 'حسینیه' in d_lower):
        cc_id = 'CC-906'
        reason = 'حمایت، بهسازی و نگهداری کالبدی مساجد، حسینیه‌ها و مجموعه فرهنگی-مذهبی گلستان شهدا.'
    elif cc_id is None and (topic == 'گردشگری' or 'گردشگری' in d_lower or 'تسهیلات سفر' in d_lower or 'اماکن گردشگری' in d_lower):
        cc_id = 'CC-905'
        reason = 'توسعه زیرساخت‌ها، ستاد اسکان مسافرین و نگهداری جاذبه‌های گردشگری شهرداری.'
    elif cc_id is None and ('آسیب های اجتماعی' in d_lower or 'کارتن خواب' in d_lower or 'متکدیان' in d_lower or 'گرمخانه' in d_lower or 'توانمندسازی شهروندی' in d_lower or 'حمایت اجتماعی' in d_lower):
        cc_id = 'CC-904'
        reason = 'پیشگیری از آسیب‌های اجتماعی، ساماندهی متکدیان و بی‌خانمان‌ها و توانمندسازی اقشار آسیب‌پذیر.'
    elif cc_id is None and ('ورزش' in d_lower or 'اماکن ورزشی' in d_lower or 'زمین چمن' in d_lower or 'استخر' in d_lower or 'سالن ورزشی' in d_lower or topic == 'شهر دوستدار کودک'):
        cc_id = 'CC-903'
        reason = 'توسعه ورزش همگانی، احداث و تجهیز زمین‌های ورزشی محلات و نگهداری سالن‌های ورزشی.'
    elif cc_id is None and ('فرهنگسرا' in d_lower or 'کتابخانه' in d_lower or 'سالن مطالعه' in d_lower or topic == 'مراکز فرهنگی، اجتماعی، ورزشی'):
        cc_id = 'CC-901'
        reason = 'احداث، تجهیز و تعمیرات اساسی فرهنگسراها، کتابخانه‌ها و مراکز فرهنگی محلات.'
    elif cc_id is None and (topic == 'برنامه های فرهنگی اجتماعی' or 'برنامه های جشنواره' in d_lower or 'همایش ها' in d_lower or 'رویدادها' in d_lower or 'استعداد یابی' in d_lower or 'فرهنگی و هنری' in d_lower or 'برنامه های هویتی' in d_lower or 'نوجوانی فصل تحول' in d_lower):
        cc_id = 'CC-902'
        reason = 'اجرای جشنواره‌ها، جشن‌های مناسبتی، مسابقات قرآن و برنامه‌های فرهنگی-هنری محلات و شهر.'

    # 12. REAL ESTATE & LEGAL (املاک و حقوقی)
    elif cc_id is None and (topic == 'بازگشایی گلوگاه های محلی' or 'تملک' in d_lower or 'آزادسازی' in d_lower or topic == 'تامین معوض' or 'غرامت' in d_lower):
        cc_id = 'CC-1001'
        reason = 'تملک املاک واقع در طرح‌های عمرانی، پرداخت غرامت عرصه و اعیان و بازگشایی گلوگاه‌ها.'
    elif cc_id is None and ('بانک اطلاعات املاک' in d_lower or 'ثبت اسناد و املاک' in d_lower or 'سند تک برگ' in d_lower):
        cc_id = 'CC-1002'
        reason = 'تثبیت مالکیت شهرداری بر اراضی و املاک، اخذ اسناد رسمی و ساماندهی سامانه جامع املاک.'
    elif cc_id is None and ('حقوقی' in d_lower or 'دعاوی' in d_lower or 'دادگاه' in d_lower or 'قراردادها' in d_lower or 'حق الوکاله' in d_lower or 'تمبر' in d_lower):
        cc_id = 'CC-1003'
        reason = 'دفاع حقوقی از شهرداری در محاکم قضایی، تنظیم لوایح و قراردادها و احکام دادگاه.'

    # 13. IT & SMART CITY (فناوری اطلاعات)
    elif cc_id is None and ('نرم افزار' in d_lower or 'سامانه' in d_lower or 'سامانه ها' in d_lower or 'لایسنس' in d_lower):
        cc_id = 'CC-1101'
        reason = 'توسعه، لایسنس، نگهداری و ارتقای نرم‌افزارهای یکپارچه اداری، مالی و خدماتی شهرداری.'
    elif cc_id is None and (topic == 'فناوری اطلاعات' or 'مرکز داده' in d_lower or 'سرور' in d_lower or 'شبکه' in d_lower or 'دیتاسنتر' in d_lower or 'کابل کشی شبکه' in d_lower):
        cc_id = 'CC-1102'
        reason = 'تأمین و نگهداری سرورها، امنیت سایبری، مراکز داده و زیرساخت ارتباطی شبکه شهرداری.'
    elif cc_id is None and (topic == 'هوشمندسازی و خدمات الکترونیکی' or 'خدمات الکترونیک' in d_lower or 'شهر هوشمند' in d_lower or 'درگاه' in d_lower or 'اپلیکیشن' in d_lower):
        cc_id = 'CC-1103'
        reason = 'درگاه‌های یکپارچه خدمات الکترونیک شهروندی، اپلیکیشن «اصفهان من» و پنجره واحد شهر هوشمند.'

    # 14. FINANCIAL MANAGEMENT & DEBTS (مالی، درآمد و دیون)
    elif cc_id is None and ('ممیزی املاک' in d_lower or 'قبوض عوارض' in d_lower or 'وصول درآمد' in d_lower or 'تشخیص درآمد' in d_lower):
        cc_id = 'CC-1201'
        reason = 'ممیزی منابع درآمدی، محاسبه و وصول عوارض نوسازی، خودرو و مطالبات پایدار شهرداری.'
    elif cc_id is None and (topic == 'هزینه های مالی' or 'سود تسهیلات' in d_lower or 'کارمزد وام' in d_lower or 'اوراق مشارکت' in d_lower or 'تسهیلات بانکی' in d_lower):
        cc_id = 'CC-1204'
        reason = 'کارمزد و سود تسهیلات بانکی، عاملیت انتشار اوراق مشارکت و ساختارهای تأمین مالی طرح‌ها.'
    elif cc_id is None and (topic == 'دیون' or 'دیون با محل' in d_lower or 'دیون بلامحل' in d_lower or 'بدهی سنواتی' in d_lower):
        cc_id = 'CC-1205'
        reason = 'تسویه و پرداخت دیون قطعی سنوات گذشته به پیمانکاران، اشخاص، دارایی و تأمین اجتماعی.'
    elif cc_id is None and ('حسابرسی' in d_lower or 'حسابal' in d_lower or 'کنترل داخلی' in d_lower):
        cc_id = 'CC-1203'
        reason = 'حق‌الزحمه حسابرسان رسمی مستقل منتخب شورای اسلامی شهر و ارزیابی انطباق مالی.'
    elif cc_id is None and ('مالی، خزانه و حسابداری' in old_l2 or 'خزانه داری' in d_lower or 'بستن حسابها' in d_lower or 'کارمزد خدمات بانکی' in d_lower):
        cc_id = 'CC-1202'
        reason = 'امور مالی، دریافت و پرداخت‌های خزانه‌داری، حسابداری تعهدی و تنظیم صورت‌های مالی تلفیقی.'

    # 15. HUMAN RESOURCES (منابع انسانی و رفاه)
    elif cc_id is None and (topic == 'حقوق و مزایا' or 'حقوق ثابت' in d_lower or 'دستمزد کارگران' in d_lower or 'حقوق کارمندان' in d_lower or 'حقوق مامورین' in d_lower or 'مزایای کارمندان' in d_lower):
        cc_id = 'CC-1301'
        reason = 'جبران خدمات ماهانه، حقوق مبنا، دستمزد کارگری و مزایای قانونی پرسنل رسمی و قراردادی.'
    elif cc_id is None and (topic == 'رفاهیات' or 'بازنشستگان' in d_lower or 'بیمه تکمیلی' in d_lower or 'بن کارکنان' in d_lower or 'رفاهی کارکنان' in d_lower or 'مهد کودک' in d_lower):
        cc_id = 'CC-1302'
        reason = 'بیمه درمان تکمیلی، بن‌کارت‌های معیشتی و خدمات رفاهی کارکنان شاغل و بازنشستگان.'
    elif cc_id is None and ('آموزش' in d_lower and 'کارکنان' in d_lower):
        cc_id = 'CC-1303'
        reason = 'نیازسنجی و اجرای دوره‌های آموزشی ضمن خدمت و توسعه شایستگی‌های کارکنان شهرداری.'
    elif cc_id is None and ('طب کار' in d_lower or 'بهداشت و سلامت کارکنان' in d_lower or 'لباس کارکنان' in d_lower or 'حفاظت فردی' in d_lower):
        cc_id = 'CC-1304'
        reason = 'معاینات ادواری طب کار، بهداشت حرفه‌ای و البسه کار و حفاظت ایمنی پرسنل.'
    elif cc_id is None and ('تشکیلات' in d_lower or 'بایگانی الکترونیک' in d_lower or 'دبیرخانه' in d_lower or 'نظام تحول' in d_lower):
        cc_id = 'CC-1305'
        reason = 'اصلاح ساختار سازمانی، پست‌های اداری، دبیرخانه مرکزی و آرشیو اسناد شهرداری.'

    # 16. GENERAL LOGISTICS & SECURITY (پشتیبانی، تدارکات و حراست)
    elif cc_id is None and (topic == 'ساختمان های شهرداری' or 'ساختمان اداری' in d_lower or 'ساختمان شهرداری' in d_lower or 'نگهداری ساختمان و تاسیسات' in old_l2):
        cc_id = 'CC-1401'
        reason = 'احداث، تعمیرات و نوسازی کالبدی ابنیه اداری شهرداری مرکزی و مناطق.'
    elif cc_id is None and ('کرایه وسایل نقلیه' in d_lower or 'ترابری' in d_lower or 'وسائط نقلیه' in d_lower or 'سوخت' in d_lower):
        cc_id = 'CC-1402'
        reason = 'ترابری اداری، خودروهای استیجاری مأموریت‌های ستادی و سوخت ناوگان خدمات اداری.'
    elif cc_id is None and ('حفاظت فیزیکی' in d_lower or 'انتظامات' in d_lower or topic == 'حراست' or 'نگهبان' in d_lower):
        cc_id = 'CC-1405'
        reason = 'برون‌سپاری نگهبانی و حفاظت فیزیکی اماکن، گیت‌های ورودی و حراست ابنیه شهرداری.'
    elif cc_id is None and ('ناهار و پذیرایی' in d_lower or 'بهای آب اماکن' in d_lower or 'بهای برق مصرفی اماکن' in d_lower or 'خدمات شهری اتباع' in d_lower or 'اپراتورهای بی سیم' in d_lower or 'خدمات امور شهر' in d_lower):
        cc_id = 'CC-1404'
        reason = 'هزینه‌های مصارف آب، برق، پذیرایی جلسات و خدمات عمومی و سرایداری دفاتر شهرداری.'
    elif cc_id is None and ('لوازم مصرفی' in d_lower or 'کارپردازی' in d_lower or topic == 'تامین مصالح' or 'خرید اثاثیه' in d_lower):
        cc_id = 'CC-1403'
        reason = 'خرید ملزومات مصرفی، اثاثیه اداری، کاغذ و لوازم بایگانی و انبارداری عمومی.'

    # 17. COMMUNICATIONS & MEDIA (ارتباطات و روابط عمومی)
    elif cc_id is None and (topic == 'رسانه و ارتباطات مردمی' or 'انتشار محتوا' in d_lower or 'نشریات' in d_lower or 'تولیدات رسانه ای' in d_lower or 'روزنامه' in d_lower):
        cc_id = 'CC-1501'
        reason = 'تولید اخبار، برنامه‌های چندرسانه‌ای صدای شهر و سیمای شهر و نشریات مکتوب شهرداری.'
    elif cc_id is None and ('۱۳۷' in d_lower or '137' in d_lower or 'ارتباطات مردمی' in d_lower or 'پاسخگویی به شکایات' in d_lower):
        cc_id = 'CC-1502'
        reason = 'سامانه ۱۳۷ ارتباطات مردمی، ثبت فوریت‌های شهری و نظارت همگانی بازرسی.'
    elif cc_id is None and ('تشریفات' in d_lower or topic == 'شورای شهر' or 'امور شورا' in d_lower or 'بین الملل' in d_lower):
        cc_id = 'CC-1503'
        reason = 'امور تشریفات، میزبانی هیئت‌های دیپلماتیک و پشتیبانی جلسات شورای اسلامی شهر.'

    # 18. PLANNING & INNOVATION (برنامه‌ریزی و پژوهش)
    elif cc_id is None and ('بودجه' in d_lower or 'کنترل عملیات و مدیریت پروژه' in d_lower or 'ارزیابی عملکرد' in d_lower or 'داده کاوی' in d_lower):
        cc_id = 'CC-1601'
        reason = 'تنظیم بودجه، تحلیل داده‌ها و کنترل پیشرفت پروژه‌های شهری در سامانه کنترل پروژه.'
    elif cc_id is None and (topic == 'آموزش، مطالعات و ارزیابی' or 'مدیریت دانش' in d_lower or 'پژوهش' in d_lower or 'دانش بنیان' in d_lower or 'خلاقیت' in d_lower or 'نظام پیشنهاد' in d_lower):
        cc_id = 'CC-1602'
        reason = 'پژوهش‌های کاربردی شهری، حمایت از استارتاپ‌ها، نظام پیشنهادات و مدیریت دانش.'

    # 19. FALLBACK BASED ON EXISTING LEVEL 2
    if cc_id is None:
        # Fallback mapping
        if 'ورزش' in old_l2: cc_id = 'CC-903'
        elif 'فضای سبز' in old_l2: cc_id = 'CC-201'
        elif 'پیاده‌رو' in old_l2 or 'پیاده رو' in old_l2: cc_id = 'CC-503'
        elif 'خیابان' in old_l2 or 'معبر' in old_l2: cc_id = 'CC-501'
        elif 'تقاطع' in old_l2: cc_id = 'CC-501'
        elif 'آسفالت' in old_l2: cc_id = 'CC-502'
        elif 'کانال' in old_l2 or 'قنوات' in old_l2: cc_id = 'CC-304'
        elif 'آبیاری' in old_l2: cc_id = 'CC-301'
        elif 'روشنایی' in old_l2: cc_id = 'CC-601'
        elif 'مبلمان' in old_l2: cc_id = 'CC-602'
        elif 'رویداد' in old_l2 or 'همایش' in old_l2: cc_id = 'CC-902'
        elif 'فرهنگسرا' in old_l2 or 'کتابخانه' in old_l2: cc_id = 'CC-901'
        elif 'تدارکات' in old_l2:
            cc_id = 'CC-1403'
            reason = 'تدارکات و پشتیبانی عمومی؛ سرفصل نیازمند بررسی ریزاقلام قرارداد.'
            confidence = 'متوسط'
            status = 'اصلاح مرکز هزینه'
        elif 'جبران خدمات' in old_l2: cc_id = 'CC-1301'
        elif 'رفاه' in old_l2: cc_id = 'CC-1302'
        elif 'ساختمان' in old_l2: cc_id = 'CC-1401'
        elif 'پارکینگ' in old_l2: cc_id = 'CC-405'
        elif 'دوربین' in old_l2: cc_id = 'CC-401'
        elif 'آتش‌نشانی' in old_l2: cc_id = 'CC-801'
        elif 'گلوگاه' in old_l2: cc_id = 'CC-1001'
        elif 'مناسب‌سازی' in old_l2: cc_id = 'CC-504'
        elif 'بافت' in old_l2: cc_id = 'CC-704'
        elif 'نورپردازی' in old_l2: cc_id = 'CC-601'
        elif 'آبنما' in old_l2: cc_id = 'CC-603'
        elif 'قطار شهری' in old_l2: cc_id = 'CC-407'
        elif 'هوشمندسازی' in old_l2: cc_id = 'CC-1103'
        elif 'نرم‌افزار' in old_l2: cc_id = 'CC-1101'
        elif 'پایانه' in old_l2: cc_id = 'CC-403'
        elif 'رسانه‌ای' in old_l2: cc_id = 'CC-1501'
        elif 'پل عابر' in old_l2: cc_id = 'CC-404'
        elif 'ارتباطات مردمی' in old_l2: cc_id = 'CC-1502'
        elif 'ایمن‌سازی' in old_l2: cc_id = 'CC-804'
        elif 'مالی' in old_l2: cc_id = 'CC-1202'
        elif 'دیوان' in old_l2 or 'حقوقی' in old_l2: cc_id = 'CC-1003'
        else:
            cc_id = 'CC-1403'
            confidence = 'پایین'
            status = 'ابهام نیازمند بررسی انسانی'
            reason = 'شرح کلی؛ انتساب پیش‌فرض به تدارکات اداری تا بررسی قرارداد.'

    # Determine status & reason if not set
    meta = cc_catalog[cc_id]
    if not reason:
        reason = f'تطبیق مفهومی شرح ردیف با ماهیت عملیاتی «{meta["name"]}».'

    # Compare with existing
    if status == 'تأیید نگاشت موجود':
        old_l2_clean = old_l2.replace('\u200c', ' ').strip()
        new_name_clean = meta['name'].replace('\u200c', ' ').strip()
        if old_l2_clean != new_name_clean:
            if 'تدارکات و پشتیبانی عمومی' in old_l2:
                status = 'اصلاح مرکز هزینه (تفکیک از سبد متفرقه)'
            elif old_l1 != meta['domain']:
                status = 'اصلاح مرکز هزینه و حوزه خدمت'
            else:
                status = 'اصلاح و تدقیق مرکز هزینه'

    # Check owner alignment
    owner_str = meta['owner_l1']
    if old_owner and old_owner not in ['nan', 'نامشخص', 'مشترک', 'نیازمند تصمیم کارشناسی']:
        if old_owner.replace('\u200c', ' ') != owner_str.replace('\u200c', ' '):
            if 'فرهنگی' in old_owner and 'فرهنگی' in owner_str:
                pass # synonymous
            elif 'عمران' in old_owner and 'عمران' in owner_str:
                pass
            else:
                if 'اصلاح' not in status:
                    status = 'اصلاح متولی سازمانی'
                else:
                    status += ' و اصلاح متولی'

    return {
        'cc_id': cc_id,
        'domain': meta['domain'],
        'activity': meta['activity'],
        'subactivity': meta['subactivity'],
        'cc_name': meta['name'],
        'owner_l1': meta['owner_l1'],
        'owner_l2': meta['owner_l2'],
        'owner_l3': meta['owner_l3'],
        'driver': meta['driver'],
        'output': meta['output'],
        'unit': meta['unit'],
        'confidence': confidence,
        'status': status,
        'reason': reason
    }

print("Classification function configured successfully.")

# 2. Iterate and classify all 2830 rows
mapped_rows = []
for idx, r in df_cc.iterrows():
    c = r['کد بودجه']
    d = r['شرح ردیف']
    ct = r['نوع هزینه']
    old_l1 = r['عنوان سطح 1']
    old_l2 = r['عنوان سطح 2']
    old_owner = r['متولی']
    
    # Retrieve topic and subtopic
    key = (c, d.replace('\u200c', ' ').strip())
    if key in b_lookup:
        topic, subtopic = b_lookup[key]
    elif c in code_to_topic.index:
        topic, subtopic = code_to_topic.loc[c, 'موضوع'], code_to_topic.loc[c, 'زیرموضوع']
    else:
        topic, subtopic = 'نامشخص (ردیف کارشناسی)', 'نامشخص'
        
    res = classify_row(c, d, ct, old_l1, old_l2, old_owner, topic, subtopic)
    
    mapped_rows.append({
        'ردیف': idx + 1,
        'کد بودجه': c,
        'شرح ردیف': d,
        'نوع هزینه': ct,
        'موضوع موجود': topic,
        'زیرموضوع موجود': subtopic,
        'مرکز هزینه فعلی سطح 1': old_l1,
        'مرکز هزینه فعلی سطح 2': old_l2,
        'حوزه خدمت پیشنهادی': res['domain'],
        'فعالیت اصلی پیشنهادی': res['activity'],
        'زیر‌فعالیت پیشنهادی': res['subactivity'],
        'مرکز هزینه پیشنهادی': res['cc_name'],
        'کد مرکز هزینه پیشنهادی': res['cc_id'],
        'معاونت متولی': res['owner_l1'],
        'واحد سازمانی متولی': f"{res['owner_l2']} - {res['owner_l3']}",
        'محرک هزینه پیشنهادی': res['driver'],
        'خروجی فعالیت': res['output'],
        'واحد سنجش': res['unit'],
        'درجه اطمینان': res['confidence'],
        'وضعیت نگاشت': res['status'],
        'دلیل/توضیح تحلیل': res['reason']
    })

df_mapped = pd.DataFrame(mapped_rows)
print(f"Successfully classified all {len(df_mapped)} rows!")
print("\nMapping Status Distribution:")
print(df_mapped['وضعیت نگاشت'].value_counts())
print("\nConfidence Distribution:")
print(df_mapped['درجه اطمینان'].value_counts())
print("\nTop 10 Proposed Cost Centers by Row Count:")
print(df_mapped['مرکز هزینه پیشنهادی'].value_counts().head(10))

# 3. Cleanliness Pilot Extraction
pilot_mask = (df_mapped['حوزه خدمت پیشنهادی'] == 'پاکیزگی و مدیریت پسماند شهری') | \
             (df_mapped['مرکز هزینه فعلی سطح 1'].str.contains('پاکیزگی|پسماند', na=False)) | \
             (df_mapped['شرح ردیف'].str.contains('نظافت|پسماند|بازیافت|جاروب|نخاله|سرویس بهداشتی|سکوی زباله|لندفیل|هاضم', na=False))

df_pilot = df_mapped[pilot_mask].copy()
print(f"\nPilot Cleanliness dataset extracted: {len(df_pilot)} rows.")

# 4. Master Dictionary Dataframe
dict_rows = []
for cc_id, m in cc_catalog.items():
    cnt = (df_mapped['کد مرکز هزینه پیشنهادی'] == cc_id).sum()
    sample_lines = df_mapped[df_mapped['کد مرکز هزینه پیشنهادی'] == cc_id]['شرح ردیف'].head(3).tolist()
    dict_rows.append({
        'کد مرکز هزینه': m['id'],
        'حوزه خدمت': m['domain'],
        'فعالیت اصلی': m['activity'],
        'زیرفعالیت عملیاتی': m['subactivity'],
        'عنوان مرکز هزینه': m['name'],
        'تعداد ردیف‌های تخصیص‌یافته': cnt,
        'تعریف دقیق عملیاتی': m['def'],
        'معیارهای شمول (Inclusion)': m['inc'],
        'معیارهای عدم شمول (Exclusion)': m['exc'],
        'نمونه شرح ردیف‌ها': ' | '.join(sample_lines),
        'محرک هزینه (Cost Driver)': m['driver'],
        'خروجی فعالیت': m['output'],
        'واحد سنجش': m['unit'],
        'معاونت متولی (سطح ۱)': m['owner_l1'],
        'اداره کل / مدیریت (سطح ۲)': m['owner_l2'],
        'اداره / واحد عملیاتی (سطح ۳)': m['owner_l3'],
        'درجه اطمینان استاندارد': 'بالا'
    })
df_dict = pd.DataFrame(dict_rows)

# 5. Activity Taxonomy Dataframe
tax_rows = []
for cc_id, m in cc_catalog.items():
    tax_rows.append({
        'کد حوزه': cc_id.split('-')[0] + '-' + cc_id.split('-')[1][:1],
        'سطح A - حوزه خدمت شهری': m['domain'],
        'سطح B - فعالیت اصلی': m['activity'],
        'سطح C - زیرفعالیت عملیاتی': m['subactivity'],
        'مرکز هزینه متناظر': m['name'],
        'کد مرکز هزینه': m['id'],
        'معاونت مسئول': m['owner_l1']
    })
df_taxonomy = pd.DataFrame(tax_rows).drop_duplicates()

# 6. Organizational Responsibility Matrix
org_matrix = []
for cc_id, m in cc_catalog.items():
    cnt = (df_mapped['کد مرکز هزینه پیشنهادی'] == cc_id).sum()
    org_matrix.append({
        'حوزه خدمت': m['domain'],
        'فعالیت اصلی': m['activity'],
        'کد مرکز هزینه': m['id'],
        'مرکز هزینه': m['name'],
        'معاونت متولی (Owner L1)': m['owner_l1'],
        'اداره کل / مدیریت متولی (Owner L2)': m['owner_l2'],
        'واحد / اداره عملیاتی (Owner L3)': m['owner_l3'],
        'تعداد خطوط مالی تحت مسئولیت': cnt
    })
df_org = pd.DataFrame(org_matrix)

# 7. Exceptions and Human Review Queue
df_exceptions = df_mapped[
    (df_mapped['درجه اطمینان'] != 'بالا') | 
    (df_mapped['وضعیت نگاشت'].str.contains('ابهام|بررسی انسانی', na=False)) |
    (df_mapped['نوع هزینه'] == 'نیازمند تصمیم کارشناسی') |
    (df_mapped['موضوع موجود'].str.contains('نامشخص', na=False))
].copy()
print(f"Exceptions and Human Review Queue: {len(df_exceptions)} rows.")

# 8. Create Multi-Tab Excel Workbook with Openpyxl
wb = openpyxl.Workbook()
wb.remove(wb.active) # Remove default sheet

# Helper styling
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="B Nazanin", size=11, bold=True, color="FFFFFF")
cell_font = Font(name="B Nazanin", size=10)
bold_font = Font(name="B Nazanin", size=10, bold=True)
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

def add_dataframe_to_sheet(wb, title, df, is_rtl=True):
    ws = wb.create_sheet(title=title)
    ws.views.sheetView[0].rightToLeft = is_rtl
    
    # Headers
    headers = list(df.columns)
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(1, col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    # Data rows
    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        ws.append(list(row))
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(r_idx, col_idx)
            c.font = cell_font
            c.border = border_thin
            if isinstance(c.value, (int, float)):
                c.alignment = Alignment(horizontal='center', vertical='center')
            else:
                c.alignment = Alignment(horizontal='right', vertical='center')
                
    # Adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col[:50])
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)
    
    ws.row_dimensions[1].height = 28
    return ws

# Add all 7 sheets
print("Writing sheets to Excel workbook...")
# Sheet 1: Diagnostic Report Summary (خلاصه آسیب‌شناسی)
diag_summary = [
    {'شاخص': 'تعداد کل خطوط مالی در بودجه مصوب (df_b)', 'مقدار': len(df_b), 'توضیحات': 'فایل مبنای بودجه شهرداری اصفهان'},
    {'شاخص': 'تعداد کل خطوط در فایل مراکز هزینه (df_cc)', 'مقدار': len(df_cc), 'توضیحات': 'شامل ۲۴۴۸ ردیف بودجه + ۳۸۲ ردیف نیازمند تصمیم کارشناسی'},
    {'شاخص': 'تعداد ردیف‌های نیازمند تصمیم کارشناسی', 'مقدار': 382, 'توضیحات': 'ردیف‌های بلاتکلیف و غربال‌شده از بودجه اصلی'},
    {'شاخص': 'تعداد کدهای بودجه منحصر‌به‌فرد', 'مقدار': df_cc['کد بودجه'].nunique(), 'توضیحات': 'تکرار کدها ناشی از تفکیک مناطق و پروژه‌هاست'},
    {'شاخص': 'تعداد مراکز هزینه سطح ۱ در مدل قبلی', 'مقدار': 17, 'توضیحات': 'دسته‌بندی‌های کلی و بعضاً همپوشان'},
    {'شاخص': 'تعداد مراکز هزینه سطح ۲ در مدل قبلی', 'مقدار': df_b['مرکز هزینه سطح 2'].nunique(), 'توضیحات': 'شامل سبدهای انباشت ناهمگن مثل تدارکات و پشتیبانی عمومی'},
    {'شاخص': 'تعداد فعالیت‌های تعریف‌شده در activity.xlsx', 'مقدار': len(df_act), 'توضیحات': 'نگاشت مکانیکی ۱ به ۱ موضوع و زیرموضوع بودجه بدون توجه به شرح ردیف'},
    {'شاخص': 'تعداد حوزه‌های خدمت پیشنهادی (سطح A)', 'مقدار': 16, 'توضیحات': 'پوشش کامل خدمات صف و پشتیبانی شهرداری اصفهان'},
    {'شاخص': 'تعداد مراکز هزینه نرمال‌شده پیشنهادی (Master CC)', 'مقدار': len(cc_catalog), 'توضیحات': 'مراکز هزینه یکنواخت با رفتار هزینه و متولی معین'},
    {'شاخص': 'ردیف‌های نیازمند اصلاح مرکز هزینه', 'مقدار': (df_mapped['وضعیت نگاشت'].str.contains('اصلاح', na=False)).sum(), 'توضیحات': 'ردیف‌هایی که در مدل قبلی در مرکز هزینه نامتجانس قرار داشتند'},
    {'شاخص': 'ردیف‌های با انتساب صحیح و مورد تأیید', 'مقدار': (df_mapped['وضعیت نگاشت'] == 'تأیید نگاشت موجود').sum(), 'توضیحات': 'تطابق کامل معنایی با اصول حسابداری مدیریت'}
]
df_diag_sum = pd.DataFrame(diag_summary)
add_dataframe_to_sheet(wb, 'خلاصه_آسیب_شناسی', df_diag_sum)

# Sheet 2: Cleanliness Pilot
add_dataframe_to_sheet(wb, 'پایلوت_پاکیزگی_و_پسماند', df_pilot)

# Sheet 3: Activity Taxonomy
add_dataframe_to_sheet(wb, 'درخت_فعالیت_ها', df_taxonomy)

# Sheet 4: Cost Center Master Dictionary
add_dataframe_to_sheet(wb, 'شناسنامه_مراکز_هزینه', df_dict)

# Sheet 5: Financial Line Mapping (All 2830 Rows)
add_dataframe_to_sheet(wb, 'نگاشت_ردیف_های_مالی', df_mapped)

# Sheet 6: Organizational Responsibility Matrix
add_dataframe_to_sheet(wb, 'ماتریس_مسئولیت_سازمانی', df_org)

# Sheet 7: Exceptions Queue
add_dataframe_to_sheet(wb, 'صف_بررسی_انسانی', df_exceptions)

# Save workbook
wb.save(output_excel)
print(f"=== FULL WORKBOOK SAVED SUCCESSFULLY AT: {output_excel} ===")
