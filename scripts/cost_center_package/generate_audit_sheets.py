import sys
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import numpy as np

print("=== BUILDING AUDIT WORKBOOK WITH 11 SHEETS ===")

output_path = r'C:\Users\sadra\Downloads\ممیزی_نسخه_اول_مدل_مراکز_هزینه.xlsx'

# 1. Load data
df_mapped = pd.read_excel(r'C:\Users\sadra\Downloads\تحلیل_جامع_مراکز_هزینه_و_فعالیت_های_شهرداری_اصفهان.xlsx', sheet_name='نگاشت_ردیف_های_مالی')
df_dict = pd.read_excel(r'C:\Users\sadra\Downloads\تحلیل_جامع_مراکز_هزینه_و_فعالیت_های_شهرداری_اصفهان.xlsx', sheet_name='شناسنامه_مراکز_هزینه')
wb_org = openpyxl.load_workbook(r'C:\Users\sadra\Downloads\آمار پست های معاونتها-1405.xlsx', data_only=True)

# Build org lookup
org_cells = []
for sname in wb_org.sheetnames:
    ws = wb_org[sname]
    for r in range(1, ws.max_row+1):
        for c in range(1, ws.max_column+1):
            val = ws.cell(r, c).value
            if val is not None:
                v_str = str(val).strip()
                v_norm = v_str.replace('\u200c', ' ').replace('ي', 'ی').replace('ك', 'ک').strip()
                if v_norm:
                    org_cells.append((sname, r, c, v_str, v_norm))

print(f"Loaded {len(org_cells)} organizational cells from source workbook.")

wb_out = openpyxl.Workbook()
wb_out.remove(wb_out.active)

# Styling helpers
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_fill_red = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
header_fill_amber = PatternFill(start_color="ED7D31", end_color="ED7D31", fill_type="solid")
header_font = Font(name="B Nazanin", size=11, bold=True, color="FFFFFF")
cell_font = Font(name="B Nazanin", size=10)
bold_font = Font(name="B Nazanin", size=10, bold=True)
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

def add_sheet(wb, title, df, fill=header_fill):
    ws = wb.create_sheet(title=title)
    ws.views.sheetView[0].rightToLeft = True
    headers = list(df.columns)
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(1, col_idx)
        c.fill = fill
        c.font = header_font
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        ws.append(list(row))
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(r_idx, col_idx)
            c.font = cell_font
            c.border = border_thin
            if isinstance(c.value, (int, float)):
                c.alignment = Alignment(horizontal='center', vertical='center')
            else:
                c.alignment = Alignment(horizontal='right', vertical='center')
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col[:50])
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
    ws.row_dimensions[1].height = 28
    return ws

print("Helper functions ready.")
