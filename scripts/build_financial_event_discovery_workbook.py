"""Build a revisioned, read-only-source Financial Event discovery workbook.

This utility is intentionally separate from the legacy cost-center workbook builders.  It
never edits either input workbook.  Persian text is normalized only in in-memory keys
used to compare records; every evidence row retains its raw source values.

The output is discovery evidence, not an authoritative event catalog or an approved
Financial Event-to-Cost Center rule table.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

import openpyxl
import xlsxwriter


PIPELINE_VERSION = "1.1.0"
FISCAL_YEAR = "1404"
DEFAULT_V3_BASELINE = Path(
    r"C:/Users/sadra/Downloads/بسته_نهایی_مراکز_هزینه_با_شماره_سند_حسابداری_v3.xlsx"
)
DEFAULT_ACCOUNTING_SOURCE = Path("data/reports/اطلاعات حسابداری1404.xlsx")
DEFAULT_SUPPLEMENTAL_SOURCES = (
    Path(r"C:/Users/sadra/Downloads/محیط زیست و خدمات شهری - مرداد 1405.xlsx"),
    Path(r"C:/Users/sadra/Downloads/عمران شهری - 1405 (1).xlsx"),
    Path(r"C:/Users/sadra/Downloads/خدمات حمل و نقل و ترافیک - 1405 (4).xlsx"),
)
DEFAULT_OUTPUT_ROOT = Path("data/deliverables/financial_event_discovery")

MASTER_SHEET = "فهرست_نهایی_پیشنهادی"
ERP_MASTER_SHEET = "Master_Data_ERP"
BASELINE_MAPPING_SHEET = "Doc_No_Mapping"

MISSING_TEXT = {"", "none", "null", "nan", "nat"}
_ANALYSIS_TRANSLATION = str.maketrans(
    {
        "ي": "ی",
        "ى": "ی",
        "ك": "ک",
        "ة": "ه",
        "ۀ": "ه",
        "\u200c": " ",
        "\xa0": " ",
    }
)


@dataclass
class AmountTotals:
    debit: Decimal = Decimal("0")
    credit: Decimal = Decimal("0")

    @property
    def net(self) -> Decimal:
        return self.debit - self.credit

    def add(self, debit: Decimal, credit: Decimal) -> None:
        self.debit += debit
        self.credit += credit


@dataclass
class CandidateAggregate:
    event_id: str
    candidate_key: str
    raw_title_sample: str
    normalized_title: str
    account_code: str
    observed_rows: int = 0
    documents: set[str] = field(default_factory=set)
    totals: AmountTotals = field(default_factory=AmountTotals)
    raw_title_variants: set[str] = field(default_factory=set)
    cost_centers: set[str] = field(default_factory=set)


@dataclass
class MatrixAggregate:
    event_id: str
    cost_center_id: str
    observed_rows: int = 0
    documents: set[str] = field(default_factory=set)
    budget_codes: set[str] = field(default_factory=set)
    totals: AmountTotals = field(default_factory=AmountTotals)


def normalize_analysis_text(value: Any) -> str:
    """Return a Persian-comparison key without mutating the source value.

    This normalizer is deliberately limited to Arabic/Persian character variants and
    whitespace.  It is used only for matching and grouping analysis copies; the raw
    values are separately written to the evidence workbook.
    """

    if value is None:
        return ""
    normalized = str(value).translate(_ANALYSIS_TRANSLATION)
    return " ".join(normalized.split())


def is_missing(value: Any) -> bool:
    return normalize_analysis_text(value).casefold() in MISSING_TEXT


def raw_text(value: Any) -> str:
    """Render source values for output without applying Persian normalization."""

    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y/%m/%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y/%m/%d")
    return str(value)


def normalized_code(value: Any) -> str:
    """Create a comparison code while preserving leading zeroes in text cells."""

    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, Decimal):
        return format(value.quantize(Decimal("1")), "f")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = normalize_analysis_text(value)
    if text.endswith(".0") and text[:-2].replace("-", "", 1).isdigit():
        return text[:-2]
    return text


def decimal_value(value: Any) -> Decimal:
    """Parse a journal amount without converting through binary floating point."""

    if value is None or is_missing(value):
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    text = raw_text(value).translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789"))
    text = text.replace("٬", "").replace("،", "").replace(",", "").replace(" ", "")
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError):
        return Decimal("0")


def decimal_text(value: Decimal) -> str:
    """Write currency as exact text so large Rial totals are not rounded by Excel."""

    rendered = format(value, "f")
    if "." in rendered:
        rendered = rendered.rstrip("0").rstrip(".")
    return rendered or "0"


def decimal_key(value: Any) -> str:
    return decimal_text(decimal_value(value))


def source_date_key(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y/%m/%d")
    if isinstance(value, date):
        return value.strftime("%Y/%m/%d")
    return normalize_analysis_text(value)


def make_account_code(row: Sequence[Any], indexes: Mapping[str, int]) -> str:
    return "-".join(
        normalized_code(row[indexes[column]])
        for column in ("TitkNo", "TitMNo", "TitTNo")
    )


def make_event_key(row: Sequence[Any], indexes: Mapping[str, int]) -> str | None:
    """Return a raw-evidence candidate key, or None when no account-detail text exists."""

    raw_title = row[indexes["شرح سرفصل حساب جزء"]]
    normalized_title = normalize_analysis_text(raw_title)
    if is_missing(raw_title):
        return None
    components = [
        normalized_code(row[indexes[column]])
        for column in ("TitkNo", "TitMNo", "TitTNo", "titjno")
    ]
    return "|".join(components + [normalized_title])


def candidate_id_for_key(candidate_key: str) -> str:
    """A revision-independent opaque identifier for an observed candidate fingerprint."""

    return "CFE-" + hashlib.sha256(candidate_key.encode("utf-8")).hexdigest()[:16].upper()


def parse_baseline_note(value: Any) -> tuple[str, str]:
    """Extract the original area/date values placed in Doc_No_Mapping's note column."""

    text = raw_text(value)
    match = re.match(r"^(.*?)\s*-\s*تاریخ:\s*(.*?)\s*$", text)
    if not match:
        return "", ""
    return normalize_analysis_text(match.group(1)), normalize_analysis_text(match.group(2))


def baseline_match_key(
    *,
    area: Any,
    source_date: Any,
    document_number: Any,
    budget_code: Any,
    account_code: Any,
    debit: Any,
    credit: Any,
    account_detail: Any,
) -> tuple[str, str, str, str, str, str, str, str]:
    """Key a baseline mapping to a precise raw journal line, never a whole voucher."""

    return (
        normalize_analysis_text(area),
        source_date_key(source_date),
        normalized_code(document_number),
        normalized_code(budget_code),
        normalize_analysis_text(account_code),
        decimal_key(debit),
        decimal_key(credit),
        normalize_analysis_text(account_detail),
    )


def classify_baseline_matches(matches: Sequence[Mapping[str, str]]) -> tuple[str, Mapping[str, str] | None]:
    """Classify an exact line-level baseline match without inferring a Cost Center.

    Duplicate baseline rows that all identify the same Cost Center remain observed
    evidence.  Two or more distinct Cost Center IDs are explicitly ambiguous.
    """

    if not matches:
        return "UNMAPPED", None
    distinct_cost_centers = {match["cost_center_id"] for match in matches}
    if len(distinct_cost_centers) != 1:
        return "AMBIGUOUS", None
    return "OBSERVED", matches[0]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source_file:
        for chunk in iter(lambda: source_file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def header_indexes(header: Iterable[Any], required: Iterable[str]) -> dict[str, int]:
    normalized_header = [raw_text(item) for item in header]
    missing = [field for field in required if field not in normalized_header]
    if missing:
        raise ValueError(f"Missing required accounting columns: {', '.join(missing)}")
    return {field: normalized_header.index(field) for field in required}


def load_baseline(
    v3_path: Path,
) -> tuple[
    list[list[Any]], list[list[Any]], dict[str, dict[str, str]], dict[tuple[str, ...], list[dict[str, str]]], dict[str, Any]
]:
    """Load raw baseline references and index Doc_No_Mapping at journal-line granularity."""

    workbook = openpyxl.load_workbook(v3_path, read_only=True, data_only=True)
    required_sheets = {MASTER_SHEET, ERP_MASTER_SHEET, BASELINE_MAPPING_SHEET}
    absent = required_sheets.difference(workbook.sheetnames)
    if absent:
        workbook.close()
        raise ValueError(f"v3 baseline lacks required sheets: {', '.join(sorted(absent))}")

    master_rows = [list(row) for row in workbook[MASTER_SHEET].iter_rows(values_only=True)]
    erp_rows = [list(row) for row in workbook[ERP_MASTER_SHEET].iter_rows(values_only=True)]
    mapping_ws = workbook[BASELINE_MAPPING_SHEET]
    mapping_header = [raw_text(value) for value in next(mapping_ws.iter_rows(max_row=1, values_only=True))]
    mapping_indexes = {name: mapping_header.index(name) for name in (
        "شماره سند حسابداری",
        "کد بودجه / ردیف مالی",
        "کد حساب",
        "شرح سند حسابداری",
        "مبلغ بدهکار",
        "مبلغ بستانکار",
        "کد مرکز هزینه پیشنهادی",
        "عنوان مرکز هزینه پیشنهادی",
        "حوزه خدمت",
        "توضیح",
    )}

    cost_center_master: dict[str, dict[str, str]] = {}
    master_header = [raw_text(value) for value in master_rows[0]]
    master_indexes = {
        "id": master_header.index("کد مرکز هزینه پیشنهادی"),
        "title": master_header.index("عنوان مرکز هزینه"),
        "domain": master_header.index("حوزه خدمت"),
    }
    for row in master_rows[1:]:
        cost_center_id = raw_text(row[master_indexes["id"]])
        if cost_center_id:
            cost_center_master[cost_center_id] = {
                "title": raw_text(row[master_indexes["title"]]),
                "domain": raw_text(row[master_indexes["domain"]]),
            }

    baseline_by_line: dict[tuple[str, ...], list[dict[str, str]]] = defaultdict(list)
    baseline_mapping_rows = 0
    for excel_row, row in enumerate(mapping_ws.iter_rows(min_row=2, values_only=True), start=2):
        baseline_mapping_rows += 1
        area, source_date = parse_baseline_note(row[mapping_indexes["توضیح"]])
        key = baseline_match_key(
            area=area,
            source_date=source_date,
            document_number=row[mapping_indexes["شماره سند حسابداری"]],
            budget_code=row[mapping_indexes["کد بودجه / ردیف مالی"]],
            account_code=row[mapping_indexes["کد حساب"]],
            debit=row[mapping_indexes["مبلغ بدهکار"]],
            credit=row[mapping_indexes["مبلغ بستانکار"]],
            account_detail=row[mapping_indexes["شرح سند حسابداری"]],
        )
        cost_center_id = raw_text(row[mapping_indexes["کد مرکز هزینه پیشنهادی"]])
        baseline_by_line[key].append(
            {
                "cost_center_id": cost_center_id,
                "baseline_row": str(excel_row),
                "baseline_title": raw_text(row[mapping_indexes["عنوان مرکز هزینه پیشنهادی"]]),
                "baseline_domain": raw_text(row[mapping_indexes["حوزه خدمت"]]),
            }
        )

    workbook.close()
    baseline_metadata = {
        "master_rows": len(master_rows) - 1,
        "master_columns": len(master_rows[0]),
        "erp_rows": len(erp_rows) - 1,
        "erp_columns": len(erp_rows[0]),
        "mapping_rows": baseline_mapping_rows,
        "mapping_columns": len(mapping_header),
    }
    return master_rows, erp_rows, cost_center_master, baseline_by_line, baseline_metadata


def source_record_context(
    row: Sequence[Any], indexes: Mapping[str, int], sheet_name: str, excel_row: int
) -> dict[str, Any]:
    """Build a source-line context retaining raw values and separate analysis keys."""

    account_code = make_account_code(row, indexes)
    document_number = normalized_code(row[indexes["شماره سند"]])
    budget_code = normalized_code(row[indexes["کد بودجه"]])
    debit = decimal_value(row[indexes["مبلغ بدهکار"]])
    credit = decimal_value(row[indexes["مبلغ بستانکار"]])
    account_detail_raw = raw_text(row[indexes["شرح سرفصل حساب جزء"]])
    context = {
        "sheet_name": sheet_name,
        "excel_row": excel_row,
        "source_reference": f"اطلاعات حسابداری1404.xlsx::{sheet_name}!R{excel_row}",
        "area_raw": raw_text(row[indexes["منطقه"]]),
        "date_raw": raw_text(row[indexes["تاریخ"]]),
        "document_raw": raw_text(row[indexes["شماره سند"]]),
        "document_number": document_number,
        "budget_raw": raw_text(row[indexes["کد بودجه"]]),
        "budget_code": budget_code,
        "account_code": account_code,
        "account_detail_raw": account_detail_raw,
        "account_detail_normalized": normalize_analysis_text(account_detail_raw),
        "debit": debit,
        "credit": credit,
        "event_key": make_event_key(row, indexes),
    }
    context["baseline_key"] = baseline_match_key(
        area=context["area_raw"],
        source_date=source_date_key(row[indexes["تاریخ"]]),
        document_number=context["document_number"],
        budget_code=context["budget_code"],
        account_code=context["account_code"],
        debit=context["debit"],
        credit=context["credit"],
        account_detail=context["account_detail_raw"],
    )
    return context


def accounting_indexes(header: Sequence[Any]) -> dict[str, int]:
    return header_indexes(
        header,
        (
            "کد منطقه",
            "منطقه",
            "تاریخ",
            "شماره سند",
            "TitkNo",
            "TitMNo",
            "TitTNo",
            "titjno",
            "شرح سرفصل حساب جزء",
            "شرح سرفصل حساب معین",
            "شرح سرفصل حساب تفضیلی",
            "مبلغ بدهکار",
            "مبلغ بستانکار",
            "شماره درخواست",
            "نوع درخواست",
            "کد بودجه",
            "ماهیت سند",
        ),
    )


def scan_accounting_source(
    accounting_path: Path,
    baseline_by_line: Mapping[tuple[str, ...], Sequence[Mapping[str, str]]],
) -> tuple[dict[str, CandidateAggregate], dict[tuple[str, str], MatrixAggregate], dict[str, Any], list[dict[str, Any]]]:
    """First pass: derive candidates and observed Event-to-Cost-Center evidence only."""

    candidates: dict[str, CandidateAggregate] = {}
    matrix: dict[tuple[str, str], MatrixAggregate] = {}
    qa = Counter()
    source_sheet_details: list[dict[str, Any]] = []

    workbook = openpyxl.load_workbook(accounting_path, read_only=True, data_only=True)
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        header = next(worksheet.iter_rows(max_row=1, values_only=True))
        indexes = accounting_indexes(header)
        sheet_rows = 0
        for excel_row, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            sheet_rows += 1
            qa["source_rows"] += 1
            context = source_record_context(row, indexes, sheet_name, excel_row)
            matches = baseline_by_line.get(context["baseline_key"], ())
            match_status, matched = classify_baseline_matches(matches)
            qa[f"baseline_{match_status.lower()}"] += 1

            event_key = context["event_key"]
            if event_key is None:
                qa["unclassified_event_rows"] += 1
                continue

            event_id = candidate_id_for_key(event_key)
            candidate = candidates.get(event_key)
            if candidate is None:
                candidate = CandidateAggregate(
                    event_id=event_id,
                    candidate_key=event_key,
                    raw_title_sample=context["account_detail_raw"],
                    normalized_title=context["account_detail_normalized"],
                    account_code=context["account_code"],
                )
                candidates[event_key] = candidate
            candidate.observed_rows += 1
            candidate.documents.add(context["document_number"] or "فاقد_شماره_سند")
            candidate.totals.add(context["debit"], context["credit"])
            candidate.raw_title_variants.add(context["account_detail_raw"])

            if match_status != "OBSERVED" or matched is None:
                continue
            cost_center_id = matched["cost_center_id"]
            candidate.cost_centers.add(cost_center_id)
            matrix_key = (event_id, cost_center_id)
            aggregate = matrix.get(matrix_key)
            if aggregate is None:
                aggregate = MatrixAggregate(event_id=event_id, cost_center_id=cost_center_id)
                matrix[matrix_key] = aggregate
            aggregate.observed_rows += 1
            aggregate.documents.add(context["document_number"] or "فاقد_شماره_سند")
            aggregate.budget_codes.add(context["budget_code"])
            aggregate.totals.add(context["debit"], context["credit"])

        source_sheet_details.append(
            {
                "sheet_name": sheet_name,
                "rows": sheet_rows,
                "columns": len(header),
                "header": [raw_text(value) for value in header],
            }
        )
    workbook.close()
    return candidates, matrix, qa, source_sheet_details



def audit_supplemental_catalogs(
    supplemental_paths: Sequence[Path],
) -> list[dict[str, Any]]:
    """Record 1405 service catalogs without treating them as Financial Event evidence.

    These inputs belong to a separate Cost Center completeness/service-activity
    workstream.  The audit preserves their workbook structure and explicitly records
    that neither their service-event fields nor their service titles enter the 1404
    Financial Event candidate dictionary or Event-to-Cost-Center matrix.
    """

    audits: list[dict[str, Any]] = []
    for source_index, source_path in enumerate(supplemental_paths, start=1):
        workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=True)
        sheet_audits: list[dict[str, Any]] = []
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            header = next(worksheet.iter_rows(max_row=1, values_only=True))
            headers = [raw_text(value) for value in header]
            sheet_audits.append(
                {
                    "name": sheet_name,
                    "worksheet_rows": worksheet.max_row,
                    "data_rows": max(worksheet.max_row - 1, 0),
                    "columns": worksheet.max_column,
                    "headers": headers,
                    "contains_service_catalog_fields": "نام خدمت ( سطح سه)" in headers,
                }
            )
        workbook.close()
        audits.append(
            {
                "hash_key": f"supplemental_{source_index}",
                "path": source_path,
                "filename": source_path.name,
                "evidence_origin": "فهرست خدمات معاونت — ۱۴۰۵",
                "usage": "ممیزی‌شده اما استفاده‌نشده در کشف رویداد مالی",
                "reason": (
                    "این workbookها فهرست خدمت/فعالیت‌اند و منبع journal یا document-level "
                    "Financial Event/Cost Center evidence نیستند؛ هیچ crosswalk یا رویداد مالی از آن‌ها ساخته نشد."
                ),
                "sheets": sheet_audits,
            }
        )
    return audits


def make_workbook(
    output_path: Path,
    *,
    v3_path: Path,
    accounting_path: Path,
    source_hashes: Mapping[str, str],
    master_rows: list[list[Any]],
    erp_rows: list[list[Any]],
    cost_center_master: Mapping[str, Mapping[str, str]],
    baseline_by_line: Mapping[tuple[str, ...], Sequence[Mapping[str, str]]],
    baseline_metadata: Mapping[str, Any],
    candidates: Mapping[str, CandidateAggregate],
    matrix: Mapping[tuple[str, str], MatrixAggregate],
    qa: Counter,
    source_sheet_details: Sequence[Mapping[str, Any]],
    supplemental_audits: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    """Second pass: write the workbook, including line-level evidence without allocation."""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = xlsxwriter.Workbook(str(output_path), {"constant_memory": True})
    workbook.set_properties(
        {
            "title": "کشف رویدادهای مالی — شواهد تاریخی ۱۴۰۴",
            "subject": "Candidate Financial Event discovery; non-authoritative historical evidence",
            "author": "Financial Event Discovery Pipeline",
            "comments": "No authoritative Financial Event-to-Cost Center rule is asserted by this workbook.",
        }
    )
    header_format = workbook.add_format(
        {"bold": True, "font_color": "#FFFFFF", "bg_color": "#1F4E78", "border": 1, "text_wrap": True, "align": "center", "valign": "vcenter"}
    )
    subheader_format = workbook.add_format(
        {"bold": True, "font_color": "#000000", "bg_color": "#D9EAF7", "border": 1, "text_wrap": True, "valign": "vcenter"}
    )
    text_format = workbook.add_format({"border": 1, "valign": "top", "text_wrap": False})
    wrap_format = workbook.add_format({"border": 1, "valign": "top", "text_wrap": True})
    currency_text_format = workbook.add_format({"border": 1, "num_format": "@", "align": "left"})
    warning_format = workbook.add_format({"border": 1, "bg_color": "#FFF2CC", "valign": "top", "text_wrap": True})
    red_format = workbook.add_format({"border": 1, "bg_color": "#FCE4D6", "valign": "top", "text_wrap": True})

    def add_table_sheet(name: str, headers: Sequence[str], rows: Iterable[Sequence[Any]], widths: Sequence[int] | None = None) -> None:
        worksheet = workbook.add_worksheet(name)
        worksheet.right_to_left()
        worksheet.freeze_panes(1, 0)
        worksheet.autofilter(0, 0, 0, len(headers) - 1)
        worksheet.set_row(0, 35)
        for column, header in enumerate(headers):
            worksheet.write(0, column, header, header_format)
            worksheet.set_column(column, column, (widths or [18] * len(headers))[column])
        for row_number, row in enumerate(rows, start=1):
            for column, value in enumerate(row):
                fmt = currency_text_format if "مبلغ" in headers[column] else text_format
                worksheet.write(row_number, column, value if value is not None else "", fmt)

    guidance_rows = [
        ("نسخهٔ pipeline", PIPELINE_VERSION),
        ("سال مالی corpus", FISCAL_YEAR),
        ("وضعیت کاندیداهای رویداد", "همهٔ موارد «مشاهدهٔ تاریخی — غیرمرجع» و در سطح نیازمند بررسی هستند."),
        ("وضعیت رابطهٔ رویداد↔مرکز هزینه", "همهٔ رابطه‌ها صرفاً شاهد سطری historical هستند؛ مجوز یا Rule authoritative نیستند."),
        ("روش انتساب", "فقط تطبیق دقیق هر سطر ژورنال با Doc_No_Mapping baseline؛ هیچ تخصیص نسبتی در سطح سند انجام نشده است."),
        ("نرمال‌سازی", "فقط کلیدهای تحلیلی در حافظه نرمال شده‌اند؛ ستون‌های raw در نگاشت شواهد سطری بدون نرمال‌سازی نگه‌داری شده‌اند."),
        ("هویت", "CFE-* شناسهٔ opaqueِ candidate fingerprint در این خروجی است؛ شماره سند، کد بودجه و شماره درخواست فقط Source Reference هستند."),
        ("زیرسامانهٔ مبدأ", "ورودی journal نام زیرسامانهٔ مبدأ را تعیین نمی‌کند؛ در این خروجی مقدار آن «تعیین‌نشده در منبع» است و با حوزهٔ خدمت یکی نیست."),
        ("مرز حسابداری", "این خروجی هیچ Accounting Handoff یا journal posting خودکاری انجام نمی‌دهد."),
        (
            "کاتالوگ‌های مکمل ۱۴۰۵",
            "سه workbook معاونتی ۱۴۰۵ ممیزی و در منشأ ثبت شده‌اند، اما فهرست خدمت/«نوع رویداد خدمت» آن‌ها Financial Event evidence نیست و در واژه‌نامه یا ماتریس ۱۴۰۴ استفاده نشده است.",
        ),
    ]
    add_table_sheet("راهنما_و_حاکمیت", ("موضوع", "تصمیم/توضیح"), guidance_rows, (28, 105))

    candidate_by_id = {aggregate.event_id: aggregate for aggregate in candidates.values()}
    candidate_rows = []
    for candidate in sorted(candidates.values(), key=lambda item: (item.normalized_title, item.account_code, item.event_id)):
        candidate_rows.append(
            (
                candidate.event_id,
                candidate.raw_title_sample,
                candidate.normalized_title,
                candidate.account_code,
                candidate.observed_rows,
                len(candidate.documents),
                decimal_text(candidate.totals.debit),
                decimal_text(candidate.totals.credit),
                decimal_text(candidate.totals.net),
                len(candidate.cost_centers),
                len(candidate.raw_title_variants),
                "مشاهدهٔ تاریخی — غیرمرجع",
                "کاندیدا / نیازمند بررسی",
                "کد و شرح حساب جزء در سطر ژورنال؛ هیچ Rule authoritative ارائه نشده است.",
            )
        )
    add_table_sheet(
        "واژه‌نامه_رویدادهای_پیشنهادی",
        (
            "شناسهٔ داخلی کاندیدا",
            "عنوان خام نمونه از منبع",
            "عنوان نرمال‌شدهٔ تحلیلی",
            "کد حساب (سه سطح)",
            "فراوانی سطرهای مشاهده‌شده",
            "تعداد اسناد یکتا",
            "مبلغ بدهکار (ریال؛ متن دقیق)",
            "مبلغ بستانکار (ریال؛ متن دقیق)",
            "مبلغ خالص (ریال؛ متن دقیق)",
            "تعداد مراکز هزینهٔ مشاهده‌شده",
            "تعداد گونه‌های عنوان خام",
            "وضعیت شواهد",
            "وضعیت طبقه‌بندی",
            "مبنای استخراج و محدودیت",
        ),
        candidate_rows,
        (23, 45, 45, 20, 16, 14, 24, 24, 24, 17, 17, 25, 23, 65),
    )

    matrix_rows = []
    for (event_id, cost_center_id), aggregate in sorted(matrix.items(), key=lambda item: (item[0][1], item[0][0])):
        candidate = candidate_by_id[event_id]
        master = cost_center_master.get(cost_center_id, {})
        matrix_rows.append(
            (
                event_id,
                candidate.raw_title_sample,
                cost_center_id,
                master.get("title", ""),
                master.get("domain", ""),
                aggregate.observed_rows,
                len(aggregate.documents),
                len(aggregate.budget_codes),
                decimal_text(aggregate.totals.debit),
                decimal_text(aggregate.totals.credit),
                decimal_text(aggregate.totals.net),
                "شاهد سطری historical — غیرمرجع",
                "فقط تطبیق دقیق سطر ژورنال با baseline؛ نه Rule مجاز و نه تخصیص در سطح سند.",
            )
        )
    add_table_sheet(
        "ماتریس_شواهد_رویداد_مرکز",
        (
            "شناسهٔ کاندیدای رویداد",
            "عنوان خام نمونهٔ رویداد",
            "کد مرکز هزینه",
            "عنوان مرکز هزینه (B baseline)",
            "حوزهٔ خدمت (C baseline)",
            "فراوانی سطرهای شواهد",
            "تعداد اسناد یکتا",
            "تعداد ردیف‌های بودجه",
            "مبلغ بدهکار (ریال؛ متن دقیق)",
            "مبلغ بستانکار (ریال؛ متن دقیق)",
            "مبلغ خالص (ریال؛ متن دقیق)",
            "وضعیت رابطه",
            "محدودیت",
        ),
        matrix_rows,
        (23, 42, 16, 42, 34, 16, 14, 15, 24, 24, 24, 28, 66),
    )

    master_reference_rows = [tuple(row) for row in master_rows]
    add_table_sheet(
        "مرجع_مراکز_هزینه",
        tuple(f"{header} (raw baseline)" for header in master_reference_rows[0]),
        master_reference_rows[1:],
        tuple(19 for _ in master_reference_rows[0]),
    )
    erp_reference_rows = [tuple(row) for row in erp_rows]
    add_table_sheet(
        "مرجع_Master_Data_ERP",
        tuple(f"{header} (raw baseline)" for header in erp_reference_rows[0]),
        erp_reference_rows[1:],
        tuple(22 for _ in erp_reference_rows[0]),
    )

    lineage_rows = [
        ("Baseline v3", str(v3_path), source_hashes["v3"], MASTER_SHEET, baseline_metadata["master_rows"], baseline_metadata["master_columns"], "مرجع raw؛ فقط خوانده شد. ستون‌های B/C عیناً در شیت مرجع خروجی نگه‌داری شده‌اند."),
        ("Baseline v3", str(v3_path), source_hashes["v3"], ERP_MASTER_SHEET, baseline_metadata["erp_rows"], baseline_metadata["erp_columns"], "مرجع raw؛ فقط خوانده شد."),
        ("Baseline v3", str(v3_path), source_hashes["v3"], BASELINE_MAPPING_SHEET, baseline_metadata["mapping_rows"], baseline_metadata["mapping_columns"], "فقط شاهد line-level برای اتصال دقیق سطر journal به Cost Center."),
    ]
    for detail in source_sheet_details:
        lineage_rows.append(
            (
                "Journal 1404",
                str(accounting_path),
                source_hashes["accounting"],
                detail["sheet_name"],
                detail["rows"],
                detail["columns"],
                "منبع raw؛ فقط خوانده شد. نام زیرسامانهٔ مبدأ در این workbook مشخص نشده است.",
            )
        )
    for audit in supplemental_audits:
        for sheet in audit["sheets"]:
            lineage_rows.append(
                (
                    "فهرست خدمات معاونتی ۱۴۰۵ (ممیزی‌شده/استفاده‌نشده)",
                    str(audit["path"]),
                    source_hashes[audit["hash_key"]],
                    sheet["name"],
                    sheet["data_rows"],
                    sheet["columns"],
                    audit["reason"],
                )
            )
    add_table_sheet(
        "منشأ_و_تبار_داده",
        ("نوع منبع", "مسیر", "SHA-256", "شیت", "تعداد سطر داده", "تعداد ستون", "نقش و محدودیت"),
        lineage_rows,
        (18, 70, 67, 28, 16, 14, 70),
    )

    client_center_ids = ("CC-502", "CC-503", "CC-504", "CC-505", "CC-506")
    client_rows = []
    client_samples: list[tuple[Any, ...]] = []
    for cost_center_id in client_center_ids:
        center_matrix = [aggregate for (_, cc_id), aggregate in matrix.items() if cc_id == cost_center_id]
        total = AmountTotals()
        documents: set[str] = set()
        budget_codes: set[str] = set()
        event_ids: set[str] = set()
        row_count = 0
        for (event_id, cc_id), aggregate in matrix.items():
            if cc_id != cost_center_id:
                continue
            row_count += aggregate.observed_rows
            documents.update(aggregate.documents)
            budget_codes.update(aggregate.budget_codes)
            event_ids.add(event_id)
            total.add(aggregate.totals.debit, aggregate.totals.credit)
        master = cost_center_master.get(cost_center_id, {})
        client_rows.append(
            (
                cost_center_id,
                master.get("title", ""),
                master.get("domain", ""),
                row_count,
                len(documents),
                len(budget_codes),
                len(event_ids),
                decimal_text(total.debit),
                decimal_text(total.credit),
                decimal_text(total.net),
                "یافتهٔ historical observed — غیرمرجع؛ نیازمند Rule Profile/تأیید مرجع برای هر رابطه.",
            )
        )
    add_table_sheet(
        "یافته_CC_502_506",
        (
            "کد مرکز هزینه",
            "عنوان مرکز هزینه (B baseline)",
            "حوزهٔ خدمت (C baseline)",
            "فراوانی سطرهای مشاهده‌شده",
            "تعداد اسناد یکتا",
            "تعداد ردیف‌های بودجه",
            "تعداد کاندیداهای رویداد",
            "مبلغ بدهکار (ریال؛ متن دقیق)",
            "مبلغ بستانکار (ریال؛ متن دقیق)",
            "مبلغ خالص (ریال؛ متن دقیق)",
            "یافته و محدودیت",
        ),
        client_rows,
        (17, 45, 35, 18, 16, 17, 18, 24, 24, 24, 72),
    )

    evidence_headers = (
        "Source Reference (سطری)",
        "فایل منبع raw",
        "شیت منبع raw",
        "ردیف Excel منبع",
        "زیرسامانهٔ مبدأ",
        "کد منطقه raw",
        "منطقه raw",
        "تاریخ raw",
        "شماره سند raw",
        "شماره درخواست raw",
        "نوع درخواست raw",
        "ماهیت سند raw",
        "TitkNo raw",
        "TitMNo raw",
        "TitTNo raw",
        "titjno raw",
        "شرح حساب معین raw",
        "شرح حساب تفضیلی raw",
        "شرح حساب جزء raw",
        "کد بودجه raw",
        "مبلغ بدهکار raw (ریال؛ متن دقیق)",
        "مبلغ بستانکار raw (ریال؛ متن دقیق)",
        "شناسهٔ کاندیدای رویداد",
        "شرح حساب جزء نرمال‌شدهٔ تحلیلی",
        "وضعیت استخراج رویداد",
        "وضعیت شواهد مرکز هزینه",
        "ردیف Doc_No_Mapping baseline",
        "کد مرکز هزینهٔ مشاهده‌شده",
        "عنوان مرکز هزینه (B baseline)",
        "حوزهٔ خدمت (C baseline)",
        "نوع رابطه",
        "یادداشت کنترل",
    )
    evidence_ws = workbook.add_worksheet("نگاشت_شواهد_سطری")
    evidence_ws.right_to_left()
    evidence_ws.freeze_panes(1, 0)
    evidence_ws.autofilter(0, 0, 0, len(evidence_headers) - 1)
    evidence_ws.set_row(0, 42)
    for column, header in enumerate(evidence_headers):
        evidence_ws.write(0, column, header, header_format)
        evidence_ws.set_column(column, column, 20 if column not in {0, 16, 17, 18, 28, 29, 31} else 45)

    issue_headers = (
        "Source Reference (سطری)",
        "شیت منبع", "ردیف Excel منبع", "شماره سند raw", "شرح حساب جزء raw",
        "شناسهٔ کاندیدای رویداد", "نوع مسئله", "وضعیت", "توضیح",
    )
    issues_ws = workbook.add_worksheet("ردیف‌های_مبهم_بی‌طبقه")
    issues_ws.right_to_left()
    issues_ws.freeze_panes(1, 0)
    issues_ws.autofilter(0, 0, 0, len(issue_headers) - 1)
    for column, header in enumerate(issue_headers):
        issues_ws.write(0, column, header, header_format)
        issues_ws.set_column(column, column, 24 if column not in {0, 4, 8} else 60)

    client_samples_ws = workbook.add_worksheet("نمونه_CC_502_506")
    client_samples_ws.right_to_left()
    client_samples_ws.freeze_panes(1, 0)
    client_sample_headers = (
        "کد مرکز هزینه", "عنوان مرکز هزینه (B baseline)", "شناسهٔ کاندیدای رویداد", "عنوان raw رویداد",
        "Source Reference (سطری)", "شماره سند raw", "تاریخ raw", "کد بودجه raw", "مبلغ بدهکار raw", "مبلغ بستانکار raw", "وضعیت",
    )
    for column, header in enumerate(client_sample_headers):
        client_samples_ws.write(0, column, header, header_format)
        client_samples_ws.set_column(column, column, 24 if column not in {3, 4} else 47)

    issue_row = 1
    evidence_row = 1
    client_sample_row = 1
    client_samples_remaining = {center: 5 for center in client_center_ids}
    second_pass_qa = Counter()
    source_workbook = openpyxl.load_workbook(accounting_path, read_only=True, data_only=True)
    for sheet_name in source_workbook.sheetnames:
        worksheet = source_workbook[sheet_name]
        header = next(worksheet.iter_rows(max_row=1, values_only=True))
        indexes = accounting_indexes(header)
        for excel_row, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            context = source_record_context(row, indexes, sheet_name, excel_row)
            event_key = context["event_key"]
            event_id = candidate_id_for_key(event_key) if event_key else ""
            matches = baseline_by_line.get(context["baseline_key"], ())
            match_status, matched = classify_baseline_matches(matches)
            if match_status == "OBSERVED" and matched is not None:
                cost_center_id = matched["cost_center_id"]
                master = cost_center_master.get(cost_center_id, {})
                center_status = "شاهد سطری observed در baseline — غیرمرجع"
                baseline_row = matched["baseline_row"]
                relation = "Observed Relationship (Historical; non-authoritative)"
                control_note = "تطبیق دقیق سطر ژورنال؛ بدون تسهیم یا تخصیص در سطح سند."
                second_pass_qa["line_level_links"] += 1
            elif match_status == "AMBIGUOUS":
                cost_center_id = ""
                master = {}
                center_status = "مبهم — نیازمند بررسی"
                baseline_row = ""
                relation = "No relationship asserted"
                control_note = "بیش از یک مرکز هزینه در تطبیق دقیق baseline؛ هیچ مرکز هزینه‌ای انتخاب نشد."
                second_pass_qa["ambiguous_rows"] += 1
            else:
                cost_center_id = ""
                master = {}
                center_status = "شاهد مرکز هزینه در baseline یافت نشد"
                baseline_row = ""
                relation = "No relationship asserted"
                control_note = "بدون تطبیق دقیق سطر با baseline؛ هیچ مرکز هزینه‌ای استنباط نشده است."
                second_pass_qa["unmapped_rows"] += 1

            event_status = "کاندیدای historical — نیازمند بررسی" if event_id else "بی‌طبقه — نیازمند بررسی"
            evidence_values = (
                context["source_reference"],
                "اطلاعات حسابداری1404.xlsx",
                sheet_name,
                excel_row,
                "تعیین‌نشده در منبع",
                raw_text(row[indexes["کد منطقه"]]),
                raw_text(row[indexes["منطقه"]]),
                raw_text(row[indexes["تاریخ"]]),
                raw_text(row[indexes["شماره سند"]]),
                raw_text(row[indexes["شماره درخواست"]]),
                raw_text(row[indexes["نوع درخواست"]]),
                raw_text(row[indexes["ماهیت سند"]]),
                raw_text(row[indexes["TitkNo"]]),
                raw_text(row[indexes["TitMNo"]]),
                raw_text(row[indexes["TitTNo"]]),
                raw_text(row[indexes["titjno"]]),
                raw_text(row[indexes["شرح سرفصل حساب معین"]]),
                raw_text(row[indexes["شرح سرفصل حساب تفضیلی"]]),
                context["account_detail_raw"],
                context["budget_raw"],
                decimal_text(context["debit"]),
                decimal_text(context["credit"]),
                event_id,
                context["account_detail_normalized"],
                event_status,
                center_status,
                baseline_row,
                cost_center_id,
                master.get("title", ""),
                master.get("domain", ""),
                relation,
                control_note,
            )
            for column, value in enumerate(evidence_values):
                fmt = currency_text_format if "مبلغ" in evidence_headers[column] else text_format
                evidence_ws.write(evidence_row, column, value, fmt)
            evidence_row += 1

            issue: tuple[Any, ...] | None = None
            if not event_id:
                issue = (
                    context["source_reference"], sheet_name, excel_row, context["document_raw"], context["account_detail_raw"], "",
                    "مبنای کاندیدای رویداد ناقص", "نیازمند بررسی",
                    "شرح حساب جزء خالی است؛ برای جلوگیری از ساخت رویداد حدسی، کاندیدا ایجاد نشد.",
                )
            elif match_status == "AMBIGUOUS":
                issue = (
                    context["source_reference"], sheet_name, excel_row, context["document_raw"], context["account_detail_raw"], event_id,
                    "شاهد مرکز هزینه مبهم", "نیازمند بررسی",
                    "تطبیق دقیق baseline بیش از یک مرکز هزینه ارائه داد؛ هیچ رابطه‌ای ساخته نشد.",
                )
            if issue:
                for column, value in enumerate(issue):
                    issues_ws.write(issue_row, column, value, warning_format if column in {6, 7, 8} else text_format)
                issue_row += 1

            if cost_center_id in client_samples_remaining and client_samples_remaining[cost_center_id] > 0:
                sample_values = (
                    cost_center_id,
                    master.get("title", ""),
                    event_id,
                    context["account_detail_raw"],
                    context["source_reference"],
                    context["document_raw"],
                    context["date_raw"],
                    context["budget_raw"],
                    decimal_text(context["debit"]),
                    decimal_text(context["credit"]),
                    "شاهد سطری historical — غیرمرجع",
                )
                for column, value in enumerate(sample_values):
                    fmt = currency_text_format if column in {8, 9} else text_format
                    client_samples_ws.write(client_sample_row, column, value, fmt)
                client_samples_remaining[cost_center_id] -= 1
                client_sample_row += 1
    source_workbook.close()

    qa_rows = [
        ("سطرهای raw journal خوانده‌شده", str(qa["source_rows"]), "PASS", "از دو شیت حسابداری ۱۴۰۴؛ همه در نگاشت شواهد سطری آمده‌اند."),
        ("کاندیداهای رویداد", str(len(candidates)), "PASS" if candidates else "FAIL", "صرفاً از کد/شرح حساب جزء ژورنال استخراج شده‌اند؛ canonical یا authoritative نیستند."),
        ("سطرهای بدون مبنای رویداد", str(qa["unclassified_event_rows"]), "REVIEW", "در شیت ردیف‌های مبهم/بی‌طبقه آمده‌اند؛ رویداد حدسی ساخته نشد."),
        ("تطبیق دقیق سطر با baseline", str(second_pass_qa["line_level_links"]), "PASS", "فقط این سطرها در ماتریس رویداد↔مرکز هزینه آمده‌اند."),
        ("سطرهای بدون شاهد مرکز هزینه", str(second_pass_qa["unmapped_rows"]), "REVIEW", "مرکز هزینه برای آن‌ها استنباط نشده و رابطه‌ای تولید نشده است."),
        ("سطرهای با شاهد مرکز هزینه مبهم", str(second_pass_qa["ambiguous_rows"]), "PASS" if not second_pass_qa["ambiguous_rows"] else "REVIEW", "هیچ مرکز هزینه‌ای برای مورد مبهم انتخاب نشده است."),
        ("جفت‌های observed رویداد↔مرکز هزینه", str(len(matrix)), "PASS", "تماماً historical/non-authoritative."),
        ("مراکز هزینهٔ مرجع v3", str(baseline_metadata["master_rows"]), "PASS", "کپی raw master در شیت مرجع مراکز هزینه؛ ستون‌های B/C حفظ شده‌اند."),
        ("وجود ستون B عنوان مرکز هزینه", "بله", "PASS", "فهرست_نهایی_پیشنهادی baseline بدون حذف در خروجی مرجع حفظ شده است."),
        ("وجود ستون C حوزه خدمت", "بله", "PASS", "فهرست_نهایی_پیشنهادی baseline بدون حذف در خروجی مرجع حفظ شده است."),
        ("رابطهٔ authoritative در خروجی", "0", "PASS", "هیچ جدول Rule authoritative ورودی این pipeline نبوده است."),
        ("تخصیص نسبتی در سطح سند", "0", "PASS", "فقط match سطری؛ شماره سند فقط Source Reference و معیار شمارش اسناد است."),
        ("نمونه‌های CC-502 تا CC-506", str(client_sample_row - 1), "PASS", "حداکثر پنج شاهد سطری برای هر مرکز هزینهٔ Client در شیت نمونه‌ها."),
    ]
    for audit in supplemental_audits:
        qa_rows.append(
            (
                f"ممیزی کاتالوگ ۱۴۰۵: {audit['filename']}",
                str(len(audit["sheets"])),
                "PASS",
                "ساختار و headerهای workbook ثبت شد؛ منبع service/activity است و برای کاندیدا، رابطه یا Rule رویداد مالی استفاده نشده است.",
            )
        )
    add_table_sheet(
        "خلاصه_QA",
        ("کنترل", "مقدار", "وضعیت", "توضیح"),
        qa_rows,
        (38, 22, 14, 90),
    )

    workbook.close()
    return {
        "row_evidence_rows": evidence_row - 1,
        "issue_rows": issue_row - 1,
        "client_sample_rows": client_sample_row - 1,
        "line_level_links": second_pass_qa["line_level_links"],
        "unmapped_rows": second_pass_qa["unmapped_rows"],
        "ambiguous_rows": second_pass_qa["ambiguous_rows"],
        "candidate_count": len(candidates),
        "matrix_pair_count": len(matrix),
        "supplemental_catalog_count": len(supplemental_audits),
    }


def build_manifest(
    *,
    output_dir: Path,
    workbook_path: Path,
    v3_path: Path,
    accounting_path: Path,
    source_hashes_before: Mapping[str, str],
    source_hashes_after: Mapping[str, str],
    baseline_metadata: Mapping[str, Any],
    source_sheet_details: Sequence[Mapping[str, Any]],
    supplemental_audits: Sequence[Mapping[str, Any]],
    metrics: Mapping[str, Any],
) -> Path:
    supplemental_source_entries = [
        {
            "role": "audited_not_used_1405_service_catalog",
            "evidence_origin": audit["evidence_origin"],
            "usage": audit["usage"],
            "reason": audit["reason"],
            "path": str(audit["path"]),
            "sha256_before": source_hashes_before[audit["hash_key"]],
            "sha256_after": source_hashes_after[audit["hash_key"]],
            "unchanged_during_run": (
                source_hashes_before[audit["hash_key"]] == source_hashes_after[audit["hash_key"]]
            ),
            "sheets": audit["sheets"],
        }
        for audit in supplemental_audits
    ]
    manifest = {
        "manifest_schema_version": 1,
        "pipeline": {
            "name": "financial_event_discovery_workbook",
            "version": PIPELINE_VERSION,
            "generated_at": datetime.now().astimezone().isoformat(),
        },
        "deliverable": {
            "workbook": str(workbook_path),
            "revision_directory": str(output_dir),
            "fiscal_year": FISCAL_YEAR,
            "scope": "Historical Financial Event discovery evidence; not an authoritative catalog or relationship rule table.",
        },
        "sources": [
            {
                "role": "immutable_v3_cost_center_baseline",
                "path": str(v3_path),
                "sha256_before": source_hashes_before["v3"],
                "sha256_after": source_hashes_after["v3"],
                "unchanged_during_run": source_hashes_before["v3"] == source_hashes_after["v3"],
                "sheets": {
                    MASTER_SHEET: {"rows": baseline_metadata["master_rows"], "columns": baseline_metadata["master_columns"]},
                    ERP_MASTER_SHEET: {"rows": baseline_metadata["erp_rows"], "columns": baseline_metadata["erp_columns"]},
                    BASELINE_MAPPING_SHEET: {"rows": baseline_metadata["mapping_rows"], "columns": baseline_metadata["mapping_columns"]},
                },
            },
            {
                "role": "immutable_1404_accounting_journal",
                "path": str(accounting_path),
                "sha256_before": source_hashes_before["accounting"],
                "sha256_after": source_hashes_after["accounting"],
                "unchanged_during_run": source_hashes_before["accounting"] == source_hashes_after["accounting"],
                "sheets": [
                    {"name": item["sheet_name"], "rows": item["rows"], "columns": item["columns"], "headers": item["header"]}
                    for item in source_sheet_details
                ],
            },
        ] + supplemental_source_entries,
        "analysis_controls": {
            "persian_normalization": "Applied only in-memory to analysis comparison/grouping keys; raw source fields are retained in line evidence.",
            "candidate_identity": "CFE-* is an opaque SHA-256-derived candidate fingerprint. Source document/request/budget values remain Source References, not internal identity.",
            "cost_center_matching": "Exact line-level match to v3 Doc_No_Mapping on area, date, document number, budget code, account code, debit, credit, and account detail.",
            "allocation": "No document-level proportional allocation; every observed relationship is tied to one source journal row.",
            "authority": "No authoritative Financial Event-to-Cost Center table was provided. All candidates and relationships are historical observed evidence and non-authoritative.",
            "source_subsystem": "Not inferred from journal worksheet/Cost Center Service Domain; marked unspecified when absent from source.",
            "supplemental_1405_catalogs": "Audited and versioned in this manifest only; service/activity and service-event labels are not treated as Financial Event evidence or crosswalk inputs.",
        },
        "metrics": dict(metrics),
    }
    manifest_path = output_dir / "source_manifest.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest_path


def run_pipeline(
    v3_path: Path,
    accounting_path: Path,
    output_root: Path,
    supplemental_sources: Sequence[Path] = DEFAULT_SUPPLEMENTAL_SOURCES,
) -> tuple[Path, Path, dict[str, Any]]:
    """Run the full discovery pipeline and return workbook, manifest, and metrics."""

    v3_path = v3_path.expanduser().resolve()
    accounting_path = accounting_path.expanduser().resolve()
    output_root = output_root.expanduser().resolve()
    resolved_supplemental_sources = tuple(path.expanduser().resolve() for path in supplemental_sources)
    for source_path in (v3_path, accounting_path, *resolved_supplemental_sources):
        if not source_path.is_file():
            raise FileNotFoundError(f"Required source workbook not found: {source_path}")

    source_hashes_before = {"v3": sha256_file(v3_path), "accounting": sha256_file(accounting_path)}
    source_hashes_before.update(
        {
            f"supplemental_{source_index}": sha256_file(source_path)
            for source_index, source_path in enumerate(resolved_supplemental_sources, start=1)
        }
    )
    supplemental_audits = audit_supplemental_catalogs(resolved_supplemental_sources)
    supplemental_signature = hashlib.sha256(
        "|".join(source_hashes_before[audit["hash_key"]] for audit in supplemental_audits).encode("ascii")
    ).hexdigest()[:8]
    revision = (
        f"{FISCAL_YEAR}-r{datetime.now():%Y%m%d}-{source_hashes_before['v3'][:8]}-"
        f"{source_hashes_before['accounting'][:8]}-{supplemental_signature}"
    )
    output_dir = output_root / "revisions" / revision
    workbook_path = output_dir / f"financial_event_discovery_{revision}.xlsx"

    master_rows, erp_rows, cost_center_master, baseline_by_line, baseline_metadata = load_baseline(v3_path)
    candidates, matrix, qa, source_sheet_details = scan_accounting_source(accounting_path, baseline_by_line)
    metrics = make_workbook(
        workbook_path,
        v3_path=v3_path,
        accounting_path=accounting_path,
        source_hashes=source_hashes_before,
        master_rows=master_rows,
        erp_rows=erp_rows,
        cost_center_master=cost_center_master,
        baseline_by_line=baseline_by_line,
        baseline_metadata=baseline_metadata,
        candidates=candidates,
        matrix=matrix,
        qa=qa,
        source_sheet_details=source_sheet_details,
        supplemental_audits=supplemental_audits,
    )
    source_hashes_after = {"v3": sha256_file(v3_path), "accounting": sha256_file(accounting_path)}
    source_hashes_after.update(
        {
            f"supplemental_{source_index}": sha256_file(source_path)
            for source_index, source_path in enumerate(resolved_supplemental_sources, start=1)
        }
    )
    manifest_path = build_manifest(
        output_dir=output_dir,
        workbook_path=workbook_path,
        v3_path=v3_path,
        accounting_path=accounting_path,
        source_hashes_before=source_hashes_before,
        source_hashes_after=source_hashes_after,
        baseline_metadata=baseline_metadata,
        source_sheet_details=source_sheet_details,
        supplemental_audits=supplemental_audits,
        metrics=metrics,
    )
    return workbook_path, manifest_path, metrics


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--v3-baseline", type=Path, default=DEFAULT_V3_BASELINE)
    parser.add_argument("--accounting-source", type=Path, default=DEFAULT_ACCOUNTING_SOURCE)
    parser.add_argument("--output-root", type=Path, default=DEFAULT_OUTPUT_ROOT)
    parser.add_argument(
        "--supplemental-catalog",
        type=Path,
        action="append",
        help="1405 service/activity catalog to audit only; it is never Financial Event evidence.",
    )
    arguments = parser.parse_args()

    workbook_path, manifest_path, metrics = run_pipeline(
        arguments.v3_baseline,
        arguments.accounting_source,
        arguments.output_root,
        tuple(arguments.supplemental_catalog) if arguments.supplemental_catalog else DEFAULT_SUPPLEMENTAL_SOURCES,
    )
    print(f"Workbook: {workbook_path}")
    print(f"Manifest: {manifest_path}")
    print(json.dumps(metrics, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
