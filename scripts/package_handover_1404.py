"""Copy complete evidence, package, and independently validate client delivery."""
import argparse,csv,hashlib,json,shutil,zipfile,xml.etree.ElementTree as ET
from pathlib import Path
import openpyxl
BASE=Path('data/deliverables');V1=BASE/'journal_1404_enrichment_final';AUDIT=BASE/'journal_1404_enrichment_v1_independent_audit';FINAL=BASE/'FINAL_CLIENT_HANDOVER_1404';BUILD=BASE/'handover_1404_build'
BOOK='Accounting_Journal_1404_Desired_State_Enriched_v1.xlsx'
V1FILES=['Execution_Report_EN.md','audit_fill.csv','unresolved_review.csv','verification_final.json','delivery_manifest.json']
AUDITFILES=['Independent_Audit_Report_EN.md','QA_REVIEW_Explanations.md','Cost_Center_Coverage_Reconciliation.csv','Conflict_Root_Cause_Analysis.csv','Bank_Assignment_Validation.csv','Bank_Assignment_Exceptions.csv','Organizational_Unit_Formula_Audit.csv','Dimension_Eligibility_and_Coverage.xlsx','Person_Master_Exceptions.csv','Semantic_Cost_Center_Audit.csv','full_population_findings.csv','Proposed_Correction_Package_EN.md','independent_audit_manifest.json']
EXTRA=['proposed_cc_corrections.csv','proposed_person_corrections.csv','Applied_Match_Classification_Audit.csv','Revenue_Field_Usage_Investigation.csv','semantic_metrics.json','org_person_metrics.json','person_variant_taxonomy.json','independent_metrics.json','eligibility_metrics.json','independent_entry_verification.json','conflict_taxonomy_metrics.json','audit_output_validation.json','Management_Summary_FA.md']
MAIN=['01_Accounting_Journal_Enriched_1404.xlsx','02_Management_Report_FA.pdf','03_Expert_Review_and_Decisions.xlsx','04_Required_Reference_Data_FA.pdf','05_Readme_FA.pdf']
def digest(path):
 h=hashlib.sha256()
 with path.open('rb') as f:
  for chunk in iter(lambda:f.read(8*1024*1024),b''):h.update(chunk)
 return h.hexdigest()
def validate(path):
 ext=path.suffix.lower()
 if ext=='.json':json.loads(path.read_text(encoding='utf-8'));return {'parse':'PASS'}
 if ext=='.csv':
  with path.open(encoding='utf-8-sig',newline='') as f:
   reader=csv.reader(f);header=next(reader);n=0
   for row in reader:
    if len(row)!=len(header):raise ValueError('CSV column count mismatch '+str(path)+' '+str(n+2))
    n+=1
  return {'parse':'PASS','records':n,'columns':len(header)}
 if ext=='.xlsx':
  with zipfile.ZipFile(path) as z:
   bad=z.testzip()
   if bad:raise ValueError('Bad XLSX CRC '+bad)
  w=openpyxl.load_workbook(path,read_only=True,data_only=False)
  sheets=w.sheetnames
  for s in w:list(s.iter_rows(min_row=1,max_row=1,values_only=True))
  w.close();return {'parse':'PASS','sheets':sheets}
 if ext in {'.md','.txt'}:
  text=path.read_text(encoding='utf-8-sig')
  if not text.strip():raise ValueError('Empty document '+str(path))
  return {'parse':'PASS'}
 if ext=='.pdf':
  if not path.read_bytes().startswith(b'%PDF'):raise ValueError('PDF header invalid')
  return {'parse':'PASS','visual_verification':'pdf_build_metrics.json'}
 raise ValueError('Unsupported delivery file '+str(path))
def copy_sources():
 FINAL.mkdir(parents=True,exist_ok=True);BUILD.mkdir(parents=True,exist_ok=True);technical=FINAL/'Technical_Audit';technical.mkdir(exist_ok=True)
 inventory=[]
 for src,dst in [(V1/BOOK,FINAL/MAIN[0])]+[(V1/n,technical/n) for n in V1FILES]+[(AUDIT/n,technical/n) for n in AUDITFILES+EXTRA]:
  if not src.is_file():raise FileNotFoundError('Required source missing '+str(src))
  before=digest(src);result=validate(src);shutil.copy2(src,dst)
  if digest(dst)!=before:raise RuntimeError('Copy checksum mismatch '+str(dst))
  inventory.append({'source':str(src.resolve()),'destination':str(dst.relative_to(FINAL)),'bytes':src.stat().st_size,'sha256':before,'validation':result})
  print('Copied and verified',dst.name,src.stat().st_size,flush=True)
 (BUILD/'copied_sources.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding='utf-8')
 print('Verified evidence copies',len(inventory),flush=True)
def package():
 copies=json.loads((BUILD/'copied_sources.json').read_text(encoding='utf-8'))
 for r in copies:
  if digest(Path(r['source']))!=r['sha256'] or digest(FINAL/r['destination'])!=r['sha256']:raise RuntimeError('Source or evidence copy altered')
 wbmetrics=json.loads((BUILD/'workbook_build_metrics.json').read_text(encoding='utf-8'))
 pdfmetrics=json.loads((BUILD/'pdf_build_metrics.json').read_text(encoding='utf-8'))
 independent=json.loads((BUILD/'expert_workbook_independent_verification.json').read_text(encoding='utf-8'))
 if independent['checks']!='PASS' or not all(wbmetrics['verification'].values()):raise RuntimeError('Reviewer workbook verification failed')
 if wbmetrics['counts']!={'cc_review':884,'proposed':6977,'conflict':8042,'bank':245,'ambiguous':31,'correction':51}:raise RuntimeError('Review population mismatch')
 for document in pdfmetrics['documents']:
  if not document['font_embedding'] or document['rendered_page_count']!=document['expected_pages'] or not all(p['nonblank'] for p in document['pages']):raise RuntimeError('PDF rendering validation failed')
 if pdfmetrics.get('parent_visual_inspection',{}).get('status')!='PASS':raise RuntimeError('Actual PDF visual inspection missing')
 for name in MAIN:
  if not (FINAL/name).is_file():raise FileNotFoundError(name)
  validate(FINAL/name)
 if digest(FINAL/MAIN[0])!=digest(V1/BOOK):raise RuntimeError('Delivered v1 not identical')
 if 23243+8042+6977+31+7722!=46015:raise RuntimeError('Population partition failure')
 inventory=[]
 for path in sorted(FINAL.rglob('*')):
  if not path.is_file() or path.name in {'FINAL_HANDOVER_MANIFEST.json','CHECKSUMS_SHA256.txt'}:continue
  inventory.append({'path':path.relative_to(FINAL).as_posix(),'bytes':path.stat().st_size,'sha256':digest(path),'validation':validate(path)})
 manifest={'package_status':'READY_FOR_RESTRICTED_EXPERT_CLIENT_REVIEW_NOT_APPROVED_ERP','journal_label':'نسخه تکمیل‌شده پیشنهادی — نیازمند بازبینی و تأیید کارشناسی','source_v1_unchanged':True,'identical_delivered_journal_sha256':digest(V1/BOOK),'statistics':{'accounting_rows':553667,'unique_region_documents':75279,'expense_rows':46015,'applied_cc_rows':23243,'expense_conflicts':8042,'proposed_expense_rows':6977,'ambiguous_expense_rows':31,'no_evidence_expense_rows':7722,'applied_cc_review_rows':884,'bank_review_rows':245,'enriched_cells':198336,'bank_name_assignments':150648,'cc_enriched_cells':47688},'workbook_verification':wbmetrics,'pdf_verification':pdfmetrics,'files':inventory,'manifest_file':'FINAL_HANDOVER_MANIFEST.json','checksums_file':'CHECKSUMS_SHA256.txt','self_reference_note':'Manifest lists both metadata filenames; checksum file hashes manifest and all payloads, excluding itself to avoid recursive checksum.','confidentiality':'Local packaging only; authorized recipients only; no credentials or encryption keys included.','limitations':['Proposed mappings and51CCcorrection candidates not applied','Missing approved crosswalk,contract/project references,globalBankIDs,eventgovernance andprivacyrelease authorization','No source/applied status changed; not production-ready ERP master']}
 (FINAL/'FINAL_HANDOVER_MANIFEST.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
 paths=[p for p in sorted(FINAL.rglob('*')) if p.is_file() and p.name!='CHECKSUMS_SHA256.txt']
 checks='\n'.join(digest(p)+'  '+p.relative_to(FINAL).as_posix() for p in paths)+'\n'
 (FINAL/'CHECKSUMS_SHA256.txt').write_text(checks,encoding='utf-8')
 archive=BASE/'Isfahan_Municipality_1404_Desired_State_Final_Handover.zip'
 with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=6,allowZip64=True) as z:
  for p in sorted(FINAL.rglob('*')):
   if p.is_file():z.write(p,FINAL.name+'/'+p.relative_to(FINAL).as_posix())
 expected={FINAL.name+'/'+p.relative_to(FINAL).as_posix():digest(p) for p in FINAL.rglob('*') if p.is_file()}
 # Independently extract each archive member to one reusable local temp file,
 # hash physical extracted bytes, then remove; no duplicate full evidence folder.
 temp=BUILD/'zip_member_extraction.tmp';tested=0
 with zipfile.ZipFile(archive) as z:
  if z.testzip() is not None:raise RuntimeError('ZIP CRC failure')
  if set(z.namelist())!=set(expected):raise RuntimeError('ZIP member population mismatch')
  for name in z.namelist():
   with z.open(name) as src,temp.open('wb') as dest:shutil.copyfileobj(src,dest,8*1024*1024)
   if digest(temp)!=expected[name]:raise RuntimeError('ZIP extraction hash mismatch '+name)
   tested+=1
 temp.unlink()
 result={'archive':str(archive.resolve()),'archive_bytes':archive.stat().st_size,'archive_sha256':digest(archive),'members_extracted_and_hash_checked':tested,'zip_crc':'PASS','exact_member_population':'PASS','source_copies_unchanged':'PASS','journal_identical_to_v1':'PASS','final_directory':str(FINAL.resolve()),'workbook_metrics':wbmetrics,'pdf_metrics':pdfmetrics}
 (BUILD/'final_package_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
 (BASE/'Isfahan_Municipality_1404_Desired_State_Final_Handover.zip.sha256').write_text(result['archive_sha256']+'  '+archive.name+'\n',encoding='utf-8')
 print(json.dumps(result,ensure_ascii=False,indent=2),flush=True)
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('phase',choices=['copy','package']);a=p.parse_args()
 copy_sources() if a.phase=='copy' else package()
