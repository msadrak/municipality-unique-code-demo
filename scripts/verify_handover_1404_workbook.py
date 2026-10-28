"""Independently reconcile expert workbook populations and blank authority fields."""
import csv,json,sqlite3
from pathlib import Path
from collections import Counter
import openpyxl
ROOT=Path('data/deliverables');FINAL=ROOT/'FINAL_CLIENT_HANDOVER_1404';AUDIT=ROOT/'journal_1404_enrichment_v1_independent_audit';BUILD=ROOT/'handover_1404_build'
SHEETS=['راهنما_و_خلاصه','تصمیمات_گروه‌بندی‌شده','بازبینی_مراکز_هزینه','نگاشت‌های_پیشنهادی','تعارض‌های_مراکز_هزینه','بازبینی_بانک','چندتطبیقی‌ها','اصلاحات_پیشنهادی','مراجع_مفقود','ثبت_تصمیمات']
def main():
 wb=openpyxl.load_workbook(FINAL/'03_Expert_Review_and_Decisions.xlsx',read_only=False,data_only=False)
 missing=set(SHEETS)-set(wb.sheetnames)
 if missing:raise RuntimeError('Required sheet missing '+str(missing))
 expected={'بازبینی_مراکز_هزینه':884,'نگاشت‌های_پیشنهادی':6977,'تعارض‌های_مراکز_هزینه':8042,'بازبینی_بانک':245,'چندتطبیقی‌ها':31,'اصلاحات_پیشنهادی':51}
 metrics={};fails=[]
 for name,n in expected.items():
  ws=wb[name];observed=ws.max_row-1
  metrics[name]={'data_rows':observed,'physical_rows':ws.max_row,'expected':n,'rtl':ws.sheet_view.rightToLeft,'freeze':str(ws.freeze_panes),'filter':ws.auto_filter.ref,'tables':len(ws.tables)}
  if observed!=n:fails.append([name,'population',observed,n])
  if not ws.sheet_view.rightToLeft or not ws.freeze_panes:fails.append([name,'missing RTL/freeze'])
 groups=wb['تصمیمات_گروه‌بندی‌شده'];metrics[groups.title]={'data_rows':groups.max_row-1,'physical_rows':groups.max_row}
 if groups.max_column<20:fails.append(['groups','less than20fields'])
 for row in groups.iter_rows(min_row=2):
  for cell in row[14:20]:
   if cell.value is not None:fails.append([cell.coordinate,'prefilled authoritative group field',cell.value])
 register=wb['ثبت_تصمیمات']
 # IDs may be preallocated; all approval/implementation data must remain blank.
 for row in register.iter_rows(min_row=2):
  for cell in row[2:]:
   if cell.value is not None:fails.append([register.title,cell.coordinate,'prefilled committee authority'])
 metrics['registered_sheets']={s.title:{'physical_rows':s.max_row,'columns':s.max_column} for s in wb}
 metrics['decision_dropdown_count']=sum(len(s.data_validations.dataValidation) for s in wb)
 if metrics['decision_dropdown_count']==0:fails.append(['dropdowns','missing'])
 wb.close()
 result={'checks':'PASS' if not fails else 'FAIL','sheets':metrics,'failures':fails}
 (BUILD/'expert_workbook_independent_verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps(result,ensure_ascii=False,indent=2))
 if fails:raise RuntimeError('Expert workbook verification failed')
if __name__=='__main__':main()
