"""Yekta issuance for statement-level financial events."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import (
    BudgetItem,
    CivilProject,
    ContinuousAction,
    CostCenterRef,
    EntityDocument,
    FinancialEventRef,
    FinancialMatrix,
    OrgUnit,
    Transaction,
    WorkflowStep,
)
from app.services.yekta_codes import build_yekta_prefix, build_yekta_segments, prefix_matches


STATEMENT_UPLOAD_DOCUMENT_TYPES = {
    "APPENDIX_5_TIMESHEET",
    "MONTHLY_TIMESHEET",
    "CONTRACTOR_STATEMENT_UPLOAD",
}
VOLUME_FINANCIAL_MATRIX_FORM = "FINANCIAL_MATRIX_LINK"
VOLUME_BASE_FORM = "VOLUME_CONTRACT_BASE"
VOLUME_YEKTA_FORM = "YEKTA_CODE_ISSUANCE"
VOLUME_BASE_STEP_NUMBER = 101
VOLUME_STATEMENT_STEP_NUMBER = 102
VOLUME_YEKTA_STEP_NUMBER = 103
VOLUME_GENERATION_STEP_NUMBER = 104
READY_FOR_PAYMENT_QUEUE = "READY_FOR_PAYMENT_QUEUE"

PERSIAN_ARABIC_DIGITS = str.maketrans(
    "\u06f0\u06f1\u06f2\u06f3\u06f4\u06f5\u06f6\u06f7\u06f8\u06f9"
    "\u0660\u0661\u0662\u0663\u0664\u0665\u0666\u0667\u0668\u0669",
    "01234567890123456789",
)


@dataclass(frozen=True)
class MunicipalityFinancialEvent:
    code: str
    label: str
    group: str
    direction: str
    rate: Decimal | None = None
    amount_key: str | None = None
    beneficiary: str | None = None


MUNICIPALITY_FINANCIAL_EVENTS: tuple[MunicipalityFinancialEvent, ...] = (
    MunicipalityFinancialEvent("GROSS_STATEMENT", "Gross statement value", "STATEMENT", "DEBIT", Decimal("1")),
    MunicipalityFinancialEvent("VALUE_ADDED_TAX", "Value added tax", "TAX", "DEBIT", Decimal("0.10"), beneficiary="Tax authority"),
    MunicipalityFinancialEvent("TAX_WITHHOLDING", "Tax withholding", "DEDUCTION", "CREDIT", Decimal("0.03"), beneficiary="Tax authority"),
    MunicipalityFinancialEvent(
        "EMPLOYEE_INSURANCE",
        "Employee insurance share",
        "DEDUCTION",
        "CREDIT",
        Decimal("0.07"),
        beneficiary="Social security",
    ),
    MunicipalityFinancialEvent(
        "EMPLOYER_INSURANCE",
        "Employer insurance share",
        "INSURANCE",
        "DEBIT",
        Decimal("0.23"),
        beneficiary="Social security",
    ),
    MunicipalityFinancialEvent(
        "SOCIAL_SECURITY_DEPOSIT",
        "Social security deposit",
        "DEDUCTION",
        "CREDIT",
        Decimal("0.05"),
        beneficiary="Social security",
    ),
    MunicipalityFinancialEvent(
        "RETENTION_GUARANTEE",
        "Retention guarantee",
        "GUARANTEE",
        "CREDIT",
        Decimal("0.10"),
        beneficiary="Municipality retention account",
    ),
    MunicipalityFinancialEvent(
        "PERFORMANCE_GUARANTEE",
        "Good performance guarantee",
        "GUARANTEE",
        "CREDIT",
        Decimal("0.05"),
        beneficiary="Municipality guarantee account",
    ),
    MunicipalityFinancialEvent(
        "OTHER_DEDUCTIONS",
        "Other approved deductions",
        "DEDUCTION",
        "CREDIT",
        None,
        "other_deductions",
        beneficiary="Municipality",
    ),
    MunicipalityFinancialEvent("NET_PAYABLE", "Net payable to contractor", "PAYABLE", "DEBIT", None),
)


DEFAULT_FINANCIAL_EVENT_CODES: dict[str, str] = {
    "GROSS_STATEMENT": "04",
    "VALUE_ADDED_TAX": "20",
    "TAX_WITHHOLDING": "20",
    "EMPLOYEE_INSURANCE": "20",
    "EMPLOYER_INSURANCE": "07",
    "SOCIAL_SECURITY_DEPOSIT": "07",
    "RETENTION_GUARANTEE": "08",
    "PERFORMANCE_GUARANTEE": "08",
    "OTHER_DEDUCTIONS": "99",
    "NET_PAYABLE": "10",
}


class YektaService:
    """Issues one Yekta code per approved financial event row."""

    def __init__(self, db: Session):
        self.db = db

    def issue_for_approved_volume_statement(self, project_id: int) -> list[FinancialMatrix]:
        project = self._project_or_error(project_id)
        statement_document = self._latest_statement_document(project.id)
        if statement_document is None:
            raise ValueError("Step 102 contractor upload was not found.")

        base_payload, yekta_payload = self._volume_context_payloads(project.id)
        yekta_context = self._normalize_yekta_payload(
            project=project,
            base_payload=base_payload,
            statement_payload=statement_document.payload or {},
            yekta_payload=yekta_payload,
        )

        preview = self.build_volume_statement_preview(project_id, statement_document=statement_document)
        rows: list[FinancialMatrix] = []
        for event_payload in preview["financial_events"]:
            event_code = event_payload["event_code"]
            event_context = self._event_yekta_context(yekta_context, event_payload)
            matrix_row = self._upsert_financial_matrix_row(
                project=project,
                statement_document=statement_document,
                yekta_context=event_context,
                event_payload=event_payload,
            )
            if not matrix_row.yekta_code or not prefix_matches(
                matrix_row.yekta_code,
                self._yekta_prefix(event_context),
            ):
                matrix_row.yekta_code = self._next_available_yekta_code(
                    yekta_context=event_context,
                    current_matrix_id=matrix_row.id,
                )

            transaction = self._upsert_payment_queue_transaction(
                project=project,
                statement_document=statement_document,
                matrix_row=matrix_row,
                yekta_context=event_context,
            )
            matrix_row.transaction_id = transaction.id
            matrix_row.payment_queue_status = READY_FOR_PAYMENT_QUEUE
            rows.append(matrix_row)

        project.yekta_code = None
        self._sync_statement_payload(statement_document, rows)
        self.db.flush()
        return rows

    def build_volume_statement_preview(
        self,
        project_id: int,
        *,
        statement_document: EntityDocument | None = None,
    ) -> dict[str, Any]:
        project = self._project_or_error(project_id)
        statement_document = statement_document or self._latest_statement_document(project.id)
        if statement_document is None:
            raise ValueError("Step 102 contractor upload was not found.")

        base_payload, _yekta_payload = self._volume_context_payloads(project.id, require_yekta=False)
        statement_payload = statement_document.payload or {}
        gross_amount = self._extract_gross_statement_amount(statement_payload)
        period_key = self._statement_period_key(statement_payload, statement_document)
        existing_rows = self._financial_rows_for_statement(project.id, statement_document.id)

        if existing_rows:
            return {
                "project_id": project.id,
                "project_code": project.project_code,
                "statement_document_id": statement_document.id,
                "statement_number": statement_payload.get("statement_number"),
                "work_period": period_key,
                "gross_amount": float(gross_amount),
                "financial_events": [self._matrix_row_to_payload(row) for row in existing_rows],
            }

        event_payloads = self._build_standard_event_payloads(
            gross_amount=gross_amount,
            statement_payload=statement_payload,
            period_key=period_key,
            activity_name=str(base_payload.get("contract_title") or project.title or ""),
        )
        return {
            "project_id": project.id,
            "project_code": project.project_code,
            "statement_document_id": statement_document.id,
            "statement_number": statement_payload.get("statement_number"),
            "work_period": period_key,
            "gross_amount": float(gross_amount),
            "financial_events": event_payloads,
        }

    def statement_rows_for_project(self, project_id: int) -> list[FinancialMatrix]:
        return (
            self.db.query(FinancialMatrix)
            .filter(FinancialMatrix.project_id == project_id)
            .order_by(
                FinancialMatrix.statement_document_id.desc(),
                FinancialMatrix.event_sequence.asc(),
                FinancialMatrix.id.asc(),
            )
            .all()
        )

    def _upsert_financial_matrix_row(
        self,
        *,
        project: CivilProject,
        statement_document: EntityDocument,
        yekta_context: dict[str, Any],
        event_payload: dict[str, Any],
    ) -> FinancialMatrix:
        matrix_row = (
            self.db.query(FinancialMatrix)
            .filter(
                FinancialMatrix.project_id == project.id,
                FinancialMatrix.statement_document_id == statement_document.id,
                FinancialMatrix.financial_event == event_payload["event_code"],
            )
            .first()
        )
        if matrix_row is None:
            matrix_row = FinancialMatrix(
                project_id=project.id,
                statement_document_id=statement_document.id,
                financial_event=event_payload["event_code"],
            )
            self.db.add(matrix_row)
            self.db.flush()

        matrix_row.budget_line_level_5 = str(yekta_context["budget_code"])
        matrix_row.cost_center = str(yekta_context["cost_center_code"])
        matrix_row.financial_event_label = event_payload["event_label"]
        matrix_row.event_group = event_payload["event_group"]
        matrix_row.event_sequence = int(event_payload["event_sequence"])
        matrix_row.event_direction = event_payload["event_direction"]
        matrix_row.activity_name = str(yekta_context["special_activity"])
        matrix_row.period_key = event_payload["period_key"]
        matrix_row.amount = self._money(event_payload["amount"])
        matrix_row.source_payload = {
            **event_payload,
            "yekta_payload": yekta_context,
        }
        matrix_row.is_active = True
        return matrix_row

    def _upsert_payment_queue_transaction(
        self,
        *,
        project: CivilProject,
        statement_document: EntityDocument,
        matrix_row: FinancialMatrix,
        yekta_context: dict[str, Any],
    ) -> Transaction:
        tracking_number = self._tracking_number(project, statement_document, matrix_row)
        transaction = (
            self.db.query(Transaction)
            .filter(Transaction.tracking_number == tracking_number)
            .first()
        )
        amount = int(self._money(matrix_row.amount or 0).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
        budget_item = self._budget_item(str(yekta_context["budget_code"]))
        beneficiary_name = str(yekta_context["beneficiary_name"])

        if transaction is None:
            transaction = Transaction(
                tracking_number=tracking_number,
                current_approval_level=0,
                created_by_id=statement_document.created_by_id,
            )
            self.db.add(transaction)
            self.db.flush()

        transaction.unique_code = matrix_row.yekta_code
        transaction.status = READY_FOR_PAYMENT_QUEUE
        transaction.zone_id = yekta_context["zone_id"]
        transaction.department_id = yekta_context["department_id"]
        transaction.section_id = yekta_context["section_id"]
        transaction.activity_id = project.activity_id
        transaction.budget_item_id = budget_item.id if budget_item else None
        transaction.cost_center_id = yekta_context.get("cost_center_id")
        transaction.continuous_action_id = yekta_context.get("continuous_action_id")
        transaction.financial_event_id = yekta_context.get("financial_event_id")
        transaction.amount = amount
        transaction.beneficiary_name = beneficiary_name
        transaction.contract_number = project.project_code
        transaction.special_activity = str(yekta_context["special_activity"])
        transaction.description = (
            f"Volume contract statement event {matrix_row.financial_event}: "
            f"{matrix_row.financial_event_label or matrix_row.financial_event}"
        )
        transaction.form_data = {
            "source": "VOLUME_CONTRACT_STATEMENT",
            "civil_project_id": project.id,
            "civil_project_code": project.project_code,
            "entity_document_id": statement_document.id,
            "financial_matrix_id": matrix_row.id,
            "financial_event": matrix_row.financial_event,
            "financial_event_label": matrix_row.financial_event_label,
            "event_group": matrix_row.event_group,
            "event_direction": matrix_row.event_direction,
            "period_key": matrix_row.period_key,
            "yekta_code": matrix_row.yekta_code,
            "yekta_payload": yekta_context,
        }
        transaction.fiscal_year = str(yekta_context["fiscal_year"])
        return transaction

    def _build_standard_event_payloads(
        self,
        *,
        gross_amount: Decimal,
        statement_payload: dict[str, Any],
        period_key: str,
        activity_name: str,
    ) -> list[dict[str, Any]]:
        gross_amount = self._money(gross_amount)
        amounts: dict[str, Decimal] = {}
        for event in MUNICIPALITY_FINANCIAL_EVENTS:
            if event.code == "NET_PAYABLE":
                continue
            if event.amount_key:
                amount = self._to_decimal(statement_payload.get(event.amount_key, 0))
            elif event.rate is not None:
                amount = gross_amount * event.rate
            else:
                amount = Decimal("0")
            amounts[event.code] = self._money(amount)

        deductions = (
            amounts["TAX_WITHHOLDING"]
            + amounts["EMPLOYEE_INSURANCE"]
            + amounts["SOCIAL_SECURITY_DEPOSIT"]
            + amounts["RETENTION_GUARANTEE"]
            + amounts["PERFORMANCE_GUARANTEE"]
            + amounts["OTHER_DEDUCTIONS"]
        )
        amounts["NET_PAYABLE"] = self._money(
            gross_amount + amounts["VALUE_ADDED_TAX"] - deductions
        )

        payloads: list[dict[str, Any]] = []
        for sequence, event in enumerate(MUNICIPALITY_FINANCIAL_EVENTS, start=1):
            amount = amounts[event.code]
            payloads.append(
                {
                    "event_sequence": sequence,
                    "event_code": event.code,
                    "event_label": event.label,
                    "event_group": event.group,
                    "event_direction": event.direction,
                    "rate": str(event.rate) if event.rate is not None else None,
                    "amount": float(amount),
                    "period_key": period_key,
                    "activity_name": activity_name,
                    "beneficiary": event.beneficiary,
                    "payment_queue_status": READY_FOR_PAYMENT_QUEUE,
                    "yekta_code": None,
                }
            )
        return payloads

    def _extract_gross_statement_amount(self, payload: dict[str, Any]) -> Decimal:
        for key in (
            "statement_total_amount",
            "declared_total_amount",
            "gross_amount",
            "total_amount",
            "invoice_amount",
        ):
            value = payload.get(key)
            amount = self._to_decimal(value)
            if amount > 0:
                return self._money(amount)

        excel_amounts: list[Decimal] = []
        for attachment in self._payload_attachments(payload):
            file_url = str(attachment.get("file_url") or attachment.get("url") or "")
            file_name = str(attachment.get("file_name") or file_url)
            suffix = Path(file_name or file_url).suffix.lower()
            if suffix not in {".xlsx", ".xlsm", ".xls"}:
                continue
            excel_amounts.extend(self._extract_amounts_from_excel(file_url))

        candidates = [amount for amount in excel_amounts if amount >= Decimal("10000")]
        if candidates:
            return self._money(max(candidates))

        raise ValueError(
            "Could not infer the gross statement amount from the uploaded Excel files. "
            "Upload a filled 6-column statement or include statement_total_amount in the payload."
        )

    def _extract_amounts_from_excel(self, file_url: str) -> list[Decimal]:
        path = self._resolve_upload_path(file_url)
        if path is None or not path.exists():
            return []

        suffix = path.suffix.lower()
        if suffix == ".xls":
            return self._extract_amounts_from_legacy_excel(path)

        try:
            workbook = load_workbook(path, data_only=True, read_only=True)
        except Exception:
            return []

        amounts: list[Decimal] = []
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                for value in row:
                    amount = self._to_decimal(value)
                    if amount > 0:
                        amounts.append(amount)
        return amounts

    def _extract_amounts_from_legacy_excel(self, path: Path) -> list[Decimal]:
        try:
            import pandas as pd
        except Exception:
            return []

        try:
            sheets = pd.read_excel(path, sheet_name=None, header=None)
        except Exception:
            return []

        amounts: list[Decimal] = []
        for frame in sheets.values():
            for value in frame.to_numpy().ravel():
                amount = self._to_decimal(value)
                if amount > 0:
                    amounts.append(amount)
        return amounts

    def _payload_attachments(self, payload: dict[str, Any]) -> list[dict[str, Any]]:
        attachments: list[dict[str, Any]] = []
        for key in (
            "_attachments",
            "attachments",
            "statement_file",
            "attendance_18_column_file",
            "statement_6_column_file",
        ):
            value = payload.get(key)
            if isinstance(value, list):
                attachments.extend(item for item in value if isinstance(item, dict))
            elif isinstance(value, dict):
                attachments.append(value)
        return attachments

    def _resolve_upload_path(self, file_url: str) -> Path | None:
        if not file_url:
            return None
        normalized = file_url.split("?", 1)[0].lstrip("/")
        root = Path(__file__).resolve().parents[2]
        candidate = (root / normalized).resolve()
        uploads_root = (root / "uploads").resolve()
        try:
            candidate.relative_to(uploads_root)
        except ValueError:
            return None
        return candidate

    def _next_available_yekta_code(
        self,
        *,
        yekta_context: dict[str, Any],
        current_matrix_id: int,
    ) -> str:
        sequence = 1
        while True:
            candidate = build_yekta_segments(
                zone_code=yekta_context["zone_code"],
                department_code=yekta_context["department_code"],
                section_code=yekta_context["section_code"],
                budget_code=yekta_context["budget_code"],
                cost_center_code=yekta_context["cost_center_code"],
                continuous_action_code=yekta_context["continuous_activity_code"],
                special_activity=yekta_context["special_activity"],
                beneficiary_name=yekta_context["beneficiary_name"],
                financial_event_code=yekta_context["financial_event_code"],
                fiscal_year=yekta_context["fiscal_year"],
                sequence=sequence,
            ).code()
            matrix_exists = (
                self.db.query(FinancialMatrix.id)
                .filter(
                    FinancialMatrix.id != current_matrix_id,
                    FinancialMatrix.yekta_code == candidate,
                )
                .first()
            )
            transaction_exists = (
                self.db.query(Transaction.id)
                .filter(Transaction.unique_code == candidate)
                .first()
            )
            project_exists = (
                self.db.query(CivilProject.id)
                .filter(CivilProject.yekta_code == candidate)
                .first()
            )
            if matrix_exists is None and transaction_exists is None and project_exists is None:
                return candidate
            sequence += 1

    def _yekta_prefix(self, yekta_context: dict[str, Any]) -> str:
        return build_yekta_prefix(
            zone_code=yekta_context["zone_code"],
            department_code=yekta_context["department_code"],
            section_code=yekta_context["section_code"],
            budget_code=yekta_context["budget_code"],
            cost_center_code=yekta_context["cost_center_code"],
            continuous_action_code=yekta_context["continuous_activity_code"],
            special_activity=yekta_context["special_activity"],
            beneficiary_name=yekta_context["beneficiary_name"],
            financial_event_code=yekta_context["financial_event_code"],
            fiscal_year=yekta_context["fiscal_year"],
        )

    def _financial_rows_for_statement(self, project_id: int, statement_document_id: int) -> list[FinancialMatrix]:
        return (
            self.db.query(FinancialMatrix)
            .filter(
                FinancialMatrix.project_id == project_id,
                FinancialMatrix.statement_document_id == statement_document_id,
            )
            .order_by(FinancialMatrix.event_sequence.asc(), FinancialMatrix.id.asc())
            .all()
        )

    def _latest_statement_document(self, project_id: int) -> EntityDocument | None:
        return (
            self.db.query(EntityDocument)
            .join(WorkflowStep, EntityDocument.step_id == WorkflowStep.id)
            .filter(
                EntityDocument.project_id == project_id,
                EntityDocument.document_type.in_(STATEMENT_UPLOAD_DOCUMENT_TYPES),
                WorkflowStep.order_index == VOLUME_STATEMENT_STEP_NUMBER,
            )
            .order_by(EntityDocument.created_at.desc(), EntityDocument.id.desc())
            .first()
        )

    def _volume_context_payloads(
        self,
        project_id: int,
        *,
        require_yekta: bool = True,
    ) -> tuple[dict[str, Any], dict[str, Any]]:
        base_document = self._latest_step_order_document(
            project_id,
            VOLUME_BASE_STEP_NUMBER,
            VOLUME_BASE_FORM,
        )
        yekta_document = self._latest_step_order_document(
            project_id,
            VOLUME_YEKTA_STEP_NUMBER,
            VOLUME_YEKTA_FORM,
        )
        if base_document is None:
            raise ValueError("Step 101 form is required before statement Yekta generation.")
        if require_yekta and yekta_document is None:
            raise ValueError("Step 103 Yekta sub-process payload is required before Yekta generation.")

        yekta_payload = dict(yekta_document.payload or {}) if yekta_document else {}
        if yekta_document is not None:
            yekta_payload["_document_id"] = yekta_document.id
            yekta_payload["_submitted_at"] = (
                yekta_document.created_at.isoformat() if yekta_document.created_at else None
            )
        return base_document.payload or {}, yekta_payload

    def _latest_step_order_document(
        self,
        project_id: int,
        step_number: int,
        document_type: str,
    ) -> EntityDocument | None:
        return (
            self.db.query(EntityDocument)
            .join(WorkflowStep, EntityDocument.step_id == WorkflowStep.id)
            .filter(
                EntityDocument.project_id == project_id,
                EntityDocument.document_type == document_type,
                WorkflowStep.order_index == step_number,
            )
            .order_by(EntityDocument.created_at.desc(), EntityDocument.id.desc())
            .first()
        )

    def _normalize_yekta_payload(
        self,
        *,
        project: CivilProject,
        base_payload: dict[str, Any],
        statement_payload: dict[str, Any],
        yekta_payload: dict[str, Any],
    ) -> dict[str, Any]:
        required_fields = [
            "zone_id",
            "department_id",
            "section_id",
            "budget_code",
            "cost_center_code",
            "continuous_activity_code",
            "special_activity",
            "beneficiary_name",
            "fiscal_year",
        ]
        missing = [
            field
            for field in required_fields
            if yekta_payload.get(field) is None
            or (isinstance(yekta_payload.get(field), str) and not yekta_payload.get(field).strip())
        ]
        if missing:
            raise ValueError(
                "Step 103 Yekta payload is missing required fields: " + ", ".join(missing)
            )

        if "sequence" in yekta_payload or "occurrence_sequence" in yekta_payload:
            raise ValueError("Step 103 Yekta payload must not include the system-generated sequence segment.")

        zone = self._org_unit(int(yekta_payload["zone_id"]), "zone_id")
        department = self._org_unit(int(yekta_payload["department_id"]), "department_id")
        section = self._org_unit(int(yekta_payload["section_id"]), "section_id")
        budget_code = str(yekta_payload.get("budget_code") or "").strip()
        budget_item_id = yekta_payload.get("budget_item_id")
        if budget_item_id and not budget_code:
            budget_item = self.db.query(BudgetItem).filter(BudgetItem.id == int(budget_item_id)).first()
            budget_code = budget_item.budget_code if budget_item else ""
        if not budget_code:
            raise ValueError("Step 103 Yekta payload is missing budget_code.")

        cost_center_code = str(yekta_payload.get("cost_center_code") or "").strip()
        continuous_activity_code = str(yekta_payload.get("continuous_activity_code") or "").strip()
        financial_event_defaults = yekta_payload.get("financial_event_defaults") or {}
        if not isinstance(financial_event_defaults, dict):
            raise ValueError("financial_event_defaults must be an object keyed by statement event code.")
        for event_code in [event.code for event in MUNICIPALITY_FINANCIAL_EVENTS]:
            if not str(financial_event_defaults.get(event_code) or "").strip():
                raise ValueError(f"financial_event_defaults is missing {event_code}.")

        event_overrides = yekta_payload.get("event_overrides") or []
        if not isinstance(event_overrides, list):
            raise ValueError("event_overrides must be a list.")

        default_beneficiary = (
            str(yekta_payload.get("beneficiary_name") or "").strip()
            or str(base_payload.get("contractor_name") or "").strip()
            or str(statement_payload.get("contractor_name") or "").strip()
            or project.title
        )
        special_activity = (
            str(yekta_payload.get("special_activity") or "").strip()
            or str(base_payload.get("contract_title") or "").strip()
            or str(project.title or "").strip()
        )

        cost_center = self._cost_center(cost_center_code)
        continuous_action = self._continuous_action(continuous_activity_code)

        return {
            "source_document_id": yekta_payload.get("_document_id"),
            "zone_id": zone.id,
            "zone_code": zone.code or zone.id,
            "zone_title": zone.title,
            "department_id": department.id,
            "department_code": department.code or department.id,
            "department_title": department.title,
            "section_id": section.id,
            "section_code": section.code or section.id,
            "section_title": section.title,
            "budget_item_id": int(budget_item_id) if budget_item_id else None,
            "budget_code": budget_code,
            "cost_center_id": cost_center.id if cost_center else None,
            "cost_center_code": cost_center_code,
            "cost_center_title": cost_center.title if cost_center else None,
            "continuous_action_id": continuous_action.id if continuous_action else None,
            "continuous_activity_code": continuous_activity_code,
            "continuous_activity_title": continuous_action.title if continuous_action else None,
            "special_activity": special_activity,
            "beneficiary_name": default_beneficiary,
            "financial_event_defaults": financial_event_defaults,
            "event_overrides": event_overrides,
            "fiscal_year": str(yekta_payload.get("fiscal_year") or self._current_persian_year()).strip(),
            "raw_payload": {
                key: value
                for key, value in yekta_payload.items()
                if key not in {"_document_id", "_submitted_at"}
            },
        }

    def _event_yekta_context(
        self,
        base_context: dict[str, Any],
        event_payload: dict[str, Any],
    ) -> dict[str, Any]:
        event_code = str(event_payload["event_code"])
        override = self._event_override(base_context.get("event_overrides"), event_code)
        default_mapping = base_context.get("financial_event_defaults") or {}
        event_financial_code = (
            (override or {}).get("financial_event_code")
            or default_mapping.get(event_code)
            or DEFAULT_FINANCIAL_EVENT_CODES.get(event_code)
            or "000"
        )
        beneficiary_name = (
            (override or {}).get("beneficiary_name")
            or event_payload.get("beneficiary")
            or base_context["beneficiary_name"]
        )
        financial_event_ref = self._financial_event(str(event_financial_code))
        return {
            **base_context,
            "statement_event_code": event_code,
            "statement_event_label": event_payload.get("event_label"),
            "financial_event_code": str(event_financial_code),
            "financial_event_id": financial_event_ref.id if financial_event_ref else None,
            "financial_event_title": financial_event_ref.title if financial_event_ref else None,
            "beneficiary_name": str(beneficiary_name),
            "event_override": override or None,
        }

    def _event_override(self, overrides: Any, event_code: str) -> dict[str, Any] | None:
        if not isinstance(overrides, list):
            return None
        for item in overrides:
            if not isinstance(item, dict):
                continue
            if str(item.get("event_code") or "") == event_code:
                return item
        return None

    def _org_unit(self, org_unit_id: int, field_name: str) -> OrgUnit:
        org_unit = self.db.query(OrgUnit).filter(OrgUnit.id == org_unit_id).first()
        if org_unit is None:
            raise ValueError(f"Step 103 Yekta payload references unknown {field_name}: {org_unit_id}")
        return org_unit

    def _cost_center(self, code: str) -> CostCenterRef | None:
        if not code:
            return None
        return self.db.query(CostCenterRef).filter(CostCenterRef.code == code).first()

    def _continuous_action(self, code: str) -> ContinuousAction | None:
        if not code:
            return None
        return self.db.query(ContinuousAction).filter(ContinuousAction.code == code).first()

    def _financial_event(self, code: str) -> FinancialEventRef | None:
        if not code:
            return None
        return self.db.query(FinancialEventRef).filter(FinancialEventRef.code == code).first()

    def _sync_statement_payload(self, statement_document: EntityDocument, rows: Iterable[FinancialMatrix]) -> None:
        payload = dict(statement_document.payload or {})
        payload["financial_matrix_rows"] = [self._matrix_row_to_payload(row) for row in rows]
        payload["yekta_issued_at"] = datetime.utcnow().isoformat()
        statement_document.payload = payload

    def _matrix_row_to_payload(self, row: FinancialMatrix) -> dict[str, Any]:
        return {
            "id": row.id,
            "transaction_id": row.transaction_id,
            "event_sequence": row.event_sequence,
            "event_code": row.financial_event,
            "event_label": row.financial_event_label or row.financial_event,
            "event_group": row.event_group,
            "event_direction": row.event_direction,
            "amount": float(row.amount or 0),
            "yekta_code": row.yekta_code,
            "payment_queue_status": row.payment_queue_status,
            "period_key": row.period_key,
            "source_payload": row.source_payload,
        }

    def _tracking_number(
        self,
        project: CivilProject,
        statement_document: EntityDocument,
        matrix_row: FinancialMatrix,
    ) -> str:
        normalized_project_code = re.sub(r"[^A-Za-z0-9]+", "-", project.project_code).strip("-")
        if not normalized_project_code:
            normalized_project_code = f"PROJECT-{project.id}"
        return f"CIVIL-{normalized_project_code}-DOC-{statement_document.id}-EVT-{matrix_row.event_sequence:02d}"

    def _beneficiary_name(self, matrix_row: FinancialMatrix, contractor_name: str) -> str:
        payload = matrix_row.source_payload or {}
        yekta_payload = payload.get("yekta_payload") if isinstance(payload, dict) else None
        if isinstance(yekta_payload, dict) and yekta_payload.get("beneficiary_name"):
            return str(yekta_payload["beneficiary_name"])
        explicit = payload.get("beneficiary")
        if explicit:
            return str(explicit)
        if matrix_row.financial_event == "NET_PAYABLE":
            return contractor_name
        if matrix_row.event_group == "TAX":
            return "Tax authority"
        if matrix_row.event_group == "INSURANCE":
            return "Social security"
        if matrix_row.event_group == "GUARANTEE":
            return "Municipality guarantee account"
        return contractor_name or "Municipality"

    def _statement_period_key(self, payload: dict[str, Any], document: EntityDocument) -> str:
        for key in ("work_period", "work_month", "period", "statement_period"):
            value = payload.get(key)
            if value:
                return str(value)
        if document.created_at:
            return document.created_at.strftime("%Y-%m")
        return self._current_persian_year()

    def _project_or_error(self, project_id: int) -> CivilProject:
        project = self.db.query(CivilProject).filter(CivilProject.id == project_id).first()
        if project is None:
            raise ValueError(f"Civil project {project_id} was not found")
        return project

    def _budget_item(self, budget_line_level_5: str) -> BudgetItem | None:
        if not budget_line_level_5:
            return None
        return self.db.query(BudgetItem).filter(BudgetItem.budget_code == budget_line_level_5).first()

    def _to_decimal(self, value: Any) -> Decimal:
        if value is None:
            return Decimal("0")
        if isinstance(value, Decimal):
            return value
        if isinstance(value, bool):
            return Decimal("0")
        if isinstance(value, int):
            return Decimal(value)
        if isinstance(value, float):
            return Decimal(str(value))

        normalized = str(value).translate(PERSIAN_ARABIC_DIGITS).strip()
        normalized = normalized.replace(",", "").replace("\u066c", "").replace(" ", "")
        normalized = re.sub(r"[^\d.\-]", "", normalized)
        if normalized in {"", "-", ".", "-."}:
            return Decimal("0")
        try:
            return Decimal(normalized)
        except InvalidOperation:
            return Decimal("0")

    def _money(self, value: Any) -> Decimal:
        return self._to_decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    def _current_persian_year(self) -> str:
        now = datetime.utcnow()
        if (now.month, now.day) < (3, 21):
            return str(now.year - 622)
        return str(now.year - 621)


__all__ = [
    "MUNICIPALITY_FINANCIAL_EVENTS",
    "READY_FOR_PAYMENT_QUEUE",
    "STATEMENT_UPLOAD_DOCUMENT_TYPES",
    "YektaService",
]
