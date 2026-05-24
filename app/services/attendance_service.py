import io
import logging
import re
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Callable, Optional

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.models.attendance import (
    AttendanceRecord,
    ContractorStatement,
    ContractorStatementItem,
    PenaltyRecord,
    RegionalAttendanceSheet,
    WorkPeriod,
)
from app.models.enums import AttendanceSheetStatus, WorkFlowStatus
from app.services.notification_service import create_notification
from app.schemas.attendance import (
    AttendanceAggregationItem,
    AttendanceAggregationOut,
    AttendanceReconciliationOut,
    PenaltyRecordCreate,
    PenaltyRecordOut,
    SheetMetaSummary,
    WorkPeriodApprovalOut,
    WorkPeriodStatusOut,
    WorkflowStatusLabel,
)
from app.utils.constants import AUTHORIZED_REGIONS
from app.utils.pricing_engine import calculate_record_legal_value

EXCEL_SKIPROWS = 2
EXCEL_HEADER_ROW_NUMBER = 3
EXCEL_DATA_START_ROW_NUMBER = EXCEL_HEADER_ROW_NUMBER + 1
logger = logging.getLogger(__name__)
UNIQUE_CODE_CONNECTION_ERROR_DETAIL = (
    "خطا در برقراری ارتباط با سرور مالی جهت صدور کد یکتا. لطفاً دقایقی دیگر تلاش کنید."
)

ATTENDANCE_FINANCE_ACTIVITY_NAME = "تسویه کارکرد پیمانکاران"

ATTENDANCE_FIELD_ORDER = [
    "location",
    "full_name",
    "father_name",
    "national_id",
    "marital_status",
    "children_count",
    "start_date",
    "average_per_capita",
    "min_paid_hours",
    "max_paid_hours",
    "monthly_performance_days",
    "monthly_absence_days",
    "overtime_replacement",
    "overtime_hours",
    "total_average_per_capita",
    "total_overtime_given",
    "performance_balance",
]

ATTENDANCE_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "location": ("موقعیت",),
    "full_name": ("نام و نام خانوادگی", "نام ونام خانوادگی", "full_name"),
    "father_name": ("نام پدر",),
    "national_id": ("شماره ملی",),
    "marital_status": ("وضعیت تاهل", "وضعیت تأهل"),
    "children_count": ("تعداد فرزند",),
    "start_date": ("تاریخ شروع به کار",),
    "average_per_capita": ("میانگین سرانه", "amount"),
    "min_paid_hours": ("حداقل ساعت پرداختی",),
    "max_paid_hours": ("حداکثر ساعت پرداختی",),
    "monthly_performance_days": ("کارکرد ماهانه روز",),
    "monthly_absence_days": ("غیبت ماهانه روز",),
    "overtime_replacement": ("اضافه کار جایگزین مرخصی",),
    "overtime_hours": ("اضافه کار", "worked"),
    "total_average_per_capita": ("جمع کل میانگین سرانه", "total_sum", "total"),
    "total_overtime_given": ("جمع کل اضافه کار داده شده",),
    "performance_balance": ("مانده کارکرد", "description"),
}

SERIAL_COLUMN_KEYS = {
    "ردیف",
    "row",
    "index",
    "شماره",
}

SUMMARY_ROW_MARKERS = {
    "جمعکل",
    "جمع",
    "تنظیمکننده",
}

_ARABIC_CHAR_TRANSLATION = str.maketrans(
    {
        "ي": "ی",
        "ك": "ک",
        "ة": "ه",
        "ۀ": "ه",
    }
)
_PERSIAN_DIGIT_TRANSLATION = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
    "01234567890123456789",
)


def _normalize_text(value: str) -> str:
    normalized = value.translate(_ARABIC_CHAR_TRANSLATION)
    normalized = normalized.replace("\u200c", " ").replace("\xa0", " ")
    normalized = normalized.replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", normalized).strip()


def _compact_text(value: str) -> str:
    normalized = _normalize_text(value).lower()
    return re.sub(r"[\s\-_()/\\:؛،,.]+", "", normalized)


def _is_missing(value) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        cleaned = _normalize_text(value).lower()
        return cleaned in {"", "nan", "none", "null", "nat"}
    try:
        return value != value
    except Exception:
        return False


def _safe_string(value) -> Optional[str]:
    if _is_missing(value):
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    cleaned = _normalize_text(str(value))
    return cleaned or None


def _safe_decimal(value) -> Optional[Decimal]:
    if _is_missing(value):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return Decimal(str(value))
        except (InvalidOperation, ValueError):
            return None

    normalized = (_safe_string(value) or "").translate(_PERSIAN_DIGIT_TRANSLATION)
    normalized = normalized.replace("−", "-").replace("–", "-").replace("—", "-")
    normalized = normalized.replace("٬", "").replace("،", "").replace(",", "")
    normalized = normalized.replace("٫", ".").replace(" ", "")
    if normalized in {"", "-", ".", "-."}:
        return None

    try:
        return Decimal(normalized)
    except (InvalidOperation, ValueError):
        return None


def _safe_int(value) -> Optional[int]:
    decimal_value = _safe_decimal(value)
    if decimal_value is None:
        return None
    try:
        return int(decimal_value)
    except (ValueError, OverflowError):
        return None


FIELD_PARSERS: dict[str, Callable[[object], object]] = {
    "location": _safe_string,
    "full_name": _safe_string,
    "father_name": _safe_string,
    "national_id": _safe_string,
    "marital_status": _safe_string,
    "children_count": _safe_int,
    "start_date": _safe_string,
    "average_per_capita": _safe_decimal,
    "min_paid_hours": _safe_decimal,
    "max_paid_hours": _safe_decimal,
    "monthly_performance_days": _safe_decimal,
    "monthly_absence_days": _safe_decimal,
    "overtime_replacement": _safe_decimal,
    "overtime_hours": _safe_decimal,
    "total_average_per_capita": _safe_decimal,
    "total_overtime_given": _safe_decimal,
    "performance_balance": _safe_string,
}


def _build_column_mapping(columns) -> dict[str, object]:
    original_columns = list(columns)
    normalized_lookup: dict[str, object] = {}

    for column_name in original_columns:
        normalized = _compact_text(str(column_name))
        if normalized and normalized not in normalized_lookup:
            normalized_lookup[normalized] = column_name

    mapping: dict[str, object] = {}
    used_columns: set[object] = set()

    for field_name in ATTENDANCE_FIELD_ORDER:
        for alias in ATTENDANCE_COLUMN_ALIASES[field_name]:
            matched_column = normalized_lookup.get(_compact_text(alias))
            if matched_column is not None and matched_column not in used_columns:
                mapping[field_name] = matched_column
                used_columns.add(matched_column)
                break

    remaining_fields = [field for field in ATTENDANCE_FIELD_ORDER if field not in mapping]
    remaining_columns = [column for column in original_columns if column not in used_columns]

    if remaining_columns and _compact_text(str(remaining_columns[0])) in {
        _compact_text(value) for value in SERIAL_COLUMN_KEYS
    }:
        remaining_columns = remaining_columns[1:]

    if remaining_columns:
        # Fallback: map by column order when file provides a subset (for example 6-column templates).
        for field_name, column_name in zip(remaining_fields, remaining_columns):
            mapping[field_name] = column_name

    if "full_name" not in mapping:
        raise HTTPException(
            status_code=422,
            detail="Excel headers do not match the expected Regional Attendance template.",
        )

    return mapping


def _get_row_value(row, mapping: dict[str, object], field_name: str):
    column_name = mapping.get(field_name)
    if column_name is None:
        return None
    return row.get(column_name)


def _should_skip_row(row, mapping: dict[str, object]) -> bool:
    full_name = _safe_string(_get_row_value(row, mapping, "full_name"))
    if not full_name:
        return True

    compact_full_name = _compact_text(full_name)
    if compact_full_name in SUMMARY_ROW_MARKERS:
        return True
    if compact_full_name.startswith("جمعکل"):
        return True
    if compact_full_name == _compact_text("نام و نام خانوادگی"):
        return True

    return False


def _read_excel_rows(contents: bytes) -> tuple[list[object], list[dict[object, object]]]:
    try:
        workbook = load_workbook(io.BytesIO(contents), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Cannot parse Excel file: {exc}") from exc

    worksheet = workbook.active
    all_rows = list(worksheet.iter_rows(values_only=True))
    if len(all_rows) <= EXCEL_SKIPROWS:
        raise HTTPException(
            status_code=422,
            detail="Excel file does not contain any attendance rows.",
        )

    header_row = list(all_rows[EXCEL_SKIPROWS])
    data_rows = all_rows[EXCEL_SKIPROWS + 1 :]

    last_non_empty_index = -1
    for idx, value in enumerate(header_row):
        if not _is_missing(value):
            last_non_empty_index = idx

    if last_non_empty_index < 0:
        raise HTTPException(
            status_code=422,
            detail="Excel file does not contain any attendance headers.",
        )

    headers = header_row[: last_non_empty_index + 1]
    records: list[dict[object, object]] = []
    for row_values in data_rows:
        trimmed_values = list(row_values[: len(headers)])
        if all(_is_missing(value) for value in trimmed_values):
            continue
        records.append({header: value for header, value in zip(headers, trimmed_values)})

    if not records:
        raise HTTPException(
            status_code=422,
            detail="Excel file does not contain any attendance rows.",
        )

    return headers, records


def get_or_create_work_period(db: Session, year: int, month: int) -> WorkPeriod:
    work_period = (
        db.query(WorkPeriod)
        .filter(
            WorkPeriod.year == year,
            WorkPeriod.month == month,
        )
        .first()
    )
    if work_period:
        return work_period

    work_period = WorkPeriod(
        year=year,
        month=month,
        status=WorkFlowStatus.DRAFT_REGIONS,
    )
    db.add(work_period)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        work_period = (
            db.query(WorkPeriod)
            .filter(
                WorkPeriod.year == year,
                WorkPeriod.month == month,
            )
            .first()
        )
        if work_period:
            return work_period
        raise

    return work_period


def _get_work_period_by_period(
    db: Session,
    *,
    year: int,
    month: int,
) -> Optional[WorkPeriod]:
    return (
        db.query(WorkPeriod)
        .filter(
            WorkPeriod.year == year,
            WorkPeriod.month == month,
        )
        .first()
    )


def _get_work_period_or_404(
    db: Session,
    *,
    year: int,
    month: int,
) -> WorkPeriod:
    work_period = _get_work_period_by_period(db, year=year, month=month)
    if not work_period:
        raise HTTPException(
            status_code=404,
            detail=f"No work period found for year={year}, month={month}.",
        )
    return work_period


def get_work_period_upload_status(
    db: Session,
    *,
    year: int,
    month: int,
) -> WorkPeriodStatusOut:
    work_period = _get_work_period_by_period(db, year=year, month=month)
    normalized_authorized_regions = {
        _normalize_text(region): region for region in AUTHORIZED_REGIONS
    }

    uploaded_region_set: set[str] = set()
    if work_period:
        uploaded_region_rows = (
            db.query(RegionalAttendanceSheet.region)
            .filter(
                RegionalAttendanceSheet.work_period_id == work_period.id,
                RegionalAttendanceSheet.row_count > 0,
                RegionalAttendanceSheet.region.isnot(None),
            )
            .all()
        )
        uploaded_region_set = {
            normalized_authorized_regions[_normalize_text(region)]
            for (region,) in uploaded_region_rows
            if region and _normalize_text(region) in normalized_authorized_regions
        }

    uploaded_regions = [
        region for region in AUTHORIZED_REGIONS if region in uploaded_region_set
    ]
    missing_regions = [
        region for region in AUTHORIZED_REGIONS if region not in uploaded_region_set
    ]

    return WorkPeriodStatusOut(
        status=(work_period.status if work_period else WorkFlowStatus.DRAFT_REGIONS),
        total_uploaded=len(uploaded_regions),
        uploaded_regions=uploaded_regions,
        missing_regions=missing_regions,
    )


def approve_work_period_aggregation(
    db: Session,
    *,
    year: int,
    month: int,
) -> WorkPeriodApprovalOut:
    work_period = _get_work_period_or_404(db, year=year, month=month)
    if work_period.status in {
        WorkFlowStatus.PENDING_CONTRACTOR_STATEMENT,
        WorkFlowStatus.PENDING_SECURITY_AUDIT,
        WorkFlowStatus.FINAL_APPROVED,
    }:
        raise HTTPException(
            status_code=400,
            detail="This work period has already passed the aggregation approval stage.",
        )

    work_period.status = WorkFlowStatus.PENDING_CONTRACTOR_STATEMENT

    try:
        db.commit()
        db.refresh(work_period)
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail=f"Failed to approve aggregation: {exc}",
        ) from exc

    return WorkPeriodApprovalOut(
        success=True,
        status=work_period.status,
        message="Aggregation approved successfully.",
    )


def _get_latest_contractor_statement_for_work_period(
    db: Session,
    *,
    work_period_id: int,
) -> Optional[ContractorStatement]:
    return (
        db.query(ContractorStatement)
        .filter(ContractorStatement.work_period_id == work_period_id)
        .order_by(ContractorStatement.created_at.desc(), ContractorStatement.id.desc())
        .first()
    )


def _generate_finance_tracking_number() -> str:
    """Generate the finance tracking number used before a unique code exists."""
    now = datetime.utcnow()
    persian_year = now.year - 622 if (now.month, now.day) < (3, 21) else now.year - 621
    return f"TRX-{persian_year}-{uuid.uuid4().hex[:8]}"


def _extract_attendance_work_period_id(transaction) -> Optional[int]:
    form_data = transaction.form_data if isinstance(transaction.form_data, dict) else {}
    raw_work_period_id = form_data.get("attendance_work_period_id", form_data.get("work_period_id"))
    if raw_work_period_id is None:
        return None

    try:
        return int(raw_work_period_id)
    except (TypeError, ValueError):
        return None


def _get_attendance_finance_activity(
    db: Session,
    *,
    required: bool = True,
):
    from app.models import Activity

    activity = (
        db.query(Activity)
        .filter(Activity.name == ATTENDANCE_FINANCE_ACTIVITY_NAME)
        .first()
    )
    if activity is None and required:
        raise HTTPException(
            status_code=500,
            detail=(
                "Attendance finance handoff activity "
                f"'{ATTENDANCE_FINANCE_ACTIVITY_NAME}' was not found in the Financial Rules Matrix."
            ),
        )
    return activity


def _find_attendance_financial_transaction(
    db: Session,
    *,
    work_period_id: int,
    activity_id: Optional[int] = None,
):
    from app.models import Transaction

    resolved_activity_id = activity_id
    if resolved_activity_id is None:
        activity = _get_attendance_finance_activity(db, required=False)
        if activity is None:
            return None
        resolved_activity_id = activity.id

    candidates = (
        db.query(Transaction)
        .filter(Transaction.activity_id == resolved_activity_id)
        .order_by(Transaction.created_at.desc(), Transaction.id.desc())
        .all()
    )
    for transaction in candidates:
        if _extract_attendance_work_period_id(transaction) == work_period_id:
            return transaction
    return None


def _ensure_attendance_financial_transaction(
    db: Session,
    *,
    work_period: WorkPeriod,
    year: int,
    month: int,
    final_payable_amount: Decimal,
    initiated_by,
):
    from app.models import Transaction, User

    activity = _get_attendance_finance_activity(db)
    existing_transaction = _find_attendance_financial_transaction(
        db,
        work_period_id=work_period.id,
        activity_id=activity.id,
    )
    if existing_transaction:
        return existing_transaction

    contractor_statement = _get_latest_contractor_statement_for_work_period(
        db,
        work_period_id=work_period.id,
    )

    contractor_id = None
    beneficiary_name = "پیمانکار"
    if contractor_statement and contractor_statement.sheet and contractor_statement.sheet.created_by_id:
        contractor_id = contractor_statement.sheet.created_by_id
        contractor_user = db.query(User).filter(User.id == contractor_id).first()
        if contractor_user and contractor_user.full_name:
            beneficiary_name = contractor_user.full_name

    amount_value = int(final_payable_amount)
    transaction = Transaction(
        tracking_number=_generate_finance_tracking_number(),
        unique_code=None,
        status="PENDING_L1",
        current_approval_level=0,
        created_by_id=initiated_by.id,
        zone_id=getattr(initiated_by, "default_zone_id", None),
        department_id=getattr(initiated_by, "default_dept_id", None),
        section_id=getattr(initiated_by, "default_section_id", None),
        activity_id=activity.id,
        amount=amount_value,
        beneficiary_name=beneficiary_name,
        contract_number=f"ATT-{year}-{str(month).zfill(2)}",
        description=f"Attendance reconciliation payout for {year}/{str(month).zfill(2)}",
        fiscal_year=str(year),
        form_data={
            "source_module": "ATTENDANCE",
            "attendance_work_period_id": work_period.id,
            "work_period_id": work_period.id,
            "year": year,
            "month": month,
            "contractor_user_id": contractor_id,
            "contractor_name": beneficiary_name,
            "final_payable_amount": amount_value,
        },
    )
    db.add(transaction)
    db.flush()
    return transaction


def get_work_period_reconciliation(
    db: Session,
    *,
    year: int,
    month: int,
) -> AttendanceReconciliationOut:
    work_period = _get_work_period_or_404(db, year=year, month=month)

    attendance_records = (
        db.query(AttendanceRecord)
        .filter(AttendanceRecord.work_period_id == work_period.id)
        .all()
    )

    region_total = Decimal("0")
    for record in attendance_records:
        region_total += calculate_record_legal_value(record)

    contractor_statement = _get_latest_contractor_statement_for_work_period(
        db,
        work_period_id=work_period.id,
    )
    contractor_total = _as_decimal(
        contractor_statement.declared_total_amount if contractor_statement else Decimal("0")
    )

    penalties = (
        db.query(PenaltyRecord)
        .filter(PenaltyRecord.work_period_id == work_period.id)
        .order_by(PenaltyRecord.id.desc())
        .all()
    )
    total_penalties_amount = sum(
        (_as_decimal(penalty.total_amount) for penalty in penalties),
        Decimal("0"),
    )
    final_payable_amount = contractor_total - total_penalties_amount
    financial_transaction = _find_attendance_financial_transaction(
        db,
        work_period_id=work_period.id,
    )

    return AttendanceReconciliationOut(
        year=year,
        month=month,
        status=work_period.status,
        region_total=region_total,
        contractor_total=contractor_total,
        total_penalties_amount=total_penalties_amount,
        final_payable_amount=final_payable_amount,
        financial_transaction_id=financial_transaction.id if financial_transaction else None,
        financial_transaction_status=financial_transaction.status if financial_transaction else None,
        financial_tracking_number=financial_transaction.tracking_number if financial_transaction else None,
        financial_unique_code=work_period.financial_unique_code,
        penalties=[PenaltyRecordOut.model_validate(penalty) for penalty in penalties],
    )


def create_penalty_record_for_work_period(
    db: Session,
    *,
    year: int,
    month: int,
    payload: PenaltyRecordCreate,
) -> PenaltyRecordOut:
    work_period = _get_work_period_or_404(db, year=year, month=month)
    if work_period.status == WorkFlowStatus.FINAL_APPROVED:
        raise HTTPException(
            status_code=400,
            detail="Cannot add penalties to a finalized work period.",
        )

    total_amount = _as_decimal(payload.unit_amount) * payload.count
    penalty_record = PenaltyRecord(
        work_period_id=work_period.id,
        description=_normalize_text(payload.description),
        count=payload.count,
        unit_amount=_as_decimal(payload.unit_amount),
        total_amount=total_amount,
    )
    db.add(penalty_record)

    try:
        db.commit()
        db.refresh(penalty_record)
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail=f"Failed to save penalty record: {exc}",
        ) from exc

    return PenaltyRecordOut.model_validate(penalty_record)


def finalize_work_period(
    db: Session,
    *,
    year: int,
    month: int,
    initiated_by,
) -> WorkPeriodApprovalOut:
    work_period = _get_work_period_or_404(db, year=year, month=month)
    activity = _get_attendance_finance_activity(db)
    existing_transaction = _find_attendance_financial_transaction(
        db,
        work_period_id=work_period.id,
        activity_id=activity.id,
    )

    if work_period.status in {WorkFlowStatus.PENDING_UNIQUE_CODE, WorkFlowStatus.SENT_TO_FINANCE}:
        return WorkPeriodApprovalOut(
            success=True,
            status=work_period.status,
            message="Work period has already been handed off to finance.",
            financial_transaction_id=existing_transaction.id if existing_transaction else None,
            financial_transaction_status=existing_transaction.status if existing_transaction else None,
            financial_tracking_number=existing_transaction.tracking_number if existing_transaction else None,
        )

    # Resolve reconciliation before mutating workflow state so failures in the
    # financial handoff do not leave the period in a partially transitioned state.
    reconciliation = get_work_period_reconciliation(db, year=year, month=month)
    final_payable_amount = reconciliation.final_payable_amount

    logger.info("Preparing finance handoff for work period %s/%s.", year, month)
    contractor_statement = _get_latest_contractor_statement_for_work_period(
        db,
        work_period_id=work_period.id,
    )
    contractor_id = None
    if contractor_statement and contractor_statement.sheet and contractor_statement.sheet.created_by_id:
        contractor_id = contractor_statement.sheet.created_by_id

    # Hook into the financial queue by creating a real Transaction record.
    from app.models import Transaction, User

    tracking_number = _generate_finance_tracking_number()
    beneficiary_name = "پیمانکار"
    if contractor_id:
        contractor_user = db.query(User).filter(User.id == contractor_id).first()
        if contractor_user and contractor_user.full_name:
            beneficiary_name = contractor_user.full_name

    payload = {
        "tracking_number": tracking_number,
        "unique_code": None,
        "status": "DRAFT",
        "current_approval_level": 0,
        "activity_id": activity.id,
        "activity_name": activity.name,
        "amount": int(final_payable_amount),
        "beneficiary_name": beneficiary_name,
        "description": f"تسویه حساب کارکرد منابع انسانی - سال {year} ماه {month}",
        "created_by_id": initiated_by.id,
        "fiscal_year": str(year),
        "form_data": {
            "source_module": "ATTENDANCE",
            "attendance_work_period_id": work_period.id,
            "work_period_id": work_period.id,
            "year": year,
            "month": month,
            "contractor_user_id": contractor_id,
            "contractor_name": beneficiary_name,
            "final_payable_amount": int(final_payable_amount),
        },
    }
    logger.info("Attendance finance handoff payload prepared: %s", payload)

    transaction = Transaction(
        tracking_number=payload["tracking_number"],
        unique_code=payload["unique_code"],  # explicitly NULL until Transaction reaches APPROVED
        status=payload["status"],
        current_approval_level=payload["current_approval_level"],
        zone_id=getattr(initiated_by, "default_zone_id", None),
        department_id=getattr(initiated_by, "default_dept_id", None),
        section_id=getattr(initiated_by, "default_section_id", None),
        activity_id=payload["activity_id"],
        amount=payload["amount"],
        description=f"تسویه حساب کارکرد منابع انسانی - سال {year} ماه {month}",
        beneficiary_name=payload["beneficiary_name"],
        contract_number=f"ATT-{year}-{str(month).zfill(2)}",
        created_by_id=payload["created_by_id"],
        fiscal_year=payload["fiscal_year"],
        form_data=payload["form_data"],
    )
    try:
        db.add(transaction)
        db.flush()
        work_period.status = WorkFlowStatus.PENDING_UNIQUE_CODE
        db.commit()
        db.refresh(transaction)
        db.refresh(work_period)
    except Exception as exc:
        db.rollback()
        logger.exception("Failed to hand off work period %s/%s to finance.", year, month)
        raise HTTPException(
            status_code=500,
            detail=f"Failed to finalize work period: {exc}",
        ) from exc

    return WorkPeriodApprovalOut(
        success=True,
        status=work_period.status,
        message="Work period finalized and handed off to the finance workflow.",
        financial_transaction_id=transaction.id,
        financial_transaction_status=transaction.status,
        financial_tracking_number=transaction.tracking_number,
    )


def generate_unique_code_for_period(
    db: Session,
    *,
    year: int,
    month: int,
) -> WorkPeriodApprovalOut:
    raise HTTPException(
        status_code=410,
        detail="Direct unique-code generation is removed. Continue the workflow in the finance module.",
    )

    work_period = _get_work_period_or_404(db, year=year, month=month)

    if work_period.status == WorkFlowStatus.SENT_TO_FINANCE:
        return WorkPeriodApprovalOut(
            success=True,
            status=work_period.status,
            message="Unique code has already been generated and sent to finance.",
        )

    if work_period.status != WorkFlowStatus.PENDING_UNIQUE_CODE:
        raise HTTPException(
            status_code=400,
            detail=f"Work period must be in PENDING_UNIQUE_CODE status to generate a unique code. Current status: {work_period.status}",
        )

    # Find the Transaction created for this work period
    from app.models import Transaction

    activity = _get_attendance_finance_activity(db)
    transaction = (
        db.query(Transaction)
        .filter(
            Transaction.activity_id == activity.id,
            Transaction.status == "DRAFT",
        )
        .order_by(Transaction.created_at.desc())
        .first()
    )

    if not transaction:
        # Fallback search in form_data if simple filter fails
        all_drafts = db.query(Transaction).filter(Transaction.status == "DRAFT").all()
        for t in all_drafts:
            if t.form_data and t.form_data.get("work_period_id") == work_period.id:
                transaction = t
                break

    if not transaction:
        raise HTTPException(
            status_code=404,
            detail="No draft transaction found for this work period. Please finalize the period first.",
        )

    # Logic to generate the unique code
    # Pattern: HR-{year}{month}-{tracking_suffix}
    tracking_suffix = (
        transaction.tracking_number.split("-")[-1]
        if "-" in transaction.tracking_number
        else uuid.uuid4().hex[:8]
    )
    generated_code = f"HR-{year}{str(month).zfill(2)}-{tracking_suffix.upper()}"

    try:
        transaction.unique_code = generated_code
        transaction.status = "APPROVED"
        transaction.current_approval_level = 4

        work_period.financial_unique_code = generated_code
        work_period.status = WorkFlowStatus.SENT_TO_FINANCE

        db.commit()
        db.refresh(work_period)
    except Exception as exc:
        db.rollback()
        logger.error(f"Failed to generate unique code: {str(exc)}")
        raise HTTPException(
            status_code=500,
            detail=f"Failed to generate unique code: {exc}",
        ) from exc

    return WorkPeriodApprovalOut(
        success=True,
        status=work_period.status,
        message=f"کد یکتا {generated_code} با موفقیت صادر و پرونده به امور مالی ارجاع شد.",
    )


def ingest_excel(
    file: UploadFile,
    title: str,
    year: int,
    month: int,
    region_name: str,
    creator_id: int,
    db: Session,
) -> RegionalAttendanceSheet:
    contents = file.file.read()
    columns, rows = _read_excel_rows(contents)
    template_kind = _infer_template_kind(columns)
    if template_kind == "statement_6":
        raise HTTPException(
            status_code=422,
            detail=(
                "Uploaded file appears to be the 6-column contractor statement template. "
                "Please upload it via /api/attendance/upload/statement."
            ),
        )
    if template_kind == "unknown":
        raise HTTPException(
            status_code=422,
            detail=(
                "Unsupported Excel template. Only 18-column regional attendance and "
                "6-column contractor statement templates are accepted."
            ),
        )

    column_mapping = _build_column_mapping(columns)

    try:
        work_period = get_or_create_work_period(db, year=year, month=month)
        sheet = db.query(RegionalAttendanceSheet).filter(
            RegionalAttendanceSheet.year == year,
            RegionalAttendanceSheet.month == month,
            RegionalAttendanceSheet.region == region_name,
        ).first()

        had_existing_submission_for_period = bool(
            sheet
            and sheet.work_period_id == work_period.id
            and (sheet.row_count or 0) > 0
        )

        if not sheet:
            sheet = RegionalAttendanceSheet(
                title=title,
                year=year,
                month=month,
                region=region_name,
                original_filename=file.filename or "unknown.xlsx",
                row_count=0,
                status=AttendanceSheetStatus.DRAFT,
                created_by_id=creator_id,
                work_period_id=work_period.id,
            )
            db.add(sheet)
            db.flush()
        else:
            db.query(AttendanceRecord).filter(
                AttendanceRecord.sheet_id == sheet.id
            ).delete(synchronize_session=False)
            sheet.title = title
            sheet.original_filename = file.filename or "unknown.xlsx"
            sheet.created_by_id = creator_id
            sheet.status = AttendanceSheetStatus.DRAFT
            sheet.rejection_reason = None
            sheet.row_count = 0
            sheet.work_period_id = work_period.id

        records: list[AttendanceRecord] = []
        for idx, row in enumerate(rows):
            if _should_skip_row(row, column_mapping):
                continue

            record_kwargs = {
                "sheet": sheet,
                "row_number": EXCEL_DATA_START_ROW_NUMBER + idx,
                "work_period_id": work_period.id,
            }

            for field_name, parser in FIELD_PARSERS.items():
                record_kwargs[field_name] = parser(_get_row_value(row, column_mapping, field_name))

            if not record_kwargs["full_name"]:
                continue

            records.append(AttendanceRecord(**record_kwargs))

        sheet.row_count = len(records)
        if not had_existing_submission_for_period:
            work_period.total_regions_uploaded = (work_period.total_regions_uploaded or 0) + 1
        db.add_all(records)
        db.commit()
        db.refresh(sheet)
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Database write failed: {str(exc)}") from exc

    return sheet


def parse_attendance_excel_to_db(
    file: UploadFile,
    title: str,
    year: int,
    month: int,
    region_name: str,
    creator_id: int,
    db: Session,
) -> RegionalAttendanceSheet:
    return ingest_excel(
        file=file,
        title=title,
        year=year,
        month=month,
        region_name=region_name,
        creator_id=creator_id,
        db=db,
    )


STATEMENT_FIELD_PARSERS: dict[str, Callable[[object], object]] = {
    "national_id": _safe_string,
    "full_name": _safe_string,
    "gross_amount": _safe_decimal,
    "tax_deduction": _safe_decimal,
    "insurance_deduction": _safe_decimal,
    "net_payable_amount": _safe_decimal,
}

STATEMENT_FIELD_ORDER = [
    "national_id",
    "full_name",
    "gross_amount",
    "tax_deduction",
    "insurance_deduction",
    "net_payable_amount",
]

STATEMENT_COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "national_id": ("شماره ملی", "کد ملی", "national_id"),
    "full_name": ("نام و نام خانوادگی", "نام ونام خانوادگی", "full_name"),
    "gross_amount": ("مبلغ ناخالص", "ناخالص", "gross_amount"),
    "tax_deduction": ("مالیات", "کسورات مالیات", "tax"),
    "insurance_deduction": ("بیمه", "کسورات بیمه", "insurance"),
    "net_payable_amount": ("مبلغ قابل پرداخت", "خالص پرداختی", "مبلغ خالص", "net_amount"),
}

STATEMENT_TOTAL_ROW_MARKERS = {
    _compact_text("جمع کل"),
    _compact_text("جمع کل ریال"),
    _compact_text("جمع کل (ریال)"),
    _compact_text("جمع‌کل"),
}


def _count_alias_matches(columns, aliases_map: dict[str, tuple[str, ...]]) -> int:
    normalized_columns = {
        _compact_text(str(column))
        for column in columns
        if not _is_missing(column)
    }
    matches = 0
    for aliases in aliases_map.values():
        if any(_compact_text(alias) in normalized_columns for alias in aliases):
            matches += 1
    return matches


def _infer_template_kind(columns) -> str:
    """Return one of: attendance_18, statement_6, unknown."""
    effective_columns = [column for column in columns if not _is_missing(column)]
    column_count = len(effective_columns)

    attendance_hits = _count_alias_matches(effective_columns, ATTENDANCE_COLUMN_ALIASES)
    statement_hits = _count_alias_matches(effective_columns, STATEMENT_COLUMN_ALIASES)

    if statement_hits >= 4 and column_count <= 8:
        return "statement_6"
    if attendance_hits >= 6 and column_count >= 10:
        return "attendance_18"

    # Fallback by column count for header variants that still follow known templates.
    if column_count <= 8:
        return "statement_6"
    if column_count >= 12:
        return "attendance_18"
    return "unknown"


def _build_statement_column_mapping(columns) -> dict[str, object]:
    original_columns = list(columns)
    normalized_lookup: dict[str, object] = {}

    for column_name in original_columns:
        normalized = _compact_text(str(column_name))
        if normalized and normalized not in normalized_lookup:
            normalized_lookup[normalized] = column_name

    mapping: dict[str, object] = {}
    used_columns: set[object] = set()

    for field_name in STATEMENT_FIELD_ORDER:
        for alias in STATEMENT_COLUMN_ALIASES[field_name]:
            matched_column = normalized_lookup.get(_compact_text(alias))
            if matched_column is not None and matched_column not in used_columns:
                mapping[field_name] = matched_column
                used_columns.add(matched_column)
                break

    remaining_fields = [field for field in STATEMENT_FIELD_ORDER if field not in mapping]
    remaining_columns = [column for column in original_columns if column not in used_columns]

    if remaining_columns and _compact_text(str(remaining_columns[0])) in {
        _compact_text(value) for value in SERIAL_COLUMN_KEYS
    }:
        remaining_columns = remaining_columns[1:]

    if remaining_columns:
        # Fallback for 6-column contractor statement files with non-standard headers.
        for field_name, column_name in zip(remaining_fields, remaining_columns):
            mapping[field_name] = column_name

    if "full_name" not in mapping:
        raise HTTPException(
            status_code=422,
            detail="Excel headers do not match the expected Contractor Statement template.",
        )

    return mapping


def ingest_statement_excel(
    file: UploadFile,
    title: str,
    year: int,
    month: int,
    region_name: str,
    declared_total_amount: float,
    creator_id: int,
    db: Session,
) -> RegionalAttendanceSheet:
    contents = file.file.read()
    columns, rows = _read_excel_rows(contents)
    template_kind = _infer_template_kind(columns)
    if template_kind == "attendance_18":
        raise HTTPException(
            status_code=422,
            detail=(
                "Uploaded file appears to be the 18-column regional attendance template. "
                "Please upload it via /api/attendance/upload."
            ),
        )
    if template_kind == "unknown":
        raise HTTPException(
            status_code=422,
            detail=(
                "Unsupported Excel template. Only 18-column regional attendance and "
                "6-column contractor statement templates are accepted."
            ),
        )

    column_mapping = _build_statement_column_mapping(columns)
    declared_total_decimal = _as_decimal(declared_total_amount)

    if declared_total_decimal <= 0:
        raise HTTPException(
            status_code=422,
            detail="Declared total amount must be greater than zero.",
        )

    try:
        work_period = get_or_create_work_period(db, year=year, month=month)
        if work_period.status == WorkFlowStatus.DRAFT_REGIONS:
            logger.warning(
                "Contractor statement uploaded while work period is still in DRAFT_REGIONS "
                "(year=%s, month=%s, work_period_id=%s).",
                year,
                month,
                work_period.id,
            )

        sheet = db.query(RegionalAttendanceSheet).filter(
            RegionalAttendanceSheet.year == year,
            RegionalAttendanceSheet.month == month,
            RegionalAttendanceSheet.region == region_name,
        ).first()

        if not sheet:
            sheet = RegionalAttendanceSheet(
                title=title,
                year=year,
                month=month,
                region=region_name,
                original_filename=file.filename or "statement.xlsx",
                row_count=0,
                status=AttendanceSheetStatus.DRAFT,
                declared_total_amount=declared_total_decimal,
                created_by_id=creator_id,
                work_period_id=work_period.id,
            )
            db.add(sheet)
            db.flush()
        else:
            # Update declared amount on existing sheet if being re-uploaded
            sheet.declared_total_amount = declared_total_decimal
            sheet.work_period_id = work_period.id

        statement = ContractorStatement(
            sheet=sheet,
            year=year,
            month=month,
            region_name=region_name,
            declared_total_amount=declared_total_decimal,
            work_period_id=work_period.id,
        )
        db.add(statement)
        db.flush()

        items: list[ContractorStatementItem] = []
        for idx, row in enumerate(rows):
            if _should_skip_row(row, column_mapping):
                continue

            item_kwargs = {
                "statement": statement,
                "sheet": sheet,
                "row_number": EXCEL_DATA_START_ROW_NUMBER + idx,
            }
            if hasattr(ContractorStatementItem, "work_period_id"):
                item_kwargs["work_period_id"] = work_period.id

            for field_name, parser in STATEMENT_FIELD_PARSERS.items():
                item_kwargs[field_name] = parser(_get_row_value(row, column_mapping, field_name))

            if not item_kwargs["full_name"]:
                continue

            items.append(ContractorStatementItem(**item_kwargs))

        db.add_all(items)
        db.commit()
        db.refresh(sheet)
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Database write failed for statement: {str(exc)}") from exc

    return sheet


def parse_contractor_statement_excel(
    file: UploadFile,
    title: str,
    year: int,
    month: int,
    region_name: str,
    declared_total_amount: float,
    creator_id: int,
    db: Session,
) -> RegionalAttendanceSheet:
    return ingest_statement_excel(
        file=file,
        title=title,
        year=year,
        month=month,
        region_name=region_name,
        declared_total_amount=declared_total_amount,
        creator_id=creator_id,
        db=db,
    )


def reject_attendance_sheet(
    db: Session,
    *,
    sheet_id: int,
    reason: str,
    actor_id: int,
) -> RegionalAttendanceSheet:
    normalized_reason = _normalize_text(reason)
    if not normalized_reason:
        raise HTTPException(status_code=422, detail="Rejection reason is required.")

    sheet = (
        db.query(RegionalAttendanceSheet)
        .filter(RegionalAttendanceSheet.id == sheet_id)
        .first()
    )
    if not sheet:
        raise HTTPException(status_code=404, detail="Sheet not found.")

    sheet.status = AttendanceSheetStatus.REJECTED
    sheet.rejection_reason = normalized_reason
    sheet.company_rep_id = actor_id
    sheet.work_period_id = None

    create_notification(
        db,
        user_id=sheet.created_by_id,
        message=f"فایل کارکرد شما به دلیل: {normalized_reason} رد شد. لطفاً اصلاح و مجدداً آپلود کنید.",
    )

    try:
        db.commit()
        db.refresh(sheet)
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail=f"Failed to reject attendance sheet: {exc}",
        ) from exc

    return sheet


_STATUS_MAP: dict[str, str] = {
    AttendanceSheetStatus.DRAFT.value: WorkflowStatusLabel.DRAFT.value,
    AttendanceSheetStatus.CREATOR_APPROVED.value: WorkflowStatusLabel.CREATOR_APPROVED.value,
    AttendanceSheetStatus.COMPANY_REP_APPROVED.value: WorkflowStatusLabel.COMPANY_REP_APPROVED.value,
    AttendanceSheetStatus.SECURITY_APPROVED.value: WorkflowStatusLabel.SECURITY_APPROVED.value,
    AttendanceSheetStatus.GM_APPROVED.value: WorkflowStatusLabel.GM_APPROVED.value,
    AttendanceSheetStatus.NEEDS_CORRECTION.value: WorkflowStatusLabel.NEEDS_CORRECTION.value,
    AttendanceSheetStatus.REJECTED.value: WorkflowStatusLabel.REJECTED.value,
}


def _resolve_status_label(raw_status: str) -> str:
    return _STATUS_MAP.get(raw_status, raw_status)


def _as_float(value: Decimal | int | float | None) -> float:
    if value is None:
        return 0.0
    return float(value)


def _as_int(value: Decimal | int | float | None) -> int:
    if value is None:
        return 0
    return int(value)


def _as_decimal(value: Decimal | int | float | None) -> Decimal:
    if value is None:
        return Decimal("0")
    if isinstance(value, Decimal):
        return value
    return Decimal(str(value))


def _is_statement_total_row(item: ContractorStatementItem) -> bool:
    label = _compact_text(item.full_name or "")
    if not label:
        return False
    return label in STATEMENT_TOTAL_ROW_MARKERS or "جمعکل" in label


def _get_statement_total_candidate(item: ContractorStatementItem) -> Optional[Decimal]:
    for value in (
        item.gross_amount,
        item.net_payable_amount,
        item.tax_deduction,
        item.insurance_deduction,
    ):
        if value is not None:
            return value
    return None


def get_attendance_aggregation(
    db: Session,
    *,
    region: Optional[str] = None,
    status: Optional[str] = None,
) -> AttendanceAggregationOut:
    totals_subquery = (
        db.query(
            AttendanceRecord.sheet_id.label("sheet_id"),
            func.count(AttendanceRecord.id).label("total_personnel"),
            func.coalesce(func.sum(AttendanceRecord.overtime_hours), 0).label(
                "total_worked_hours"
            ),
            func.coalesce(func.sum(AttendanceRecord.total_average_per_capita), 0).label(
                "total_amount"
            ),
        )
        .group_by(AttendanceRecord.sheet_id)
        .subquery()
    )

    query = (
        db.query(
            RegionalAttendanceSheet,
            func.coalesce(totals_subquery.c.total_personnel, 0).label("total_personnel"),
            func.coalesce(totals_subquery.c.total_worked_hours, 0).label(
                "total_worked_hours"
            ),
            func.coalesce(totals_subquery.c.total_amount, 0).label("total_amount"),
        )
        .outerjoin(totals_subquery, totals_subquery.c.sheet_id == RegionalAttendanceSheet.id)
        .filter(RegionalAttendanceSheet.row_count > 0)
        .order_by(RegionalAttendanceSheet.created_at.desc())
    )

    if region:
        query = query.filter(RegionalAttendanceSheet.region.ilike(f"%{region.strip()}%"))

    if status:
        normalized_status = status.strip().upper()
        try:
            status_enum = AttendanceSheetStatus(normalized_status)
        except ValueError as exc:
            raise HTTPException(status_code=422, detail="وضعیت حضور و غیاب نامعتبر است.") from exc
        query = query.filter(RegionalAttendanceSheet.status == status_enum)

    rows = query.all()
    items: list[AttendanceAggregationItem] = []

    for sheet_obj, total_personnel, total_worked_hours, total_amount in rows:
        raw_status = (
            sheet_obj.status.value
            if isinstance(sheet_obj.status, AttendanceSheetStatus)
            else str(sheet_obj.status)
        )
        sheet_meta = SheetMetaSummary(
            sheet_id=sheet_obj.id,
            title=sheet_obj.title,
            year=sheet_obj.year,
            month=sheet_obj.month,
            region_name=sheet_obj.region,
            original_filename=sheet_obj.original_filename,
            row_count=sheet_obj.row_count,
            created_at=sheet_obj.created_at,
            creator_approved_at=sheet_obj.creator_approved_at,
            company_rep_approved_at=sheet_obj.company_rep_approved_at,
            security_approved_at=sheet_obj.security_approved_at,
            gm_approved_at=sheet_obj.gm_approved_at,
        )
        items.append(
            AttendanceAggregationItem(
                sheet=sheet_meta,
                workflow_status=raw_status,
                workflow_status_label=_resolve_status_label(raw_status),
                total_personnel=_as_int(total_personnel),
                total_worked_hours=_as_float(total_worked_hours),
                total_amount=_as_int(total_amount),
            )
        )

    return AttendanceAggregationOut(count=len(items), items=items)


def update_attendance_record(
    db: Session,
    *,
    record_id: int,
    payload: "app.schemas.attendance.AttendanceRecordUpdate", # type: ignore
    user_id: int,
) -> app.models.attendance.AttendanceRecord: # type: ignore
    from app.models.attendance import AttendanceRecord, AttendanceAuditLog
    
    record = db.query(AttendanceRecord).filter(AttendanceRecord.id == record_id).first()
    if not record:
        raise HTTPException(status_code=404, detail="Attendance record not found.")

    update_data = payload.model_dump(exclude_unset=True)
    if not update_data:
        return record  # Nothing to update

    logs = []
    
    for field_name, new_value in update_data.items():
        old_value = getattr(record, field_name, None)
        
        # Compare as strings or float, because Decimal vs float vs string
        # For simple difference check, converting both to strings handles Decimals/Floats
        str_old = str(old_value) if old_value is not None else None
        str_new = str(new_value) if new_value is not None else None

        if str_old != str_new:
            setattr(record, field_name, new_value)
            logs.append(
                AttendanceAuditLog(
                    record_id=record.id,
                    changed_by_user_id=user_id,
                    field_name=field_name,
                    old_value=str_old,
                    new_value=str_new,
                )
            )

    if logs:
        db.add_all(logs)
        try:
            db.commit()
            db.refresh(record)
        except Exception as exc:
            db.rollback()
            raise HTTPException(status_code=500, detail=f"Failed to update record: {exc}") from exc

    return record


def get_attendance_audit_logs(
    db: Session,
    *,
    record_id: int,
) -> list["app.schemas.attendance.AttendanceAuditLogOut"]: # type: ignore
    from app.models.attendance import AttendanceAuditLog
    from app.models import User
    from app.schemas.attendance import AttendanceAuditLogOut
    
    logs = (
        db.query(AttendanceAuditLog, User)
        .outerjoin(User, AttendanceAuditLog.changed_by_user_id == User.id)
        .filter(AttendanceAuditLog.record_id == record_id)
        .order_by(AttendanceAuditLog.created_at.desc())
        .all()
    )

    out_logs = []
    for log, user in logs:
        out_log = AttendanceAuditLogOut.model_validate(log)
        out_log.changed_by_full_name = user.full_name if user else None
        out_logs.append(out_log)

    return out_logs


def get_work_period_region_records(
    db: Session,
    *,
    year: int,
    month: int,
    region_name: str,
) -> list[app.models.attendance.AttendanceRecord]: # type: ignore
    from app.models.attendance import AttendanceRecord, RegionalAttendanceSheet
    
    records = (
        db.query(AttendanceRecord)
        .join(RegionalAttendanceSheet, AttendanceRecord.sheet_id == RegionalAttendanceSheet.id)
        .filter(
            RegionalAttendanceSheet.year == year,
            RegionalAttendanceSheet.month == month,
            RegionalAttendanceSheet.region == region_name,
        )
        .order_by(AttendanceRecord.row_number.asc())
        .all()
    )
    return records
